"""tests/test_build.py · tests/test_operations.py가 공유하는 작업 공간 도우미(openpyxl 생성 문서 + 실제 Reader 격리)."""

from __future__ import annotations

import copy
import types

from openpyxl import Workbook

from schema import service as service_module
from schema.service import Service

SCHEMA = {
    "format": "parsing-schema",
    "schema_version": "3.0",
    "schema_key": "process_standard",
    "schema_name": "공정 데이터 표준",
    "fields": [
        {"field_key": "process", "name": "공정 정보", "type": "group", "level": 1},
        {"field_key": "lot", "name": "배치", "type": "text", "level": 2, "parents": "process", "aliases": ["LOT"]},
        {"field_key": "temperature", "name": "온도", "type": "decimal", "unit": "°C", "level": 2, "parents": ["process"]},
        {"field_key": "pressure", "name": "압력", "type": "decimal", "unit": "bar", "level": 2, "parents": ["process"]},
        {"field_key": "note", "name": "비고", "type": "text"},
    ],
}

UNITS_YAML = """
version: "t"
aliases:
  "℃": "°C"
  "degC": "°C"
dimensions:
  temperature:
    "°C": 1.0
    "K": {factor: 1.0, offset: -273.15}
    "°F": {factor: 0.5555555555555556, offset: -17.77777777777778}
  pressure:
    "bar": 1.0
    "kPa": 0.01
"""


def profile_definition(**over):
    """온도(list, °C 단위 영역)·배치(list)·서명자(scalar) 규칙 3개. sheet 이름 '공정 기록'."""
    base = {
        "format": "parsing-profile",
        "schema_version": "3.0",
        "profile_name": "공정데이터_A양식",
        "schema_key": "process_standard",
        "sheet_roles": {"main": {"cardinality": "one", "match": {"name": "공정 기록"}}},
        "anchors": {
            "hdr_temp": {"sheet_role": "main", "find": {"texts": ["온도"], "within": "A1:Z60"}},
            "hdr_lot": {"sheet_role": "main", "find": {"regex": "^(배치|LOT)$", "within": "A1:Z60"}},
            "sign": {"sheet_role": "main", "find": {"texts": ["서명"], "within": "A1:Z60"}},
        },
        "rules": [
            {
                "rule_key": "lot",
                "rule_name": "배치",
                "field_key": "lot",
                "selector": {
                    "key": {"areas": [{"sheet_role": "main", "anchor": "hdr_lot"}]},
                    "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 20, "cols": 1}}], "cardinality": "list", "axis": "down", "stop": {"kind": "blank_run", "count": 2}},
                },
                "value_spec": {"type": "text"},
                "record_spec": {"scope": ["process-table"], "key": "physical_row"},
            },
            {
                "rule_key": "temperature",
                "rule_name": "온도",
                "field_key": "temperature",
                "selector": {
                    "key": {"areas": [{"sheet_role": "main", "anchor": "hdr_temp"}]},
                    "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 20, "cols": 1}}], "cardinality": "list", "axis": "down", "stop": {"kind": "blank_run", "count": 2}},
                    "unit": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1, "anchor": "hdr_temp"}}]},
                },
                "value_spec": {"type": "decimal", "unit": "°C"},
                "record_spec": {"scope": ["process-table"], "key": "physical_row"},
                "relations": [{"to_rule": "lot", "kind": "same_row"}],
            },
            {
                "rule_key": "signer",
                "rule_name": "서명자",
                "field_key": "note",
                "selector": {"key": {"areas": [{"sheet_role": "main", "anchor": "sign"}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1}}]}},
                "value_spec": {"type": "text"},
            },
        ],
    }
    base.update(over)
    return base


def kelvin_profile():
    """온도를 K 단위로 읽는 두 번째 프로파일(단위 셀이 'K'인 문서용; 단위 불일치 큐)."""
    definition = copy.deepcopy(profile_definition(profile_name="공정데이터_K양식"))
    definition["rules"][1]["value_spec"] = {"type": "decimal", "unit": "K"}
    return definition


def variant_profile(name="공정데이터_B양식", sheet="공정 기록"):
    """서명자를 C55(서명 오른쪽 두 번째 칸)에서 읽는 변형 프로파일(같은 필드에 다른 출처 → 빌드 충돌)."""
    definition = copy.deepcopy(profile_definition(profile_name=name))
    definition["sheet_roles"] = {"main": {"cardinality": "one", "match": {"name": sheet}}}
    definition["rules"][2]["selector"]["value"] = {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 2}}]}
    return definition


def build_workbook(path, shift=0, temps=(10.5, 20, 30), lots=("L1", "L2", "L3"), signer="홍길동", unit="°C", sheet="공정 기록", second_signer="이순신", temp_label="온도"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = "공정 A"
    ws.cell(3 + shift, 2, "LOT")
    ws.cell(3 + shift, 3, temp_label)
    ws.cell(3 + shift, 4, unit)
    for n, (lot, temp) in enumerate(zip(lots, temps)):
        ws.cell(4 + shift + n, 2, lot)
        ws.cell(4 + shift + n, 3, temp)
    ws["A55"], ws["B55"], ws["C55"] = "서명", signer, second_signer
    wb.save(path)
    return path


class World:
    """작업 공간 + Reader 호출 기록. 렌더 클라이언트는 무효화 기록만 남기는 대역."""

    def __init__(self, root, monkeypatch):
        (root / "data/raw").mkdir(parents=True)
        self.root = root
        self.calls = []
        self.invalidated = []
        original_result, original_events = service_module.reader_result, service_module.reader_events

        def counted_result(root_, provider, principal, operation, payload, checkpoint=lambda: None):
            self.calls.append(operation)
            return original_result(root_, provider, principal, operation, payload, checkpoint)

        def counted_events(root_, provider, principal, operation, payload, checkpoint=lambda: None):
            self.calls.append(operation)
            return original_events(root_, provider, principal, operation, payload, checkpoint)

        monkeypatch.setattr(service_module, "reader_result", counted_result)
        monkeypatch.setattr(service_module, "reader_events", counted_events)
        self.service = Service(root)
        self.service._render = types.SimpleNamespace(invalidate=self.invalidated.append, close=lambda: None)

    def file(self, name, **kw):
        return build_workbook(self.root / "data/raw" / name, **kw)

    def register(self, name):
        job = self.service.register_documents([name], wait=60)
        assert job["state"] == "succeeded", job
        return job["result"]["documents"][0]

    def status(self, document_id):
        return self.service.document(document_id)["status"]

    def units(self):
        (self.root / "config").mkdir(exist_ok=True)
        (self.root / "config/units.yaml").write_text(UNITS_YAML, encoding="utf-8")

    def approve_profile_with(self, name="a.xlsx", definition=None, **file_kw):
        """스키마 + 프로파일(draft) → 문서 등록 → 수동 적용 → approve_all(발행) → 프로파일 승인(+rematch). → {profile, document, application}."""
        s = self.service
        if s.schema_fields("process_standard") is None:
            s.import_schema(SCHEMA)
        profile = s.import_profile("process_standard", definition or profile_definition())
        self.file(name, **file_kw)
        doc = self.register(name)
        app = s.apply_profile(doc["snapshot"]["snapshot_id"], profile["profile_id"])
        result = s.approve_all(app["application_id"], reason="검수 완료")
        assert result["extraction"]["state"] == "succeeded", result
        approved = s.approve_profile(profile["profile_id"], app["application_id"])
        rematch = s.jobs.wait(approved["reparse_job"]["job_id"], 60)
        assert rematch["state"] == "succeeded", rematch
        self.calls.clear()
        return {"profile": profile, "document": doc, "application": app}
