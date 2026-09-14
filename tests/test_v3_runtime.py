"""v3 API 통합 시나리오(계약 §11 `test_v3_runtime.py`) + 성능 회귀(문서 2,000건 목록 페이지 < 50ms).

한 작업 공간 위에서 순서대로 진행한다(모듈 fixture `world`, 테스트 함수는 파일 순서대로 실행된다):
스키마·초안 프로파일 → 등록(Reader 프로세스 1회, approved 프로파일만 자동 적용 → 초안뿐이면 unmatched) → 수동 적용(draft) →
검수 승인(CAS·반려·approve_all→추출·발행) → 프로파일 승인(+rematch 소급) → 자동 승인·추출(identical) → 값/역방향 조회 →
빌드(candidates·preview·csv/xlsx/sqlite·manifest·INVALID_HEADER) → 큐 묶음·묶음 처리 → 새 snapshot 승계(proposed) → approve_all → 발행 →
프로파일 테스트 → 검색 → 문서 상태 전이.
Reader 호출은 `kg.v3.service.reader_result/reader_events`를 감싸 연산 이름을 기록한다(격리 프로세스는 실제로 뜬다)."""

from __future__ import annotations

import copy
import csv
import io
import sqlite3
import time
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from examples.schema_v3 import demo
from kg.v3 import service as service_module
from kg.v3.api import create_app
from kg.v3.db import now
from kg.v3.render.client import RenderClient
from tests.test_v3_api import inprocess_source

DOC_REF = "공정데이터_2024_01.xlsx"  # 대표 문서(수동 적용 → 검수 → 프로파일 승인)
DOC_SAME = "공정데이터_2024_02.xlsx"  # 같은 양식 → 자동 승인·추출·발행
DOC_SHIFTED = "공정데이터_2024_04_양식이동.xlsx"  # 표 이동 → compatible → 검수
DOC_OTHER = "품질검사_2024_05.xlsx"  # 다른 양식 → unmatched
DOC_LOCKED = "공정데이터_2024_06_잠김.xlsx"  # DRM_READER_REQUIRED → locked
DOC_KELVIN = "공정데이터_2024_07_단위K.xlsx"  # 같은 양식이지만 단위 셀이 K → 자동 승인 뒤 추출 실패(UNIT_MISMATCH)
SCHEMA_KEY = demo.SCHEMA_KEY


class World:
    def __init__(self, root, client, service):
        self.root, self.client, self.service = root, client, service
        self.calls = []
        self.state = {"transitions": {}}  # document_name -> [status...]

    # ---- HTTP -------------------------------------------------------------------------------
    def get(self, path, expect=200, **kw):
        response = self.client.get("/api/v3" + path, **kw)
        assert response.status_code == expect, (path, response.status_code, response.text[:500])
        return response.json() if response.content else None

    def post(self, path, body=None, expect=200, **kw):
        response = self.client.post("/api/v3" + path, json=body, **kw)
        assert response.status_code == expect, (path, response.status_code, response.text[:500])
        return response.json()

    # ---- 시나리오 도우미 ------------------------------------------------------------------------
    def register(self, name, document_id=None, expect_state="succeeded"):
        self.calls.clear()
        body = {"source_refs": [name]}
        if document_id:
            body["document_id"] = document_id
        job = self.post("/documents/register?wait=60", body, expect=200)
        assert job["state"] == expect_state and job["kind"] == "register", job
        doc = job["result"]["documents"][0] if job.get("result") else None
        if doc and doc.get("document_id"):
            self.remember(doc["document_id"], doc["document_name"])
        return job, doc

    def remember(self, document_id, name):
        status = self.get(f"/documents/{document_id}")["status"]
        self.state["transitions"].setdefault(name, [])
        if not self.state["transitions"][name] or self.state["transitions"][name][-1] != status:
            self.state["transitions"][name].append(status)
        return status

    def status_of(self, name):
        return self.remember(self.state["documents"][name]["document_id"], name)

    def doc(self, name):
        return self.state["documents"][name]


@pytest.fixture(scope="module")
def world(tmp_path_factory):
    root = tmp_path_factory.mktemp("v3-runtime")
    raw = root / "data/raw"
    raw.mkdir(parents=True)
    demo.build_process_workbook(raw / DOC_REF, 1)
    demo.build_process_workbook(raw / DOC_SAME, 2, with_image=True)
    demo.build_process_workbook(raw / DOC_SHIFTED, 4, table_shift=3)
    demo.build_other_workbook(raw / DOC_OTHER)
    demo.build_locked_file(raw / DOC_LOCKED)
    kelvin = demo.build_process_workbook(raw / DOC_KELVIN, 7)
    wb = load_workbook(kelvin)
    wb[demo.COMMON_SHEET]["B1"] = "K"
    wb.save(kelvin)

    calls = []
    original_result, original_events = service_module.reader_result, service_module.reader_events

    def counted_result(root_, provider, principal, operation, payload, checkpoint=lambda: None):
        calls.append(operation)
        return original_result(root_, provider, principal, operation, payload, checkpoint)

    def counted_events(root_, provider, principal, operation, payload, checkpoint=lambda: None):
        calls.append(operation)
        return original_events(root_, provider, principal, operation, payload, checkpoint)

    patch = pytest.MonkeyPatch()
    patch.setattr(service_module, "reader_result", counted_result)
    patch.setattr(service_module, "reader_events", counted_events)
    app = create_app(root, start_worker=False)
    try:
        with TestClient(app) as client:
            service = app.state.v3
            service._render = RenderClient(root, event_source=inprocess_source)
            env = World(root, client, service)
            env.calls = calls
            yield env
    finally:
        patch.undo()


# ---------------------------------------------------------------------------- 1. 정의


def test_schema_and_draft_profile(world):
    schema = world.post("/schemas", {"definition": demo.SCHEMA}, expect=201)
    assert schema["schema_key"] == SCHEMA_KEY and schema["current_rev"] == 1 and schema["fields"] == 14
    profile = world.post("/profiles", {"schema_key": SCHEMA_KEY, "definition": demo.PROFILE}, expect=201)
    assert profile["status"] == "draft" and profile["current_rev"] == 1 and profile["report"]["format_detected"] == "parsing-profile-3.0"
    assert profile["rules"] == 11 and profile["deprecated_rules"] in (0, [])
    world.state["profile_id"] = profile["profile_id"]
    counts = world.get("/status")["counts"]
    assert counts == {"documents": 0, "profiles": 1, "schemas": 1, "jobs_running": 0, "review": 0}


# ---------------------------------------------------------------------------- 2. 등록


def test_register_uses_one_reader_process_and_auto_applies_approved_profiles_only(world):
    world.state["documents"] = {}
    job, doc = world.register(DOC_REF)
    # 등록 = describe 1회(같은 프로세스에서 approved 프로파일 매치까지). 초안 프로파일은 자동 적용 대상이 아니다.
    assert world.calls == ["describe"], world.calls
    assert doc["status"] == "unmatched" and doc["applied"] == [] and doc["snapshot"]["revision_no"] == 1 and not doc["snapshot"]["unchanged"]
    assert job["label"] == DOC_REF and job["target_kind"] in ("workspace", "document") and job["completed"] == job["total"] == 1
    world.state["documents"][DOC_REF] = doc
    detail = world.get(f"/documents/{doc['document_id']}")
    assert detail["status"] == "unmatched" and detail["current_snapshot"]["snapshot_id"] == doc["snapshot"]["snapshot_id"] and detail["profiles"] == []
    sheets = world.get(f"/snapshots/{doc['snapshot']['snapshot_id']}/sheets")["items"]
    assert [s["sheet_name"] for s in sheets] == [demo.MAIN_SHEET, demo.COMMON_SHEET] and sheets[0]["ordinal"] == 0

    # 같은 파일을 다시 등록하면 snapshot이 늘지 않는다(token 동일).
    job, again = world.register(DOC_REF)
    assert world.calls == ["describe"] and again["snapshot"]["unchanged"] and again["snapshot"]["snapshot_id"] == doc["snapshot"]["snapshot_id"]
    assert len(world.get(f"/documents/{doc['document_id']}/snapshots")["items"]) == 1

    # 잠긴 파일: 작업 failed(DRM_READER_REQUIRED), 문서 행은 locked로 남는다.
    job, locked = world.register(DOC_LOCKED, expect_state="failed")
    assert job["error_code"] == "DRM_READER_REQUIRED" and world.calls == ["describe"]
    listed = world.get(f"/documents?status=locked")["items"]
    assert [d["document_name"] for d in listed] == [DOC_LOCKED] and listed[0]["current_snapshot"] is None
    world.state["documents"][DOC_LOCKED] = {"document_id": listed[0]["document_id"], "document_name": DOC_LOCKED, "snapshot": None}
    assert world.status_of(DOC_LOCKED) == "locked"
    assert world.get(f"/documents/{listed[0]['document_id']}")["last_error"]


# ---------------------------------------------------------------------------- 3. 수동 적용(draft) + 검수


def test_manual_apply_draft_then_review_cas_reject_and_approve_all(world):
    ref = world.doc(DOC_REF)
    sid, pid = ref["snapshot"]["snapshot_id"], world.state["profile_id"]
    world.calls.clear()
    app = world.post(f"/snapshots/{sid}/applications", {"profile_id": pid}, expect=201)
    assert world.calls == ["match"]
    assert app["origin"] == "manual" and app["compatibility"] == "manual" and not app["published"]
    assert app["heads_total"] == 11 and app["heads_approved"] == 0 and app["document"]["document_id"] == ref["document_id"]
    assert all(m["status"] == "proposed" and m["origin"] == "profile" and m["edit_seq"] == 1 and m["revision_no"] == 1 and m["field"] for m in app["mappings"])
    assert {r["role"] for m in app["mappings"] for r in m["regions"]} >= {"key", "value"}
    assert sorted(s["sheet_name"] for s in app["sheets"]) == sorted([demo.MAIN_SHEET, demo.COMMON_SHEET])
    assert world.post(f"/snapshots/{sid}/applications", {"profile_id": pid}, expect=409)["error"]["code"] == "ALREADY_APPLIED"
    assert world.status_of(DOC_REF) == "review"
    aid = app["application_id"]
    world.state["reference_application_id"] = aid
    by_rule = {m["rule_key"]: m for m in app["mappings"]}

    # 리비전 하나 승인: 헤드가 전부 approved가 아니므로 추출은 없다.
    temp = by_rule["temperature"]
    approved = world.post(f"/mappings/{temp['mapping_id']}/revisions", {"expected_seq": 1, "status": "approved", "reason": "확인"}, expect=201)
    assert approved["status"] == "approved" and approved["edit_seq"] == 2 and approved["revision_no"] == 2 and approved["extraction"] is None
    # CAS: 같은 expected_seq로 다시 쓰면 409, 리비전 수는 그대로.
    conflict = world.post(f"/mappings/{temp['mapping_id']}/revisions", {"expected_seq": 1, "status": "approved"}, expect=409)
    assert conflict["error"]["code"] == "EDIT_CONFLICT"
    assert len(world.get(f"/mappings/{temp['mapping_id']}/revisions")["items"]) == 2
    # 반려 → 다시 승인(반려 헤드는 추출을 막는다).
    lot = by_rule["lot"]
    rejected = world.post(f"/mappings/{lot['mapping_id']}/revisions", {"expected_seq": 1, "status": "rejected", "reason": "위치 확인 필요"}, expect=201)
    assert rejected["status"] == "rejected" and rejected["edit_seq"] == 2
    gate = world.post(f"/applications/{aid}/extract", expect=422)["error"]
    assert gate["code"] == "REVIEW_REQUIRED" and {f["status"] for f in gate["fields"]} == {"proposed", "rejected"} and len(gate["fields"]) == 10
    fixed = world.post(f"/mappings/{lot['mapping_id']}/revisions", {"expected_seq": 2, "status": "approved", "reason": "확인"}, expect=201)
    assert fixed["status"] == "approved" and fixed["edit_seq"] == 3
    history = world.get(f"/mappings/{lot['mapping_id']}/revisions")["items"]
    assert [h["status"] for h in history] == ["approved", "rejected", "proposed"] and history[0]["reason"] == "확인" and "mapping_revision_id" in history[0]

    # approve_all: 남은 proposed 9개 → approved + 추출·발행이 요청 1회.
    world.calls.clear()
    done = world.post(f"/applications/{aid}/approve-all?wait=60", {"reason": "대표 문서 검수"})
    assert done["approved"] == 9 and done["skipped"] == [] and done["extraction"]["state"] == "succeeded" and done["extraction"]["kind"] == "extract"
    assert world.calls == ["extract"]
    summary = world.get(f"/applications/{aid}")
    assert summary["published"] and summary["heads_approved"] == summary["heads_total"] == 11 and "published_run_id" not in summary
    assert all(m["value"] is not None for m in summary["mappings"])
    assert world.status_of(DOC_REF) == "normal"
    assert world.get("/status")["counts"]["review"] == 0


# ---------------------------------------------------------------------------- 4. 프로파일 승인 + rematch 소급


def test_profile_approval_then_rematch_backfills_unmatched_documents(world):
    pid = world.state["profile_id"]
    # 승인 전에 등록된 문서: 초안은 자동 적용되지 않으므로 둘 다 unmatched.
    for name in (DOC_SHIFTED, DOC_OTHER):
        job, doc = world.register(name)
        assert world.calls == ["describe"] and doc["status"] == "unmatched" and doc["applied"] == []
        world.state["documents"][name] = doc
    assert world.post(f"/profiles/{pid}/reparse", {"mode": "fill"}, expect=422)["error"]["code"] == "PROFILE_NOT_APPROVED"
    assert world.post(f"/profiles/{pid}/approve", {"application_id": "nope"}, expect=404)["error"]["code"] == "NOT_FOUND"

    approved = world.post(f"/profiles/{pid}/approve", {"application_id": world.state["reference_application_id"]})
    assert approved["status"] == "approved" and approved["reference_application_id"] == world.state["reference_application_id"]
    assert approved["reference_profile_rev"] == 1 and approved["reference_signature"] and approved["reparse_job"]["state"] == "queued"
    world.calls.clear()
    rematch = world.post(f"/documents/register?wait=0", {"source_refs": []}, expect=422)  # 큐만 확인: 검증 오류는 작업을 만들지 않는다
    assert rematch["error"]["code"] == "VALIDATION_ERROR"
    job = world.service.jobs.wait(approved["reparse_job"]["job_id"], 60)
    assert job["state"] == "succeeded" and job["kind"] == "reparse" and job["target_id"] == pid, job
    # rematch: 대표 문서(발행됨) up_to_date, 표 이동 문서 compatible → proposed, 다른 양식 → incompatible.
    reasons = {s["document_name"]: s["reason"] for s in job["result"]["skipped"]}
    assert job["result"]["queued"] == 1 and reasons == {DOC_OTHER: "incompatible", DOC_REF: "up_to_date"}, job["result"]
    assert "match" in world.calls
    assert world.status_of(DOC_SHIFTED) == "review" and world.status_of(DOC_OTHER) == "unmatched"
    shifted_apps = world.get(f"/snapshots/{world.doc(DOC_SHIFTED)['snapshot']['snapshot_id']}/applications")["items"]
    assert len(shifted_apps) == 1 and shifted_apps[0]["origin"] == "auto" and shifted_apps[0]["compatibility"] == "compatible" and shifted_apps[0]["heads_approved"] == 0
    world.state["shifted_application_id"] = shifted_apps[0]["application_id"]
    # 헤드가 승인되지 않은 적용 건으로는 승인할 수 없다.
    gate = world.post(f"/profiles/{pid}/approve", {"application_id": shifted_apps[0]["application_id"]}, expect=422)
    assert gate["error"]["code"] == "REVIEW_REQUIRED"

    detail = world.get(f"/profiles/{pid}")
    assert detail["status"] == "approved" and detail["auto_approval_active"] and detail["reference"]["document_name"] == DOC_REF
    assert detail["reference"]["snapshot"]["is_current"] and detail["document_count"] == 2
    docs = {d["document_name"]: d for d in world.get(f"/profiles/{pid}/documents")["items"]}
    assert docs[DOC_REF]["is_reference"] and docs[DOC_REF]["published"] and docs[DOC_SHIFTED]["compatibility"] == "compatible" and not docs[DOC_SHIFTED]["published"]


# ---------------------------------------------------------------------------- 5. 자동 승인·추출·발행


def test_identical_document_is_auto_approved_extracted_and_published(world):
    job, doc = world.register(DOC_SAME)
    # describe(매치 포함) 1회 + 추출 1회: 사람 개입 없이 발행까지.
    assert world.calls == ["describe", "extract"], world.calls
    assert doc["status"] == "normal" and [(a["compatibility"], a["state"]) for a in doc["applied"]] == [("identical", "published")]
    world.state["documents"][DOC_SAME] = doc
    apps = world.get(f"/snapshots/{doc['snapshot']['snapshot_id']}/applications")["items"]
    assert apps[0]["origin"] == "auto" and apps[0]["published"] and apps[0]["last_run"] == "succeeded"
    summary = world.get(f"/applications/{apps[0]['application_id']}")
    assert all(m["status"] == "approved" and m["origin"] == "auto" for m in summary["mappings"])
    assert world.status_of(DOC_SAME) == "normal"

    # 단위 셀이 K인 문서: 구조는 identical → 자동 승인 → 추출 실패(UNIT_MISMATCH) → failed(auto_approved 표시).
    job, kelvin = world.register(DOC_KELVIN)
    assert job["state"] == "succeeded" and world.calls == ["describe", "extract"]
    assert kelvin["status"] == "failed" and kelvin["applied"][0]["compatibility"] == "identical"
    world.state["documents"][DOC_KELVIN] = kelvin
    failed = world.get(f"/documents/{kelvin['document_id']}")
    assert failed["status"] == "failed" and failed["last_error"].startswith("UNIT_MISMATCH")
    apps = world.get(f"/snapshots/{kelvin['snapshot']['snapshot_id']}/applications")["items"]
    assert apps[0]["last_run"] == "failed" and apps[0]["last_error"].startswith("UNIT_MISMATCH") and not apps[0]["published"]
    assert world.status_of(DOC_KELVIN) == "failed"


# ---------------------------------------------------------------------------- 6. 값 · 역방향 조회


def test_values_and_reverse_lookup(world):
    same = world.doc(DOC_SAME)
    sid = same["snapshot"]["snapshot_id"]
    page1 = world.get(f"/snapshots/{sid}/values?field_key=temperature&limit=5")
    assert len(page1["items"]) == 5 and page1["has_more"] and page1["next_cursor"]
    page2 = world.get(f"/snapshots/{sid}/values?field_key=temperature&limit=50&cursor={page1['next_cursor']}")
    assert len(page2["items"]) == 7 and not page2["has_more"]
    values = page1["items"] + page2["items"]
    assert [v["value_text"] for v in values] == ["150", "152.5", "155", "157.5", "160"] * 2 + ["150", "152.5"]
    first = values[0]
    assert first["rule_key"] == "temperature" and first["field"]["key"] == "temperature" and first["unit_normalized"] == "°C" and first["value_state"] == "present"
    assert first["source"]["sheet_name"] == demo.MAIN_SHEET and first["source"]["range"] == "C9" and first["record_key"] == values[0]["record_key"]
    assert first["derivation_key"] == "relation:lot:same_row"
    roles = {r["role"] for r in first["regions"]}
    assert roles >= {"key", "value", "unit", "record_key"}
    by_rule = world.get(f"/applications/{first['application_id']}/values?rule_key=lot&limit=200")["items"]
    assert [v["value_text"] for v in by_rule][:2] == ["LOT-02-001", "LOT-02-002"] and len(by_rule) == 12

    value = world.get(f"/values/{first['value_id']}")
    assert value["document"]["document_id"] == same["document_id"] and value["snapshot"]["snapshot_id"] == sid and "run_id" not in value and value["mapping_id"]
    region = next(r for r in first["regions"] if r["role"] == "value")
    reverse = world.get(f"/regions/{region['region_id']}/values")
    assert reverse["region"]["range"] == "C9" and first["value_id"] in {v["value_id"] for v in reverse["values"]}
    assert any(m["rule_key"] == "temperature" and m["is_head"] and m["status"] == "approved" for m in reverse["mappings"])
    regions = world.get(f"/sheets/{region['sheet_id']}/regions?range=C8:C20")
    assert regions["sheet"]["sheet_name"] == demo.MAIN_SHEET and regions["range"] == "C8:C20"
    ranges = {r["range"]: r for r in regions["items"]}
    assert "C8" in ranges and ranges["C8"]["mapping_count"] >= 1 and ranges["C9"]["value_count"] >= 1
    assert world.get("/values/nope", expect=404)["error"]["code"] == "NOT_FOUND"
    world.state["value_id"] = first["value_id"]


# ---------------------------------------------------------------------------- 7. 데이터 빌드


def test_builds_candidates_preview_three_formats_manifest_and_header_validation(world):
    ids = [world.doc(n)["document_id"] for n in (DOC_REF, DOC_SAME, DOC_SHIFTED, DOC_OTHER, DOC_LOCKED, DOC_KELVIN)]
    candidates = world.post("/builds/candidates", {"document_ids": ids + ["missing-doc"], "schema_key": SCHEMA_KEY})
    assert candidates["summary"] == {"total": 7, "usable": 2, "excluded": 5}
    reasons = {d["document_name"]: d.get("reason") for d in candidates["documents"] if not d["usable"]}
    assert reasons[DOC_SHIFTED] == "review_required" and reasons[DOC_OTHER] == "unmatched" and reasons[DOC_LOCKED] == "locked" and reasons[DOC_KELVIN] == "failed"
    missing = next(d for d in candidates["documents"] if d["document_id"] == "missing-doc")
    assert missing["reason"] == "not_found" and missing["document_name"] is None and not missing["usable"]
    fields = {f["field_key"]: f for f in candidates["fields"]}
    assert fields["temperature"]["document_count"] == 2 and fields["temperature"]["unit"] == "°C" and "basic" not in fields
    usable = [d["document_id"] for d in candidates["documents"] if d["usable"]]
    assert usable == ids[:2]

    columns = [
        {"field_key": "product_name", "header": "제품"},
        {"field_key": "lot", "header": "LOT"},
        {"field_key": "temperature", "header": "온도"},
        {"field_key": "verdict", "header": "판정"},
    ]
    preview = world.post("/builds/preview", {"document_ids": usable, "schema_key": SCHEMA_KEY, "columns": columns, "row_mode": "record"})
    assert preview["row_count"] == 24 and len(preview["rows"]) == 24 and preview["headers"][:4] == ["제품", "LOT", "온도", "판정"] or preview["headers"][-4:] == ["제품", "LOT", "온도", "판정"]
    assert preview["excluded"] == [] and preview["conflicts"] == [] and preview["document_count"] == 2
    row = preview["rows"][0]
    assert row["document"]["document_name"] in (DOC_REF, DOC_SAME) and row["record_key"]
    cells = {c["rule_key"]: c for c in row["cells"] if c}
    assert cells["temperature"]["text"] == "150" and cells["temperature"]["range"] == "C9" and cells["temperature"]["sheet_name"] == demo.MAIN_SHEET
    assert cells["product_name"]["text"].startswith("제품-") and cells["product_name"]["value_id"]
    duplicate = world.post("/builds/preview", {"document_ids": usable, "schema_key": SCHEMA_KEY, "columns": columns + [{"field_key": "pressure", "header": "온도"}]}, expect=422)
    assert duplicate["error"]["code"] == "INVALID_HEADER" and duplicate["error"]["fields"][0]["reason"] == "duplicate"
    blank = world.post("/builds", {"document_ids": usable, "schema_key": SCHEMA_KEY, "columns": [{"field_key": "lot", "header": "  "}], "format": "csv"}, expect=422)
    assert blank["error"]["code"] == "INVALID_HEADER" and blank["error"]["fields"][0]["reason"] == "empty"
    unknown = world.post("/builds", {"document_ids": usable, "schema_key": SCHEMA_KEY, "columns": [{"field_key": "nope", "header": "x"}]}, expect=422)
    assert unknown["error"]["code"] == "UNKNOWN_FIELD"
    assert world.post("/builds", {"document_ids": usable, "schema_key": SCHEMA_KEY, "columns": columns, "format": "pdf"}, expect=422)["error"]["code"] == "VALIDATION_ERROR"

    request = {"document_ids": usable, "schema_key": SCHEMA_KEY, "columns": columns, "row_mode": "record"}
    built = {}
    for fmt in ("csv", "xlsx", "sqlite"):
        result = world.post("/builds?wait=30", {**request, "format": fmt})
        assert result["build_key"] == preview["build_key"] and result["download_url"] == f"/api/v3/builds/{result['build_key']}/download?format={fmt}"
        assert result["manifest"]["row_count"] == 24 and result["format"] == fmt and "path" not in result and result["job_id"]
        built[fmt] = result
    key = preview["build_key"]
    manifest = world.get(f"/builds/{key}/manifest")
    assert manifest["schema"] == {"key": SCHEMA_KEY, "rev": 1, "name": "공정 데이터 표준"} and set(manifest["formats"]) == {"csv", "xlsx", "sqlite"}
    assert [c["header"] for c in manifest["columns"]] == ["제품", "LOT", "온도", "판정"] and manifest["conflicts"] == [] and len(manifest["sources"]) == 2
    assert {s["document_name"] for s in manifest["sources"]} == {DOC_REF, DOC_SAME} and all(s["profile"]["name"] == demo.PROFILE_NAME for s in manifest["sources"])
    # 같은 입력은 재사용(새 작업 없음).
    again = world.post("/builds", {**request, "format": "csv"})
    assert again["reused"] and again["build_key"] == key

    download = world.client.get(f"/api/v3/builds/{key}/download?format=csv")
    assert download.status_code == 200 and download.content.startswith(b"\xef\xbb\xbf") and "attachment" in download.headers["content-disposition"]
    text = download.content.decode("utf-8-sig")
    assert "\r\n" in text
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][:4] == ["제품", "LOT", "온도", "판정"] and len(rows) == 25 and rows[1][2] == "150"
    xlsx = world.client.get(f"/api/v3/builds/{key}/download?format=xlsx")
    assert xlsx.status_code == 200 and xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(xlsx.content), read_only=True).active
    sheet_rows = list(ws.iter_rows(values_only=True))
    assert list(sheet_rows[0][:4]) == ["제품", "LOT", "온도", "판정"] and len(sheet_rows) == 25 and sheet_rows[1][2] == 150
    sqlite_file = world.client.get(f"/api/v3/builds/{key}/download?format=sqlite")
    assert sqlite_file.status_code == 200
    path = world.root / "build-download.sqlite"
    path.write_bytes(sqlite_file.content)
    with sqlite3.connect(path) as conn:
        columns_in_db = [r[1] for r in conn.execute("PRAGMA table_info(data)")]
        assert {"제품", "LOT", "온도", "판정", "_source_온도", "_document", "_snapshot", "_record_key"} <= set(columns_in_db)
        assert conn.execute("SELECT count(*) FROM data").fetchone()[0] == 24
        assert conn.execute('SELECT "_source_온도" FROM data ORDER BY rowid LIMIT 1').fetchone()[0] == f"{demo.MAIN_SHEET}!C9"
    assert world.get(f"/builds/{key}/download?format=pdf", expect=422)["error"]["code"] == "VALIDATION_ERROR"
    assert world.get("/builds/0123456789abcdef/manifest", expect=404)["error"]["code"] == "NOT_FOUND"
    assert world.get("/builds/../etc/download", expect=404)
    jobs = world.get("/jobs?kind=build")["items"]
    assert len(jobs) == 3 and all(j["state"] == "succeeded" and j["target_kind"] == "build" and j["target_id"] == key for j in jobs)
    world.state["build_key"] = key


# ---------------------------------------------------------------------------- 8. 검수 큐 묶음·묶음 처리


def test_queue_groups_and_group_action(world):
    pid = world.state["profile_id"]
    queues = world.get("/queues")
    assert queues["summary"] == {"unmatched": 1, "review": 1, "failed": 2, "changed": 0, "conflict": 0}
    review = queues["groups"]["review"][0]
    assert review["group_key"] == pid and review["count"] == 1 and review["representative"]["document_name"] == DOC_SHIFTED
    assert "approve_all" in review["actions"] and "open_review" in review["actions"] and "UUID" not in review["label"]
    failed = {g["group_key"]: g for g in queues["groups"]["failed"]}
    assert failed["UNIT_MISMATCH"]["count"] == 1 and failed["locked"]["count"] == 1 and "reparse" in failed["UNIT_MISMATCH"]["actions"]
    assert failed["UNIT_MISMATCH"]["impact"].get("auto_approved") == 1
    unmatched = queues["groups"]["unmatched"][0]
    assert unmatched["count"] == 1 and set(unmatched["actions"]) >= {"create_profile", "assign_profile"}

    listed = world.get("/queues/review?limit=10")
    assert listed["kind"] == "review" and [g["group_key"] for g in listed["items"]] == [pid] and not listed["has_more"]
    members = world.get(f"/queues/review/groups/{pid}/members")
    assert members["group_key"] == pid and [m["document_name"] for m in members["items"]] == [DOC_SHIFTED]
    assert members["items"][0]["application_id"] == world.state["shifted_application_id"] and members["items"][0]["compatibility"] == "compatible"
    assert world.get("/queues/nope", expect=404)["error"]["code"] == "GROUP_NOT_FOUND"
    assert world.get("/queues/review/groups/none/members", expect=404)["error"]["code"] == "GROUP_NOT_FOUND"
    assert world.post(f"/queues/review/groups/{pid}/actions", {"action": "delete"}, expect=422)["error"]["code"] == "VALIDATION_ERROR"
    not_allowed = world.post(f"/queues/unmatched/groups/{unmatched['group_key']}/actions", {"action": "reparse"}, expect=422)
    assert not_allowed["error"]["code"] == "ACTION_NOT_ALLOWED" and not_allowed["error"]["fields"]["allowed"] == ["assign_profile"]
    assert world.post(f"/queues/review/groups/{pid}/actions", {"action": "approve_all"}, expect=202)["state"] == "queued"
    # 방금 넣은 작업은 워커 없이도 wait가 직접 실행한다(멱등: 같은 묶음의 두 번째 요청).
    world.calls.clear()
    job = world.post(f"/queues/review/groups/{pid}/actions?wait=60", {"action": "approve_all", "extract": True})
    assert job["state"] == "succeeded" and job["kind"] == "queue_action" and job["target_kind"] == "queue_group" and job["target_id"] == f"review/{pid}"
    assert job["result"] == {"queued": 1, "skipped": []} and job["total"] == 1 and job["completed"] == 1 and world.calls == ["extract"]
    assert world.status_of(DOC_SHIFTED) == "normal"
    summary = world.get(f"/applications/{world.state['shifted_application_id']}")
    assert summary["published"] and summary["heads_approved"] == 11
    after = world.get("/queues")
    assert after["summary"]["review"] == 0 and after["groups"]["review"] == []
    assert world.post(f"/queues/review/groups/{pid}/actions", {"action": "approve_all"}, expect=404)["error"]["code"] == "GROUP_NOT_FOUND"
    # 표가 이동한 문서의 값도 원본 위치가 다르다(C12 = 헤더 11행 + 1).
    shifted_values = world.get(f"/snapshots/{world.doc(DOC_SHIFTED)['snapshot']['snapshot_id']}/values?rule_key=temperature&limit=1")["items"]
    assert shifted_values[0]["source"]["range"] == "C12"


# ---------------------------------------------------------------------------- 9. 새 snapshot 승계 → approve_all → 발행


def test_new_snapshot_inherits_proposed_then_approve_all_publishes(world):
    ref = world.doc(DOC_REF)
    old_sid = ref["snapshot"]["snapshot_id"]
    old_values = [v["value_text"] for v in world.get(f"/snapshots/{old_sid}/values?field_key=temperature&limit=200")["items"]]
    demo.mutate_first_document(world.root, temp_offset=0.5)
    job, doc = world.register(DOC_REF, document_id=ref["document_id"])
    # 승계 = describe 1회 + 이전 헤드 spec으로 match_specs 1회. 자동 승인 없음.
    assert world.calls == ["describe", "match_specs"], world.calls
    assert doc["document_id"] == ref["document_id"] and doc["snapshot"]["revision_no"] == 2 and doc["snapshot"]["snapshot_id"] != old_sid
    # applied[].state는 application 상태 어휘(approved/review/changed …); 승계 리비전 자체는 proposed다(아래 mappings 검사).
    assert doc["status"] == "changed" and [(a["compatibility"], a["state"]) for a in doc["applied"]] == [("identical", "changed")]
    world.state["documents"][DOC_REF] = doc
    new_sid = doc["snapshot"]["snapshot_id"]
    snapshots = world.get(f"/documents/{ref['document_id']}/snapshots")["items"]
    assert [(s["revision_no"], s["current"]) for s in snapshots] == [(2, True), (1, False)]
    apps = world.get(f"/snapshots/{new_sid}/applications")["items"]
    assert apps[0]["origin"] == "inherited" and apps[0]["compatibility"] == "identical" and not apps[0]["published"]
    aid = apps[0]["application_id"]
    summary = world.get(f"/applications/{aid}")
    assert all(m["status"] == "proposed" and m["origin"] == "inherited" and m["field"] for m in summary["mappings"])
    history = world.get(f"/mappings/{summary['mappings'][0]['mapping_id']}/revisions")["items"]
    assert history[0]["evidence"]["compatibility"] == "identical" and history[0]["evidence"]["previous_application_id"] == world.state["reference_application_id"]
    queues = world.get("/queues")
    changed = queues["groups"]["changed"]
    assert queues["summary"]["changed"] == 1 and changed[0]["group_key"] == world.state["profile_id"] and "approve_all" in changed[0]["actions"]
    # 이전 snapshot의 값은 그대로(불변); 현재 snapshot은 아직 발행 전.
    assert [v["value_text"] for v in world.get(f"/snapshots/{old_sid}/values?field_key=temperature&limit=200")["items"]] == old_values
    assert world.get(f"/snapshots/{new_sid}/values?field_key=temperature")["items"] == []
    excluded = world.post("/builds/candidates", {"document_ids": [ref["document_id"]], "schema_key": SCHEMA_KEY})["documents"][0]
    assert not excluded["usable"] and excluded["reason"] == "review_required"

    world.calls.clear()
    done = world.post(f"/applications/{aid}/approve-all?wait=60", {"reason": "변경 확인"})
    assert done["approved"] == 11 and done["extraction"]["state"] == "succeeded" and world.calls == ["extract"]
    assert world.get(f"/applications/{aid}")["published"] and world.status_of(DOC_REF) == "normal"
    new_values = [v["value_text"] for v in world.get(f"/snapshots/{new_sid}/values?field_key=temperature&limit=200")["items"]]
    assert new_values[:2] == ["150.5", "153"] and new_values != old_values
    assert world.get("/queues")["summary"]["changed"] == 0
    rebuilt = world.post("/builds/preview", {"document_ids": [ref["document_id"]], "schema_key": SCHEMA_KEY, "columns": [{"field_key": "temperature", "header": "온도"}]})
    assert rebuilt["build_key"] != world.state["build_key"] and rebuilt["rows"][0]["cells"][0]["text"] == "150.5"


# ---------------------------------------------------------------------------- 10. 프로파일 테스트(dry-run)


def test_profile_test_dry_run_leaves_no_application(world):
    pid = world.state["profile_id"]
    sid = world.doc(DOC_SAME)["snapshot"]["snapshot_id"]
    before = len(world.get(f"/snapshots/{sid}/applications")["items"])
    result = world.post(f"/profiles/{pid}/test", {"snapshot_id": sid})
    assert result["compatibility"] == "identical" and result["errors"] == [] and result["missing"] == [] and result["job_id"]
    groups = {g["rule_key"]: g for g in result["groups"]}
    assert len(groups) == 11 and groups["temperature"]["count"] == 12 and len(groups["temperature"]["values"]) == 12
    assert groups["temperature"]["field"]["key"] == "temperature" and groups["temperature"]["regions"]["value"][0]["sheet_name"] == demo.MAIN_SHEET
    assert groups["product_name"]["values"][0]["value_text"] == "제품-02"
    definition = copy.deepcopy(demo.PROFILE)
    definition["rules"] = definition["rules"][:5]
    partial = world.post("/profiles/test", {"schema_key": SCHEMA_KEY, "definition": definition, "snapshot_id": sid})
    assert partial["compatibility"] == "compatible" and len(partial["groups"]) == 5 and partial["errors"] == []
    other = world.post(f"/profiles/{pid}/test", {"snapshot_id": world.doc(DOC_OTHER)["snapshot"]["snapshot_id"]})
    assert other["compatibility"] == "incompatible" and other["groups"] == [] and other["errors"] and other["errors"][0]["code"]
    assert len(world.get(f"/snapshots/{sid}/applications")["items"]) == before
    tests = world.get("/jobs?kind=test")["items"]
    assert len(tests) == 3 and tests[-1]["label"].startswith(f"{demo.PROFILE_NAME} r1 · ") and tests[-1]["state"] == "succeeded"


# ---------------------------------------------------------------------------- 11. 검색 · 문서 상태 전이


def test_search_and_document_status_transitions(world):
    found = world.get("/search?q=공정")["items"]
    kinds = {i["kind"] for i in found}
    assert kinds == {"document", "profile", "schema", "field"} and all(i["route"].startswith("?screen=") and i["label"] for i in found)
    assert sum(1 for i in found if i["kind"] == "document") == 5
    assert [i["label"] for i in world.get("/search?q=품질")["items"]] == [DOC_OTHER]
    assert [i["kind"] for i in world.get("/search?q=온도")["items"]] == ["field"]
    assert world.get("/search?q=")["items"] == []

    transitions = world.state["transitions"]
    assert transitions[DOC_REF] == ["unmatched", "review", "normal", "changed", "normal"]
    assert transitions[DOC_SHIFTED] == ["unmatched", "review", "normal"]
    assert transitions[DOC_SAME] == ["normal"] and transitions[DOC_OTHER] == ["unmatched"]
    assert transitions[DOC_LOCKED] == ["locked"] and transitions[DOC_KELVIN] == ["failed"]
    listed = {d["document_name"]: d for d in world.get("/documents?limit=50")["items"]}
    assert {n: d["status"] for n, d in listed.items()} == {
        DOC_REF: "normal",
        DOC_SAME: "normal",
        DOC_SHIFTED: "normal",
        DOC_OTHER: "unmatched",
        DOC_LOCKED: "locked",
        DOC_KELVIN: "failed",
    }
    assert listed[DOC_REF]["current_snapshot"]["revision_no"] == 2 and listed[DOC_REF]["profiles"][0]["state"] == "published"
    assert listed[DOC_KELVIN]["last_error"].startswith("UNIT_MISMATCH") and listed[DOC_LOCKED]["status_detail"]["locked"]["code"] == "DRM_READER_REQUIRED"
    assert [d["document_name"] for d in world.get("/documents?status=failed")["items"]] == [DOC_KELVIN]
    assert [d["document_name"] for d in world.get(f"/documents?profile_id={world.state['profile_id']}&sort=document_name")["items"]] == [DOC_REF, DOC_SAME, DOC_SHIFTED, DOC_KELVIN]
    assert world.get("/status")["counts"] == {"documents": 6, "profiles": 1, "schemas": 1, "jobs_running": 0, "review": 0}
    # 발행된 값(현재 snapshot)만 필드 최근 값에 보인다: 대표 문서는 새 snapshot 값이다.
    recent = world.get(f"/schemas/{SCHEMA_KEY}/fields/temperature/values?limit=100")["items"]
    assert {v["document_name"] for v in recent} == {DOC_REF, DOC_SAME, DOC_SHIFTED} and all("sheet_name" in v and "range" in v for v in recent)


# ---------------------------------------------------------------------------- 12. 성능: 문서 2,000건 목록 페이지 < 50ms


def test_document_list_page_under_50ms_with_2000_documents(tmp_path):
    app = create_app(tmp_path / "ws", start_worker=False)
    with TestClient(app) as client:
        service = app.state.v3
        stamp = now()
        docs, snaps = [], []
        statuses = ("normal", "review", "unmatched", "changed", "failed", "not_extracted")
        for n in range(2000):
            did, sid = str(uuid.uuid4()), str(uuid.uuid4())
            docs.append((did, f"문서_{n:04d}.xlsx", "local-xlsx", f"docs/{n:04d}.xlsx", "xlsx", sid, statuses[n % len(statuses)], f"2024-01-{1 + n % 28:02d}T00:00:{n % 60:02d}+00:00", stamp, stamp))
            snaps.append((sid, did, 1, f"tok{n}", f"sha{n}", f"문서_{n:04d}.xlsx", stamp))
        # 문서 → snapshot 순환 FK는 COMMIT 시점에 검사된다(DEFERRABLE).
        with service.db.connect(write=True) as conn:
            conn.executemany(
                "INSERT INTO document(document_id, document_name, provider, source_path, file_type, current_snapshot_id, status, last_processed_at, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                docs,
            )
            conn.executemany(
                "INSERT INTO document_snapshot(snapshot_id, document_id, revision_no, change_token, content_sha256, filename, captured_at) VALUES (?,?,?,?,?,?,?)",
                snaps,
            )
        first = client.get("/api/v3/documents?limit=50").json()
        assert len(first["items"]) == 50 and first["has_more"] and first["next_cursor"]
        assert first["items"][0]["last_processed_at"] >= first["items"][-1]["last_processed_at"]
        paths = [
            "/api/v3/documents?limit=50",
            f"/api/v3/documents?limit=50&cursor={first['next_cursor']}",
            "/api/v3/documents?limit=50&status=review",
            "/api/v3/documents?limit=50&sort=document_name",
            "/api/v3/documents?limit=50&q=문서_19",
        ]
        client.get(paths[0])  # 워밍업(연결·컴파일)
        elapsed = []
        for path in paths:
            started = time.perf_counter()
            response = client.get(path)
            elapsed.append(time.perf_counter() - started)
            assert response.status_code == 200 and len(response.json()["items"]) == 50
        average = sum(elapsed) / len(elapsed)
        assert average < 0.05, f"평균 {average * 1000:.1f}ms: {[round(e * 1000, 1) for e in elapsed]}"
        assert client.get(f"/api/v3/documents?limit=50&status=review").json()["items"][0]["status"] == "review"
