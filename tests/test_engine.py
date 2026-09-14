"""실행 엔진: regex find·이름 앵커·composite·AMBIGUOUS_ANCHOR·relative.anchor·relations·list/matrix/merged/stop·
match_profile(within 밖 앵커, missing, identical/compatible/incompatible)·Reader describe(profiles)/extract/render."""

import copy
import importlib.util

import pytest
from openpyxl import Workbook, load_workbook

from schema import engine
from schema.db import Problem
from schema.profile import compile_profile, validate_profile
from schema.readers import XlsxReader, make_reader

FIELDS = {
    "lot": {"field_id": "f1", "value_type": "text", "status": "active"},
    "temperature": {"field_id": "f2", "value_type": "decimal", "status": "active"},
    "pressure": {"field_id": "f3", "value_type": "decimal", "status": "active"},
    "note": {"field_id": "f4", "value_type": "text", "status": "active"},
}


def build_workbook(path, duplicate_temp=False, shift=0, drop_common=False):
    """공정 기록 시트: B3 LOT / C3 온도 / D3 °C 머리, 4~6행 값, 7·8행 빈칸, G12~J13 가로 표, 55행 서명 앵커."""
    wb = Workbook()
    ws = wb.active
    ws.title = "공정 기록"
    ws["A1"] = "공정 A"
    ws.cell(3 + shift, 2, "LOT")
    ws.cell(3 + shift, 3, "온도")
    ws.cell(3 + shift, 4, "°C")
    ws.cell(3 + shift, 5, "비고")
    for n, (lot, temp, note) in enumerate([("L1", "10.5 / 3", "a"), ("L2", "20 / 4", None), ("L3", 30, "c")]):
        ws.cell(4 + shift + n, 2, lot)
        ws.cell(4 + shift + n, 3, temp)
        if note:
            ws.cell(4 + shift + n, 5, note)
    ws.cell(9 + shift, 2, "L9")  # 빈칸 2줄 뒤의 값: blank_run 2로는 읽지 않는다.
    ws.merge_cells(start_row=4 + shift, start_column=5, end_row=5 + shift, end_column=5)
    ws["G12"], ws["H12"], ws["I12"], ws["J12"] = "항목", "a", "b", "c"
    ws["G13"], ws["H13"], ws["I13"], ws["J13"] = "값", 1, 2, 3
    ws["G15"], ws["H15"], ws["I15"] = "합계", 4, 5
    ws["G16"], ws["H16"] = "오류", "abc"
    ws["A55"] = "서명"
    ws["B55"] = "홍길동"
    if duplicate_temp:
        ws["H20"] = "온도"
    if not drop_common:
        common = wb.create_sheet("공통")
        common["A2"] = "온도 단위"
        common["B2"] = "°C"
    for name in ("10C", "20C"):
        exp = wb.create_sheet(name)
        exp["A1"] = "실험"
        exp["B1"] = name
    wb.save(path)
    return path


def definition(**over):
    base = {
        "format": "parsing-profile",
        "schema_version": "3.0",
        "profile_name": "공정",
        "schema_key": "ps",
        "sheet_roles": {
            "main": {"cardinality": "one", "match": {"name": "공정 기록"}},
            "common": {"cardinality": "one", "match": {"any_of": [{"name_regex": "^공통"}, {"contains_text": {"texts": ["온도 단위"], "within": "A1:AD30", "mode": "any"}}]}},
            "exp": {"cardinality": "many", "match": {"name_regex": "^\\d+C$"}},
        },
        "anchors": {
            "hdr_temp": {"sheet_role": "main", "find": {"texts": ["온도"], "within": "A1:Z60"}},
            "hdr_lot": {"sheet_role": "main", "find": {"regex": "^(배치|LOT)$", "within": "A1:Z60"}},
            "table": {"all_of": ["hdr_lot", "hdr_temp"]},
            "sign": {"sheet_role": "main", "find": {"texts": ["서명"], "within": "A1:Z60"}},
        },
        "rules": [
            {
                "rule_key": "lot",
                "field_key": "lot",
                "selector": {
                    "key": {"areas": [{"sheet_role": "main", "anchor": "hdr_lot"}]},
                    "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 65, "cols": 1}}], "cardinality": "list", "axis": "down", "stop": {"kind": "blank_run", "count": 2}},
                },
                "value_spec": {"type": "text"},
                "record_spec": {"scope": ["process-table"], "key": "physical_row"},
            },
            {
                "rule_key": "temperature",
                "field_key": "temperature",
                "selector": {
                    "key": {"areas": [{"sheet_role": "main", "anchor": "hdr_temp"}]},
                    "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 65, "cols": 1}}], "cardinality": "list", "axis": "down", "stop": {"kind": "blank_run", "count": 2}},
                    "unit": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1, "anchor": "hdr_temp"}}]},
                    "context": {"areas": [{"sheet_role": "main", "range": "A1"}]},
                },
                "value_spec": {
                    "type": "decimal",
                    "unit": "°C",
                    "normalization": {"operation": "pipeline", "steps": [{"op": "trim_text"}, {"op": "split_delimiter", "delimiter": "/", "index": 0}, {"op": "split_unit_suffix"}]},
                },
                "record_spec": {"scope": ["process-table"], "key": "physical_row"},
                "relations": [{"to_rule": "lot", "kind": "same_row"}],
            },
            {
                "rule_key": "signer",
                "field_key": "note",
                "selector": {"key": {"areas": [{"sheet_role": "main", "anchor": "sign"}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1}}]}},
            },
        ],
    }
    base.update(over)
    return base


@pytest.fixture
def workbook(tmp_path):
    path = build_workbook(tmp_path / "sample.xlsx")
    raw = load_workbook(path, data_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, keep_links=False)
    yield {"path": path, "raw": raw, "cached": cached}
    raw.close()
    cached.close()


def run(workbook, canonical, bindings=None, rules=None):
    specs = compile_profile(canonical)
    if rules:
        specs = {k: specs[k] for k in rules}
    bindings = bindings or {"main": ["공정 기록"], "common": ["공통"], "exp": ["10C", "20C"]}
    groups, values = {}, {}
    for event in engine.extract(workbook["raw"], workbook["cached"], bindings, specs.values()):
        if event["type"] == "group":
            groups[event["group_key"]] = event
            values[event["group_key"]] = []
        else:
            assert event["type"] == "values" and len(event["items"]) <= 200
            values[event["group_key"]].extend(event["items"])
    return groups, values


def only(groups, values, rule_key):
    keys = [k for k in groups if groups[k]["rule_key"] == rule_key]
    assert len(keys) == 1, keys
    return groups[keys[0]], values[keys[0]]


# ---------------------------------------------------------------------------- find / anchors


def test_find_regex_and_texts(workbook):
    ws = workbook["cached"]["공정 기록"]
    ws["H30"] = "lot"
    assert [h["range"] for h in engine.find_cells(ws, {"regex": "^(배치|LOT)$", "within": "A1:Z60"})] == ["B3"]
    assert [h["range"] for h in engine.find_cells(ws, {"texts": ["LOT"], "within": "A1:Z60"})] == ["B3", "H30"]
    assert [h["range"] for h in engine.find_cells(ws, {"texts": ["LOT"], "within": "A1:Z60", "occurrence": 1})] == ["H30"]
    assert engine.find_cells(ws, {"texts": ["LOT"], "within": "A1:Z60", "occurrence": 5}) == []
    assert [h["range"] for h in engine.find_cells(ws, {"regex": "(?i)^lot$", "within": "A1:Z60"})] == ["B3", "H30"]
    # 병합 영역은 한 region으로 한 번만 잡힌다.
    assert [h["range"] for h in engine.find_cells(ws, {"texts": ["a"], "within": "A1:Z60"})] == ["E4:E5", "H12"]


def test_named_anchor_relative_anchor_and_split_delimiter(workbook):
    canonical = validate_profile(definition(), FIELDS)
    groups, values = run(workbook, canonical)
    group, items = only(groups, values, "temperature")
    assert group["group_key"] == "temperature|공정 기록|C3" and group["anchor_locator"] == "C3"
    assert group["observed_key"] == "온도" and group["primary_sheet"] == "공정 기록"
    assert group["regions"]["key"][0]["range"] == "C3"
    assert group["regions"]["unit"][0]["range"] == "D3"  # relative.anchor 기준 (0, +1)
    assert group["regions"]["context"][0]["range"] == "A1"
    assert [(i["item_index"], i["value_state"], i["value_text"], i["unit_raw"], i["unit_normalized"]) for i in items] == [
        (0, "present", "10.5", "°C", "°C"),
        (1, "present", "20", "°C", "°C"),
        (2, "present", "30", "°C", "°C"),
    ]
    assert items[0]["raw_text"] == "10.5 / 3" and items[0]["display_text"] == "10.5 / 3" and items[0]["value_type"] == "decimal"
    assert items[0]["formula_state"] == "none"
    assert {k for k in items[0]["regions"]} == {"key", "value", "unit", "context", "record_key"}
    group, items = only(groups, values, "signer")
    assert group["regions"]["key"][0]["range"] == "A55"
    assert items[0]["value_text"] == "홍길동" and items[0]["record_key"] == "B55"


def test_composite_anchor_bounding_box_key_and_relative(workbook):
    spec = definition(
        rules=[
            {
                "rule_key": "table",
                "selector": {
                    "key": {"areas": [{"sheet_role": "main", "anchor": "table"}], "separator": "|"},
                    "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 2, "cols": 2}}], "cardinality": "matrix", "axis": "row_major"},
                },
            },
            {
                "rule_key": "from_composite",
                "selector": {
                    "key": {"areas": [{"sheet_role": "main", "range": "A1"}]},
                    "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 2, "anchor": "table"}}]},
                },
            },
        ]
    )
    canonical = validate_profile(spec, FIELDS)
    compiled = compile_profile(canonical)
    box = engine.resolve_anchor(workbook["cached"], compiled["table"]["selector"]["key"]["areas"][0], "공정 기록")
    assert (box["range"], [m["range"] for m in box["members"]]) == ("B3:C3", ["B3", "C3"])
    groups, values = run(workbook, canonical)
    group, items = only(groups, values, "table")
    assert group["group_key"] == "table|공정 기록|B3:C3"
    assert group["observed_key"] == "LOT|온도"
    assert [r["range"] for r in group["regions"]["key"]] == ["B3", "C3"]  # composite 키는 멤버별 기록
    assert [(i["item_index"], i["regions"]["value"][0]["range"], i["display_text"]) for i in items] == [
        (0, "B4", "L1"),
        (1, "C4", "10.5 / 3"),
        (2, "B5", "L2"),
        (3, "C5", "20 / 4"),
    ]
    group, items = only(groups, values, "from_composite")
    assert items[0]["regions"]["value"][0]["range"] == "D3" and items[0]["value_text"] == "°C"


def test_ambiguous_anchor(tmp_path):
    path = build_workbook(tmp_path / "dup.xlsx", duplicate_temp=True)
    raw, cached = load_workbook(path), load_workbook(path, data_only=True)
    canonical = validate_profile(definition(), FIELDS)
    with pytest.raises(Problem) as exc:
        list(engine.extract(raw, cached, {"main": ["공정 기록"], "common": ["공통"], "exp": ["10C"]}, compile_profile(canonical).values()))
    assert exc.value.code == "AMBIGUOUS_ANCHOR"
    match = engine.match_profile(canonical, cached)
    assert match["compatibility"] == "incompatible"
    assert {"rule_key": "temperature", "role": "key", "code": "AMBIGUOUS_ANCHOR"} in match["missing"]
    # occurrence를 지정하면 해소된다.
    fixed = definition()
    fixed["anchors"]["hdr_temp"]["find"]["occurrence"] = 0
    canonical = validate_profile(fixed, FIELDS)
    assert engine.match_profile(canonical, cached)["compatibility"] == "compatible"
    # 일반 find 영역은 모든 히트를 돌려준다.
    spec = definition(rules=[{"rule_key": "all", "selector": {"key": {"areas": [{"sheet_role": "main", "range": "A1"}]}, "value": {"areas": [{"sheet_role": "main", "find": {"texts": ["온도"], "within": "A1:Z60"}}], "cardinality": "list", "axis": "down"}}}])
    groups, values = run({"raw": raw, "cached": cached}, validate_profile(spec, FIELDS), {"main": ["공정 기록"], "common": ["공통"], "exp": ["10C"]})
    _, items = only(groups, values, "all")
    assert [i["regions"]["value"][0]["range"] for i in items] == ["C3", "H20"]


def test_anchor_not_found_and_missing_sheet_binding(workbook):
    spec = definition()
    spec["anchors"]["sign"]["find"]["texts"] = ["없는 라벨"]
    canonical = validate_profile(spec, FIELDS)
    with pytest.raises(Problem) as exc:
        run(workbook, canonical)
    assert exc.value.code == "ANCHOR_NOT_FOUND"
    with pytest.raises(Problem) as exc:
        run(workbook, validate_profile(definition(), FIELDS), {"common": ["공통"]})
    assert exc.value.code == "INVALID_SHEET_BINDING"


# ---------------------------------------------------------------------------- relations


def test_relation_same_row_derives_record_key(workbook):
    canonical = validate_profile(definition(), FIELDS)
    groups, values = run(workbook, canonical)
    _, lots = only(groups, values, "lot")
    assert [(i["record_key"], i["value_text"]) for i in lots] == [("r4", "L1"), ("r5", "L2"), ("r6", "L3")]
    assert all(i["derivation"] is None for i in lots)
    _, temps = only(groups, values, "temperature")
    assert [i["record_key"] for i in temps] == ["r4", "r5", "r6"]
    assert all(i["derivation"] == "relation:lot:same_row" for i in temps)
    assert [i["regions"]["record_key"][0]["range"] for i in temps] == ["B4", "B5", "B6"]


def test_relation_topological_order_offset_and_fallback(workbook):
    spec = definition()
    # temperature가 먼저 선언되어도 to_rule(lot)이 먼저 실행된다. 짝이 없는 항목은 record_spec.key로 돌아간다.
    spec["rules"] = [spec["rules"][1], spec["rules"][0]]
    spec["rules"][0]["relations"] = [{"to_rule": "lot", "kind": "offset", "row": 0, "col": -1}]
    spec["rules"][0]["record_spec"] = {"scope": ["t"], "key": "coordinate"}
    spec["rules"][1]["selector"]["value"]["areas"][0]["relative"]["rows"] = 2  # lot은 2개만
    spec["rules"][1]["selector"]["value"]["stop"] = {"kind": "explicit_areas"}
    spec["rules"][1]["record_spec"] = {"scope": ["t"], "key": {"column": "B"}}
    canonical = validate_profile(spec, FIELDS)
    groups, values = run(workbook, canonical)
    _, temps = only(groups, values, "temperature")
    assert [(i["record_key"], i["derivation"]) for i in temps] == [("L1", "relation:lot:offset"), ("L2", "relation:lot:offset"), ("C6", None)]
    assert temps[0]["regions"]["record_key"][0]["range"] == "B4"


def test_relation_same_column_with_horizontal_lists(workbook):
    spec = definition(
        rules=[
            {
                "rule_key": "head",
                "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["항목"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1, "cols": 3}}], "cardinality": "list", "axis": "right"}},
                "record_spec": {"scope": ["h"], "key": "physical_column"},
            },
            {
                "rule_key": "val",
                "field_key": "pressure",
                "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["값"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1, "cols": 3}}], "cardinality": "list", "axis": "right"}},
                "value_spec": {"type": "decimal"},
                "record_spec": {"scope": ["h"], "key": None},
                "relations": [{"to_rule": "head", "kind": "same_column"}],
            },
            {
                "rule_key": "orphan",
                "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["합계"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1, "cols": 4}}], "cardinality": "list", "axis": "right"}},
                "record_spec": {"scope": ["h"], "key": None},
                "relations": [{"to_rule": "head", "kind": "same_row"}],
            },
        ]
    )
    canonical = validate_profile(spec, FIELDS)
    groups, values = run(workbook, canonical)
    _, heads = only(groups, values, "head")
    assert [(i["record_key"], i["value_text"]) for i in heads] == [("c8", "a"), ("c9", "b"), ("c10", "c")]
    _, vals = only(groups, values, "val")
    assert [(i["record_key"], i["value_text"], i["derivation"]) for i in vals] == [("c8", "1", "relation:head:same_column"), ("c9", "2", "relation:head:same_column"), ("c10", "3", "relation:head:same_column")]
    _, orphans = only(groups, values, "orphan")
    # 짝도 record_spec.key도 없으면 NULL(실패 아님). explicit_areas라 빈 셀은 empty로 남는다.
    assert [(i["record_key"], i["value_state"]) for i in orphans] == [(None, "present"), (None, "present"), (None, "empty"), (None, "empty")]


# ---------------------------------------------------------------------------- list / matrix / merged / stop / states


def test_list_blank_run_merged_and_states(workbook):
    spec = definition(
        rules=[
            {
                "rule_key": "notes",
                "field_key": "note",
                "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["비고"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 4}}], "cardinality": "list", "axis": "down", "stop": {"kind": "explicit_areas"}}},
            },
            {
                "rule_key": "lots",
                "field_key": "lot",
                "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["LOT"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 20}}], "cardinality": "list", "axis": "down", "stop": {"kind": "blank_run", "count": 3}}},
            },
            {
                "rule_key": "bad_number",
                "field_key": "pressure",
                "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["오류"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1}}]}},
                "value_spec": {"type": "decimal"},
            },
            {
                "rule_key": "no_fragment",
                "field_key": "note",
                "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["오류"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1}}]}},
                "value_spec": {"type": "text", "normalization": {"operation": "pipeline", "steps": [{"op": "split_delimiter", "delimiter": "/", "index": 3}]}},
            },
        ]
    )
    canonical = validate_profile(spec, FIELDS)
    groups, values = run(workbook, canonical)
    _, notes = only(groups, values, "notes")
    # E4:E5 병합은 한 항목, E6 값, E7 빈칸(preserve).
    assert [(i["item_index"], i["regions"]["value"][0]["range"], i["value_state"], i["value_text"]) for i in notes] == [
        (0, "E4:E5", "present", "a"),
        (1, "E6", "present", "c"),
        (2, "E7", "empty", None),
    ]
    assert notes[2]["value_type"] == "text" and notes[2]["display_text"] == ""
    _, lots = only(groups, values, "lots")
    # blank_run 3: 7·8행 빈칸 뒤 9행 값까지 읽고 10~12행 빈칸에서 멈춘다(빈칸은 항목으로 남는다).
    assert [(i["item_index"], i["value_state"]) for i in lots] == [(0, "present"), (1, "present"), (2, "present"), (3, "empty"), (4, "empty"), (5, "present")]
    _, bad = only(groups, values, "bad_number")
    assert (bad[0]["value_state"], bad[0]["value_text"], bad[0]["raw_text"], bad[0]["value_type"]) == ("error", None, "abc", "decimal")
    _, none = only(groups, values, "no_fragment")
    assert (none[0]["value_state"], none[0]["value_text"], none[0]["raw_text"]) == ("null", None, "abc")


def test_matrix_orders_and_scalar_combine(workbook):
    spec = definition(
        rules=[
            {"rule_key": "grid", "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["항목"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1, "rows": 2, "cols": 3}}], "cardinality": "matrix", "axis": "column_major"}}},
            {"rule_key": "total", "field_key": "pressure", "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["합계"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1, "cols": 2}}], "combine": "sum"}}, "value_spec": {"type": "decimal"}},
            {"rule_key": "joined", "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["항목"]}}]}, "value": {"areas": [{"sheet_role": "main", "range": "H12"}, {"sheet_role": "main", "range": "I12"}], "combine": "concat", "separator": "-"}}},
            {"rule_key": "multi", "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["항목"]}}]}, "value": {"areas": [{"sheet_role": "main", "range": "H12:I12"}]}}},
        ]
    )
    canonical = validate_profile(spec, FIELDS)
    with pytest.raises(Problem) as exc:
        run(workbook, canonical, rules=["multi"])
    assert exc.value.code == "MULTIPLE_SCALAR_VALUES"
    groups, values = run(workbook, canonical, rules=["grid", "total", "joined"])
    _, grid = only(groups, values, "grid")
    assert [i["display_text"] for i in grid] == ["a", "1", "b", "2", "c", "3"]
    assert [i["record_key"] for i in grid][:2] == ["H12", "H13"]
    _, total = only(groups, values, "total")
    assert (total[0]["value_text"], total[0]["value_state"], total[0]["value_type"]) == ("9", "present", "decimal")
    assert [r["range"] for r in total[0]["regions"]["input"]] == ["H15", "I15"] and "value" not in total[0]["regions"]
    _, joined = only(groups, values, "joined")
    assert joined[0]["value_text"] == "a-b"


def test_values_events_batch_at_200(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "big"
    ws["A1"] = "값"
    for r in range(2, 452):
        ws.cell(r, 1, r)
    path = tmp_path / "big.xlsx"
    wb.save(path)
    raw, cached = load_workbook(path), load_workbook(path, data_only=True)
    spec = definition(
        sheet_roles={"main": {"match": {"name": "big"}}},
        anchors={},
        rules=[{"rule_key": "v", "selector": {"key": {"areas": [{"sheet_role": "main", "range": "A1"}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 600}}], "cardinality": "list", "axis": "down", "stop": {"kind": "blank_run", "count": 1}}}}],
    )
    events = list(engine.extract(raw, cached, {"main": ["big"]}, compile_profile(validate_profile(spec, FIELDS)).values()))
    sizes = [len(e["items"]) for e in events if e["type"] == "values"]
    assert sizes == [200, 200, 50]
    assert events[-1]["items"][-1]["item_index"] == 449


def test_key_each_and_many_role(workbook):
    spec = definition(
        rules=[
            {
                "rule_key": "exp_name",
                "selector": {"key": {"areas": [{"sheet_role": "exp", "find": {"texts": ["실험"]}}]}, "value": {"areas": [{"sheet_role": "exp", "relative": {"row": 0, "col": 1}}]}},
            },
            {
                "rule_key": "letters",
                "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["a", "b"], "within": "A1:Z60"}}], "repeat": "each"}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0}}]}},
            },
        ]
    )
    canonical = validate_profile(spec, FIELDS)
    groups, values = run(workbook, canonical)
    keys = sorted(k for k in groups if groups[k]["rule_key"] == "exp_name")
    assert keys == ["exp_name|10C|A1", "exp_name|20C|A1"]
    assert [values[k][0]["value_text"] for k in keys] == ["10C", "20C"]
    keys = [k for k in groups if groups[k]["rule_key"] == "letters"]
    assert keys == ["letters|공정 기록|E4:E5", "letters|공정 기록|H12", "letters|공정 기록|I12"]
    assert [values[k][0]["display_text"] for k in keys] == ["a", "1", "2"]  # E5는 병합 E4:E5 안


# ---------------------------------------------------------------------------- match_profile / match_specs


def test_match_profile_resolves_anchor_outside_signature_window(workbook):
    canonical = validate_profile(definition(), FIELDS)
    match = engine.match_profile(canonical, workbook["cached"])
    assert match["compatibility"] == "compatible" and match["missing"] == []
    assert match["bindings"] == {"main": ["공정 기록"], "common": ["공통"], "exp": ["10C", "20C"]}
    assert match["resolved"]["signer"] == {"key": "공정 기록!A55", "value": "공정 기록!B55"}
    assert match["resolved"]["temperature"] == {"key": "공정 기록!C3", "value": "공정 기록!C4:C68", "unit": "공정 기록!D3", "context": "공정 기록!A1"}
    rows = match["match_signature_json"]
    assert [r[0] for r in rows] == ["lot", "signer", "temperature"]
    assert rows[0] == ["lot", "main", "공정 기록!B3", "공정 기록!B4:B68", None, None]
    assert len(match["match_signature"]) == 64
    # 같은 워크북·같은 프로파일이면 서명이 결정적이다.
    assert engine.match_profile(canonical, workbook["cached"])["match_signature"] == match["match_signature"]


def test_match_profile_identical_compatible_incompatible(tmp_path, workbook):
    canonical = validate_profile(definition(), FIELDS)
    first = engine.match_profile(canonical, workbook["cached"], profile_rev=3)
    reference = {"signature": first["match_signature"], "profile_rev": 3}
    assert engine.match_profile(canonical, workbook["cached"], None, reference, 3)["compatibility"] == "identical"
    # 프로파일 rev가 바뀌면 같은 서명이어도 identical이 아니다.
    assert engine.match_profile(canonical, workbook["cached"], None, reference, 4)["compatibility"] == "compatible"
    # 같은 양식이지만 표가 한 행 아래로 밀린 문서: 서명이 달라져 compatible.
    shifted = load_workbook(build_workbook(tmp_path / "shifted.xlsx", shift=1), data_only=True)
    moved = engine.match_profile(canonical, shifted, None, reference, 3)
    assert moved["compatibility"] == "compatible" and moved["resolved"]["lot"]["key"] == "공정 기록!B4"
    assert moved["match_signature"] != first["match_signature"]
    # 시트가 없으면 incompatible + missing[{role}].
    without = load_workbook(build_workbook(tmp_path / "nocommon.xlsx", drop_common=True), data_only=True)
    lost = engine.match_profile(canonical, without, None, reference, 3)
    assert lost["compatibility"] == "incompatible"
    assert lost["missing"] == [{"role": "common", "code": "SHEET_NOT_FOUND"}]
    assert "common" not in lost["bindings"] and set(lost["resolved"]) == {"lot", "signer", "temperature"}
    # 규칙의 영역이 없으면 missing[{rule_key, role}].
    broken = definition()
    broken["rules"][2]["selector"]["value"]["areas"] = [{"sheet_role": "main", "find": {"texts": ["없음"]}}]
    broken["anchors"]["hdr_lot"]["find"]["regex"] = "^배치$"
    lost = engine.match_profile(validate_profile(broken, FIELDS), workbook["cached"])
    assert lost["compatibility"] == "incompatible"
    assert lost["missing"] == [
        {"rule_key": "lot", "role": "key", "code": "ANCHOR_NOT_FOUND"},
        {"rule_key": "signer", "role": "value", "code": "VALUE_NOT_FOUND"},
    ]
    # one 역할에 시트가 둘이면 AMBIGUOUS_SHEET, match가 없는 역할은 MATCH_REQUIRED.
    ambiguous = definition()
    ambiguous["sheet_roles"]["exp"]["cardinality"] = "one"
    ambiguous["sheet_roles"]["manual"] = {"cardinality": "one"}
    lost = engine.match_profile(validate_profile(ambiguous, FIELDS), workbook["cached"])
    assert lost["missing"] == [{"role": "exp", "code": "AMBIGUOUS_SHEET"}, {"role": "manual", "code": "MATCH_REQUIRED"}]


def test_sheet_matchers(workbook):
    wb = workbook["cached"]
    sheets = [{"name": n, "ordinal": i} for i, n in enumerate(wb.sheetnames)]
    roles = {
        "by_ordinal": {"cardinality": "one", "match": {"ordinal": 1}},
        "by_text_all": {"cardinality": "one", "match": {"contains_text": {"texts": ["온도 단위", "°c"], "within": "A1:AD30", "mode": "all"}}},
        "by_text_missing": {"cardinality": "one", "match": {"contains_text": {"texts": ["온도 단위", "없음"], "within": "A1:AD30", "mode": "all"}}},
        "many": {"cardinality": "many", "match": {"contains_text": {"texts": ["실험"], "within": "A1:AD30", "mode": "any"}}},
        "outside": {"cardinality": "one", "match": {"contains_text": {"texts": ["서명"], "within": "A1:AD30", "mode": "any"}}},
    }
    bindings, missing = engine.bind_sheets(roles, wb, sheets)
    assert bindings == {"by_ordinal": ["공통"], "by_text_all": ["공통"], "many": ["10C", "20C"]}
    assert missing == [{"role": "by_text_missing", "code": "SHEET_NOT_FOUND"}, {"role": "outside", "code": "SHEET_NOT_FOUND"}]


def test_match_specs_with_bindings_hint(tmp_path, workbook):
    canonical = validate_profile(definition(), FIELDS)
    specs = list(compile_profile(canonical).values())
    hint = {"main": ["공정 기록"], "common": ["공통"], "exp": ["10C", "20C"]}
    result = engine.match_specs(specs, workbook["cached"], None, hint)
    assert result["compatibility"] == "compatible" and result["missing"] == []
    assert result["match_signature"] == engine.match_profile(canonical, workbook["cached"])["match_signature"]
    without = load_workbook(build_workbook(tmp_path / "nocommon.xlsx", drop_common=True), data_only=True)
    lost = engine.match_specs(specs, without, None, hint)
    assert lost["compatibility"] == "incompatible" and lost["missing"] == [{"role": "common", "code": "SHEET_NOT_FOUND"}]
    assert lost["bindings"] == {"main": ["공정 기록"], "exp": ["10C", "20C"]}


# ---------------------------------------------------------------------------- Reader


@pytest.fixture
def reader_root(tmp_path):
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    build_workbook(raw / "sample.xlsx")
    return tmp_path


def test_reader_describe_with_profiles_returns_matches(reader_root):
    reader = make_reader(reader_root, "local-xlsx", "tester")
    assert isinstance(reader, XlsxReader)
    plain = reader.describe("sample.xlsx")
    assert "matches" not in plain and [s["name"] for s in plain["sheets"]] == ["공정 기록", "공통", "10C", "20C"]
    assert plain["signature"]["sheets"][0]["name"] == "공정 기록"
    canonical = validate_profile(definition(), FIELDS)
    other = validate_profile(definition(sheet_roles={"main": {"match": {"name": "없는 시트"}}}, anchors={}, rules=[{"rule_key": "x", "selector": {"key": {"areas": [{"sheet_role": "main", "range": "A1"}]}, "value": {"areas": [{"sheet_role": "main", "range": "B1"}]}}}]), FIELDS)
    first = reader.describe("sample.xlsx", [{"profile_id": "p1", "profile_rev": 2, "canonical": canonical, "reference": None}])
    assert first["token"] == plain["token"]
    (m,) = first["matches"]
    assert (m["profile_id"], m["profile_rev"], m["compatibility"]) == ("p1", 2, "compatible")
    assert m["bindings"]["main"] == ["공정 기록"] and m["missing"] == [] and "temperature" in m["resolved"]
    profiles = [
        {"profile_id": "p1", "profile_rev": 2, "canonical": canonical, "reference": {"signature": m["match_signature"], "profile_rev": 2}},
        {"profile_id": "p2", "profile_rev": 1, "canonical": other, "reference": None},
    ]
    described = reader.describe("sample.xlsx", profiles)
    assert [(x["profile_id"], x["compatibility"]) for x in described["matches"]] == [("p1", "identical"), ("p2", "incompatible")]
    assert described["matches"][1]["missing"] == [{"role": "main", "code": "SHEET_NOT_FOUND"}]
    matched = reader.match("sample.xlsx", plain["token"], profiles)
    assert [(x["profile_id"], x["compatibility"]) for x in matched] == [("p1", "identical"), ("p2", "incompatible")]
    specs = list(compile_profile(canonical).values())
    inherited = reader.match_specs("sample.xlsx", plain["token"], specs, m["bindings"])
    assert inherited["match_signature"] == m["match_signature"] and inherited["compatibility"] == "compatible"
    with pytest.raises(Problem) as exc:
        reader.match("sample.xlsx", "stale", profiles)
    assert exc.value.code == "SOURCE_VERSION_CHANGED"


def test_reader_extract_stream_ends_with_verified(reader_root):
    reader = make_reader(reader_root, "local-xlsx", "tester")
    token = reader.describe("sample.xlsx")["token"]
    canonical = validate_profile(definition(), FIELDS)
    specs = list(compile_profile(canonical).values())
    events = list(reader.extract("sample.xlsx", token, specs, {"main": ["공정 기록"], "common": ["공통"], "exp": ["10C"]}))
    assert [e["type"] for e in events] == ["group", "values", "group", "values", "group", "values", "verified"]
    assert events[-1] == {"type": "verified", "token": token}
    assert events[3]["items"][0]["value_text"] == "10.5"
    with pytest.raises(Problem) as exc:
        list(reader.extract("sample.xlsx", "stale", specs, {"main": ["공정 기록"]}))
    assert exc.value.code == "SOURCE_VERSION_CHANGED"


def test_make_reader_requires_registered_factory(reader_root, monkeypatch):
    monkeypatch.delenv("SCHEMA_READER_FACTORY", raising=False)
    with pytest.raises(Problem) as exc:
        make_reader(reader_root, "drm-x", "tester")
    assert exc.value.code == "DRM_READER_REQUIRED"
    monkeypatch.setenv("SCHEMA_READER_FACTORY", "tests.reader_fixture:factory")
    assert make_reader(reader_root, "revocable-xlsx", "tester").__class__.__name__ == "Revocable"


@pytest.mark.skipif(importlib.util.find_spec("schema.render") is None or importlib.util.find_spec("schema.render.renderer") is None, reason="render module not present yet")
def test_reader_render_stream(reader_root):
    reader = make_reader(reader_root, "local-xlsx", "tester")
    token = reader.describe("sample.xlsx")["token"]
    events = list(reader.render("sample.xlsx", token, "공정 기록", 1, 1, 100, 20))
    assert events[0]["type"] == "meta" and events[-1] == {"type": "verified", "token": token}
    assert any(e["type"] == "band" for e in events)
