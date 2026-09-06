"""watcher 이관 회귀 — 현행 파이프라인(kg.cli watch)이 raw를 감시해 자동
등록하고, 잠긴 파일은 우회 없이 건너뛴다. src.watch 경로는 호환 셔틀."""
from __future__ import annotations

import shutil

from openpyxl import Workbook

from tests.conftest import FIXTURES

ROOT = FIXTURES.parent.parent
FIN_CONFIG = ROOT / "domains" / "financier" / "config"


def _xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "190도"
    ws["A1"] = "오븐온도"
    ws["A2"] = 190
    wb.save(path)


def test_legacy_import_path_still_works():
    from src.watch.watcher import FileEventWatcher as legacy

    from kg.watch import FileEventWatcher
    assert legacy is FileEventWatcher     # 셔틀 re-export — 같은 구현


def test_kg_watch_ingests_new_files_and_skips_locked(tmp_path, capsys):
    ws = tmp_path / "ws"
    (ws / "config").mkdir(parents=True)
    raw = ws / "data" / "raw"
    raw.mkdir(parents=True)
    for f in ("domain_kg.yaml", "units.yaml"):
        shutil.copy(FIN_CONFIG / f, ws / "config" / f)

    from kg.cli import main
    assert main(["--ws", str(ws), "seed"]) == 0

    _xlsx(raw / "새실험.xlsx")
    (raw / "잠김.xlsx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    # 안정화 판정(2회 연속 동일)까지 포함해 3회 스캔이면 등록이 끝난다
    assert main(["--ws", str(ws), "watch", "--raw", str(raw),
                 "--interval", "0", "--scans", "3", "--no-map"]) == 0

    from kg.store import KgStore
    store = KgStore(ws / "data" / "kg" / "kg.db")
    docs = {r[0] for r in store.conn.execute("SELECT filename FROM document")}
    assert "새실험.xlsx" in docs                      # 새 파일 자동 등록
    assert "잠김.xlsx" not in docs                    # 잠긴 파일은 우회하지 않는다
    assert "잠김" in capsys.readouterr().err          # 스킵을 드러낸다
    store.close()
