"""서비스(§4.1–§4.9, §4.12): 정의 가져오기 projection, 등록(Reader 1회), 수동 적용·검수·승인·rematch 소급,
자동 승인·추출·발행, compatible 검수, CAS 충돌, 역방향 조회, 새 snapshot 승계, 잠긴 파일, dry-run, 목록·검색."""

from __future__ import annotations

import copy
import types

import pytest
from openpyxl import Workbook

from schema import service as service_module
from schema.db import Problem
from schema.service import Service

SCHEMA = {
    "format": "parsing-schema",
    "schema_version": "3.0",
    "schema_key": "process_standard",
    "schema_name": "공정 데이터 표준",
    "fields": [
        {"field_key": "process", "name": "공정 정보", "type": "group", "level": 1},
        {"field_key": "lot", "name": "배치", "type": "text", "level": 2, "parents": "process", "aliases": ["LOT", "배치번호"]},
        {"field_key": "temperature", "name": "온도", "type": "decimal", "unit": "°C", "level": 2, "parents": ["process"], "aliases": ["온도값", "Temp"], "related": ["pressure"]},
        {"field_key": "pressure", "name": "압력", "type": "decimal", "unit": "bar", "level": 2, "parents": ["process"]},
        {"field_key": "note", "name": "비고", "type": "text"},
    ],
}


def profile_definition(**over):
    base = {
        "format": "parsing-profile",
        "schema_version": "3.0",
        "profile_name": "공정데이터_A양식",
        "schema_key": "process_standard",
        "sheet_roles": {"main": {"cardinality": "one", "match": {"name": "공정 기록"}}},
        "anchors": {
            "hdr_temp": {"sheet_role": "main", "find": {"texts": ["온도"], "within": "A1:Z60"}},
            "hdr_lot": {"sheet_role": "main", "find": {"regex": "^(배치|LOT)$", "within": "A1:Z60"}},
            "sign": {"sheet_role": "main", "find": {"texts": ["서명"], "within": "A1:Z60"}},
        },
        "rules": [
            {
                "rule_key": "lot",
                "rule_name": "배치",
                "field_key": "lot",
                "selector": {
                    "key": {"areas": [{"sheet_role": "main", "anchor": "hdr_lot"}]},
                    "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 20, "cols": 1}}], "cardinality": "list", "axis": "down", "stop": {"kind": "blank_run", "count": 2}},
                },
                "value_spec": {"type": "text"},
                "record_spec": {"scope": ["process-table"], "key": "physical_row"},
            },
            {
                "rule_key": "temperature",
                "rule_name": "온도",
                "field_key": "temperature",
                "selector": {
                    "key": {"areas": [{"sheet_role": "main", "anchor": "hdr_temp"}]},
                    "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 20, "cols": 1}}], "cardinality": "list", "axis": "down", "stop": {"kind": "blank_run", "count": 2}},
                    "unit": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1, "anchor": "hdr_temp"}}]},
                },
                "value_spec": {"type": "decimal", "unit": "°C"},
                "record_spec": {"scope": ["process-table"], "key": "physical_row"},
                "relations": [{"to_rule": "lot", "kind": "same_row"}],
            },
            {
                "rule_key": "signer",
                "rule_name": "서명자",
                "field_key": "note",
                "selector": {"key": {"areas": [{"sheet_role": "main", "anchor": "sign"}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1}}]}},
                "value_spec": {"type": "text"},
            },
        ],
    }
    base.update(over)
    return base


def build_workbook(path, shift=0, temps=(10.5, 20, 30), lots=("L1", "L2", "L3"), signer="홍길동"):
    wb = Workbook()
    ws = wb.active
    ws.title = "공정 기록"
    ws["A1"] = "공정 A"
    ws.cell(3 + shift, 2, "LOT")
    ws.cell(3 + shift, 3, "온도")
    ws.cell(3 + shift, 4, "°C")
    for n, (lot, temp) in enumerate(zip(lots, temps)):
        ws.cell(4 + shift + n, 2, lot)
        ws.cell(4 + shift + n, 3, temp)
    ws["A55"], ws["B55"] = "서명", signer
    wb.save(path)
    return path


class World:
    """작업 공간 + Reader 호출 기록 + 렌더 무효화 기록."""

    def __init__(self, root, monkeypatch):
        (root / "data/raw").mkdir(parents=True)
        self.root = root
        self.calls = []
        self.invalidated = []
        original_result, original_events = service_module.reader_result, service_module.reader_events

        def counted_result(root_, provider, principal, operation, payload, checkpoint=lambda: None):
            self.calls.append(operation)
            return original_result(root_, provider, principal, operation, payload, checkpoint)

        def counted_events(root_, provider, principal, operation, payload, checkpoint=lambda: None):
            self.calls.append(operation)
            return original_events(root_, provider, principal, operation, payload, checkpoint)

        monkeypatch.setattr(service_module, "reader_result", counted_result)
        monkeypatch.setattr(service_module, "reader_events", counted_events)
        self.service = Service(root)
        self.service._render = types.SimpleNamespace(invalidate=self.invalidated.append, close=lambda: None)

    def file(self, name, **kw):
        return build_workbook(self.root / "data/raw" / name, **kw)

    def register(self, name):
        job = self.service.register_documents([name], wait=60)
        assert job["state"] == "succeeded", job
        return job["result"]["documents"][0]

    def status(self, document_id):
        return self.service.document(document_id)["status"]

    def mapping_ids(self, application_id):
        return {m["rule_key"]: m for m in self.service.application_mappings(application_id)}


@pytest.fixture
def world(tmp_path, monkeypatch):
    w = World(tmp_path, monkeypatch)
    yield w
    w.service.close()


@pytest.fixture
def approved(world):
    """스키마·프로파일(draft) → 문서 A 등록(unmatched) → 수동 적용 → approve_all(추출·발행) → 프로파일 승인 + rematch."""
    s = world.service
    s.import_schema(SCHEMA)
    profile = s.import_profile("process_standard", profile_definition())
    world.file("a.xlsx")
    doc = world.register("a.xlsx")
    assert world.calls == ["describe"] and doc["status"] == "unmatched" and doc["applied"] == []
    app = s.apply_profile(doc["snapshot"]["snapshot_id"], profile["profile_id"])
    assert app["origin"] == "manual" and app["compatibility"] == "manual" and app["heads_total"] == 3 and app["heads_approved"] == 0
    assert all(m["status"] == "proposed" and m["origin"] == "profile" for m in app["mappings"])
    assert world.status(doc["document_id"]) == "review"
    result = s.approve_all(app["application_id"], reason="검수 완료")
    assert result["approved"] == 3 and result["skipped"] == [] and result["extraction"]["state"] == "succeeded", result
    assert world.status(doc["document_id"]) == "normal"
    approved = s.approve_profile(profile["profile_id"], app["application_id"])
    assert approved["status"] == "approved" and approved["reference_signature"]
    rematch = s.jobs.wait(approved["reparse_job"]["job_id"], 60)
    assert rematch["state"] == "succeeded" and rematch["result"]["queued"] == 0
    assert rematch["result"]["skipped"] == [{"document_id": doc["document_id"], "document_name": "a.xlsx", "reason": "up_to_date"}]
    world.calls.clear()
    return {"profile": profile, "document": doc, "application": app}


# ---------------------------------------------------------------------------- 정의 가져오기


def test_schema_import_projection_and_reimport_deprecates(world):
    s = world.service
    first = s.import_schema(SCHEMA)
    assert first["current_rev"] == 1 and first["fields"] == {"total": 5, "added": 5, "updated": 0, "deprecated": 0}
    fields = s.schema_fields("process_standard")
    assert fields["temperature"]["aliases"] == ["온도값", "Temp"] and fields["process"]["value_type"] == "group"
    assert (world.root / "schemas/process_standard/r0001.json").is_file() and (world.root / "schemas/process_standard/current.json").is_file()
    with s.db.connect() as conn:
        edges = conn.execute("SELECT relation, count(*) FROM parsing_field_edge GROUP BY relation").fetchall()
    assert {tuple(e) for e in edges} == {("parent_of", 3), ("related_to", 1)}
    assert s.import_schema(SCHEMA)["unchanged"] is True
    smaller = copy.deepcopy(SCHEMA)
    smaller["fields"] = [f for f in smaller["fields"] if f["field_key"] != "pressure"]
    smaller["fields"][2].pop("related")
    smaller["fields"][2]["aliases"] = ["온도값"]
    second = s.import_schema(smaller)
    assert second["current_rev"] == 2 and second["fields"] == {"total": 4, "added": 0, "updated": 4, "deprecated": 1}
    fields = s.schema_fields("process_standard")
    assert fields["pressure"]["status"] == "deprecated" and fields["temperature"]["aliases"] == ["온도값"]
    with s.db.connect() as conn:
        assert conn.execute("SELECT count(*) FROM parsing_field").fetchone()[0] == 5  # 삭제 없음
    bad = copy.deepcopy(SCHEMA)
    bad["fields"][1]["level"] = 3
    with pytest.raises(Problem) as exc:
        s.import_schema(bad)
    assert exc.value.code == "INVALID_SCHEMA"
    patched = s.patch_field("process_standard", "lot", aliases=["로트"], description="배치 식별자")
    assert patched["current_rev"] == 3 and s.schema_fields("process_standard")["lot"]["aliases"] == ["로트"]


def test_import_schema_modes_and_delete_paths(world):
    """§4.2 create/revision 구분 · §4.2.1 스키마 삭제 · §4.2.2 필드 삭제(참조 검사·새 리비전)."""
    s = world.service
    created = s.import_schema(SCHEMA, mode="create")
    assert created["current_rev"] == 1 and created["schema_name"] == "공정 데이터 표준"

    # 같은 키 create → 409, 파일도 리비전도 늘지 않는다.
    with pytest.raises(Problem) as exc:
        s.import_schema(SCHEMA, mode="create")
    assert exc.value.code == "SCHEMA_EXISTS" and exc.value.status == 409
    assert sorted(p.name for p in (world.root / "schemas/process_standard").iterdir()) == ["current.json", "r0001.json"]

    # 없는 키 revision → 404, 폴더를 만들지 않는다.
    absent = copy.deepcopy(SCHEMA)
    absent["schema_key"], absent["schema_name"] = "absent", "없는 스키마"
    with pytest.raises(Problem) as exc:
        s.import_schema(absent, mode="revision")
    assert exc.value.code == "UNKNOWN_SCHEMA" and exc.value.status == 404
    assert not (world.root / "schemas/absent").exists()

    # 필드 추가 → 새 리비전, 삭제 → 그 필드를 뺀 새 리비전 + projection 행 제거.
    s.create_field("process_standard", "note2", "비고2", parent_field_key="process")
    assert s.schema_fields("process_standard")["note2"]["field_id"]
    removed = s.delete_field("process_standard", "note2")
    assert removed == {"schema_key": "process_standard", "field_key": "note2", "name": "비고2", "current_rev": 3, "fields_remaining": 5}
    assert "note2" not in s.schema_fields("process_standard")
    assert [f["field_key"] for f in s.schema_definition("process_standard")["fields"]] == ["process", "lot", "temperature", "pressure", "note"]

    # 자식이 있는 필드는 거부.
    with pytest.raises(Problem) as exc:
        s.delete_field("process_standard", "process")
    assert exc.value.code == "FIELD_HAS_CHILDREN" and exc.value.detail["count"] == 3

    # 프로파일이 쓰면 스키마 삭제 거부, 지우고 나면 폴더까지 사라진다.
    s.import_profile("process_standard", profile_definition())
    with pytest.raises(Problem) as exc:
        s.delete_schema("process_standard")
    assert exc.value.code == "SCHEMA_IN_USE" and exc.value.detail["profile_count"] == 1 and exc.value.detail["document_count"] == 0
    assert (world.root / "schemas/process_standard").is_dir()

    # 아무도 쓰지 않는 스키마는 행과 정의 폴더가 함께 사라진다.
    other = {**copy.deepcopy(SCHEMA), "schema_key": "spare", "schema_name": "여분 스키마"}
    s.import_schema(other, mode="create")
    assert s.delete_schema("spare") == {
        "schema_key": "spare",
        "schema_name": "여분 스키마",
        "deleted": {"fields": 5, "aliases": 4, "edges": 4, "revisions": 1},
    }
    assert not (world.root / "schemas/spare").exists()
    with pytest.raises(Problem) as exc:
        s.delete_schema("spare")
    assert exc.value.status == 404


def test_profile_import_30_and_v2_format_and_rule_projection(world):
    s = world.service
    s.import_schema(SCHEMA)
    created = s.import_profile("process_standard", profile_definition())
    assert created["current_rev"] == 1 and created["status"] == "draft" and created["rules"] == 3
    assert created["report"]["format_detected"] == "parsing-profile-3.0"
    with s.db.connect() as conn:
        rules = {r["rule_key"]: dict(r) for r in conn.execute("SELECT r.*, f.field_key FROM parsing_rule r LEFT JOIN parsing_field f ON f.field_id=r.default_field_id")}
    assert rules["temperature"]["field_key"] == "temperature" and rules["temperature"]["status"] == "active"
    updated = s.import_profile("process_standard", profile_definition(rules=profile_definition()["rules"][:2]), profile_id=created["profile_id"])
    assert updated["current_rev"] == 2 and updated["deprecated_rules"] == 1
    with s.db.connect() as conn:
        assert conn.execute("SELECT status FROM parsing_rule WHERE rule_key='signer'").fetchone()[0] == "deprecated"
    assert (world.root / "profiles" / created["profile_id"] / "r0002.json").is_file()
    v2 = {
        "kg_revision_id": "kg-1",
        "sheet_roles": {"main": {"cardinality": "one", "match": {"name": "공정 기록"}}},
        "rules": [
            {"rule_key": "temp", "concept_id": "temperature", "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["온도"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 1}}]}}, "value_spec": {"type": "decimal", "unit": "°C"}},
            {"rule_key": "free", "concept_id": None, "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["비고"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"col": 1}}]}}},
        ],
    }
    imported = s.import_profile("process_standard", v2, name="v2 이관")
    assert imported["report"]["format_detected"] == "v2-template"
    assert {w["code"] for w in imported["report"]["warnings"]} >= {"DROPPED_FIELD", "MISSING_FIELD"}
    with s.db.connect() as conn:
        free = conn.execute("SELECT default_field_id FROM parsing_rule WHERE profile_id=? AND rule_key='free'", (imported["profile_id"],)).fetchone()
    assert free["default_field_id"] is None
    with pytest.raises(Problem) as exc:
        s.import_profile("process_standard", profile_definition(), name="공정데이터_A양식")
    assert exc.value.code == "PROFILE_NAME_CONFLICT"
    preview = s.preview_profile("process_standard", {"fields": [{"name": "temperature", "cell": "C4"}]})
    assert preview["format_detected"] == "generic-keyvalue" and preview["errors"] == [] and preview["canonical"]["rules"][0]["field_key"] == "temperature"


# ---------------------------------------------------------------------------- 등록·자동 적용·검수


def test_auto_approve_identical_and_review_compatible(world, approved):
    s = world.service
    world.file("b.xlsx", temps=(1, 2, 3), lots=("X1", "X2", "X3"))
    doc_b = world.register("b.xlsx")
    assert world.calls == ["describe", "extract"], world.calls  # 등록 1회 + 자동 승인 추출 1회
    assert doc_b["status"] == "normal" and [(a["compatibility"], a["state"]) for a in doc_b["applied"]] == [("identical", "published")]
    app_b = s.application_summary(doc_b["applied"][0]["application_id"])
    assert app_b["origin"] == "auto" and app_b["published"] and app_b["heads_approved"] == 3
    reference_heads = {m["rule_key"]: m for m in s.application_mappings(approved["application"]["application_id"])}
    revisions = s.mapping_revisions(app_b["mappings"][0]["mapping_id"])
    assert revisions[0]["origin"] == "auto" and revisions[0]["evidence"]["reference_revision_id"]
    with s.db.connect() as conn:
        ref = conn.execute("SELECT current_revision_id FROM mapping WHERE mapping_id=?", (reference_heads["lot"]["mapping_id"],)).fetchone()[0]
        published_run = conn.execute("SELECT published_run_id FROM parsing_application WHERE application_id=?", (app_b["application_id"],)).fetchone()[0]
        auto = conn.execute("SELECT auto_approved FROM extraction_run WHERE run_id=?", (published_run,)).fetchone()[0]
    assert revisions[0]["evidence"]["reference_revision_id"] == ref and auto == 1
    values = s.snapshot_values(doc_b["snapshot"]["snapshot_id"], field_key="temperature")
    assert [v["value_text"] for v in values["items"]] == ["1", "2", "3"]
    assert values["items"][0]["record_key"] == "r4" and values["items"][0]["derivation_key"] == "relation:lot:same_row"
    assert values["items"][0]["source"]["range"] == "C4" and {r["role"] for r in values["items"][0]["regions"]} == {"key", "value", "unit", "record_key"}
    assert app_b["mappings"][1]["value"]["count"] == 3 and app_b["mappings"][1]["value"]["first_region"]["range"] == "C4"

    world.calls.clear()
    world.file("c.xlsx", shift=1)
    doc_c = world.register("c.xlsx")
    assert world.calls == ["describe"] and doc_c["status"] == "review"
    assert [(a["compatibility"], a["state"]) for a in doc_c["applied"]] == [("compatible", "review")]
    app_c = s.application_summary(doc_c["applied"][0]["application_id"])
    assert all(m["status"] == "proposed" and m["origin"] == "profile" and m["edit_seq"] == 1 for m in app_c["mappings"])
    lot = next(m for m in app_c["mappings"] if m["rule_key"] == "lot")
    assert lot["regions"][0] == {"role": "key", "sheet_id": app_c["sheets"][0]["sheet_id"], "sheet_name": "공정 기록", "range": "B4"}

    # CAS: 이미 edit_seq=1인 매핑에 expected_seq=0으로 쓰면 409, 리비전 수 불변.
    with pytest.raises(Problem) as exc:
        s.revise(lot["mapping_id"], expected_seq=0, status="approved")
    assert exc.value.code == "EDIT_CONFLICT" and len(s.mapping_revisions(lot["mapping_id"])) == 1
    revised = s.revise(lot["mapping_id"], expected_seq=1, status="approved", regions=[{"role": "key", "sheet_id": lot["regions"][0]["sheet_id"], "range": "B4"}, {"role": "value", "sheet_id": lot["regions"][0]["sheet_id"], "range": "B5:B24"}], reason="확인")
    assert revised["status"] == "approved" and revised["edit_seq"] == 2 and revised["extraction"] is None
    assert [r["range"] for r in revised["regions"]] == ["B4", "B5:B24"]
    with pytest.raises(Problem) as exc:
        s.revise(lot["mapping_id"], expected_seq=1, status="proposed")
    assert exc.value.code == "EDIT_CONFLICT"
    assert world.status(doc_c["document_id"]) == "review"

    world.calls.clear()
    done = s.approve_all(app_c["application_id"])
    assert done["approved"] == 2 and done["extraction"]["state"] == "succeeded" and world.calls == ["extract"]
    assert world.status(doc_c["document_id"]) == "normal"
    app_c = s.application_summary(app_c["application_id"])
    assert app_c["published"] and app_c["heads_approved"] == 3
    signer = next(m for m in app_c["mappings"] if m["rule_key"] == "signer")
    assert signer["value"]["value_text"] == "홍길동" and signer["value"]["first_region"]["range"] == "B55"

    # 역방향 조회: 값 → 영역 → 이 셀에서 나온 값과 헤드 매핑.
    detail = s.value(signer["value"]["value_id"])
    region_id = next(r["region_id"] for r in detail["regions"] if r["role"] == "value")
    reverse = s.region_values(region_id)
    assert reverse["region"]["range"] == "B55" and [v["value_id"] for v in reverse["values"]] == [signer["value"]["value_id"]]
    assert reverse["values"][0]["published"] and reverse["mappings"][0]["rule_key"] == "signer" and reverse["mappings"][0]["is_head"]
    grid = s.sheet_regions(reverse["region"]["sheet_id"], "A50:C60")
    assert [r["range"] for r in grid["items"]] == ["A55", "B55"] and grid["items"][1]["value_count"] == 1
    assert [r["range"] for r in s.sheet_regions(reverse["region"]["sheet_id"], "A1:A1")["items"]] == []

    # 값 목록 keyset 페이지.
    first = s.snapshot_values(doc_c["snapshot"]["snapshot_id"], limit=2)
    assert first["has_more"] and len(first["items"]) == 2
    second = s.snapshot_values(doc_c["snapshot"]["snapshot_id"], limit=2, cursor=first["next_cursor"])
    assert {v["value_id"] for v in first["items"]}.isdisjoint({v["value_id"] for v in second["items"]})
    total = len(s.application_values(app_c["application_id"], limit=200)["items"])
    assert total == 7 and len(s.application_values(app_c["application_id"], rule_key="lot")["items"]) == 3


def test_new_snapshot_inherits_as_proposed_and_keeps_old_values(world, approved):
    s = world.service
    doc, app = approved["document"], approved["application"]
    old_sid = doc["snapshot"]["snapshot_id"]
    old_values = s.snapshot_values(old_sid, field_key="temperature")
    assert [v["value_text"] for v in old_values["items"]] == ["10.5", "20", "30"]
    unchanged = world.register("a.xlsx")
    assert unchanged["snapshot"]["unchanged"] and unchanged["snapshot"]["snapshot_id"] == old_sid and world.calls == ["describe"]

    world.calls.clear()
    world.file("a.xlsx", temps=(11, 21, 31))
    changed = world.register("a.xlsx")
    assert world.calls == ["describe", "match_specs"], world.calls
    assert changed["snapshot"]["revision_no"] == 2 and changed["snapshot"]["snapshot_id"] != old_sid
    assert changed["status"] == "changed" and [(a["compatibility"], a["state"]) for a in changed["applied"]] == [("identical", "changed")]
    assert world.invalidated == [old_sid]
    new_app = s.application_summary(changed["applied"][0]["application_id"])
    assert new_app["origin"] == "inherited" and not new_app["published"] and all(m["status"] == "proposed" and m["origin"] == "inherited" for m in new_app["mappings"])
    evidence = s.mapping_revisions(new_app["mappings"][0]["mapping_id"])[0]["evidence"]
    assert evidence["previous_application_id"] == app["application_id"] and evidence["compatibility"] == "identical" and evidence["previous_signature"] == evidence["new_signature"]
    assert [r["range"] for r in new_app["mappings"][1]["regions"]] == ["C3", "C4:C23", "D3"]
    # 이전 snapshot의 값은 그대로, 새 snapshot은 발행 전이라 값이 없다.
    assert [v["value_text"] for v in s.snapshot_values(old_sid, field_key="temperature")["items"]] == ["10.5", "20", "30"]
    assert s.snapshot_values(changed["snapshot"]["snapshot_id"])["items"] == []
    assert [x["revision_no"] for x in s.document_snapshots(doc["document_id"])] == [2, 1]
    listed = s.document_query(status="changed")["items"]
    assert [d["document_id"] for d in listed] == [doc["document_id"]] and listed[0]["profiles"][0]["state"] == "changed"

    world.calls.clear()
    done = s.approve_all(new_app["application_id"])
    assert done["approved"] == 3 and done["extraction"]["state"] == "succeeded" and world.calls == ["extract"]
    assert world.status(doc["document_id"]) == "normal"
    assert [v["value_text"] for v in s.snapshot_values(changed["snapshot"]["snapshot_id"], field_key="temperature")["items"]] == ["11", "21", "31"]
    assert s.application_summary(new_app["application_id"])["published"]


def test_locked_file_marks_document_locked_and_fails_job(world, approved):
    s = world.service
    (world.root / "data/raw/locked.xlsx").write_bytes(b"\xd0\xcf\x11\xe0drm-protected")
    job = s.register_documents(["locked.xlsx"], wait=60)
    assert job["state"] == "failed" and job["error_code"] == "DRM_READER_REQUIRED"
    listed = s.document_query(status="locked")["items"]
    assert len(listed) == 1 and listed[0]["document_name"] == "locked.xlsx" and listed[0]["current_snapshot"] is None
    assert listed[0]["last_error"] and listed[0]["status_detail"]["locked"]["code"] == "DRM_READER_REQUIRED"
    assert s.refresh_document_status(listed[0]["document_id"])["status"] == "locked"
    mixed = s.register_documents(["locked.xlsx", "a.xlsx"], wait=60)
    assert mixed["state"] == "succeeded" and [d["status"] for d in mixed["result"]["documents"]] == ["locked", "normal"]
    assert mixed["result"]["documents"][0]["error"]["code"] == "DRM_READER_REQUIRED"


def test_profile_dry_run_returns_groups_and_records_test_job(world, approved):
    s = world.service
    sid = approved["document"]["snapshot"]["snapshot_id"]
    result = s.test_profile(sid, profile_id=approved["profile"]["profile_id"])
    assert result["compatibility"] == "identical" and result["errors"] == [] and result["bindings"] == {"main": ["공정 기록"]}
    groups = {g["rule_key"]: g for g in result["groups"]}
    assert groups["temperature"]["count"] == 3 and groups["temperature"]["field"]["key"] == "temperature" and groups["temperature"]["observed_key"] == "온도"
    assert groups["temperature"]["values"][0]["value_text"] == "10.5" and groups["temperature"]["values"][0]["range"] == "C4"
    assert groups["temperature"]["regions"]["unit"] == [{"sheet_name": "공정 기록", "range": "D3"}]
    job = s.jobs.get(result["job_id"])
    assert job["kind"] == "test" and job["state"] == "succeeded" and job["result"] == {"errors": 0, "groups": 3, "compatibility": "identical"}
    assert job["label"].startswith("공정데이터_A양식 r1 · a.xlsx")
    with s.db.connect() as conn:
        assert conn.execute("SELECT count(*) FROM parsing_application").fetchone()[0] == 1  # dry-run은 application을 만들지 않는다
    broken = profile_definition(anchors={**profile_definition()["anchors"], "sign": {"sheet_role": "main", "find": {"texts": ["없는 라벨"], "within": "A1:Z60"}}})
    result = s.test_profile(sid, definition=broken, schema_key="process_standard")
    assert result["compatibility"] == "incompatible" and result["missing"] == [{"rule_key": "signer", "role": "key", "code": "ANCHOR_NOT_FOUND"}]
    assert {g["rule_key"] for g in result["groups"]} == {"lot", "temperature"}
    assert result["errors"][0]["code"] == "ANCHOR_NOT_FOUND" and result["errors"][0]["rule_key"] is None
    assert s.jobs.get(result["job_id"])["target_id"] is None


def test_document_query_filters_sort_cursor_and_search(world, approved):
    s = world.service
    world.file("b.xlsx")
    world.file("zz.xlsx", signer=None)
    (world.root / "data/raw/zz.xlsx").rename(world.root / "data/raw/z.xlsx")
    world.register("b.xlsx")
    other = Workbook()
    other.active.title = "다른 양식"
    other.active["A1"] = "무관"
    other.save(world.root / "data/raw/other.xlsx")
    unmatched = world.register("other.xlsx")
    assert unmatched["status"] == "unmatched"
    page1 = s.document_query(sort="document_name", limit=2)
    assert [d["document_name"] for d in page1["items"]] == ["a.xlsx", "b.xlsx"] and page1["has_more"]
    page2 = s.document_query(sort="document_name", limit=2, cursor=page1["next_cursor"])
    assert [d["document_name"] for d in page2["items"]] == ["other.xlsx"] and not page2["has_more"]
    assert [d["document_name"] for d in s.document_query(sort="-document_name")["items"]] == ["other.xlsx", "b.xlsx", "a.xlsx"]
    with pytest.raises(Problem) as exc:
        s.document_query(sort="-document_name", cursor=page1["next_cursor"])
    assert exc.value.code == "INVALID_CURSOR"
    assert [d["document_name"] for d in s.document_query(status="normal", sort="document_name")["items"]] == ["a.xlsx", "b.xlsx"]
    assert [d["document_name"] for d in s.document_query(q="oth")["items"]] == ["other.xlsx"]
    pid = approved["profile"]["profile_id"]
    by_profile = s.document_query(profile_id=pid, sort="document_name")["items"]
    assert [d["document_name"] for d in by_profile] == ["a.xlsx", "b.xlsx"]
    assert by_profile[0]["profiles"][0]["profile_name"] == "공정데이터_A양식" and by_profile[0]["schemas"] == [{"schema_key": "process_standard", "schema_name": "공정 데이터 표준"}]
    assert [d["document_name"] for d in s.document_query(schema_key="process_standard", sort="document_name")["items"]] == ["a.xlsx", "b.xlsx"]
    by_time = s.document_query()["items"]
    assert by_time[0]["document_name"] in ("other.xlsx", "b.xlsx") and len(by_time) == 3
    docs = s.profile_documents(pid)["items"]
    assert [(d["document_name"], d["is_reference"], d["published"]) for d in docs] == [("a.xlsx", True, True), ("b.xlsx", False, True)]
    found = s.search("온도")
    assert [(i["kind"], i["id"]) for i in found["items"]] == [("field", "temperature")]
    kinds = {(i["kind"], i["label"]) for i in s.search("공정")["items"]}
    assert ("profile", "공정데이터_A양식 v1") in kinds and ("schema", "공정 데이터 표준 v1") in kinds
    assert [i["kind"] for i in s.search("a.xlsx")["items"]] == ["document"]
    assert s.search("Temp")["items"][0]["route"] == "?screen=schema&schema=process_standard&field=temperature"
    assert s.status()["counts"] == {"documents": 3, "profiles": 1, "schemas": 1, "jobs_running": 0, "review": 0}


def test_extract_requires_all_heads_approved_and_reparse_rules(world, approved):
    s = world.service
    world.file("c.xlsx", shift=1)
    doc_c = world.register("c.xlsx")
    app_c = doc_c["applied"][0]["application_id"]
    with pytest.raises(Problem) as exc:
        s.extract(app_c)
    assert exc.value.code == "REVIEW_REQUIRED" and [f["rule_key"] for f in exc.value.fields] == ["lot", "temperature", "signer"]
    with pytest.raises(Problem) as exc:
        s.approve_profile(approved["profile"]["profile_id"], app_c)
    assert exc.value.code == "REVIEW_REQUIRED"
    draft = s.import_profile("process_standard", profile_definition(profile_name="초안"))
    with pytest.raises(Problem) as exc:
        s.reparse(draft["profile_id"], "fill")
    assert exc.value.code == "PROFILE_NOT_APPROVED"
    # 새 프로파일 리비전 → 참조 rev가 뒤처져 자동 승인이 중단된다(compatible → proposed).
    revised = s.import_profile("process_standard", profile_definition(description="v2"), profile_id=approved["profile"]["profile_id"])
    assert revised["current_rev"] == 2
    world.file("d.xlsx")
    doc_d = world.register("d.xlsx")
    assert doc_d["status"] == "review" and doc_d["applied"][0]["compatibility"] == "compatible"
    # 대표 문서로 재승인하면 참조가 새 rev로 갱신되고 rematch가 d를 소급 승인·추출한다.
    again = s.approve_profile(approved["profile"]["profile_id"], approved["application"]["application_id"])
    assert again["reference_profile_rev"] == 2
    rematch = s.jobs.wait(again["reparse_job"]["job_id"], 120)
    assert rematch["state"] == "succeeded", rematch
    assert world.status(doc_d["document_id"]) == "normal" and s.application_summary(doc_d["applied"][0]["application_id"])["published"]
    head = s.mapping(s.application_mappings(doc_d["applied"][0]["application_id"])[0]["mapping_id"])
    assert head["status"] == "approved" and head["origin"] == "auto" and head["revision_no"] == 2
    # rollback: 과거(proposed) 리비전을 복제한 새 리비전 → 발행 무효 → 재추출 필요.
    first_rev = s.mapping_revisions(head["mapping_id"])[-1]
    rolled = s.rollback(head["mapping_id"], expected_seq=2, target_revision_id=first_rev["mapping_revision_id"])
    assert rolled["status"] == "proposed" and rolled["revision_no"] == 3
    assert world.status(doc_d["document_id"]) == "review" and not s.application_summary(doc_d["applied"][0]["application_id"])["published"]
    fill = s.jobs.wait(s.reparse(approved["profile"]["profile_id"], "fill")["job_id"], 60)
    assert fill["result"]["queued"] == 0 and {x["reason"] for x in fill["result"]["skipped"]} <= {"published", "review_required"}
