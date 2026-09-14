"""v2 워크스페이스 → v3 이관(계약 §9): 건수, 폐기 개념, NULL 개념 리비전, 값/영역 병합, 발행 유지, 재추출 동일성, dry-run."""

from __future__ import annotations

import json
import sqlite3
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from kg.v2.build import create_integration, prepare_build
from kg.v2.db import uid
from kg.v2.service import Service as V2Service
from kg.v3 import migrate as migrate_module
from kg.v3.db import Problem
from kg.v3.migrate import SCHEMA_KEY, run as migrate
from kg.v3.service import Service as V3Service

PRINCIPAL = "tester"


def build_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "세로"
    ws["B2"], ws["D2"] = "공정", "온도"
    ws.merge_cells("B2:C2")
    ws["A3"], ws["A4"], ws["A5"] = "lot-1", "lot-2", "lot-3"
    ws["B3"], ws["B4"], ws["B5"] = "123456789012345678901234567890.1250", "20.5", "30"
    ws.merge_cells("B3:C3")
    ws["F2"], ws["G2"] = "비고", "정상"
    second = wb.create_sheet("단위")
    second["A1"] = "°C"
    wb.save(path)
    return path


def v2_job(service, kind, payload, prepare=None):
    job = service.jobs.submit(kind, payload, PRINCIPAL, uid(), prepare)
    assert service.jobs.run_one()
    job = service.jobs.get(job["job_id"], PRINCIPAL)
    assert job["state"] == "succeeded", job
    return job


@pytest.fixture
def v2(tmp_path):
    """실제 XLSX 위에 v2 워크스페이스: KG 2리비전(구 개념 1개 폐기) → 템플릿(NULL 개념 규칙 포함) → 적용 → 검수 → 추출 → 통합+빌드."""
    root = tmp_path / "v2"
    (root / "data/raw").mkdir(parents=True)
    path = build_workbook(root / "data/raw/sample.xlsx")
    service = V2Service(root)
    world = {"root": root, "service": service, "path": path}
    try:
        old_kg = service.import_kg(
            {"concepts": [
                {"concept_id": "temperature", "name": "공정온도", "level": 1, "aliases": ["Temp"], "canonical_unit": "°C", "value_type": "decimal"},
                {"concept_id": "legacy_note", "name": "구비고", "level": 1},
            ]},
            PRINCIPAL,
        )
        kg = service.import_kg(
            {"concepts": [
                {"concept_id": "process", "name": "공정 정보", "level": 1},
                {"concept_id": "temperature", "name": "공정온도", "level": 2, "aliases": ["공정 온도", "Temp"], "canonical_unit": "°C", "value_type": "number"},
                {"concept_id": "remark", "name": "비고", "level": 2, "value_type": "text"},
            ], "relations": [["process", "temperature", "parent_of"], ["temperature", "remark", "related"]]},
            PRINCIPAL,
        )
        job = v2_job(service, "register", {"source_refs": ["sample.xlsx"], "provider": "local-xlsx"})
        doc = job["result"]["documents"][0]
        with service.db.connect() as conn:
            sheets = {r["name"]: r["sheet_id"] for r in conn.execute("SELECT sheet_id, name FROM sheet WHERE document_version_id=?", (doc["version_id"],))}
        definition = {
            "kg_revision_id": kg["kg_revision_id"],
            "sheet_roles": {"main": {"cardinality": "one"}, "units": {"cardinality": "one"}},
            "rules": [
                {
                    "rule_key": "temperature",
                    "concept_id": "temperature",
                    "record_spec": {"scope": ["process"], "key": {"column": "A"}},
                    "selector": {
                        "key": {"areas": [{"sheet_role": "main", "range": "B2:C2"}, {"sheet_role": "main", "range": "D2"}]},
                        "value": {"areas": [{"sheet_role": "main", "range": "B3:C5"}], "cardinality": "list", "axis": "down", "element_layout": "one_per_row"},
                        "unit": {"areas": [{"sheet_role": "units", "range": "A1"}]},
                    },
                    "value_spec": {"type": "decimal", "unit": "°C"},
                },
                {
                    "rule_key": "remark",
                    "selector": {"key": {"areas": [{"sheet_role": "main", "find": {"texts": ["비고"]}}]}, "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1}}]}},
                    "value_spec": {"type": "text"},
                },
            ],
        }
        template = service.create_template("공정온도", definition, PRINCIPAL)
        application = service.apply_template(
            doc["version_id"], template["template_version_id"], {"main": [sheets["세로"]], "units": [sheets["단위"]]}, None, True, PRINCIPAL
        )
        with service.db.connect() as conn:
            proposed = conn.execute(
                "SELECT * FROM mapping_revision WHERE application_id=? AND rule_key='remark'", (application["application_id"],)
            ).fetchone()
        assert proposed["status"] == "proposed" and proposed["concept_id"] is None
        spec = json.loads(proposed["effective_spec_json"])
        approved = service.revise(application["application_id"], proposed["mapping_revision_id"], 0, spec, "remark", "approved", "비고 개념 연결", PRINCIPAL)
        run = v2_job(service, "extract", {"application_id": application["application_id"]}, service.prepare_extraction)
        integration = create_integration(
            service,
            "온도 DB",
            {
                "kg_revision_id": kg["kg_revision_id"],
                "row_mode": "record_scope",
                "fields": [
                    {"field_key": "temp", "output_name": "공정온도", "concept_id": "temperature", "target_type": "decimal", "target_unit": "°C",
                     "sources": [{"application_id": application["application_id"], "rule_key": "temperature"}]}
                ],
            },
            PRINCIPAL,
        )
        build = v2_job(service, "build", {"integration_version_id": integration["integration_version_id"]}, prepare_build)
        with service.db.connect() as conn:
            items = [
                dict(r)
                for r in conn.execute(
                    "SELECT i.*, s.instance_key, s.mapping_revision_id, m.rule_key FROM extracted_item i JOIN extracted_series s USING(series_id) JOIN mapping_revision m USING(mapping_revision_id) WHERE s.run_id=? ORDER BY m.rule_key, i.item_index",
                    (run["result"]["run_id"],),
                )
            ]
            counts = {
                t: conn.execute(f"SELECT count(*) FROM {t}").fetchone()[0]
                for t in ("document", "document_version", "sheet", "source_region", "mapping_revision", "extracted_series", "extracted_item", "item_region", "series_region")
            }
            # series_region은 그 series의 항목마다 붙으므로 v3 영역 수 = item_region + Σ(항목별 series_region)
            counts["merged_regions"] = counts["item_region"] + conn.execute(
                "SELECT count(*) FROM series_region sr JOIN extracted_item i USING(series_id)"
            ).fetchone()[0]
            published = conn.execute("SELECT published_run_id FROM template_application WHERE application_id=?", (application["application_id"],)).fetchone()[0]
        world.update(
            old_kg=old_kg, kg=kg, doc=doc, sheets=sheets, template=template, application=application, proposed=dict(proposed), approved=approved,
            run_id=run["result"]["run_id"], build_id=build["result"]["build_id"], items=items, counts=counts, published=published,
        )
        assert published == run["result"]["run_id"]
        yield world
    finally:
        service.jobs.close()


def test_dry_run_writes_nothing(v2, tmp_path):
    ws = tmp_path / "v3"
    report = migrate(ws, v2["root"], dry_run=True)
    assert report["dry_run"] is True
    assert not ws.exists()
    counts = report["counts"]
    assert counts["document"]["migrated"] == 1
    assert counts["document_snapshot"]["migrated"] == v2["counts"]["document_version"]
    assert counts["sheet"]["migrated"] == v2["counts"]["sheet"]
    assert counts["parsing_field"]["migrated"] == 4  # process, temperature, remark + 폐기된 legacy_note
    assert counts["parsing_profile"]["migrated"] == 1
    assert counts["mapping_revision"]["migrated"] == v2["counts"]["mapping_revision"]
    assert counts["extracted_value"]["migrated"] == v2["counts"]["extracted_item"]
    assert counts["published_run"]["migrated"] == 1
    assert counts["build_manifest"]["migrated"] == 1
    assert [r["reason"] for r in report["review_required"]] == ["FIELD_REQUIRED"]
    assert report["raw"]["available"] == 1


def test_missing_v2_database(tmp_path):
    with pytest.raises(Problem) as exc:
        migrate(tmp_path / "v3", tmp_path / "nowhere")
    assert exc.value.code == "V2_DATABASE_MISSING"


@pytest.fixture
def migrated(v2, tmp_path):
    ws = tmp_path / "v3"
    report = migrate(ws, v2["root"])
    service = V3Service(ws)
    try:
        yield SimpleNamespace(ws=ws, report=report, service=service, v2=v2)
    finally:
        service.close()


def test_counts_schema_and_profile(migrated):
    report, v2 = migrated.report, migrated.v2
    counts = report["counts"]
    assert report["skipped"] == [], report["skipped"]
    assert counts["document"]["migrated"] == 1 and counts["document_snapshot"]["migrated"] == v2["counts"]["document_version"]
    assert counts["sheet"]["migrated"] == v2["counts"]["sheet"]
    assert counts["source_region"]["migrated"] == v2["counts"]["source_region"]
    assert counts["extracted_value"]["migrated"] == v2["counts"]["extracted_item"]
    assert counts["extracted_value_region"]["migrated"] == v2["counts"]["merged_regions"] == 24
    assert counts["snapshot_signature"]["migrated"] == 1
    assert report["raw"] == {"checked": 1, "available": 1, "changed": 0, "missing": 0, "copied": 1, "reader_errors": 0}
    assert (migrated.ws / "data/raw/sample.xlsx").is_file()
    # 스키마: 개념 합집합, 최신 리비전에 없는 개념은 deprecated, alias/edge는 최신 리비전에서
    fields = migrated.service.schema_fields(SCHEMA_KEY)
    assert fields["legacy_note"]["status"] == "deprecated"
    assert {k for k, f in fields.items() if f["status"] == "active"} == {"process", "temperature", "remark"}
    assert fields["temperature"]["value_type"] == "decimal" and fields["temperature"]["unit"] == "°C"
    assert fields["temperature"]["aliases"] == ["공정 온도", "Temp"]
    definition = migrated.service.schema_definition(SCHEMA_KEY)
    by_key = {f["field_key"]: f for f in definition["fields"]}
    assert by_key["temperature"]["parents"] == ["process"] and by_key["temperature"]["related"] == ["remark"]
    assert (migrated.ws / "schemas" / SCHEMA_KEY / "r0001.json").is_file()
    # 프로파일: draft, r0001 파일, concept NULL 규칙은 default_field_id NULL
    profile = report["profiles"][0]
    assert profile["template_id"] == v2["template"]["template_id"] and profile["rules"] == ["remark", "temperature"]
    assert (migrated.ws / "profiles" / profile["profile_id"] / "r0001.json").is_file()
    with migrated.service.db.connect() as conn:
        row = conn.execute("SELECT status, current_rev FROM parsing_profile WHERE profile_id=?", (profile["profile_id"],)).fetchone()
        rules = {r["rule_key"]: r["default_field_id"] for r in conn.execute("SELECT rule_key, default_field_id FROM parsing_rule WHERE profile_id=?", (profile["profile_id"],))}
    assert tuple(row) == ("draft", 1)
    assert rules["remark"] is None and rules["temperature"] == fields["temperature"]["field_id"]


def test_null_concept_revision_and_published_run(migrated):
    v2, service = migrated.v2, migrated.service
    aid = v2["application"]["application_id"]
    with service.db.connect() as conn:
        app = dict(conn.execute("SELECT * FROM parsing_application WHERE application_id=?", (aid,)).fetchone())
        revisions = [
            dict(r)
            for r in conn.execute(
                "SELECT v.*, r.rule_key FROM mapping_revision v JOIN mapping m USING(mapping_id) JOIN parsing_rule r ON r.rule_id=m.rule_id WHERE m.application_id=? ORDER BY r.rule_key, v.revision_no",
                (aid,),
            )
        ]
        run = dict(conn.execute("SELECT * FROM extraction_run WHERE run_id=?", (v2["run_id"],)).fetchone())
        regions = conn.execute(
            "SELECT count(*) FROM mapping_region mr JOIN mapping_revision v USING(mapping_revision_id) JOIN mapping m USING(mapping_id) WHERE m.application_id=?", (aid,)
        ).fetchone()[0]
    assert app["origin"] == "manual" and app["profile_rev"] == 1 and app["snapshot_id"] == v2["doc"]["version_id"]
    # 원본이 있어 match_specs로 서명을 계산했다(compatible). 파일이 없으면 'manual'/'migrated'.
    assert app["compatibility"] == "compatible" and len(app["match_signature"]) == 64
    assert app["published_run_id"] == v2["run_id"]
    assert run["status"] == "succeeded" and json.loads(run["input_manifest_json"])["migrated_from"] == "v2"
    remark = [r for r in revisions if r["rule_key"] == "remark"]
    assert [(r["revision_no"], r["status"], r["field_id"] is None, r["origin"]) for r in remark] == [(1, "proposed", True, "import"), (2, "approved", False, "import")]
    assert remark[0]["mapping_revision_id"] == v2["proposed"]["mapping_revision_id"]
    assert json.loads(remark[0]["evidence_json"])["v2"]["concept_id"] is None
    assert [r["rule_key"] for r in migrated.report["review_required"]] == ["remark"]
    assert migrated.report["re_extract_required"] == []
    assert regions > 0
    summary = service.application_summary(aid)
    assert summary["published"] and summary["heads_approved"] == summary["heads_total"] == 2
    assert service.document(v2["doc"]["document_id"])["status"] == "normal"


def test_values_and_regions_merge(migrated):
    v2, service = migrated.v2, migrated.service
    with service.db.connect() as conn:
        values = [dict(r) for r in conn.execute("SELECT * FROM extracted_value WHERE run_id=? ORDER BY value_id", (v2["run_id"],))]
        region_rows = [dict(r) for r in conn.execute("SELECT * FROM extracted_value_region WHERE value_id IN (SELECT value_id FROM extracted_value WHERE run_id=?) ORDER BY value_id, role, ordinal", (v2["run_id"],))]
    by_id = {v["value_id"]: v for v in values}
    assert set(by_id) == {i["item_id"] for i in v2["items"]}
    for item in v2["items"]:
        value = by_id[item["item_id"]]
        assert value["group_key"] == item["instance_key"]
        assert value["value_text"] == item["value_text"] and value["value_state"] == "present"
        assert value["record_key"] == item["record_key"]
        assert value["mapping_revision_id"] == item["mapping_revision_id"]
    # series_region(key/value/unit)은 같은 role 뒤에 이어 붙는다: 값 항목의 value 영역 1개 + series value 영역 1개
    temp = next(i for i in v2["items"] if i["rule_key"] == "temperature")
    roles = {}
    for r in region_rows:
        if r["value_id"] == temp["item_id"]:
            roles.setdefault(r["role"], []).append(r["ordinal"])
    assert roles["value"] == [0, 1] and roles["key"] == [0, 1] and roles["unit"] == [0, 1] and roles["record_key"] == [0]
    listing = service.snapshot_values(v2["doc"]["version_id"], rule_key="temperature")
    assert [v["value_text"] for v in listing["items"]] == [i["value_text"] for i in v2["items"] if i["rule_key"] == "temperature"]
    assert all(v["source"]["range"] for v in listing["items"])


def test_reextraction_matches_v2(migrated):
    """이관된 프로파일·매핑으로 v3 엔진이 같은 파일을 추출하면 v2 항목의 value_text 목록과 같다."""
    v2, service = migrated.v2, migrated.service
    aid = v2["application"]["application_id"]
    job = service.extract(aid, wait=120)
    assert job["state"] == "succeeded", job
    new_run = job["result"]["run_id"]
    assert new_run != v2["run_id"]
    with service.db.connect() as conn:
        assert conn.execute("SELECT published_run_id FROM parsing_application WHERE application_id=?", (aid,)).fetchone()[0] == new_run
        rows = [
            dict(r)
            for r in conn.execute(
                "SELECT v.value_text, v.value_state, v.record_key, r.rule_key FROM extracted_value v JOIN mapping_revision mv USING(mapping_revision_id) JOIN mapping m USING(mapping_id) JOIN parsing_rule r ON r.rule_id=m.rule_id WHERE v.run_id=? ORDER BY r.rule_key, v.item_index",
                (new_run,),
            )
        ]
    assert [r["value_text"] for r in rows] == [i["value_text"] for i in v2["items"]]
    assert [r["record_key"] for r in rows if r["rule_key"] == "temperature"] == ["lot-1", "lot-2", "lot-3"]
    assert all(r["value_state"] == "present" for r in rows)


def test_build_manifest_and_rerun(migrated):
    v2 = migrated.v2
    manifest_path = migrated.ws / "data/exports" / f"migrated-{v2['build_id']}" / "manifest.json"
    assert manifest_path.is_file()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["build_id"] == v2["build_id"] and manifest["status"] == "succeeded" and manifest["row_count"] == 3
    assert manifest["integration"]["name"] == "온도 DB"
    assert manifest["integration"]["fields"][0]["field_key_v3"] == "temperature"
    assert manifest["inputs"] == [{"run_id": v2["run_id"], "selection": manifest["inputs"][0]["selection"], "migrated": True}]
    assert manifest["v2_output"] == f"data/v2-builds/{v2['build_id']}.sqlite"
    # 같은 v3 워크스페이스에 다시 실행하면 이미 있는 문서·적용 건·실행은 건너뛴다.
    migrated.service.close()
    again = migrate(migrated.ws, v2["root"])
    reasons = {(s["table"], s["reason"]) for s in again["skipped"]}
    assert ("document", "EXISTS") in reasons and ("parsing_application", "EXISTS") in reasons
    assert again["counts"]["document"]["migrated"] == 0
    assert again["schema"]["unchanged"] is True


def test_cli_report(v2, tmp_path, capsys):
    args = SimpleNamespace(ws=tmp_path / "v3", from_ws=v2["root"], raw=None, dry_run=True, report=tmp_path / "report.json", principal="v2-migration")
    assert migrate_module.run_cli(args) == 0
    printed = json.loads(capsys.readouterr().out)
    assert printed["dry_run"] is True and printed["counts"]["document"]["migrated"] == 1
    written = json.loads((tmp_path / "report.json").read_text(encoding="utf-8"))
    assert written["format"] == "v2-migration-report/1"
    assert not (tmp_path / "v3").exists()
    bad = SimpleNamespace(ws=tmp_path / "v3", from_ws=tmp_path / "missing", raw=None, dry_run=False, report=None, principal="x")
    assert migrate_module.run_cli(bad) == 2
    assert "V2_DATABASE_MISSING" in capsys.readouterr().err


def test_field_key_sanitized():
    taken = set()
    assert migrate_module.field_key_of("temperature", taken) == "temperature"
    odd = migrate_module.field_key_of("공정 온도/℃", taken)
    assert migrate_module.KEY_RE.match(odd) and odd != "공정 온도/℃"
    taken.add(odd)
    assert migrate_module.field_key_of("공정 온도/℃", taken) != odd


def test_v2_db_must_be_v2_schema(tmp_path):
    root = tmp_path / "fake"
    (root / "data/kg").mkdir(parents=True)
    with sqlite3.connect(root / "data/kg/v2.db") as conn:
        conn.execute("CREATE TABLE document (document_id TEXT)")
    with pytest.raises(Problem) as exc:
        migrate(tmp_path / "v3", root, dry_run=True)
    assert exc.value.code == "V2_DATABASE_INVALID"


def test_without_raw_file_falls_back_to_manual(v2, tmp_path):
    """원본이 없으면 match_specs를 돌리지 않고 compatibility='manual', match_signature='migrated'로 옮긴다."""
    empty = tmp_path / "empty-raw"
    empty.mkdir()
    ws = tmp_path / "v3-noraw"
    report = migrate(ws, v2["root"], raw=empty)
    assert report["raw"]["missing"] == 1 and report["raw"]["available"] == 0
    assert report["counts"]["snapshot_signature"]["migrated"] == 0
    service = V3Service(ws)
    try:
        with service.db.connect() as conn:
            app = dict(conn.execute("SELECT * FROM parsing_application WHERE application_id=?", (v2["application"]["application_id"],)).fetchone())
            regions = conn.execute("SELECT count(*) FROM mapping_region").fetchone()[0]
        assert (app["compatibility"], app["match_signature"], app["match_signature_json"]) == ("manual", "migrated", None)
        assert app["published_run_id"] == v2["run_id"] and regions > 0  # 실제 위치는 series_region에서 복원된다
        assert service.document(v2["doc"]["document_id"])["status"] == "normal"
    finally:
        service.close()
