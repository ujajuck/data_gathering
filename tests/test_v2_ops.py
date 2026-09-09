"""watch 폴더 자동 등록(kg.v2.watch)과 템플릿 재크롤링(kg.v2.recrawl) 검증."""

import json
import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

from kg.v2.__main__ import main
from kg.v2.db import uid
from kg.v2.watch import Watcher
from tests.test_v2_runtime import extracted, setup  # noqa: F401


def make_xlsx(path, cell="x"):
    wb = Workbook()
    wb.active["A1"] = cell
    wb.save(path)
    return path


def modify(path, cell="changed"):
    wb = load_workbook(path)
    wb.active["A1"] = cell
    wb.save(path)


def rows(root, sql, params=()):
    conn = sqlite3.connect(root / "data/kg/v2.db")
    conn.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in conn.execute(sql, params)]
    finally:
        conn.close()


def lines(capsys):
    return [json.loads(l) for l in capsys.readouterr().out.splitlines()]


@pytest.fixture
def ws(tmp_path):
    (tmp_path / "data/raw").mkdir(parents=True)
    return tmp_path


def test_watch_once_registers_generated_xlsx_and_logs_ids(ws, capsys):
    make_xlsx(ws / "data/raw/sample.xlsx")
    assert main(["watch", "--ws", str(ws), "--once"]) == 0
    out = lines(capsys)
    assert len(out) == 1
    line = out[0]
    assert line["event"] == "created"
    assert line["source_ref"] == "sample.xlsx"
    assert line["unchanged"] is False
    assert line["revision_no"] == 1
    assert line["sheets"] == 1
    docs = rows(ws, "SELECT * FROM document")
    assert len(docs) == 1
    assert docs[0]["provider"] == "local-xlsx"
    assert docs[0]["source_ref"] == "sample.xlsx"
    assert docs[0]["document_id"] == line["document_id"]
    assert docs[0]["current_version_id"] == line["version_id"]


def test_watch_modified_file_creates_version_two_of_same_document(ws, capsys):
    path = make_xlsx(ws / "data/raw/sample.xlsx")
    assert main(["watch", "--ws", str(ws), "--once"]) == 0
    first = lines(capsys)[0]
    modify(path)
    assert main(["watch", "--ws", str(ws), "--once"]) == 0
    second = lines(capsys)[0]
    assert second["document_id"] == first["document_id"]
    assert second["version_id"] != first["version_id"]
    assert second["unchanged"] is False
    assert second["revision_no"] == 2
    versions = rows(ws, "SELECT * FROM document_version ORDER BY revision_no")
    assert [v["revision_no"] for v in versions] == [1, 2]
    assert {v["document_id"] for v in versions} == {first["document_id"]}
    docs = rows(ws, "SELECT * FROM document")
    assert len(docs) == 1
    assert docs[0]["current_version_id"] == second["version_id"]


def test_watch_unchanged_rerun_registers_nothing(ws, capsys):
    make_xlsx(ws / "data/raw/sample.xlsx")
    assert main(["watch", "--ws", str(ws), "--once"]) == 0
    first = lines(capsys)[0]
    assert main(["watch", "--ws", str(ws), "--once"]) == 0
    second = lines(capsys)[0]
    assert second["unchanged"] is True
    assert second["version_id"] == first["version_id"]
    assert second["revision_no"] == 1
    assert len(rows(ws, "SELECT * FROM document_version")) == 1


def test_watch_skips_encrypted_and_invalid_files_with_reason(ws, capsys):
    raw = ws / "data/raw"
    (raw / "ole.xlsx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
    (raw / "junk.xlsx").write_bytes(b"PK\x03\x04junk")
    (raw / "~$lock.xlsx").write_bytes(b"PK\x03\x04lock")
    make_xlsx(raw / "good.xlsx")
    assert main(["watch", "--ws", str(ws), "--once"]) == 0
    out = {l["source_ref"]: l for l in lines(capsys)}
    assert set(out) == {"ole.xlsx", "junk.xlsx", "good.xlsx"}
    assert out["ole.xlsx"]["skipped"] == "DRM_READER_REQUIRED"
    assert out["junk.xlsx"]["skipped"] == "READER_FAILED"
    assert out["good.xlsx"]["revision_no"] == 1
    assert [d["source_ref"] for d in rows(ws, "SELECT source_ref FROM document")] == [
        "good.xlsx"
    ]


def test_watch_logs_deleted_without_removing_records(ws):
    path = make_xlsx(ws / "data/raw/sample.xlsx")
    watcher = Watcher(ws)
    out = watcher.run_once()
    assert len(out) == 1 and out[0]["revision_no"] == 1
    before = (
        rows(ws, "SELECT count(*) n FROM document")[0]["n"],
        rows(ws, "SELECT count(*) n FROM document_version")[0]["n"],
    )
    path.unlink()
    out = watcher.scan()
    assert len(out) == 1
    assert out[0]["event"] == "deleted"
    assert out[0]["skipped"] == "FILE_DELETED"
    assert out[0]["source_ref"] == "sample.xlsx"
    after = (
        rows(ws, "SELECT count(*) n FROM document")[0]["n"],
        rows(ws, "SELECT count(*) n FROM document_version")[0]["n"],
    )
    assert before == after == (1, 1)


def test_watch_subdirectory_source_ref_is_relative_to_raw(ws, capsys):
    sub = ws / "data/raw/sub"
    sub.mkdir()
    make_xlsx(sub / "inner.xlsx")
    make_xlsx(ws / "data/raw/top.xlsx")
    assert main(["watch", "--ws", str(ws), "--raw", str(sub), "--once"]) == 0
    out = lines(capsys)
    assert [l["source_ref"] for l in out] == ["sub/inner.xlsx"]
    docs = rows(ws, "SELECT * FROM document")
    assert [d["source_ref"] for d in docs] == ["sub/inner.xlsx"]
    watcher = Watcher(ws)
    assert watcher.service.version(out[0]["version_id"])["source_ref"] == "sub/inner.xlsx"


def test_watch_rejects_raw_outside_workspace(ws, capsys):
    elsewhere = ws.parent / "elsewhere"
    elsewhere.mkdir(exist_ok=True)
    make_xlsx(elsewhere / "outside.xlsx")
    assert main(["watch", "--ws", str(ws), "--raw", str(elsewhere), "--once"]) == 2
    captured = capsys.readouterr()
    assert captured.out == ""
    assert json.loads(captured.err)["error"] == "INVALID_RAW_DIR"
    assert rows(ws, "SELECT count(*) n FROM document")[0]["n"] == 0


def test_serve_remains_default_cli(ws, monkeypatch):
    calls = []
    monkeypatch.setattr(
        "uvicorn.run", lambda app, **kw: calls.append((type(app).__name__, kw))
    )
    assert main(["--ws", str(ws), "--port", "8123"]) == 0
    assert main(["serve", "--ws", str(ws), "--port", "8123"]) == 0
    assert len(calls) == 2
    for name, kw in calls:
        assert name == "FastAPI"
        assert kw == {"host": "127.0.0.1", "port": 8123}
    with pytest.raises(SystemExit) as exc:
        main(["--help"])
    assert exc.value.code == 0


def add_application(s, scope_key, approved):
    return s["api"](
        "POST",
        "/applications",
        {
            "version_id": s["doc"]["version_id"],
            "template_version_id": s["template"]["template_version_id"],
            "bindings": {
                "main": [s["sheets"][0]["sheet_id"]],
                "units": [s["sheets"][1]["sheet_id"]],
            },
            "scope_key": scope_key,
            "approved": approved,
        },
    )["application_id"]


def mapping_state(s):
    with s["service"].db.connect() as conn:
        revisions = {
            r[0]
            for r in conn.execute(
                "SELECT mapping_revision_id FROM mapping_revision"
            )
        }
        heads = [
            tuple(r)
            for r in conn.execute(
                "SELECT application_id,rule_key,mapping_revision_id,edit_seq FROM mapping_head ORDER BY 1,2"
            )
        ]
    return revisions, heads


def application_row(s, aid):
    with s["service"].db.connect() as conn:
        return dict(
            conn.execute(
                "SELECT * FROM template_application WHERE application_id=?", (aid,)
            ).fetchone()
        )


def count(s, table):
    with s["service"].db.connect() as conn:
        return conn.execute(f"SELECT count(*) FROM {table}").fetchone()[0]


def population(s):
    a = s["application"]["application_id"]
    b = add_application(s, "b", approved=False)
    c = add_application(s, "c", approved=True)
    job = s["work"]("/applications/" + c + "/extract", {})
    assert job["state"] == "succeeded", job
    assert application_row(s, c)["published_run_id"] == job["result"]["run_id"]
    assert application_row(s, a)["published_run_id"] is None
    return a, b, c


def test_recrawl_fill_queues_unpublished_approved_and_skips_others(setup):
    s = setup
    a, b, c = population(s)
    tid = s["template"]["template_version_id"]
    key = uid()
    result = s["api"](
        "POST", f"/template-versions/{tid}/recrawl", {"mode": "fill", "request_key": key}
    )
    assert result["template_version_id"] == tid
    assert result["mode"] == "fill"
    assert result["request_key"] == key
    assert [q["application_id"] for q in result["queued"]] == [a]
    assert sorted(result["skipped"], key=lambda k: k["application_id"]) == sorted(
        [
            {"application_id": b, "reason": "review_required"},
            {"application_id": c, "reason": "published"},
        ],
        key=lambda k: k["application_id"],
    )
    job_id = result["queued"][0]["job_id"]
    assert s["api"]("GET", "/jobs/" + job_id)["kind"] == "extract"
    s["service"].jobs.run_one()
    job = s["api"]("GET", "/jobs/" + job_id)
    assert job["state"] == "succeeded", job
    assert application_row(s, a)["published_run_id"] == job_id


def test_recrawl_reset_auto_reextracts_published_and_never_touches_mappings(setup):
    s = setup
    a, b, c = population(s)
    previous_run = application_row(s, c)["published_run_id"]
    before = mapping_state(s)
    tid = s["template"]["template_version_id"]
    result = s["api"](
        "POST",
        f"/template-versions/{tid}/recrawl",
        {"mode": "reset_auto", "request_key": uid()},
    )
    assert {q["application_id"] for q in result["queued"]} == {a, c}
    assert result["skipped"] == [{"application_id": b, "reason": "review_required"}]
    s["service"].jobs.run_one()
    s["service"].jobs.run_one()
    for q in result["queued"]:
        job = s["api"]("GET", "/jobs/" + q["job_id"])
        assert job["state"] == "succeeded", job
    assert mapping_state(s) == before
    new_run = application_row(s, c)["published_run_id"]
    assert new_run != previous_run
    assert new_run in {q["job_id"] for q in result["queued"]}
    with s["service"].db.connect() as conn:
        assert (
            conn.execute(
                "SELECT status FROM extraction_run WHERE run_id=?", (previous_run,)
            ).fetchone()[0]
            == "succeeded"
        )


def test_recrawl_request_key_is_idempotent_per_application(setup):
    s = setup
    a = s["application"]["application_id"]
    tid = s["template"]["template_version_id"]
    body = {"mode": "fill", "request_key": uid()}
    first = s["api"]("POST", f"/template-versions/{tid}/recrawl", body)
    jobs, runs = count(s, "runtime_job"), count(s, "extraction_run")
    second = s["api"]("POST", f"/template-versions/{tid}/recrawl", body)
    assert first["queued"] == second["queued"]
    assert [q["application_id"] for q in first["queued"]] == [a]
    assert count(s, "runtime_job") == jobs
    assert count(s, "extraction_run") == runs
    with s["service"].db.connect() as conn:
        stored = conn.execute(
            "SELECT request_key,kind FROM runtime_job WHERE job_id=?",
            (first["queued"][0]["job_id"],),
        ).fetchone()
    assert tuple(stored) == (f"{body['request_key']}:fill:{a}", "extract")
    # 같은 request_key라도 mode가 다르면 별개의 배치다 — 같은 적용 건은 진행 중이면 건너뛰고 새 작업을 재사용하지 않는다.
    third = s["api"](
        "POST", f"/template-versions/{tid}/recrawl", {**body, "mode": "reset_auto"}
    )
    assert third["mode"] == "reset_auto" and third["truncated"] is False
    assert not any(q["job_id"] == first["queued"][0]["job_id"] for q in third["queued"])
    assert count(s, "runtime_job") == jobs + len(third["queued"])


def test_recrawl_status_summarizes_jobs(setup):
    s = setup
    a = s["application"]["application_id"]
    tid = s["template"]["template_version_id"]
    key = uid()
    s["api"](
        "POST", f"/template-versions/{tid}/recrawl", {"mode": "fill", "request_key": key}
    )
    status = s["api"]("GET", f"/template-versions/{tid}/recrawl-status?request_key={key}")
    assert status["summary"] == {"queued": 1}
    assert status["has_more"] is False
    assert len(status["items"]) == 1
    assert status["items"][0]["application_id"] == a
    assert status["items"][0]["state"] == "queued"
    s["service"].jobs.run_one()
    status = s["api"]("GET", f"/template-versions/{tid}/recrawl-status?request_key={key}")
    assert status["summary"] == {"succeeded": 1}
    assert status["items"][0]["state"] == "succeeded"
    other = s["api"](
        "GET", f"/template-versions/{tid}/recrawl-status?request_key={uid()}"
    )
    assert other["items"] == [] and other["summary"] == {}
    s["api"](
        "GET", f"/template-versions/{uid()}/recrawl-status?request_key={key}", code=404
    )


def test_recrawl_skips_applications_on_old_document_versions(setup):
    s = setup
    a = s["application"]["application_id"]
    wb = load_workbook(s["path"])
    wb["세로"]["B5"] = "31"
    wb.save(s["path"])
    job = s["work"]("/documents/register", {"source_refs": ["sample.xlsx"]})
    assert job["state"] == "succeeded", job
    registered = job["result"]["documents"][0]
    assert registered["document_id"] == s["doc"]["document_id"]
    assert registered["version_id"] != s["doc"]["version_id"]
    tid = s["template"]["template_version_id"]
    result = s["api"](
        "POST",
        f"/template-versions/{tid}/recrawl",
        {"mode": "reset_auto", "request_key": uid()},
    )
    assert result["queued"] == []
    assert result["skipped"] == [{"application_id": a, "reason": "stale_version"}]


def test_recrawl_validation_and_not_found(setup):
    s = setup
    tid = s["template"]["template_version_id"]
    error = s["api"](
        "POST",
        f"/template-versions/{uid()}/recrawl",
        {"mode": "fill", "request_key": uid()},
        404,
    )
    assert error["error"]["code"] == "NOT_FOUND"
    error = s["api"](
        "POST", f"/template-versions/{tid}/recrawl", {"mode": "x", "request_key": uid()}, 422
    )
    assert error["error"]["code"] == "VALIDATION_ERROR"
    error = s["api"]("POST", f"/template-versions/{tid}/recrawl", {}, 422)
    assert error["error"]["code"] == "REQUEST_KEY_REQUIRED"
    error = s["api"](
        "POST",
        f"/template-versions/{tid}/recrawl",
        {"mode": "fill", "request_key": "k" * 65},
        422,
    )
    assert error["error"]["code"] == "REQUEST_KEY_REQUIRED"
    assert count(s, "runtime_job") == 1  # setup의 register 작업만 존재


def test_watch_skips_symlink_outside_raw_and_continues(ws, capsys, tmp_path):
    outside = tmp_path / "outside" / "secret.xlsx"
    outside.parent.mkdir(parents=True)
    make_xlsx(outside)
    make_xlsx(ws / "data/raw/good.xlsx")
    (ws / "data/raw/link.xlsx").symlink_to(outside)
    assert main(["watch", "--ws", str(ws), "--once"]) == 0
    lines = [json.loads(l) for l in capsys.readouterr().out.splitlines()]
    skipped = [l for l in lines if l.get("skipped") == "OUTSIDE_RAW_DIR"]
    assert skipped and all("secret" not in json.dumps(l) for l in skipped)
    assert any(l.get("source_ref") == "good.xlsx" and l.get("version_id") for l in lines)


def test_watch_dangling_symlink_does_not_block_other_files(ws, capsys):
    make_xlsx(ws / "data/raw/good.xlsx")
    (ws / "data/raw/dangling.xlsx").symlink_to(ws / "data/raw/missing.xlsx")
    assert main(["watch", "--ws", str(ws), "--once"]) == 0
    lines = [json.loads(l) for l in capsys.readouterr().out.splitlines()]
    assert any(l.get("skipped") == "STAT_FAILED" and l.get("path") == "dangling.xlsx" for l in lines)
    assert any(l.get("source_ref") == "good.xlsx" and l.get("version_id") for l in lines)
    assert not any(l.get("skipped") == "SCAN_FAILED" for l in lines)
