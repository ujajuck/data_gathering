"""§4.10 데이터 빌드: candidates 사유·필드별 문서 수, 미리보기(record/document, scalar 복제, 중복 제거, 충돌), 헤더 검증,
단위 변환(units.yaml), CSV/XLSX/SQLite 내용, manifest, build_key 재사용, 동기 runtime_job, 비동기 임계."""

from __future__ import annotations

import csv
import json
import sqlite3

import pytest
from openpyxl import load_workbook

from kg.v3 import build as build_module
from kg.v3.build import build, candidates, manifest, output_path, preview
from kg.v3.db import Problem
from tests.v3_ops_fixture import World, variant_profile


@pytest.fixture
def world(tmp_path, monkeypatch):
    w = World(tmp_path, monkeypatch)
    yield w
    w.service.close()


@pytest.fixture
def published(world):
    """a.xlsx(수동 승인·발행, 대표) · b.xlsx(자동 승인·발행) · c.xlsx(compatible → 검수) · other.xlsx(unmatched)."""
    base = world.approve_profile_with("a.xlsx")
    world.file("b.xlsx", temps=(1, 2, 3), lots=("X1", "X2", "X3"), signer="김철수")
    doc_b = world.register("b.xlsx")
    assert doc_b["status"] == "normal"
    world.file("c.xlsx", shift=1)
    doc_c = world.register("c.xlsx")
    assert doc_c["status"] == "review"
    world.file("other.xlsx", sheet="다른 양식")
    other = world.register("other.xlsx")
    assert other["status"] == "unmatched"
    return {"a": base["document"], "b": doc_b, "c": doc_c, "other": other, "profile": base["profile"]}


def columns(*keys, **units):
    return [{"field_key": k, "header": k, "target_unit": units.get(k)} for k in keys]


def request(published, *keys, row_mode="record", format="xlsx", ids=None, **units):
    return {
        "document_ids": ids or [published["a"]["document_id"], published["b"]["document_id"], published["c"]["document_id"], published["other"]["document_id"]],
        "schema_key": "process_standard",
        "columns": columns(*keys, **units),
        "row_mode": row_mode,
        "format": format,
    }


# ---------------------------------------------------------------------------- candidates


def test_candidates_reasons_and_field_document_counts(world, published):
    s = world.service
    ids = [published["a"]["document_id"], published["b"]["document_id"], published["c"]["document_id"], published["other"]["document_id"], "missing-id"]
    result = candidates(s, ids, "process_standard")
    by_id = {d["document_id"]: d for d in result["documents"]}
    assert [d["document_id"] for d in result["documents"]] == ids  # 입력 순서 유지
    assert by_id[published["a"]["document_id"]]["usable"] and by_id[published["a"]["document_id"]]["profile"]["profile_name"] == "공정데이터_A양식"
    assert by_id[published["b"]["document_id"]]["usable"] and "reason" not in by_id[published["b"]["document_id"]]
    assert by_id[published["c"]["document_id"]]["reason"] == "review_required"
    assert by_id[published["other"]["document_id"]]["reason"] == "unmatched"
    assert by_id["missing-id"]["reason"] == "not_found" and by_id["missing-id"]["document_name"] is None
    assert result["summary"] == {"total": 5, "usable": 2, "excluded": 3}
    fields = {f["field_key"]: f for f in result["fields"]}
    assert "process" not in fields  # group 필드는 출력 후보가 아니다
    assert fields["temperature"]["document_count"] == 2 and fields["lot"]["document_count"] == 2 and fields["note"]["document_count"] == 2
    assert fields["pressure"]["document_count"] == 0 and fields["temperature"]["unit"] == "°C" and fields["temperature"]["type"] == "decimal"
    assert result["schemas"] == [{"schema_key": "process_standard", "schema_name": "공정 데이터 표준", "document_count": 2}]
    without = candidates(s, ids[:2])
    assert without["summary"]["usable"] == 2 and {f["field_key"] for f in without["fields"]} == {"lot", "temperature", "pressure", "note"}
    with pytest.raises(Problem) as exc:
        candidates(s, ids, "nope")
    assert exc.value.code == "UNKNOWN_SCHEMA" and exc.value.status == 404
    # 잠긴 문서는 locked 사유.
    (world.root / "data/raw/locked.xlsx").write_bytes(b"\xd0\xcf\x11\xe0drm")
    job = s.register_documents(["locked.xlsx"], wait=60)
    assert job["state"] == "failed"
    locked = s.document_query(status="locked")["items"][0]["document_id"]
    assert candidates(s, [locked], "process_standard")["documents"][0]["reason"] == "locked"


# ---------------------------------------------------------------------------- 미리보기


def test_preview_record_rows_broadcast_scalars_and_document_mode(world, published):
    s = world.service
    result = preview(s, request(published, "lot", "temperature", "note"))
    assert [c["header"] for c in result["columns"]] == ["lot", "temperature", "note"] and result["headers"] == ["lot", "temperature", "note", "_document", "_snapshot", "_record_key"]
    assert result["row_count"] == 6 and len(result["rows"]) == 6 and result["document_count"] == 2
    assert [e["reason"] for e in result["excluded"]] == ["review_required", "unmatched"]
    rows_a = [r for r in result["rows"] if r["document"]["document_name"] == "a.xlsx"]
    assert [r["record_key"] for r in rows_a] == ["r4", "r5", "r6"]
    assert [r["cells"][0]["text"] for r in rows_a] == ["L1", "L2", "L3"] and [r["cells"][1]["text"] for r in rows_a] == ["10.5", "20", "30"]
    # scalar(서명자)는 문서의 모든 행에 같은 값(같은 value_id)으로 복제된다.
    assert {r["cells"][2]["text"] for r in rows_a} == {"홍길동"} and len({r["cells"][2]["value_id"] for r in rows_a}) == 1
    cell = rows_a[0]["cells"][1]
    assert cell["sheet_name"] == "공정 기록" and cell["range"] == "C4" and cell["rule_key"] == "temperature" and cell["sheet_id"] and cell["application_id"]
    assert cell["unit"] == "°C" and cell["state"] == "present"
    assert [r["row_no"] for r in result["rows"]] == [1, 2, 3, 4, 5, 6] and rows_a[0]["snapshot"]["captured_at"]
    assert result["conflicts"] == []
    rows_b = [r for r in result["rows"] if r["document"]["document_name"] == "b.xlsx"]
    assert [r["cells"][2]["text"] for r in rows_b] == ["김철수"] * 3 and [r["cells"][1]["text"] for r in rows_b] == ["1", "2", "3"]

    document = preview(s, request(published, "temperature", "note", row_mode="document"))
    assert document["row_count"] == 2 and document["headers"] == ["temperature", "temperature_count", "note", "note_count", "_document", "_snapshot", "_record_key"]
    first = next(r for r in document["rows"] if r["document"]["document_name"] == "a.xlsx")
    assert first["record_key"] is None and first["cells"][0]["text"] == "10.5" and first["cells"][0]["count"] == 3 and first["cells"][1]["count"] == 1

    # 문서 순서는 입력 순서, 중복 ID는 한 번, 대상 없는 문서만이면 행 0.
    reversed_ids = [published["b"]["document_id"], published["a"]["document_id"], published["a"]["document_id"]]
    assert [r["document"]["document_name"] for r in preview(s, request(published, "lot", ids=reversed_ids))["rows"]] == ["b.xlsx"] * 3 + ["a.xlsx"] * 3
    empty = preview(s, request(published, "lot", ids=[published["c"]["document_id"]]))
    assert empty["row_count"] == 0 and empty["rows"] == [] and empty["excluded"][0]["reason"] == "review_required"
    # 값이 없는 필드만 고르면 문서당 빈 행 1개.
    blank = preview(s, request(published, "pressure"))
    assert blank["row_count"] == 2 and [r["cells"] for r in blank["rows"]] == [[None], [None]]


def test_preview_limits_to_50_rows(world, published, monkeypatch):
    monkeypatch.setattr(build_module, "PREVIEW_ROWS", 4)
    result = preview(world.service, request(published, "lot"))
    assert result["row_count"] == 6 and len(result["rows"]) == 4 and [r["row_no"] for r in result["rows"]] == [1, 2, 3, 4]


def test_dedupe_same_source_and_conflict_on_different_source(world, published):
    """같은 문서에 두 프로파일이 발행되면 같은 셀의 값은 1개(identity 같음), 다른 셀(B55/C55)은 첫 값 + conflicts."""
    s = world.service
    variant = s.import_profile("process_standard", variant_profile())
    app = s.apply_profile(published["a"]["snapshot"]["snapshot_id"], variant["profile_id"])
    done = s.approve_all(app["application_id"])
    assert done["extraction"]["state"] == "succeeded"
    result = preview(s, request(published, "lot", "temperature", "note", ids=[published["a"]["document_id"]]))
    assert result["row_count"] == 3 and len(result["sources"]) == 2
    kinds = [(c["kind"], c["field_key"], c["record_key"]) for c in result["conflicts"]]
    assert kinds == [("multiple_values", "note", "r4"), ("multiple_values", "note", "r5"), ("multiple_values", "note", "r6")]
    conflict = result["conflicts"][0]
    assert {v["range"] for v in conflict["values"]} == {"B55", "C55"} and conflict["kept"] == result["rows"][0]["cells"][2]["value_id"]
    assert conflict["document_name"] == "a.xlsx" and conflict["header"] == "note" and conflict["row"] == 1
    assert [r["cells"][1]["text"] for r in result["rows"]] == ["10.5", "20", "30"]  # 온도/배치는 같은 셀 → 중복 제거, 충돌 없음
    # 문서 모드에서는 목록이 정상이므로 충돌이 아니다(첫 값 + count).
    document = preview(s, request(published, "note", row_mode="document", ids=[published["a"]["document_id"]]))
    assert document["conflicts"] == [] and document["rows"][0]["cells"][0]["count"] == 2


# ---------------------------------------------------------------------------- 헤더 검증


def test_invalid_header_blank_duplicate_reserved_and_unknown_field(world, published):
    s = world.service
    bad = request(published, "lot", "temperature", "note")
    bad["columns"][0]["header"] = "  "
    bad["columns"][1]["header"] = "Temp"
    bad["columns"][2]["header"] = "temp"
    with pytest.raises(Problem) as exc:
        preview(s, bad)
    assert exc.value.code == "INVALID_HEADER" and exc.value.status == 422
    assert [(f["index"], f["reason"]) for f in exc.value.fields] == [(0, "empty"), (2, "duplicate")]
    reserved = request(published, "lot")
    reserved["columns"][0]["header"] = "_document"
    with pytest.raises(Problem) as exc:
        build(s, reserved)
    assert exc.value.code == "INVALID_HEADER" and exc.value.fields[0]["reason"] == "reserved"
    # 문서 모드는 <header>_count 열이 붙으므로 그 이름과 겹치는 헤더도 중복.
    clash = request(published, "lot", "temperature", row_mode="document")
    clash["columns"][1]["header"] = "lot_count"
    with pytest.raises(Problem) as exc:
        preview(s, clash)
    assert exc.value.code == "INVALID_HEADER" and exc.value.fields[0]["reason"] == "duplicate"
    with pytest.raises(Problem) as exc:
        preview(s, request(published, "lot", "nope"))
    assert exc.value.code == "UNKNOWN_FIELD" and exc.value.fields == ["nope"]
    with pytest.raises(Problem) as exc:
        preview(s, request(published, "process"))
    assert exc.value.code == "GROUP_FIELD_TARGET"
    with pytest.raises(Problem) as exc:
        preview(s, {**request(published, "lot"), "schema_key": "nope"})
    assert exc.value.code == "UNKNOWN_SCHEMA"
    with pytest.raises(Problem) as exc:
        preview(s, {**request(published, "lot"), "row_mode": "matrix"})
    assert exc.value.code == "INVALID_REQUEST"


# ---------------------------------------------------------------------------- 단위 변환


def test_target_unit_conversion_uses_units_yaml(world, published):
    s = world.service
    no_registry = preview(s, request(published, "temperature", ids=[published["a"]["document_id"]], temperature="K"))
    assert [r["cells"][0]["text"] for r in no_registry["rows"]] == ["10.5", "20", "30"]
    assert [c["kind"] for c in no_registry["conflicts"]] == ["unit"] * 3 and "units.yaml" in no_registry["conflicts"][0]["message"]
    world.units()
    kelvin = preview(s, request(published, "temperature", ids=[published["a"]["document_id"]], temperature="K"))
    assert [r["cells"][0]["text"] for r in kelvin["rows"]] == ["283.65", "293.15", "303.15"] and kelvin["conflicts"] == []
    assert kelvin["rows"][0]["cells"][0]["unit"] == "K" and kelvin["columns"][0]["unit"] == "K"
    fahrenheit = preview(s, request(published, "temperature", ids=[published["a"]["document_id"]], temperature="°F"))
    assert [r["cells"][0]["text"] for r in fahrenheit["rows"]] == ["50.9", "68", "86"]
    same = preview(s, request(published, "temperature", ids=[published["a"]["document_id"]], temperature="℃"))  # alias → 같은 단위
    assert [r["cells"][0]["text"] for r in same["rows"]] == ["10.5", "20", "30"] and same["conflicts"] == []
    wrong = preview(s, request(published, "temperature", ids=[published["a"]["document_id"]], temperature="bar"))
    assert [r["cells"][0]["text"] for r in wrong["rows"]] == ["10.5", "20", "30"] and [c["kind"] for c in wrong["conflicts"]] == ["unit"] * 3
    assert "변환식" in wrong["conflicts"][0]["message"]
    text = preview(s, request(published, "note", ids=[published["a"]["document_id"]], note="K"))
    assert "단위가 없어" in text["conflicts"][0]["message"] and text["rows"][0]["cells"][0]["text"] == "홍길동"


# ---------------------------------------------------------------------------- 산출물


def test_exports_manifest_and_build_key_reuse(world, published):
    s = world.service
    world.units()
    csv_result = build(s, request(published, "lot", "temperature", "note", format="csv", temperature="K"))
    key = csv_result["build_key"]
    assert len(key) == 16 and csv_result["download_url"] == f"/api/v3/builds/{key}/download?format=csv" and not csv_result["reused"]
    raw = open(csv_result["path"], "rb").read()
    assert raw.startswith(b"\xef\xbb\xbf") and b"\r\n" in raw and b"\n" not in raw.replace(b"\r\n", b"")
    with open(csv_result["path"], encoding="utf-8-sig", newline="") as f:
        table = list(csv.reader(f))
    assert table[0] == ["lot", "temperature", "note", "_document", "_snapshot", "_record_key"]
    assert table[1][:3] == ["L1", "283.65", "홍길동"] and table[1][3] == "a.xlsx" and len(table[1][4]) == 10 and table[1][5] == "r4"
    assert len(table) == 7 and table[4][3] == "b.xlsx"
    assert not any("-" in cell and len(cell) == 36 for row in table for cell in row)  # UUID 비노출

    m = csv_result["manifest"]
    assert m["build_key"] == key and m["schema"] == {"key": "process_standard", "rev": 1, "name": "공정 데이터 표준"} and m["row_mode"] == "record" and m["row_count"] == 6
    assert [c for c in m["columns"]] == [{"field_key": "lot", "header": "lot", "type": "text", "unit": None}, {"field_key": "temperature", "header": "temperature", "type": "decimal", "unit": "K"}, {"field_key": "note", "header": "note", "type": "text", "unit": None}]
    assert [(x["document_name"], x["profile"]["name"], x["profile"]["rev"]) for x in m["sources"]] == [("a.xlsx", "공정데이터_A양식", 1), ("b.xlsx", "공정데이터_A양식", 1)]
    assert all(x["application_id"] and x["run_id"] and x["snapshot_id"] for x in m["sources"])
    assert [e["reason"] for e in m["excluded"]] == ["review_required", "unmatched"] and m["conflicts"] == [] and m["formats"] == {"csv": "data.csv"}
    assert json.loads((world.root / "data/exports" / key / "manifest.json").read_text(encoding="utf-8"))["build_key"] == key
    assert manifest(s, key)["row_count"] == 6
    job = s.jobs.get(csv_result["job_id"])
    assert job["kind"] == "build" and job["state"] == "succeeded" and job["target_kind"] == "build" and job["target_id"] == key
    assert job["result"] == {"build_key": key, "download_url": csv_result["download_url"], "row_count": 6} and "2문서" in job["label"]

    xlsx_result = build(s, request(published, "lot", "temperature", "note", format="xlsx", temperature="K"))
    assert xlsx_result["build_key"] == key and xlsx_result["manifest"]["formats"] == {"csv": "data.csv", "xlsx": "data.xlsx"}
    wb = load_workbook(xlsx_result["path"])
    ws = wb["data"]
    assert [c.value for c in ws[1]] == ["lot", "temperature", "note", "_document", "_snapshot", "_record_key"] and all(c.font.bold for c in ws[1])
    assert ws["A2"].value == "L1" and ws["B2"].value == 283.65 and isinstance(ws["B2"].value, float) and ws["C2"].value == "홍길동"
    assert ws.column_dimensions["A"].width >= 8 and ws.column_dimensions["D"].width > ws.column_dimensions["A"].width
    assert ws.max_row == 7

    sqlite_result = build(s, request(published, "lot", "temperature", "note", format="sqlite", temperature="K"))
    db = sqlite3.connect(sqlite_result["path"])
    cols = [r[1] for r in db.execute("PRAGMA table_info(data)")]
    assert cols == ["_row_no", "lot", "temperature", "note", "_source_lot", "_source_temperature", "_source_note", "_document", "_snapshot", "_record_key"]
    first = db.execute("SELECT * FROM data ORDER BY _row_no LIMIT 1").fetchone()
    assert first[1:4] == ("L1", 283.65, "홍길동") and first[4:7] == ("공정 기록!B4", "공정 기록!C4", "공정 기록!B55") and first[7] == "a.xlsx" and first[9] == "r4"
    assert db.execute("SELECT count(*) FROM data").fetchone()[0] == 6 and json.loads(db.execute("SELECT json FROM _manifest").fetchone()[0])["build_key"] == key
    db.close()

    # 같은 입력·같은 형식 → 재사용(파일 그대로, 작업 행 없음).
    before = (world.root / "data/exports" / key / "data.csv").stat().st_mtime_ns
    again = build(s, request(published, "lot", "temperature", "note", format="csv", temperature="K"))
    assert again["reused"] and again["path"] == csv_result["path"] and (world.root / "data/exports" / key / "data.csv").stat().st_mtime_ns == before
    with s.db.connect() as conn:
        assert conn.execute("SELECT count(*) FROM runtime_job WHERE kind='build'").fetchone()[0] == 3
    # 다른 입력(열 순서)은 다른 키.
    other = build(s, request(published, "temperature", "lot", format="csv"))
    assert other["build_key"] != key
    path, filename, media_type = output_path(s, key, "xlsx")
    assert path == world.root / "data/exports" / key / "data.xlsx" and filename == f"process_standard-{key}.xlsx" and media_type.startswith("application/vnd.openxml")
    assert output_path(s, key)[1].endswith(".sqlite")  # 마지막으로 만든 형식이 기본
    for bad in ("../x", "zz", key.replace(key[0], "g")):
        with pytest.raises(Problem) as exc:
            output_path(s, bad)
        assert exc.value.code == "NOT_FOUND"
    with pytest.raises(Problem) as exc:
        output_path(s, other["build_key"], "xlsx")
    assert exc.value.code == "NOT_FOUND"


def test_document_mode_sqlite_has_count_columns(world, published):
    s = world.service
    result = build(s, request(published, "temperature", "note", row_mode="document", format="sqlite"))
    db = sqlite3.connect(result["path"])
    cols = [r[1] for r in db.execute("PRAGMA table_info(data)")]
    assert cols == ["_row_no", "temperature", "temperature_count", "note", "note_count", "_source_temperature", "_source_note", "_document", "_snapshot", "_record_key"]
    rows_ = db.execute("SELECT temperature, temperature_count, note, note_count, _document, _record_key FROM data ORDER BY _row_no").fetchall()
    assert rows_ == [(10.5, 3, "홍길동", 1, "a.xlsx", None), (1, 3, "김철수", 1, "b.xlsx", None)]
    db.close()


# ---------------------------------------------------------------------------- 비동기 임계


def test_large_builds_run_as_jobs(world, published, monkeypatch):
    s = world.service
    monkeypatch.setattr(build_module, "SYNC_DOCUMENT_LIMIT", 1)
    queued = build(s, request(published, "lot", format="csv"), wait=0)
    assert queued["manifest"] is None and queued["job"]["state"] == "queued" and queued["job"]["kind"] == "build" and queued["job"]["target_id"] == queued["build_key"]
    assert (world.root / "data/exports" / queued["build_key"] / "data.csv").exists() is False
    done = s.jobs.wait(queued["job"]["job_id"], 60)
    assert done["state"] == "succeeded" and done["result"] == {"build_key": queued["build_key"], "download_url": queued["download_url"], "row_count": 6, "format": "csv", "reused": False}
    assert done["completed"] == 6 and done["total"] == 6
    assert manifest(s, queued["build_key"])["row_count"] == 6
    # wait 안에 끝나면 산출물과 manifest를 그대로 돌려준다(같은 요청은 같은 작업).
    waited = build(s, request(published, "lot", format="xlsx"), wait=60)
    assert waited["manifest"]["formats"] == {"csv": "data.csv", "xlsx": "data.xlsx"} and waited["job"]["state"] == "succeeded"
    monkeypatch.setattr(build_module, "SYNC_DOCUMENT_LIMIT", 200)
    monkeypatch.setattr(build_module, "SYNC_ROW_LIMIT", 5)
    by_rows = build(s, request(published, "temperature", format="sqlite"), wait=0)
    assert by_rows["job"]["state"] == "queued" and by_rows["row_count"] == 6
    assert s.jobs.wait(by_rows["job"]["job_id"], 60)["state"] == "succeeded"
