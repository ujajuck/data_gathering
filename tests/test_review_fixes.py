"""적대적 리뷰 반영 회귀 테스트: 승계 헤드 보호(§4.4), 정의 파일 원자성(§1.3), 실패 큐·auto_approved, 응답 형태(§6·§7),
빌드 수식 주입 방지, 정규식 ReDoS 차단, schema_key 경로 조작, 잘못된 effective_spec, 커서·원본 참조 정규화,
렌더 서버 토큰, 작업 대기 알림, 프로파일 문서 페이지, 한국어 검증 오류, 큐 next_action, describe 1회 로드, 목록 인덱스,
옛 배치 DB 감지."""

from __future__ import annotations

import csv
import io
import json
import sqlite3
import time

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from schema import build as build_module
from schema import operations
from schema import readers as readers_module
from schema.api import create_app
from schema.db import PREVIOUS_DB_PATH, Database, Problem, encode_cursor
from schema.profile import compile_regex
from schema.render.client import RenderClient
from schema.render.server import RenderWorker, create_render_app
from schema.service import _integrity, normalize_source_ref
from tests.ops_fixture import SCHEMA, World, profile_definition


@pytest.fixture
def world(tmp_path, monkeypatch):
    w = World(tmp_path, monkeypatch)
    yield w
    w.service.close()


def heads(service, application_id):
    return {m["rule_key"]: m for m in service.application_mappings(application_id)}


# ---------------------------------------------------------------------------- §4.4 승계 헤드는 rematch가 승인하지 않는다


def test_rematch_keeps_inherited_heads_proposed(world):
    s = world.service
    base = world.approve_profile_with("a.xlsx")
    pid = base["profile"]["profile_id"]
    world.file("a.xlsx", temps=(11, 21, 31))
    changed = world.register("a.xlsx")
    assert changed["status"] == "changed" and changed["applied"][0]["compatibility"] == "identical"
    app_id = changed["applied"][0]["application_id"]
    assert all(m["status"] == "proposed" and m["origin"] == "inherited" for m in heads(s, app_id).values())

    job = s.jobs.wait(s.reparse(pid, "rematch")["job_id"], 60)
    assert job["state"] == "succeeded", job
    assert {x["reason"] for x in job["result"]["skipped"] if x["document_id"] == changed["document_id"]} == {"review_required"}
    after = heads(s, app_id)
    assert all(m["status"] == "proposed" and m["origin"] == "inherited" and m["revision_no"] == 1 for m in after.values())
    assert not s.application_summary(app_id)["published"] and world.status(changed["document_id"]) == "changed"

    # 사람이 approve_all 하면 그때 승인·추출·발행된다.
    done = s.approve_all(app_id)
    assert done["approved"] == 3 and done["extraction"]["state"] == "succeeded"
    assert world.status(changed["document_id"]) == "normal"


def test_rematch_marks_auto_approved_and_refreshes_stale_rev(world):
    s = world.service
    base = world.approve_profile_with("a.xlsx")
    pid, ref_app = base["profile"]["profile_id"], base["application"]["application_id"]
    world.file("b.xlsx", temps=(1, 2, 3))
    doc_b = world.register("b.xlsx")
    assert doc_b["status"] == "normal" and doc_b["applied"][0]["state"] == "published"
    app_b = doc_b["applied"][0]["application_id"]
    # 새 프로파일 리비전 → 참조 rev가 뒤처진다. b의 헤드를 같은 내용으로 다시 승인하면 발행이 풀리고(전부 approved, 미발행, 옛 rev).
    v2 = profile_definition(description="v2")
    v2["rules"][0]["selector"]["value"]["areas"][0]["relative"]["rows"] = 21  # 규칙 스펙이 실제로 바뀐 리비전
    revised = s.import_profile("process_standard", v2, profile_id=pid)
    assert revised["current_rev"] == 2
    lot = heads(s, app_b)["lot"]
    s.revise(lot["mapping_id"], expected_seq=lot["edit_seq"], status="approved", extract=False)
    assert not s.application_summary(app_b)["published"]
    world.file("d.xlsx")
    doc_d = world.register("d.xlsx")
    assert doc_d["status"] == "review" and doc_d["applied"][0]["compatibility"] == "compatible"
    app_d = doc_d["applied"][0]["application_id"]

    again = s.approve_profile(pid, ref_app)
    assert again["reference_profile_rev"] == 2
    job = s.jobs.wait(again["reparse_job"]["job_id"], 120)
    assert job["state"] == "succeeded", job
    # 오래된 rev의 전부-approved 미발행 건(b)은 그냥 재추출되지 않고 rev 2 스펙의 리비전을 받은 뒤 추출·발행된다.
    for app_id in (app_b, app_d):
        summary = s.application_summary(app_id)
        assert summary["published"], app_id
        for m in summary["mappings"]:
            assert m["status"] == "approved" and m["origin"] == "auto"
        # 스펙이 바뀐 규칙(lot)만 rev 2 리비전을 받는다(같은 스펙의 헤드는 그대로).
        latest = s.mapping_revisions(next(m for m in summary["mappings"] if m["rule_key"] == "lot")["mapping_id"])[0]
        assert latest["evidence"]["profile_rev"] == 2 and latest["evidence"].get("rematch")
        with s.db.connect() as conn:
            run = conn.execute(
                "SELECT r.auto_approved FROM extraction_run r JOIN parsing_application a ON a.published_run_id=r.run_id WHERE a.application_id=?", (app_id,)
            ).fetchone()
        assert run["auto_approved"] == 1


# ---------------------------------------------------------------------------- 정의 파일은 projection 커밋 뒤에만 남는다


def test_rejected_import_leaves_no_definition_file(world):
    s = world.service
    base = world.approve_profile_with("a.xlsx")
    pid = base["profile"]["profile_id"]
    schema_dir = world.root / "schemas/process_standard"
    before = (schema_dir / "current.json").read_bytes()
    bad = json.loads(json.dumps(SCHEMA))
    bad["description"] = "changed"
    next(f for f in bad["fields"] if f["field_key"] == "temperature")["type"] = "group"
    with pytest.raises(Problem) as exc:
        s.import_schema(bad)
    assert exc.value.code == "GROUP_FIELD_TARGET"
    assert not (schema_dir / "r0002.json").exists() and (schema_dir / "current.json").read_bytes() == before
    assert s.schema_definition("process_standard")["fields"][2]["type"] == "decimal"

    other = s.import_profile("process_standard", profile_definition(profile_name="다른프로파일"))
    profile_dir = world.root / "profiles" / pid
    current = (profile_dir / "current.json").read_bytes()
    with pytest.raises(Problem) as exc:
        s.import_profile("process_standard", profile_definition(profile_name="다른프로파일"), profile_id=pid)
    assert exc.value.code == "PROFILE_NAME_CONFLICT"
    assert not (profile_dir / "r0002.json").exists() and (profile_dir / "current.json").read_bytes() == current
    with s.db.connect() as conn:
        assert conn.execute("SELECT current_rev FROM parsing_profile WHERE profile_id=?", (pid,)).fetchone()[0] == 1
    # 새 프로파일이 거부되면 폴더 자체가 남지 않는다.
    with pytest.raises(Problem):
        s.import_profile("process_standard", profile_definition(profile_name=other["profile_name"]))
    assert len([p for p in (world.root / "profiles").iterdir()]) == 2


def test_schema_key_cannot_escape_schemas_folder(world):
    s = world.service
    for key in ("..", ".", "../x", "-x"):
        bad = {**SCHEMA, "schema_key": key, "schema_name": "x" + key}
        with pytest.raises(Problem) as exc:
            s.import_schema(bad)
        assert exc.value.code == "INVALID_SCHEMA", key
    assert not (world.root / "current.json").exists() and not (world.root / "schemas/current.json").exists()


# ---------------------------------------------------------------------------- 실패 큐 · 응답 형태 · 무결성 문구


def test_failed_queue_includes_published_application_with_failed_rerun(world):
    s = world.service
    world.approve_profile_with("a.xlsx")
    world.file("b.xlsx", temps=(1, 2, 3))
    doc_b = world.register("b.xlsx")
    app_b = doc_b["applied"][0]["application_id"]
    with s.db.connect(write=True) as conn:
        app = conn.execute("SELECT snapshot_id, schema_rev, profile_rev FROM parsing_application WHERE application_id=?", (app_b,)).fetchone()
        conn.execute(
            "INSERT INTO extraction_run (run_id, application_id, snapshot_id, schema_rev, profile_rev, engine_version, input_manifest_json, status, started_at, finished_at, error_summary, auto_approved)"
            " VALUES ('run-fail', ?, ?, ?, ?, 'x', '{}', 'failed', '2999-01-01T00:00:00+00:00', '2999-01-01T00:00:01+00:00', 'UNIT_MISMATCH: 단위', 1)",
            (app_b, app["snapshot_id"], app["schema_rev"], app["profile_rev"]),
        )
    assert s.refresh_document_status(doc_b["document_id"])["status"] == "failed"
    found = operations.queues(s)
    group = next(g for g in found["groups"]["failed"] if g["group_key"] == "UNIT_MISMATCH")
    assert group["count"] == 1 and group["impact"]["auto_approved"] == 1 and group["write_actions"] == ["reparse"]
    assert group["next_action"]["kind"] == "open_review" and group["next_action"]["route"] == f"?review={app_b}"
    member = operations.members(s, "failed", "UNIT_MISMATCH")["items"][0]
    assert member["document_status"] == "failed" and member["application_state"] == "published" and member["state"] == "failed"
    assert member["next_action"]["route"] == f"?review={app_b}"
    unmatched = found["groups"]["unmatched"]
    assert unmatched == [] or unmatched[0]["write_actions"] == ["assign_profile"]


def test_responses_do_not_expose_run_ids(world):
    s = world.service
    base = world.approve_profile_with("a.xlsx")
    app_id = base["application"]["application_id"]
    summary = s.application_summary(app_id)
    assert summary["published"] and "published_run_id" not in summary
    value = s.application_values(app_id)["items"][0]
    detail = s.value(value["value_id"])
    assert "run_id" not in detail and detail["mapping_id"] and detail["published"] is True
    region = detail["regions"][0]["region_id"]
    reverse = s.region_values(region)
    assert reverse["values"] and set(reverse["values"][0]) >= {"value_id", "published", "application_id", "rule_key", "field"}
    assert not ({"run_id", "mapping_revision_id", "published_run_id"} & set(reverse["values"][0]))
    assert reverse["mappings"] and not ({"mapping_revision_id"} & set(reverse["mappings"][0])) and reverse["mappings"][0]["mapping_id"]
    assert s.status()["workspace"] == world.root.name


def test_integrity_messages_are_korean_and_mapped():
    assert _integrity(sqlite3.IntegrityError("only a successful owned run can be published")).code == "RUN_NOT_PUBLISHABLE"
    assert _integrity(sqlite3.IntegrityError("extraction_run is immutable after completion")).code == "IMMUTABLE_RECORD"
    assert _integrity(sqlite3.IntegrityError("invalid extraction state transition")).code == "INVALID_RUN_STATE"
    assert _integrity(sqlite3.IntegrityError("mapping head can only advance through a new revision")).code == "IMMUTABLE_RECORD"
    fallback = _integrity(sqlite3.IntegrityError("some raw trigger text 123"))
    assert fallback.code == "INTEGRITY_ERROR" and "raw trigger" not in fallback.message


# ---------------------------------------------------------------------------- 빌드 수식 주입 · 정규식 · 잘못된 입력


def test_build_exports_escape_formula_text(world):
    s = world.service
    evil = '=HYPERLINK("http://evil","x")'
    base = world.approve_profile_with("a.xlsx", signer=evil)
    request = {"document_ids": [base["document"]["document_id"]], "schema_key": "process_standard", "columns": [{"field_key": "note", "header": "=비고"}, {"field_key": "lot", "header": "LOT"}], "row_mode": "record"}
    result = build_module.build(s, {**request, "format": "csv"})
    text = open(result["path"], encoding="utf-8-sig").read()
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][0] == "'=비고" and rows[1][0] == "'" + evil and not any(cell.startswith("=") for row in rows for cell in row)
    result = build_module.build(s, {**request, "format": "xlsx"})
    ws = load_workbook(result["path"])["data"]
    assert ws["A1"].value == "=비고" and ws["A1"].data_type == "s" and ws["A2"].value == evil and ws["A2"].data_type == "s"
    assert build_module.csv_safe("-5.5") == "-5.5" and build_module.csv_safe("+3") == "+3" and build_module.csv_safe("@x") == "'@x" and build_module.csv_safe(None) == ""


def test_regex_nested_quantifiers_rejected():
    for pattern in (r"^(a|a?)+$", r"(a+)+", r"^(\w+\s?)*$", r"(?:x{2,}y)*"):
        with pytest.raises(Problem) as exc:
            compile_regex(pattern)
        assert exc.value.code == "INVALID_REGEX", pattern
    for pattern in (r"^(배치|LOT)$", r"^\d+C$", r"(?i)^온도\s*\(.*\)$", r"[^\s]+ [^\s]+", r"^(?:ab){2,5}$"):
        assert compile_regex(pattern).pattern == pattern


def test_malformed_effective_spec_is_422(world):
    s = world.service
    base = world.approve_profile_with("a.xlsx")
    lot = heads(s, base["application"]["application_id"])["lot"]
    for spec in (
        {"selector": {"key": {"areas": [{"all_of": 5, "sheet_role": "main"}]}}},
        {"selector": {"key": {"areas": [{"find": {"texts": ["x"]}, "anchor_name": "a"}]}}},
        {"selector": {"key": {"areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "anchor": {"all_of": [1]}}}]}}},
        {"selector": {"key": {"areas": "nope"}}},
        {"selector": 5},
    ):
        with pytest.raises(Problem) as exc:
            s.revise(lot["mapping_id"], expected_seq=lot["edit_seq"], status="proposed", effective_spec=spec)
        assert exc.value.status == 422 and exc.value.code in ("INVALID_RULE", "INVALID_SELECTOR"), spec


def test_queue_cursor_with_wrong_types_is_422(world):
    s = world.service
    for values in (["abc", 1], [1, 2]):
        with pytest.raises(Problem) as exc:
            operations.queue(s, "review", cursor=encode_cursor(values, ["queue", "review"]))
        assert exc.value.code == "INVALID_CURSOR"


# ---------------------------------------------------------------------------- 원본 참조 정규화


def test_source_ref_is_normalized(world):
    s = world.service
    world.file("a.xlsx")
    first = world.register("a.xlsx")
    same = world.register("./sub/../a.xlsx")
    assert same["document_id"] == first["document_id"] and same["snapshot"]["unchanged"]
    with s.db.connect() as conn:
        assert conn.execute("SELECT count(*) FROM document").fetchone()[0] == 1
    for bad in ("../a.xlsx", "/etc/passwd", "..", "C:/x.xlsx"):
        with pytest.raises(Problem) as exc:
            normalize_source_ref(bad)
        assert exc.value.code == "INVALID_SOURCE", bad
    job = s.register_documents(["../a.xlsx"], wait=60)
    assert job["state"] == "failed" and job["error_code"] == "INVALID_SOURCE"


# ---------------------------------------------------------------------------- 렌더 서버 토큰


class AppTransport(httpx.BaseTransport):
    def __init__(self, app):
        self.client = TestClient(app, base_url="http://127.0.0.1")

    def handle_request(self, request):
        request.read()
        response = self.client.request(request.method, str(request.url), headers=dict(request.headers), content=request.content)
        return httpx.Response(response.status_code, headers=response.headers, content=response.content)


def test_render_server_requires_internal_token(tmp_path, monkeypatch):
    """렌더 서버의 bearer는 서버 대 서버 내부 토큰(SCHEMA_RENDER_TOKEN)이다 — 사용자 접근 토큰이 아니다(계약 §5)."""
    monkeypatch.setenv("SCHEMA_RENDER_TOKEN", "secret")
    worker = RenderWorker(tmp_path, event_source=lambda *a, **k: iter([]))
    app = create_render_app(tmp_path, worker=worker)
    with TestClient(app, base_url="http://127.0.0.1") as client:
        assert client.get("/render/status").status_code == 401
        assert client.get("/render/status", headers={"Authorization": "Bearer wrong"}).status_code == 401
        assert client.get("/render/status", headers={"Authorization": "Bearer secret"}).status_code == 200
        assert client.delete(f"/render/{'0' * 8}-0000-0000-0000-000000000000").status_code == 401
    # RenderClient(http)는 같은 토큰을 붙인다.
    client = RenderClient(tmp_path, url="http://render", transport=AppTransport(app))
    try:
        assert client.status()["renderer_version"]
    finally:
        client.close()
    worker.close()


# ---------------------------------------------------------------------------- 작업 대기·페이지·검증 오류·describe


def test_wait_returns_promptly_with_worker_thread(world):
    s = world.service
    world.file("a.xlsx")
    s.jobs.start()
    started = time.monotonic()
    job = s.register_documents(["a.xlsx"], wait=30)
    assert job["state"] == "succeeded" and time.monotonic() - started < 20


def test_profile_documents_paging(world):
    s = world.service
    base = world.approve_profile_with("a.xlsx")
    pid = base["profile"]["profile_id"]
    for name in ("b.xlsx", "c.xlsx"):
        world.file(name, temps=(1, 2, 3))
        world.register(name)
    page1 = s.profile_documents(pid, limit=2)
    assert [d["document_name"] for d in page1["items"]] == ["a.xlsx", "b.xlsx"] and page1["has_more"]
    assert page1["items"][0]["state"] == "published" and "status" not in page1["items"][0] and page1["items"][0]["document_status"] == "normal"
    page2 = s.profile_documents(pid, cursor=page1["next_cursor"], limit=2)
    assert [d["document_name"] for d in page2["items"]] == ["c.xlsx"] and not page2["has_more"]


def test_validation_errors_are_korean(tmp_path):
    app = create_app(tmp_path / "ws", start_worker=False)
    with TestClient(app, base_url="http://127.0.0.1") as client:
        body = client.get("/api/documents?limit=500").json()["error"]
        assert body["code"] == "VALIDATION_ERROR" and "200 이하" in body["message"] and body["fields"][0]["detail"].startswith("Input should")
        body = client.get("/api/documents?status=bogus").json()["error"]
        assert "중 하나여야" in body["message"] and "'" not in body["message"]
        body = client.post("/api/snapshots/x/applications", json={}).json()["error"]
        assert body["message"] == "profile_id은(는) 필수입니다."
        body = client.post("/api/profiles/x/reparse", json={"mode": "bad"}).json()["error"]
        assert "mode은(는)" in body["message"] and "fill" in body["message"]
        assert client.get("/api/settings").json()["workspace"] == "ws"


def test_describe_opens_workbook_once(world, monkeypatch):
    s = world.service
    s.import_schema(SCHEMA)
    profile = s.import_profile("process_standard", profile_definition())
    world.file("a.xlsx")
    loads, hashes = [], []
    original_load, original_hash = readers_module.load_workbook, readers_module.file_hash

    def counted_load(*a, **k):
        loads.append(k.get("read_only", False))
        return original_load(*a, **k)

    def counted_hash(path):
        hashes.append(str(path))
        return original_hash(path)

    monkeypatch.setattr(readers_module, "load_workbook", counted_load)
    monkeypatch.setattr(readers_module, "file_hash", counted_hash)
    reader = readers_module.XlsxReader(world.root, "tester")
    canonical = s.profile_canonical(profile["profile_id"], 1)
    result = reader.describe("a.xlsx", profiles=[{"profile_id": profile["profile_id"], "profile_rev": 1, "canonical": canonical, "reference": None}])
    assert result["sheets"][0]["name"] == "공정 기록" and result["signature"]["sheets"] and result["matches"][0]["compatibility"] == "compatible"
    assert len(loads) == 1 and len(hashes) == 2


def test_document_list_sorts_use_indexes(world):
    s = world.service
    with s.db.connect() as conn:
        for column, direction in (("coalesce(d.last_processed_at,'')", "DESC"), ("d.document_name", "ASC"), ("d.status", "ASC")):
            plan = " ".join(
                r[3]
                for r in conn.execute(
                    f"EXPLAIN QUERY PLAN SELECT d.*, {column} sort_key, s.revision_no FROM document d LEFT JOIN document_snapshot s ON s.snapshot_id=d.current_snapshot_id ORDER BY {column} {direction}, d.document_id {direction} LIMIT 51"
                )
            )
            assert "TEMP B-TREE" not in plan, (column, plan)
        for name in ("value_by_run_revision", "mapping_by_head", "value_by_field_created", "value_unit_by_run", "document_by_processed"):
            assert conn.execute("SELECT 1 FROM sqlite_master WHERE type='index' AND name=?", (name,)).fetchone()


def test_previous_workspace_db_is_detected_instead_of_creating_an_empty_one(tmp_path):
    """옛 배치의 DB가 남아 있는 작업 공간을 열면 빈 DB를 만들지 않고 멈춘다(자동 이관은 하지 않는다)."""
    ws = tmp_path / "ws"
    Database(ws)  # 새 작업 공간 — <ws>/workspace.db를 만든다
    previous = ws / PREVIOUS_DB_PATH
    previous.parent.mkdir(parents=True, exist_ok=True)
    for stray in ws.glob("workspace.db-*"):
        stray.unlink()
    (ws / "workspace.db").rename(previous)

    with pytest.raises(Problem) as exc:
        Database(ws)
    assert exc.value.code == "WORKSPACE_DB_MOVED" and exc.value.status == 409
    assert "workspace.db" in exc.value.message
    assert not (ws / "workspace.db").exists()  # 빈 DB를 만들지 않았다

    # 안내대로 옮기면 그대로 열린다.
    previous.rename(ws / "workspace.db")
    assert Database(ws).path == (ws / "workspace.db").resolve()
