"""렌더 서버(계약 §5, §11 test_v3_render.py): 밴드·창 불변식·asset 격리·202/200/304·멱등 큐·invalidate 세대·격리 중 창 응답."""

from __future__ import annotations

import io
import json
import statistics
import threading
import time
import uuid

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Font, PatternFill
from PIL import Image

from kg.v2.readers import file_hash
from kg.v3.db import Problem
from kg.v3.render import RENDERER_VERSION
from kg.v3.render.assemble import assemble
from kg.v3.render.cache import RenderCache
from kg.v3.render.client import RenderClient, RenderUnavailable
from kg.v3.render.renderer import BAND_BYTES, _split_band, render_events
from kg.v3.render.server import RenderWorker, create_render_app

SNAPSHOT_A, SNAPSHOT_B = str(uuid.uuid4()), str(uuid.uuid4())
SHEET_A, SHEET_B = str(uuid.uuid4()), str(uuid.uuid4())


def png_bytes(color="red", size=(40, 30)):
    data = io.BytesIO()
    Image.new("RGB", size, color).save(data, format="PNG")
    return data.getvalue()


def build_workbook(path, rows=250, with_image=True):
    """250행: 밴드 3개. 병합 B2:C3, 밴드 경계를 걸치는 A95:A105, 서식 셀, 이미지 D5·(먼 곳) H200."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    bold = Font(bold=True)
    fill = PatternFill(patternType="solid", fgColor="FFFFEE00")
    for r in range(1, rows + 1):
        ws.cell(r, 1, f"row{r}")
        ws.cell(r, 2, r * 1.0)
    ws["B2"] = "merged"
    ws.merge_cells("B2:C3")
    ws["A95"] = "tall"
    ws.merge_cells("A95:A105")
    ws["E1"] = "bold-1"
    ws["E1"].font = bold
    ws["E2"] = "bold-2"
    ws["E2"].font = bold
    ws["F1"].fill = fill  # 값 없는 서식 셀: 생략하지 않는다
    ws.column_dimensions["A"].width = 20
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "B2"
    if with_image:
        pic = XlImage(io.BytesIO(png_bytes()))
        ws.add_image(pic, "D5")
        far = XlImage(io.BytesIO(png_bytes("blue")))
        ws.add_image(far, "H200")
    wb.create_sheet("Other")["A1"] = "other"
    wb.save(path)
    return path


def make_source(delay_event=None, fail=None, started=None):
    """Reader 격리 프로세스 대신 같은 프로세스에서 renderer를 직접 부르는 event_source(다른 에이전트의 readers 모듈과 무관)."""

    def source(root, provider, principal, payload, checkpoint):
        path = root / "data/raw" / payload["source_ref"]
        if fail is not None:
            raise fail
        token = file_hash(path)
        if payload["expected_token"] != token:
            raise Problem("SOURCE_VERSION_CHANGED", "원본이 변경되었습니다.", 409)
        for event in render_events(
            path, payload["sheet_name"], token, payload["r1"], payload["c1"], payload["rows"], payload["cols"]
        ):
            yield event
            if started is not None:
                started.set()
            if delay_event is not None:
                delay_event.wait(10)
            checkpoint()
        yield {"type": "verified", "token": token}

    return source


def until(condition, timeout=10.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = condition()
        if value:
            return value
        time.sleep(0.02)
    raise AssertionError("timed out")


@pytest.fixture
def root(tmp_path):
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    build_workbook(raw / "doc.xlsx")
    return tmp_path


def request_body(snapshot=SNAPSHOT_A, sheet=SHEET_A, root=None, **extra):
    return {
        "snapshot_id": snapshot,
        "sheet_id": sheet,
        "sheet_name": "Data",
        "provider": "local-xlsx",
        "source_ref": "doc.xlsx",
        "expected_token": file_hash(root / "data/raw/doc.xlsx"),
        **extra,
    }


def events_of(root, sheet="Data"):
    path = root / "data/raw/doc.xlsx"
    token = file_hash(path)
    return [*render_events(path, sheet, token), {"type": "verified", "token": token}]


# ---- renderer --------------------------------------------------------------------------


def test_band_layout_and_styles(root):
    events = events_of(root)
    meta, bands, images = events[0], [e for e in events if e["type"] == "band"], [e for e in events if e["type"] == "image"]
    assert meta["type"] == "meta" and meta["mode"] == "simplified" and meta["renderer_version"] == RENDERER_VERSION
    assert meta["band_rows"] == 100 and meta["rendered_bounds"] == {"r2": 250, "c2": 26}
    assert [(b["r1"], b["r2"]) for b in bands] == [(1, 100), (101, 200), (201, 250)]
    assert meta["truncated"] is False and meta["estimated_rows"] == 250
    assert meta["freeze"] == {"rows": 1, "cols": 1}
    assert [2, 2, 3, 3] in meta["merges"] and [95, 1, 105, 1] in meta["merges"]
    assert meta["rows"][0] == {"index": 1, "y": 0, "height": 20} and meta["rows"][1]["height"] == 40
    assert meta["columns"][0]["width"] == 145 and meta["columns"][1]["x"] == 145
    first = {(c["r1"], c["c1"]): c for c in bands[0]["cells"]}
    assert first[(2, 2)] == {"r1": 2, "c1": 2, "r2": 3, "c2": 3, "text": "merged", "s": 0}
    assert (3, 3) not in first and (2, 3) not in first  # 병합 비앵커는 없다
    assert first[(1, 5)]["s"] == first[(2, 5)]["s"] != 0  # 같은 서식은 같은 인덱스
    assert meta["styles"][first[(1, 5)]["s"]]["fontWeight"] == 700
    assert first[(1, 6)]["text"] == "" and meta["styles"][first[(1, 6)]["s"]]["background"] == "#ffee00"
    assert (1, 7) not in first  # 빈 기본 셀은 생략
    assert first[(1, 2)]["text"] == "1"
    # 밴드 경계를 걸치는 병합 앵커는 두 밴드에 모두(창에서 중복 제거) 있다
    second = {(c["r1"], c["c1"]): c for c in bands[1]["cells"]}
    assert first[(95, 1)]["r2"] == 105 and second[(95, 1)] == first[(95, 1)]
    assert len(bands[0]["cells"]) + len(bands[1]["cells"]) + len(bands[2]["cells"]) == meta["cell_count"] + 1
    assert [i["ext"] for i in images] == ["png", "png"] and images[0]["x"] == 145 + 96 * 2
    assert all(len(json.dumps(e)) < 8 * 1024 * 1024 for e in events)


def test_sub_band_split():
    text = "x" * 4000
    cells = [{"r1": r, "c1": c, "r2": r, "c2": c, "text": text, "s": 0} for r in range(1, 101) for c in range(1, 21)]
    parts = _split_band(1, 100, cells, 100)
    assert [(a, b) for a, b, _, _ in parts] == [(1, 50), (51, 100)]
    assert all(len(s.encode()) <= BAND_BYTES for _, _, _, s in parts)
    assert sum(len(c) for _, _, c, _ in parts) == len(cells)
    cells = [{"r1": r, "c1": c, "r2": r, "c2": c, "text": text, "s": 0} for r in range(1, 101) for c in range(1, 41)]
    parts = _split_band(1, 100, cells, 100)
    assert [(a, b) for a, b, _, _ in parts] == [(1, 25), (26, 50), (51, 75), (76, 100)]


def test_truncation_caps(root, monkeypatch):
    path = root / "data/raw/big.xlsx"
    wb = Workbook()
    ws = wb.active
    for r in range(1, 2101):
        ws.cell(r, 1, r)
    ws.cell(1, 210, "far")
    wb.save(path)
    meta = next(render_events(path, ws.title, "t"))
    assert meta["truncated"] is True and meta["rendered_bounds"] == {"r2": 2000, "c2": 200}
    assert meta["estimated_rows"] == 2100 and meta["estimated_cols"] == 210 and len(meta["rows"]) == 2000
    monkeypatch.setattr("kg.v3.render.renderer.MAX_CELLS", 150)
    events = list(render_events(root / "data/raw/doc.xlsx", "Data", "t"))
    meta = events[0]
    assert meta["truncated"] is True and meta["rendered_bounds"]["r2"] == 100
    assert [(b["r1"], b["r2"]) for b in events if b["type"] == "band"] == [(1, 100)]
    assert meta["merges"] == [[2, 2, 3, 3], [95, 1, 100, 1]]
    with pytest.raises(Problem) as info:
        list(render_events(root / "data/raw/doc.xlsx", "Missing", "t"))
    assert info.value.code == "SHEET_NOT_FOUND" and info.value.status == 404


# ---- cache / window --------------------------------------------------------------------


@pytest.fixture
def cache(root):
    cache = RenderCache(root)
    key = (SNAPSHOT_A, SHEET_A, RENDERER_VERSION)
    assemble(iter(events_of(root)), cache, key, file_hash(root / "data/raw/doc.xlsx"))
    return cache


def test_assemble_files_and_window_invariants(root, cache):
    key = (SNAPSHOT_A, SHEET_A, RENDERER_VERSION)
    directory = cache.dir(key)
    assert sorted(p.name for p in directory.iterdir()) == ["band-1-100.json", "band-101-200.json", "band-201-250.json", "meta.json"]
    assert not list(directory.glob("*.tmp"))
    meta = json.loads((directory / "meta.json").read_text())
    assert [i["asset_id"] for i in meta["images"]] and all("bytes" not in i for i in meta["images"])
    assert meta["images"][0]["url"] == f"/api/v3/snapshots/{SNAPSHOT_A}/render-assets/{meta['images'][0]['asset_id']}"
    assert (cache.assets_dir(SNAPSHOT_A) / meta["images"][0]["asset_id"]).is_file()

    window = cache.window(key)
    assert window["range"] == "A1:Z60" and window["sheet"] == {"sheet_id": SHEET_A, "sheet_name": "Data"}
    assert len(window["rows"]) == 250 and len(window["columns"]) == 26
    y = 0
    for n, row in enumerate(window["rows"]):
        assert row["index"] == n + 1 and row["y"] == pytest.approx(y)
        y += row["height"]
    assert window["height"] == pytest.approx(y)
    x = 0
    for n, col in enumerate(window["columns"]):
        assert col["index"] == n + 1 and col["x"] == pytest.approx(x)
        x += col["width"]
    assert window["width"] == pytest.approx(x)
    anchors = [(c["r1"], c["c1"]) for c in window["cells"]]
    assert len(anchors) == len(set(anchors))
    assert all(c["r2"] >= 1 and c["r1"] <= 60 and c["c2"] >= 1 and c["c1"] <= 26 for c in window["cells"])
    assert window["merges"] == [[2, 2, 3, 3]]
    assert len(window["images"]) == 1  # H200의 이미지는 창 밖
    assert window["renderer_version"] == RENDERER_VERSION and window["freeze"] == {"rows": 1, "cols": 1}

    # 밴드 경계(100/101)를 걸치는 병합은 한 번만, 창 밖으로 나가면 clipped
    window = cache.window(key, "A90:B100")
    straddle = [c for c in window["cells"] if (c["r1"], c["c1"]) == (95, 1)]
    assert len(straddle) == 1 and straddle[0]["clipped"] is True and straddle[0]["r2"] == 105
    assert window["merges"] == [[95, 1, 105, 1]] and window["images"] == []
    window = cache.window(key, "A101:B110")
    straddle = [c for c in window["cells"] if (c["r1"], c["c1"]) == (95, 1)]
    assert len(straddle) == 1 and straddle[0]["clipped"] is True
    expected = {(95, 1), *((r, 2) for r in range(101, 111)), *((r, 1) for r in range(106, 111))}  # A101:A105는 병합에 덮임
    assert {(c["r1"], c["c1"]) for c in window["cells"]} == expected
    # 이미지는 창과 교차할 때만
    window = cache.window(key, "G195:J205")
    assert len(window["images"]) == 1 and window["images"][0]["y"] >= window["rows"][198]["y"]


def test_window_range_rules(root, cache):
    key = (SNAPSHOT_A, SHEET_A, RENDERER_VERSION)
    assert cache.window(key, "A240:Z300")["range"] == "A240:Z250"  # 캐시 범위로 잘라 range에 반영
    assert cache.window(key, "X1:AD10")["range"] == "X1:Z10"
    with pytest.raises(Problem) as info:
        cache.window(key, "A1:A121")
    assert info.value.code == "RANGE_TOO_LARGE" and info.value.status == 422
    with pytest.raises(Problem) as info:
        cache.window(key, "A1:BI1")
    assert info.value.code == "RANGE_TOO_LARGE"
    with pytest.raises(Problem) as info:
        cache.window(key, "not-a-range")
    assert info.value.code == "INVALID_RANGE" and info.value.status == 422
    with pytest.raises(Problem) as info:
        cache.window(key, "A300:Z310")
    assert info.value.code == "RANGE_OUT_OF_BOUNDS" and info.value.fields == {"rendered_bounds": {"r2": 250, "c2": 26}}
    with pytest.raises(Problem) as info:
        cache.window(key, "AA1:AB5")
    assert info.value.code == "RANGE_OUT_OF_BOUNDS"
    assert cache.etag(key, "A1:Z60").startswith('"') and len(cache.etag(key, "A1:Z60")) == 34
    assert cache.etag(key, "A1:Z60") != cache.etag(key, "A1:Z61")


def test_cache_lru_disk_budget_and_purge(root, cache):
    key = (SNAPSHOT_A, SHEET_A, RENDERER_VERSION)
    cache.lru_limit = 2000  # 바이트 상한: 큰 밴드는 머물지 못한다
    cache.lru.clear()
    cache.lru_bytes = 0
    cache.window(key, "A1:Z60")
    assert cache.lru_bytes <= 2000 or len(cache.lru) == 1
    # 디스크 예산: 두 번째 snapshot을 넣으면 오래된 snapshot이 제거된다
    cache.disk_limit = cache.cache_bytes() + 10
    assemble(iter(events_of(root)), cache, (SNAPSHOT_B, SHEET_B, RENDERER_VERSION))
    assert cache.has((SNAPSHOT_B, SHEET_B, RENDERER_VERSION)) and not cache.has(key)
    assert not cache.snapshot_dir(SNAPSHOT_A).exists()
    # 같은 시트의 다른 renderer_version 디렉터리는 put 시, 시작 시 모두 제거
    stale = cache.snapshot_dir(SNAPSHOT_B) / f"{SHEET_B}.old-renderer_1"
    stale.mkdir()
    (stale / "meta.json").write_text("{}")
    assemble(iter(events_of(root)), cache, (SNAPSHOT_B, SHEET_B, RENDERER_VERSION))
    assert not stale.exists()
    stale.mkdir()
    (stale / "meta.json").write_text("{}")
    tmp = cache.snapshot_dir(SNAPSHOT_B) / f"{SHEET_B}.{RENDERER_VERSION.replace('/', '_')}.tmp-1-1"
    tmp.mkdir()
    fresh = RenderCache(root)
    assert not stale.exists() and not tmp.exists() and fresh.has((SNAPSHOT_B, SHEET_B, RENDERER_VERSION))
    assert fresh.cache_bytes() > 0
    assert [s["sheet_id"] for s in fresh.sheets(SNAPSHOT_B)] == [SHEET_B]


def test_assemble_rejects_bad_stream(root):
    cache = RenderCache(root)
    key = (SNAPSHOT_A, SHEET_A, RENDERER_VERSION)
    events = events_of(root)
    with pytest.raises(Problem) as info:
        assemble(iter(events[:-1] + [{"type": "verified", "token": "other"}]), cache, key)
    assert info.value.code == "SOURCE_VERSION_CHANGED" and not cache.has(key)
    with pytest.raises(Problem) as info:
        assemble(iter(events), cache, key, expected_token="expected-other")
    assert info.value.code == "SOURCE_VERSION_CHANGED"
    with pytest.raises(Problem) as info:
        assemble(iter(events[:-1]), cache, key)
    assert info.value.code == "RENDER_STREAM_INVALID"
    assert not list(cache.snapshot_dir(SNAPSHOT_A).glob("*.tmp*"))
    with pytest.raises(Problem) as info:
        assemble(iter(events), cache, key, generation=99)
    assert info.value.code == "CANCELLED" and not cache.has(key)


def test_asset_isolation(root):
    cache = RenderCache(root)
    for snapshot, sheet in ((SNAPSHOT_A, SHEET_A), (SNAPSHOT_B, SHEET_B)):
        assemble(iter(events_of(root)), cache, (snapshot, sheet, RENDERER_VERSION))
    meta_a, meta_b = cache.meta((SNAPSHOT_A, SHEET_A, RENDERER_VERSION)), cache.meta((SNAPSHOT_B, SHEET_B, RENDERER_VERSION))
    asset = meta_a["images"][0]["asset_id"]
    assert asset == meta_b["images"][0]["asset_id"]  # 같은 바이트 → 같은 sha256
    assert meta_a["images"][0]["url"] != meta_b["images"][0]["url"]
    path_a, path_b = cache.asset_path(SNAPSHOT_A, asset), cache.asset_path(SNAPSHOT_B, asset)
    assert path_a != path_b and path_a.is_file() and path_b.is_file()
    assert cache.asset_path(SNAPSHOT_A, "../" + asset) is None
    assert cache.asset_path(SNAPSHOT_A, asset.replace(".png", ".svg")) is None
    assert cache.asset_path(SNAPSHOT_A, "x" * 64 + ".png") is None
    assert cache.asset_path("../" + SNAPSHOT_A[3:], asset) is None
    cache.invalidate(SNAPSHOT_A)
    assert cache.asset_path(SNAPSHOT_A, asset) is None and path_b.is_file()


# ---- server --------------------------------------------------------------------------------


@pytest.fixture
def app(root):
    worker = RenderWorker(root, event_source=make_source(), queue_limit=4)
    application = create_render_app(root, worker=worker)
    with TestClient(application) as client:
        yield client
    worker.close()


def poll(client, snapshot, sheet, range_text=None, **kw):
    def done():
        response = client.get(f"/render/{snapshot}/sheet/{sheet}", params={"range": range_text} if range_text else None, **kw)
        return response if response.status_code != 202 else None

    return until(done)


def test_post_get_202_200_304(app, root):
    body = request_body(root=root)
    first = app.post("/render", json=body)
    assert first.status_code == 202 and first.json()["status"] in ("queued", "rendering")
    assert set(first.json()) == {"status", "job_id", "position"}
    response = poll(app, SNAPSHOT_A, SHEET_A)
    assert response.status_code == 200
    window = response.json()
    assert window["range"] == "A1:Z60" and window["sheet"]["sheet_id"] == SHEET_A and window["cells"]
    etag = response.headers["etag"]
    assert etag.startswith('"') and len(etag) == 34 and response.headers["cache-control"] == "private, no-cache"
    again = app.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}", headers={"If-None-Match": etag})
    assert again.status_code == 304 and again.headers["etag"] == etag and not again.content
    other = app.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}", params={"range": "B2:D9"})
    assert other.status_code == 200 and other.headers["etag"] != etag and other.json()["range"] == "B2:D9"
    cached = app.post("/render", json={**body, "range": "A1:C3"})
    assert cached.status_code == 200 and cached.json()["status"] == "cached"
    assert cached.json()["sheet"]["range"] == "A1:C3" and cached.headers["etag"]
    assert app.post("/render", json=body, headers={"If-None-Match": etag}).status_code == 304
    sheets = app.get(f"/render/{SNAPSHOT_A}/sheets").json()
    assert [s["sheet_id"] for s in sheets["items"]] == [SHEET_A] and sheets["items"][0]["sheet_name"] == "Data"
    status = app.get("/render/status").json()
    assert status == {"queue_depth": 0, "rendering": 0, "renderer_version": RENDERER_VERSION, "cache_bytes": status["cache_bytes"]}
    assert status["cache_bytes"] > 0
    # 창 규칙은 HTTP에서도 같은 코드
    assert app.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}", params={"range": "A1:A121"}).json()["error"]["code"] == "RANGE_TOO_LARGE"
    assert app.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}", params={"range": "?"}).json()["error"]["code"] == "INVALID_RANGE"
    out = app.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}", params={"range": "A400:B410"})
    assert out.status_code == 422 and out.json()["error"]["code"] == "RANGE_OUT_OF_BOUNDS"
    assert out.json()["error"]["fields"]["rendered_bounds"] == {"r2": 250, "c2": 26}
    # asset
    asset = window["images"][0]["asset_id"]
    image = app.get(f"/render/{SNAPSHOT_A}/assets/{asset}")
    assert image.status_code == 200 and image.headers["content-type"] == "image/png" and image.content == png_bytes()
    assert image.headers["etag"] == f'"{asset}"' and image.headers["cache-control"] == "private, no-cache"
    assert app.get(f"/render/{SNAPSHOT_A}/assets/{asset}", headers={"If-None-Match": f'"{asset}"'}).status_code == 304
    assert app.get(f"/render/{SNAPSHOT_A}/assets/..%2F{asset}").status_code == 404
    assert app.get(f"/render/{SNAPSHOT_A}/assets/{asset[:-4]}.svg").status_code == 404
    assert app.get(f"/render/{SNAPSHOT_B}/assets/{asset}").status_code == 404
    assert app.get(f"/render/not-a-uuid/assets/{asset}").status_code == 404
    # 미렌더·잘못된 경로
    missing = app.get(f"/render/{SNAPSHOT_B}/sheet/{SHEET_B}")
    assert missing.status_code == 404 and missing.json()["error"]["code"] == "NOT_RENDERED"
    assert app.get(f"/render/bad/sheet/{SHEET_B}").status_code == 404
    assert app.post("/render", json={**body, "snapshot_id": "bad"}).status_code == 422
    assert app.post("/render", json={**body, "extra": 1}).json()["error"]["code"] == "INVALID_REQUEST"
    # DELETE는 asset까지 지운다
    deleted = app.delete(f"/render/{SNAPSHOT_A}")
    assert deleted.status_code == 200 and deleted.json()["snapshot_id"] == SNAPSHOT_A
    assert app.get(f"/render/{SNAPSHOT_A}/assets/{asset}").status_code == 404
    assert app.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}").status_code == 404
    assert app.get(f"/render/{SNAPSHOT_A}/sheets").json()["items"] == []


def test_idempotent_enqueue_and_queue_full(root):
    gate = threading.Event()
    worker = RenderWorker(root, event_source=make_source(delay_event=gate), queue_limit=1)
    app = create_render_app(root, worker=worker)
    body = request_body(root=root)
    with TestClient(app) as client:
        first, second = client.post("/render", json=body), client.post("/render", json=body)
        assert first.status_code == second.status_code == 202
        assert first.json()["job_id"] == second.json()["job_id"]
        until(lambda: worker.status()["rendering"] == 1)
        pending = client.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}")
        assert pending.status_code == 202 and pending.json() == {"status": "rendering", "job_id": first.json()["job_id"], "position": 0}
        assert client.get(f"/render/{SNAPSHOT_A}/sheets").json()["pending"][0]["status"] == "rendering"
        queued = client.post("/render", json=request_body(SNAPSHOT_B, SHEET_B, root))
        assert queued.status_code == 202 and queued.json()["status"] == "queued" and queued.json()["position"] == 1
        assert queued.json()["job_id"] == client.post("/render", json=request_body(SNAPSHOT_B, SHEET_B, root)).json()["job_id"]
        full = client.post("/render", json=request_body(SNAPSHOT_B, SHEET_A, root))
        assert full.status_code == 503 and full.json()["error"]["code"] == "RENDER_QUEUE_FULL" and full.headers["retry-after"] == "5"
        assert client.get("/render/status").json()["queue_depth"] == 1
        gate.set()
        assert poll(client, SNAPSHOT_A, SHEET_A).status_code == 200
        assert poll(client, SNAPSHOT_B, SHEET_B).status_code == 200
        assert client.get("/render/status").json()["queue_depth"] == 0
    worker.close()


@pytest.mark.parametrize(
    "problem, status, code",
    [
        (Problem("SOURCE_VERSION_CHANGED", "원본이 변경되었습니다.", 409), 409, "SNAPSHOT_STALE"),
        (Problem("DRM_READER_REQUIRED", "보안 읽기 어댑터가 필요합니다.", 403), 403, "DRM_READER_REQUIRED"),
        (Problem("READER_TIMEOUT", "제한 시간 초과", 408), 408, "READER_TIMEOUT"),
        (Problem("READER_MEMORY_LIMIT", "메모리 한도", 413), 413, "READER_MEMORY_LIMIT"),
        (Problem("READER_FAILED", "읽기 실패", 422), 422, "READER_FAILED"),
        (Problem("SHEET_NOT_FOUND", "시트 없음", 404), 404, "SHEET_NOT_FOUND"),
        (Problem("BOOM", "내부 오류", 500), 422, "BOOM"),
    ],
)
def test_failed_surfaces_with_retry_after(root, problem, status, code):
    worker = RenderWorker(root, event_source=make_source(fail=problem))
    with TestClient(create_render_app(root, worker=worker)) as client:
        assert client.post("/render", json=request_body(root=root)).status_code == 202
        response = poll(client, SNAPSHOT_A, SHEET_A)
        assert response.status_code == status
        assert response.json() == {"status": "failed", "error": {"code": code, "message": problem.message}, "retry_after": 60}
        repeat = client.post("/render", json=request_body(root=root))
        assert repeat.status_code == status and repeat.json()["status"] == "failed"  # 60초 보관, 다시 넣지 않음
        assert client.get("/render/status").json()["queue_depth"] == 0
        # 보관 시간이 지나면 다음 POST가 다시 큐에 넣는다
        job = worker.jobs[(SNAPSHOT_A, SHEET_A, RENDERER_VERSION)]
        job.failed_at -= 61
        assert client.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}").json()["error"]["code"] == "NOT_RENDERED"
        assert client.post("/render", json=request_body(root=root)).status_code == 202
    worker.close()


def test_unexpected_exception_is_render_failed(root):
    def broken(root, provider, principal, payload, checkpoint):
        raise RuntimeError("secret path /etc/passwd")

    worker = RenderWorker(root, event_source=broken)
    with TestClient(create_render_app(root, worker=worker)) as client:
        client.post("/render", json=request_body(root=root))
        response = poll(client, SNAPSHOT_A, SHEET_A)
        assert response.status_code == 422 and response.json()["error"]["code"] == "RENDER_FAILED"
        assert "secret" not in response.text
    worker.close()


def test_invalidate_generation_drops_inflight_result(root):
    gate, started = threading.Event(), threading.Event()
    worker = RenderWorker(root, event_source=make_source(delay_event=gate, started=started))
    with TestClient(create_render_app(root, worker=worker)) as client:
        client.post("/render", json=request_body(root=root))
        assert started.wait(10)
        generation = worker.cache.generation(SNAPSHOT_A)
        assert client.delete(f"/render/{SNAPSHOT_A}").status_code == 200
        assert worker.cache.generation(SNAPSHOT_A) == generation + 1
        assert client.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}").status_code == 404
        gate.set()
        until(lambda: worker.status()["rendering"] == 0)
        assert not worker.cache.has((SNAPSHOT_A, SHEET_A, RENDERER_VERSION))
        assert not list(worker.cache.base.glob("**/*.tmp*"))
        assert client.get(f"/render/{SNAPSHOT_A}/sheet/{SHEET_A}").status_code == 404
        # 다음 요청은 새 세대에서 정상 렌더된다
        assert client.post("/render", json=request_body(root=root)).status_code == 202
        assert poll(client, SNAPSHOT_A, SHEET_A).status_code == 200
    worker.close()


def test_cached_window_under_20ms_while_another_render_runs(root):
    gate = threading.Event()
    client = RenderClient(root, event_source=make_source(delay_event=gate))
    try:
        gate.set()
        assert client.request(SNAPSHOT_A, SHEET_A, "Data", "local-xlsx", "doc.xlsx", request_body(root=root)["expected_token"], "p")["status"] in ("queued", "rendering")
        until(lambda: client.window(SNAPSHOT_A, SHEET_A)["status"] == "cached")
        gate.clear()
        assert client.request(SNAPSHOT_B, SHEET_B, "Data", "local-xlsx", "doc.xlsx", request_body(root=root)["expected_token"], "p")["status"] == "queued"
        until(lambda: client.status()["rendering"] == 1)
        client.window(SNAPSHOT_A, SHEET_A, "A1:Z60")  # 첫 호출은 파일을 읽어 LRU에 올린다
        samples = []
        for n in range(20):
            started = time.perf_counter()
            result = client.window(SNAPSHOT_A, SHEET_A, "A%d:Z%d" % (n + 1, n + 60))
            samples.append(time.perf_counter() - started)
            assert result["status"] == "cached" and result["body"]["cells"]
        assert statistics.median(samples) < 0.02, samples
        assert client.window(SNAPSHOT_A, SHEET_A, "A20:Z79", if_none_match=result["etag"])["status"] == "not_modified"
        gate.set()
        until(lambda: client.window(SNAPSHOT_B, SHEET_B)["status"] == "cached")
    finally:
        client.close()


# ---- client ------------------------------------------------------------------------------


class AppTransport(httpx.BaseTransport):
    """httpx.Client → FastAPI 앱(동기 TestClient) 브리지. 실제 소켓 없이 http 모드를 검증한다."""

    def __init__(self, app):
        self.client = TestClient(app)

    def handle_request(self, request):
        request.read()
        response = self.client.request(
            request.method, str(request.url), headers=dict(request.headers), content=request.content
        )
        return httpx.Response(response.status_code, headers=response.headers, content=response.content)


def run_client_scenario(client, root):
    token = request_body(root=root)["expected_token"]

    def request(**kw):
        return client.request(SNAPSHOT_A, SHEET_A, "Data", "local-xlsx", "doc.xlsx", token, "principal", range="A1:E8", **kw)

    first = request()
    assert first["status"] in ("queued", "rendering") and first["http_status"] == 202 and first["body"]["job_id"]
    cached = until(lambda: next((r for r in [request()] if r["status"] == "cached"), None))
    assert cached["http_status"] == 200 and cached["body"]["status"] == "cached" and cached["body"]["sheet"]["range"] == "A1:E8"
    assert cached["etag"] and len(cached["etag"]) == 34
    same = request(if_none_match=cached["etag"])
    assert same == {"status": "not_modified", "http_status": 304, "body": None, "etag": cached["etag"]}
    window = client.window(SNAPSHOT_A, SHEET_A, "B2:D4")
    assert window["status"] == "cached" and window["body"]["range"] == "B2:D4" and window["etag"] != cached["etag"]
    bad = client.window(SNAPSHOT_A, SHEET_A, "A1:A200")
    assert bad["status"] == "failed" and bad["http_status"] == 422 and bad["body"]["error"]["code"] == "RANGE_TOO_LARGE"
    missing = client.window(SNAPSHOT_B, SHEET_B)
    assert missing["status"] == "not_rendered" and missing["http_status"] == 404
    sheets = client.sheets(SNAPSHOT_A)
    assert [s["sheet_id"] for s in sheets["items"]] == [SHEET_A]
    asset_id = cached["body"]["sheet"]["images"][0]["asset_id"]
    data, media, etag = client.asset(SNAPSHOT_A, asset_id)
    assert data == png_bytes() and media == "image/png" and etag == f'"{asset_id}"'
    assert client.asset(SNAPSHOT_A, "../" + asset_id) is None and client.asset(SNAPSHOT_B, asset_id) is None
    status = client.status()
    assert status["renderer_version"] == RENDERER_VERSION and status["queue_depth"] == 0 and status["cache_bytes"] > 0
    assert client.invalidate(SNAPSHOT_A)["snapshot_id"] == SNAPSHOT_A
    assert client.window(SNAPSHOT_A, SHEET_A)["status"] == "not_rendered"
    assert client.asset(SNAPSHOT_A, asset_id) is None


def test_render_client_http_mode(root):
    worker = RenderWorker(root, event_source=make_source())
    app = create_render_app(root, worker=worker)
    client = RenderClient(root, url="http://render-server", transport=AppTransport(app))
    assert client.mode == "http"
    try:
        run_client_scenario(client, root)
    finally:
        client.close()
        worker.close()


def test_render_client_inprocess_mode(root, monkeypatch):
    monkeypatch.delenv("KG_V3_RENDER_URL", raising=False)
    client = RenderClient(root, event_source=make_source())
    assert client.mode == "inprocess"
    try:
        run_client_scenario(client, root)
    finally:
        client.close()


def test_render_client_unreachable(root, monkeypatch):
    monkeypatch.setenv("KG_V3_RENDER_URL", "http://127.0.0.1:9")
    client = RenderClient(root)
    assert client.mode == "http"
    try:
        with pytest.raises(RenderUnavailable) as info:
            client.request(SNAPSHOT_A, SHEET_A, "Data", "local-xlsx", "doc.xlsx", "t", "p")
        assert info.value.code == "RENDER_UNAVAILABLE" and info.value.status == 503
        with pytest.raises(RenderUnavailable):
            client.status()
    finally:
        client.close()

    class Slow(httpx.BaseTransport):
        def handle_request(self, request):
            raise httpx.ReadTimeout("slow", request=request)

    client = RenderClient(root, url="http://slow", transport=Slow())
    try:
        with pytest.raises(RenderUnavailable):
            client.window(SNAPSHOT_A, SHEET_A)
    finally:
        client.close()
