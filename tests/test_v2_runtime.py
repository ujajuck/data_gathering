"""실제 XLSX → 승인/출처 → 불변 추출 → 사용자 DB API 통합 검증."""

import copy
import json
import sqlite3
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from kg.v2.api import create_app
from kg.v2.db import uid
from kg.v2.spec import decimal_text


@pytest.fixture
def setup(tmp_path):
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "세로"
    ws["B2"], ws["D2"] = "공정", "온도"
    ws.merge_cells("B2:C2")
    ws["A3"], ws["A4"], ws["A5"] = "lot-1", "lot-2", "lot-3"
    ws["B3"], ws["B4"], ws["B5"] = "123456789012345678901234567890.1250", "20.5", "30"
    ws.merge_cells("B3:C3")
    second = wb.create_sheet("단위")
    second["A1"] = "°C"
    second["B1"], second["C1"], second["D1"] = 1, 2, 3
    ws["H1"] = "=SUM(B4:B5)"
    path = raw / "sample.xlsx"
    wb.save(path)
    app = create_app(tmp_path, start_worker=False)
    with TestClient(app) as client:
        service = app.state.v2

        def request(method, url, body=None, code=200):
            response = client.request(method, "/api/v2" + url, json=body)
            assert response.status_code == code, response.text
            return response.json()

        def work(url, body):
            job = request("POST", url, {**body, "request_key": uid()}, 202)
            service.jobs.run_one()
            return request("GET", "/jobs/" + job["job_id"])

        job = work("/documents/register", {"source_refs": ["sample.xlsx"]})
        assert job["state"] == "succeeded", job
        doc = job["result"]["documents"][0]
        sheets = request("GET", "/versions/" + doc["version_id"] + "/sheets")["items"]
        kg = request(
            "POST",
            "/kg/import",
            {
                "concepts": [
                    {
                        "concept_id": "temperature",
                        "name": "공정온도",
                        "level": 1,
                        "aliases": ["공정 온도", "Temp"],
                        "canonical_unit": "°C",
                    },
                    {
                        "concept_id": "ambient",
                        "name": "주위온도",
                        "level": 1,
                        "aliases": ["Temp"],
                    },
                ]
            },
        )
        definition = {
            "kg_revision_id": kg["kg_revision_id"],
            "sheet_roles": {
                "main": {"cardinality": "one"},
                "units": {"cardinality": "one"},
            },
            "rules": [
                {
                    "rule_key": "temperature",
                    "concept_id": "temperature",
                    "record_spec": {"scope": ["process"], "key": {"column": "A"}},
                    "selector": {
                        "key": {
                            "areas": [
                                {"sheet_role": "main", "range": "B2:C2"},
                                {"sheet_role": "main", "range": "D2"},
                            ]
                        },
                        "value": {
                            "areas": [{"sheet_role": "main", "range": "B3:C5"}],
                            "cardinality": "list",
                            "axis": "down",
                            "element_layout": "one_per_row",
                        },
                        "unit": {"areas": [{"sheet_role": "units", "range": "A1"}]},
                    },
                    "value_spec": {"type": "decimal", "unit": "°C"},
                }
            ],
        }
        template = request(
            "POST", "/templates", {"name": "공정온도", "definition": definition}
        )
        application = request(
            "POST",
            "/applications",
            {
                "version_id": doc["version_id"],
                "template_version_id": template["template_version_id"],
                "bindings": {
                    "main": [sheets[0]["sheet_id"]],
                    "units": [sheets[1]["sheet_id"]],
                },
                "approved": True,
            },
        )
        yield {
            "api": request,
            "work": work,
            "service": service,
            "doc": doc,
            "sheets": sheets,
            "kg": kg,
            "definition": definition,
            "template": template,
            "application": application,
            "path": path,
            "client": client,
        }


def extracted(s):
    job = s["work"](
        "/applications/" + s["application"]["application_id"] + "/extract", {}
    )
    assert job["state"] == "succeeded", job
    series = s["api"]("GET", "/series?run_id=" + job["result"]["run_id"])["items"]
    return job, series


def integration(s, sources=None, row_mode="record_scope"):
    aid = s["application"]["application_id"]
    return s["api"](
        "POST",
        "/integrations",
        {
            "name": "온도 DB",
            "spec": {
                "kg_revision_id": s["kg"]["kg_revision_id"],
                "row_mode": row_mode,
                "fields": [
                    {
                        "field_key": "temp",
                        "output_name": "공정온도",
                        "concept_id": "temperature",
                        "target_type": "decimal",
                        "target_unit": "°C",
                        "sources": sources
                        or [{"application_id": aid, "rule_key": "temperature"}],
                    }
                ],
            },
        },
    )


def test_full_extraction_pagination_and_build(setup):
    s = setup
    job, series = extracted(s)
    sid = series[0]["series_id"]
    values = s["api"]("GET", "/series/" + sid + "/items?limit=2")
    assert values["has_more"] and len(values["items"]) == 2
    assert values["items"][0]["value_text"] == "123456789012345678901234567890.125"
    rest = s["api"](
        "GET", "/series/" + sid + "/items?limit=2&cursor=" + values["next_cursor"]
    )
    assert [i["record_key"] for i in rest["items"]] == ["lot-3"]
    regions = s["api"]("GET", "/items/" + values["items"][0]["item_id"] + "/regions")[
        "items"
    ]
    assert len({r["sheet_id"] for r in regions}) == 2
    assert any(r["role"] == "value" and r["c2"] == 3 for r in regions)
    project = integration(s)
    build = s["work"](
        "/integrations/" + project["integration_version_id"] + "/build", {}
    )
    assert build["state"] == "succeeded", build
    bid = build["result"]["build_id"]
    rows = s["api"]("GET", "/builds/" + bid + "/rows?limit=2")
    assert (
        rows["has_more"]
        and rows["items"][0]["공정온도"] == "123456789012345678901234567890.125"
    )
    result = s["client"].get("/api/v2/builds/" + bid + "/download")
    assert result.status_code == 200 and result.content.startswith(b"SQLite format 3")
    assert result.headers["cache-control"] == "no-store"


def test_alias_ambiguity_and_optimistic_edit(setup):
    s = setup
    aliases = s["api"](
        "GET", "/kg/" + s["kg"]["kg_revision_id"] + "/aliases?text=Temp"
    )["items"]
    assert len(aliases) == 2
    extracted(s)
    aid = s["application"]["application_id"]
    mapping = s["api"]("GET", "/applications/" + aid + "/mappings")["items"][0]
    detail = s["api"]("GET", "/mappings/" + mapping["mapping_revision_id"])
    body = {
        "expected_seq": detail["edit_seq"],
        "effective_spec": detail["effective_spec"],
        "concept_id": "ambient",
        "status": "approved",
    }
    changed = s["api"](
        "POST",
        "/applications/"
        + aid
        + "/mappings/"
        + mapping["mapping_revision_id"]
        + "/revisions",
        body,
    )
    assert changed["revision_no"] == 2
    assert s["api"]("GET", "/applications/" + aid)["published_run_id"] is None
    assert (
        s["api"](
            "POST",
            "/applications/"
            + aid
            + "/mappings/"
            + mapping["mapping_revision_id"]
            + "/revisions",
            body,
            409,
        )["error"]["code"]
        == "EDIT_CONFLICT"
    )


def test_viewport_is_bounded_and_never_persisted(setup):
    s = setup
    body = {
        "version_id": s["doc"]["version_id"],
        "sheet_id": s["sheets"][0]["sheet_id"],
        "r1": 1,
        "c1": 1,
        "rows": 8,
        "cols": 8,
    }
    job = s["work"]("/viewports", body)
    assert job["state"] == "succeeded", job
    view = job["result"]
    assert view["mode"] == "simplified" and len(view["cells"]) <= 64
    assert any(c["range"] == "B2:C2" for c in view["cells"])
    with s["service"].db.connect() as conn:
        persisted = conn.execute(
            "SELECT result_json FROM runtime_job WHERE job_id=?", (job["job_id"],)
        ).fetchone()[0]
        assert "cells" not in persisted and "volatile" in persisted
    s["api"]("POST", "/viewports", {**body, "rows": 101, "request_key": uid()}, 413)


def test_cancel_and_idempotency(setup):
    s = setup
    aid = s["application"]["application_id"]
    body = {"request_key": uid()}
    job = s["api"]("POST", "/applications/" + aid + "/extract", body, 202)
    assert (
        s["api"]("POST", "/applications/" + aid + "/extract", body, 202)["job_id"]
        == job["job_id"]
    )
    s["api"]("POST", "/jobs/" + job["job_id"] + "/cancel")
    s["service"].jobs.run_one()
    assert s["api"]("GET", "/jobs/" + job["job_id"])["state"] == "cancelled"


def test_source_change_and_missing_formula_cache_fail_closed(setup):
    s = setup
    aid = s["application"]["application_id"]
    mapping = s["api"]("GET", "/applications/" + aid + "/mappings")["items"][0]
    detail = s["api"]("GET", "/mappings/" + mapping["mapping_revision_id"])
    definition = detail["effective_spec"]
    definition["selector"]["value"].update(
        areas=[{"sheet_role": "main", "range": "H1"}], cardinality="scalar", axis="none"
    )
    definition["record_spec"]["key"] = "coordinate"
    s["api"](
        "POST",
        "/applications/"
        + aid
        + "/mappings/"
        + mapping["mapping_revision_id"]
        + "/revisions",
        {
            "expected_seq": detail["edit_seq"],
            "effective_spec": definition,
            "concept_id": "temperature",
        },
    )
    job = s["work"]("/applications/" + aid + "/extract", {})
    assert (
        job["state"] == "failed" and job["error_code"] == "FORMULA_CACHE_MISSING"
    ), job
    wb = load_workbook(s["path"])
    wb.active["A10"] = "changed"
    wb.save(s["path"])
    job = s["work"]("/applications/" + aid + "/extract", {})
    assert (
        job["state"] == "failed" and job["error_code"] == "SOURCE_VERSION_CHANGED"
    ), job


def test_same_source_templates_deduplicate_with_all_lineage(setup):
    s = setup
    extracted(s)
    second = s["api"](
        "POST",
        "/applications",
        {
            "version_id": s["doc"]["version_id"],
            "template_version_id": s["template"]["template_version_id"],
            "bindings": {
                "main": [s["sheets"][0]["sheet_id"]],
                "units": [s["sheets"][1]["sheet_id"]],
            },
            "approved": True,
        },
    )
    job = s["work"]("/applications/" + second["application_id"] + "/extract", {})
    assert job["state"] == "succeeded", job
    project = integration(
        s,
        [
            {"application_id": a, "rule_key": "temperature"}
            for a in [s["application"]["application_id"], second["application_id"]]
        ],
    )
    build = s["work"](
        "/integrations/" + project["integration_version_id"] + "/build", {}
    )
    assert build["state"] == "succeeded" and build["result"]["row_count"] == 3, build
    with s["service"].db.connect() as conn:
        assert (
            conn.execute(
                "SELECT count(*) FROM build_lineage WHERE build_id=?",
                (build["result"]["build_id"],),
            ).fetchone()[0]
            == 6
        )
        assert (
            conn.execute(
                "SELECT count(*) FROM build_lineage WHERE build_id=? AND contribution_role='deduplicated'",
                (build["result"]["build_id"],),
            ).fetchone()[0]
            == 3
        )


def test_decimal_text_does_not_round_to_context():
    assert (
        decimal_text("123456789012345678901234567890.1250")
        == "123456789012345678901234567890.125"
    )


def test_protected_reader_native_geometry_cache_and_revocation(tmp_path, monkeypatch):
    monkeypatch.setenv("KG_V2_READER_FACTORY", "tests.v2_reader_fixture:factory")
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    protected = raw / "protected.bin"
    protected.write_bytes(b"DRM\0opaque-test-ciphertext")
    original = protected.read_bytes()
    app = create_app(tmp_path, start_worker=False)
    with TestClient(app) as client:
        service = app.state.v2

        def work(path, body):
            response = client.post(
                "/api/v2" + path, json={**body, "request_key": uid()}
            )
            assert response.status_code == 202, response.text
            job = response.json()
            service.jobs.run_one()
            return job["job_id"], service.jobs.get(job["job_id"], "local-user")

        _, registration = work(
            "/documents/register",
            {"provider": "protected-reader", "source_refs": ["opaque:protected-1"]},
        )
        assert registration["state"] == "succeeded", registration
        vid = registration["result"]["documents"][0]["version_id"]
        sid = client.get("/api/v2/versions/" + vid + "/sheets").json()["items"][0][
            "sheet_id"
        ]
        body = {
            "version_id": vid,
            "sheet_id": sid,
            "r1": 1,
            "c1": 1,
            "rows": 10,
            "cols": 10,
        }
        jid, result = work("/viewports", body)
        assert (
            result["state"] == "succeeded" and result["result"]["mode"] == "native"
        ), result
        with service.db.connect() as conn:
            assert (
                "base64"
                not in conn.execute(
                    "SELECT result_json FROM runtime_job WHERE job_id=?", (jid,)
                ).fetchone()[0]
            )
        (tmp_path / "bad-geometry").touch()
        _, broken = work("/viewports", body)
        assert broken["error_code"] == "INVALID_RENDER_CONTRACT", broken
        (tmp_path / "revoked").touch()
        assert client.get("/api/v2/jobs/" + jid).status_code == 403
        assert protected.read_bytes() == original and list(raw.iterdir()) == [protected]


def test_build_aggregate_preserves_decimal_and_atomic_sources(setup):
    s = setup
    extracted(s)
    project = integration(s, row_mode="aggregate")
    job = s["work"]("/integrations/" + project["integration_version_id"] + "/build", {})
    assert job["state"] == "succeeded", job
    rows = s["api"]("GET", "/builds/" + job["result"]["build_id"] + "/rows")
    assert rows["items"][0]["공정온도"] == "123456789012345678901234567940.625"
    iid = s["api"](
        "GET",
        "/series?run_id="
        + s["api"]("GET", "/applications/" + s["application"]["application_id"])[
            "published_run_id"
        ],
    )["items"][0]["series_id"]
    item = s["api"]("GET", "/series/" + iid + "/items")["items"][0]
    assert (
        s["api"]("GET", "/items/" + item["item_id"])["application_id"]
        == s["application"]["application_id"]
    )


def test_version_locator_and_build_access_recheck(setup, monkeypatch):
    s = setup
    monkeypatch.setenv("KG_V2_READER_FACTORY", "tests.v2_reader_fixture:factory")
    # 새 제공자 등록으로 권한 철회를 지원하는 별도 문서 버전을 만든다.
    job = s["work"](
        "/documents/register",
        {"source_refs": ["sample.xlsx"], "provider": "revocable-xlsx"},
    )
    vid = job["result"]["documents"][0]["version_id"]
    sheets = s["api"]("GET", "/versions/" + vid + "/sheets")["items"]
    a = s["api"](
        "POST",
        "/applications",
        {
            "version_id": vid,
            "template_version_id": s["template"]["template_version_id"],
            "bindings": {
                "main": [sheets[0]["sheet_id"]],
                "units": [sheets[1]["sheet_id"]],
            },
            "approved": True,
        },
    )
    s["application"] = a
    extracted(s)
    project = integration(s)
    build = s["work"](
        "/integrations/" + project["integration_version_id"] + "/build", {}
    )
    assert build["state"] == "succeeded", build
    (s["path"].parents[2] / "revoked").touch()
    response = s["client"].get(
        "/api/v2/builds/" + build["result"]["build_id"] + "/download"
    )
    assert response.status_code == 403, response.text


def test_horizontal_and_sparse_areas_keep_physical_positions(setup):
    s = setup
    definition = copy.deepcopy(s["definition"])
    rule = definition["rules"][0]
    rule["selector"]["value"].update(
        areas=[
            {"sheet_role": "units", "range": "B1:C1"},
            {"sheet_role": "units", "range": "D1"},
        ],
        axis="right",
        element_layout="one_per_column",
    )
    rule["record_spec"]["key"] = "physical_column"
    t = s["api"]("POST", "/templates", {"name": "가로", "definition": definition})
    a = s["api"](
        "POST",
        "/applications",
        {
            "version_id": s["doc"]["version_id"],
            "template_version_id": t["template_version_id"],
            "bindings": {
                "main": [s["sheets"][0]["sheet_id"]],
                "units": [s["sheets"][1]["sheet_id"]],
            },
            "approved": True,
        },
    )
    s["application"] = a
    _, series = extracted(s)
    values = s["api"]("GET", "/series/" + series[0]["series_id"] + "/items")["items"]
    assert [(i["value_text"], i["record_key"]) for i in values] == [
        ("1", "c2"),
        ("2", "c3"),
        ("3", "c4"),
    ]


def test_sparse_blank_stop_and_matrix_anchor_once(setup):
    s = setup
    definition = copy.deepcopy(s["definition"])
    rule = definition["rules"][0]
    rule["record_spec"]["key"] = "physical_row"
    rule["selector"]["value"].update(
        areas=[{"sheet_role": "main", "range": "B4:C9"}],
        stop={"kind": "blank_run", "count": 2, "max_items": 20},
    )
    t = s["api"]("POST", "/templates", {"name": "빈칸 종료", "definition": definition})
    s["application"] = s["api"](
        "POST",
        "/applications",
        {
            "version_id": s["doc"]["version_id"],
            "template_version_id": t["template_version_id"],
            "bindings": {
                "main": [s["sheets"][0]["sheet_id"]],
                "units": [s["sheets"][1]["sheet_id"]],
            },
            "approved": True,
        },
    )
    _, series = extracted(s)
    values = s["api"]("GET", "/series/" + series[0]["series_id"] + "/items")["items"]
    assert [i["value_text"] for i in values] == ["20.5", "30"]
    rule["selector"]["value"].update(
        cardinality="matrix",
        axis="row_major",
        areas=[{"sheet_role": "main", "range": "B3:C5"}],
        stop={"kind": "explicit_areas", "max_items": 10},
    )
    t = s["api"]("POST", "/templates", {"name": "병합 행렬", "definition": definition})
    s["application"] = s["api"](
        "POST",
        "/applications",
        {
            "version_id": s["doc"]["version_id"],
            "template_version_id": t["template_version_id"],
            "bindings": {
                "main": [s["sheets"][0]["sheet_id"]],
                "units": [s["sheets"][1]["sheet_id"]],
            },
            "approved": True,
        },
    )
    _, series = extracted(s)
    values = s["api"]("GET", "/series/" + series[0]["series_id"] + "/items")["items"]
    assert len(values) == 5 and [i["value_state"] for i in values] == [
        "present",
        "present",
        "blank",
        "present",
        "blank",
    ]


def test_max_items_and_wrong_unit_fail_without_publishing(setup):
    s = setup
    aid = s["application"]["application_id"]
    mapping = s["api"]("GET", "/applications/" + aid + "/mappings")["items"][0]
    detail = s["api"]("GET", "/mappings/" + mapping["mapping_revision_id"])
    spec = detail["effective_spec"]
    spec["selector"]["value"]["stop"] = {"kind": "explicit_areas", "max_items": 2}
    changed = s["api"](
        "POST",
        "/applications/"
        + aid
        + "/mappings/"
        + mapping["mapping_revision_id"]
        + "/revisions",
        {
            "expected_seq": detail["edit_seq"],
            "effective_spec": spec,
            "concept_id": "temperature",
        },
    )
    job = s["work"]("/applications/" + aid + "/extract", {})
    assert (
        job["error_code"] == "ITEM_LIMIT"
        and s["api"]("GET", "/applications/" + aid)["published_run_id"] is None
    )
    spec["selector"]["value"]["stop"]["max_items"] = 10
    spec["value_spec"]["unit"] = "°F"
    s["api"](
        "POST",
        "/applications/"
        + aid
        + "/mappings/"
        + changed["mapping_revision_id"]
        + "/revisions",
        {
            "expected_seq": changed["edit_seq"],
            "effective_spec": spec,
            "concept_id": "temperature",
        },
    )
    assert (
        s["work"]("/applications/" + aid + "/extract", {})["error_code"]
        == "UNIT_MISMATCH"
    )


def test_old_version_locator_survives_new_original_registration(setup):
    s = setup
    old = s["doc"]["version_id"]
    wb = load_workbook(s["path"])
    second = s["path"].with_name("second.xlsx")
    wb.active["B5"] = "45"
    wb.save(second)
    doc = s["work"](
        "/documents/register",
        {"source_refs": [second.name], "document_id": s["doc"]["document_id"]},
    )
    assert doc["state"] == "succeeded", doc
    assert s["service"].version(old)["source_ref"] == "sample.xlsx"
    assert (
        s["service"].version(doc["result"]["documents"][0]["version_id"])["source_ref"]
        == "second.xlsx"
    )
    # 과거 버전은 원본이 남아 있으면 해당 위치로 계속 읽힌다.
    extracted(s)
    project = integration(s)
    s["api"](
        "POST",
        "/integrations/" + project["integration_version_id"] + "/build",
        {"request_key": uid()},
        409,
    )


def test_real_lifespan_starts_and_stops_worker(tmp_path):
    app = create_app(tmp_path)
    with TestClient(app) as client:
        assert app.state.v2.jobs.thread.is_alive()
        assert client.get("/api/v2/status").status_code == 200
    assert not app.state.v2.jobs.thread.is_alive()


def test_scope_cursor_rejected_across_series(setup):
    s = setup
    _, series = extracted(s)
    page = s["api"]("GET", "/series/" + series[0]["series_id"] + "/items?limit=1")
    response = s["client"].get(
        "/api/v2/documents", params={"cursor": page["next_cursor"]}
    )
    assert response.status_code == 422


def test_large_text_preview_is_bounded_but_download_is_complete(setup):
    s = setup
    original = "긴문자열" * 2500
    wb = load_workbook(s["path"])
    wb.active["B3"] = original
    wb.save(s["path"])
    doc = s["work"]("/documents/register", {"source_refs": ["sample.xlsx"]})["result"][
        "documents"
    ][0]
    sheets = s["api"]("GET", "/versions/" + doc["version_id"] + "/sheets")["items"]
    definition = copy.deepcopy(s["definition"])
    definition["rules"][0]["value_spec"]["type"] = "text"
    t = s["api"]("POST", "/templates", {"name": "긴 텍스트", "definition": definition})
    s["application"] = s["api"](
        "POST",
        "/applications",
        {
            "version_id": doc["version_id"],
            "template_version_id": t["template_version_id"],
            "bindings": {
                "main": [sheets[0]["sheet_id"]],
                "units": [sheets[1]["sheet_id"]],
            },
            "approved": True,
        },
    )
    _, series = extracted(s)
    item = s["api"]("GET", "/series/" + series[0]["series_id"] + "/items")["items"][0]
    assert len(item["value_text"]) == 2048 and item["preview_truncated"]
    project = s["api"](
        "POST",
        "/integrations",
        {
            "name": "텍스트 DB",
            "spec": {
                "kg_revision_id": s["kg"]["kg_revision_id"],
                "fields": [
                    {
                        "field_key": "value",
                        "output_name": "value",
                        "concept_id": "temperature",
                        "target_type": "text",
                        "target_unit": "°C",
                        "sources": [
                            {
                                "application_id": s["application"]["application_id"],
                                "rule_key": "temperature",
                            }
                        ],
                    }
                ],
            },
        },
    )
    build = s["work"](
        "/integrations/" + project["integration_version_id"] + "/build", {}
    )
    assert build["state"] == "succeeded", build
    bid = build["result"]["build_id"]
    rows = s["api"]("GET", "/builds/" + bid + "/rows")
    assert len(rows["items"][0]["value"]) == 256
    from kg.v2.build import output_path

    with sqlite3.connect(output_path(s["service"], bid)) as conn:
        assert (
            conn.execute("SELECT value FROM data WHERE _row_no=1").fetchone()[0]
            == original
        )


def test_count_deduplicates_and_uses_no_measurement_unit(setup):
    s = setup
    extracted(s)
    project = s["api"](
        "POST",
        "/integrations",
        {
            "name": "측정 개수",
            "spec": {
                "kg_revision_id": s["kg"]["kg_revision_id"],
                "row_mode": "aggregate",
                "fields": [
                    {
                        "field_key": "count",
                        "output_name": "개수",
                        "concept_id": "temperature",
                        "target_type": "decimal",
                        "target_unit": None,
                        "aggregate": "count",
                        "sources": [
                            {
                                "application_id": s["application"]["application_id"],
                                "rule_key": "temperature",
                            }
                        ],
                    }
                ],
            },
        },
    )
    build = s["work"](
        "/integrations/" + project["integration_version_id"] + "/build", {}
    )
    assert build["state"] == "succeeded", build
    rows = s["api"]("GET", "/builds/" + build["result"]["build_id"] + "/rows")
    assert rows["items"][0]["개수"] == "3" and rows["fields"][0]["unit"] is None
