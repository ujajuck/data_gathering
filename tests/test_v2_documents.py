"""GET /documents 목록 요약(템플릿·검수 대기·문서군)과 template/review 정렬 커서 검증."""

import shutil

from fastapi.testclient import TestClient
from openpyxl import Workbook

from kg.v2.api import create_app
from kg.v2.db import uid
from tests.test_v2_runtime import extracted, setup  # noqa: F401


def rows(s, q=""):
    return s["api"]("GET", "/documents" + q)["items"]


def by_name(items):
    return {i["display_name"]: i for i in items}


def head(s, aid=None):
    aid = aid or s["application"]["application_id"]
    mapping = s["api"]("GET", "/applications/" + aid + "/mappings")["items"][0]
    return s["api"]("GET", "/mappings/" + mapping["mapping_revision_id"])


def revise(s, rev, status, concept_id, reason):
    return s["api"](
        "POST",
        f"/applications/{rev['application_id']}/mappings/{rev['mapping_revision_id']}/revisions",
        {
            "expected_seq": rev["edit_seq"],
            "effective_spec": rev["effective_spec"],
            "concept_id": concept_id,
            "status": status,
            "reason": reason,
        },
    )


def propose(s):
    return revise(s, head(s), "proposed", "ambient", "재검수")


def approve(s, rev):
    current = s["api"]("GET", "/mappings/" + rev["mapping_revision_id"])
    return revise(s, current, "approved", "temperature", "검수 완료")


def register(s, name, document_id=None):
    target = s["path"].with_name(name)
    shutil.copyfile(s["path"], target)
    body = {"source_refs": [name]}
    if document_id:
        body["document_id"] = document_id
    job = s["work"]("/documents/register", body)
    assert job["state"] == "succeeded", job
    return job["result"]["documents"][0]


def apply_template(s, doc, template, approved):
    sheets = s["api"]("GET", "/versions/" + doc["version_id"] + "/sheets")["items"]
    return s["api"](
        "POST",
        "/applications",
        {
            "version_id": doc["version_id"],
            "template_version_id": template["template_version_id"],
            "bindings": {
                "main": [sheets[0]["sheet_id"]],
                "units": [sheets[1]["sheet_id"]],
            },
            "approved": approved,
        },
    )


def test_documents_rows_reflect_assignment_publish_and_review_transitions(setup):
    s = setup
    doc_id = s["doc"]["document_id"]
    template_id = s["api"]("GET", "/templates")["items"][0]["template_id"]
    row = rows(s)[0]
    assert row["document_id"] == doc_id
    assert row["templates"] == [
        {
            "application_id": s["application"]["application_id"],
            # scope_key 기본값은 적용 건 ID(service.apply). 같은 템플릿 버전을 범위를 달리해 두 번 적용할 수 있으므로
            # template_count는 적용 건 수이며 배지는 scope_key로 구분한다.
            "scope_key": s["application"]["application_id"],
            "template_id": template_id,
            "template_version_id": s["template"]["template_version_id"],
            "name": "공정온도",
            "revision_no": 1,
            "state": "pending",
            "review_pending": 0,
        }
    ]
    assert (row["template_count"], row["review_pending"]) == (1, 0)
    assert "rn" not in row["templates"][0]
    assert (row["roots"], row["root_count"]) == ([], 0)
    assert row["extraction_status"] == "pending"
    assert "templates_json" not in row and "first_template" not in row

    extracted(s)
    row = rows(s)[0]
    assert row["templates"][0]["state"] == "published"
    assert row["roots"] == [{"concept_id": "temperature", "name": "공정온도"}]
    assert row["root_count"] == 1 and row["extraction_status"] == "published"

    rev = propose(s)
    row = rows(s)[0]
    assert row["review_pending"] == 1
    assert row["templates"][0]["review_pending"] == 1
    # 헤드가 바뀌면 mapping_head_invalidates_* 트리거가 published_run_id를 비운다(schema_sqlite.sql).
    # 따라서 적용 건 state와 extraction_status는 함께 'review'가 되고 문서군도 비워진다.
    assert row["templates"][0]["state"] == "review"
    assert row["extraction_status"] == "review"
    assert (row["roots"], row["root_count"]) == ([], 0)
    assert len(s["api"]("GET", "/review-queue?status=proposed")["items"]) == 1

    approve(s, rev)
    row = rows(s)[0]
    assert row["review_pending"] == 0
    assert row["templates"][0]["state"] == "pending" and row["roots"] == []
    extracted(s)
    row = rows(s)[0]
    assert row["templates"][0]["state"] == "published" and row["root_count"] == 1

    # 새 버전이 현재 버전이 되면 요약은 그 버전 기준으로 비워진다.
    register(s, "next.xlsx", document_id=doc_id)
    row = rows(s)[0]
    assert row["document_id"] == doc_id and row["display_name"] == "next.xlsx"
    assert (row["templates"], row["roots"], row["review_pending"]) == ([], [], 0)
    assert (row["template_count"], row["root_count"]) == (0, 0)
    assert row["extraction_status"] == "unassigned"
    assert rows(s, "?sort=template")[-1]["document_id"] == doc_id


def workspace(tmp_path, concepts, relations, rules=None):
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "세로"
    ws["B2"], ws["D2"] = "공정", "온도"
    ws.merge_cells("B2:C2")
    ws["A3"], ws["A4"], ws["A5"] = "lot-1", "lot-2", "lot-3"
    ws["B3"], ws["B4"], ws["B5"] = "10.5", "20.5", "30"
    wb.create_sheet("단위")["A1"] = "°C"
    path = raw / "sample.xlsx"
    wb.save(path)
    app = create_app(tmp_path, start_worker=False)
    client = TestClient(app)
    client.__enter__()
    service = app.state.v2

    def request(method, url, body=None, code=200):
        response = client.request(method, "/api/v2" + url, json=body)
        assert response.status_code == code, response.text
        return response.json()

    def work(url, body):
        job = request("POST", url, {**body, "request_key": uid()}, 202)
        service.jobs.run_one()
        return request("GET", "/jobs/" + job["job_id"])

    job = work("/documents/register", {"source_refs": ["sample.xlsx"]})
    assert job["state"] == "succeeded", job
    doc = job["result"]["documents"][0]
    kg = request("POST", "/kg/import", {"concepts": concepts, "relations": relations})

    def rule(rule_key, concept_id):
        return {
            "rule_key": rule_key,
            "concept_id": concept_id,
            "record_spec": {"scope": ["process"], "key": {"column": "A"}},
            "selector": {
                "key": {
                    "areas": [
                        {"sheet_role": "main", "range": "B2:C2"},
                        {"sheet_role": "main", "range": "D2"},
                    ]
                },
                "value": {
                    "areas": [{"sheet_role": "main", "range": "B3:C5"}],
                    "cardinality": "list",
                    "axis": "down",
                    "element_layout": "one_per_row",
                },
                "unit": {"areas": [{"sheet_role": "units", "range": "A1"}]},
            },
            "value_spec": {"type": "decimal", "unit": "°C"},
        }

    definition = {
        "kg_revision_id": kg["kg_revision_id"],
        "sheet_roles": {
            "main": {"cardinality": "one"},
            "units": {"cardinality": "one"},
        },
        "rules": [rule(k, c) for k, c in (rules or [("temperature", "temperature")])],
    }
    template = request(
        "POST", "/templates", {"name": "공정온도", "definition": definition}
    )
    s = {
        "api": request,
        "work": work,
        "service": service,
        "doc": doc,
        "kg": kg["kg_revision_id"],
        "template": template,
        "path": path,
        "client": client,
    }
    s["application"] = apply_template(s, doc, template, True)
    return s


def test_documents_roots_follow_coverage_graph_root_rule(tmp_path):
    s = workspace(
        tmp_path,
        [
            {"concept_id": "plant", "name": "공정", "level": 1},
            {"concept_id": "quality", "name": "품질", "level": 1},
            {
                "concept_id": "temperature",
                "name": "공정온도",
                "level": 2,
                "canonical_unit": "°C",
            },
            {"concept_id": "peak", "name": "최고온도", "level": 3},
            {
                "concept_id": "old",
                "name": "폐기온도",
                "level": 2,
                "status": "deprecated",
            },
        ],
        [
            ["plant", "temperature", "parent_of"],
            ["quality", "temperature", "parent_of"],
            ["temperature", "peak", "parent_of"],
            ["plant", "old", "parent_of"],
        ],
    )
    try:
        assert rows(s)[0]["roots"] == []
        extracted(s)
        row = rows(s)[0]
        # 다부모(plant, quality) 중 concept_id가 작은 plant가 화면 부모 = coverage_graph와 동일.
        assert row["roots"] == [{"concept_id": "plant", "name": "공정"}]
        assert row["root_count"] == 1
        groups = {
            g["root_concept_id"]: g["member_document_count"]
            for g in s["api"]("GET", "/kg/" + s["kg"] + "/graph")["groups"]
        }
        assert groups == {"plant": 1, "quality": 0}

        second = register(s, "second.xlsx")
        apply_template(s, second, s["template"], approved=False)
        other = by_name(rows(s))["second.xlsx"]
        # 발행되지 않은 적용 건은 문서군에 들어가지 않는다(승인·추출 전).
        assert (other["roots"], other["root_count"]) == ([], 0)
        assert other["templates"][0]["state"] == "review"
        assert other["review_pending"] == 1 and other["extraction_status"] == "review"
        assert by_name(rows(s))["sample.xlsx"]["root_count"] == 1
    finally:
        s["client"].__exit__(None, None, None)


PLANT_TREE = [
    {"concept_id": "plant", "name": "공정", "level": 1},
    {
        "concept_id": "temperature",
        "name": "공정온도",
        "level": 2,
        "canonical_unit": "°C",
    },
]
PLANT_EDGES = [["plant", "temperature", "parent_of"]]


def graph_groups(s, kg):
    return {
        g["root_concept_id"]: g["name"]
        for g in s["api"]("GET", "/kg/" + kg + "/graph")["groups"]
    }


def test_documents_root_names_come_from_the_pinned_kg_revision(tmp_path):
    s = workspace(tmp_path, PLANT_TREE, PLANT_EDGES)
    try:
        extracted(s)
        assert rows(s)[0]["roots"] == [{"concept_id": "plant", "name": "공정"}]
        # 새 KG 리비전에서 뿌리를 개명·폐기해도 매핑은 rev1에 고정되어 있다.
        rev2 = s["api"](
            "POST",
            "/kg/import",
            {
                "concepts": [
                    {
                        "concept_id": "plant",
                        "name": "공장(폐기)",
                        "level": 1,
                        "status": "deprecated",
                    },
                    PLANT_TREE[1],
                ],
                "relations": PLANT_EDGES,
            },
        )["kg_revision_id"]
        row = rows(s)[0]
        assert row["roots"] == [{"concept_id": "plant", "name": "공정"}]
        assert row["root_count"] == 1
        assert row["roots"][0]["name"] == graph_groups(s, s["kg"])["plant"]
        assert "plant" not in graph_groups(s, rev2)
    finally:
        s["client"].__exit__(None, None, None)


def test_documents_root_spanning_two_pinned_revisions_is_one_row(tmp_path):
    s = workspace(tmp_path, PLANT_TREE, PLANT_EDGES)
    try:
        extracted(s)
        rev2 = s["api"](
            "POST",
            "/kg/import",
            {
                "concepts": [
                    {"concept_id": "plant", "name": "공정설비", "level": 1},
                    PLANT_TREE[1],
                ],
                "relations": PLANT_EDGES,
            },
        )["kg_revision_id"]
        definition = s["api"](
            "GET", "/template-versions/" + s["template"]["template_version_id"]
        )["definition"]
        second = s["api"](
            "POST",
            "/templates",
            {"name": "설비온도", "definition": {**definition, "kg_revision_id": rev2}},
        )
        application = apply_template(s, s["doc"], second, approved=True)
        job = s["work"]("/applications/" + application["application_id"] + "/extract", {})
        assert job["state"] == "succeeded", job
        row = rows(s)[0]
        # 두 발행 적용 건이 서로 다른 KG 리비전에 고정되어도 같은 뿌리는 한 행이며,
        # 이름은 가장 새 고정 리비전(rev2)의 것이다.
        assert row["roots"] == [{"concept_id": "plant", "name": "공정설비"}]
        assert row["root_count"] == 1
        assert row["roots"][0]["name"] == graph_groups(s, rev2)["plant"]
        assert [t["name"] for t in row["templates"]] == ["공정온도", "설비온도"]
        assert row["template_count"] == 2
    finally:
        s["client"].__exit__(None, None, None)


def test_documents_roots_capped_at_eight(tmp_path):
    names = [f"군{i}" for i in range(1, 10)]
    s = workspace(
        tmp_path,
        [
            {"concept_id": f"g{i}", "name": names[i - 1], "level": 1}
            for i in range(1, 10)
        ],
        [],
        rules=[(f"rule{i}", f"g{i}") for i in range(1, 10)],
    )
    try:
        extracted(s)
        row = rows(s)[0]
        assert row["root_count"] == 9 and len(row["roots"]) == 8
        assert [r["name"] for r in row["roots"]] == sorted(names)[:8]
    finally:
        s["client"].__exit__(None, None, None)


def test_documents_sort_by_template_and_review_are_cursor_stable(setup):
    s = setup
    b = register(s, "b.xlsx")
    c = register(s, "c.xlsx")
    other = s["api"](
        "POST", "/templates", {"name": "가나다", "definition": s["definition"]}
    )
    apply_template(s, b, other, approved=False)
    a_id, b_id, c_id = s["doc"]["document_id"], b["document_id"], c["document_id"]

    asc = rows(s, "?sort=template&limit=10")
    assert [r["document_id"] for r in asc] == [b_id, a_id, c_id]
    assert [r["sort_value"] for r in asc] == ["0가나다", "0공정온도", "1"]
    desc = rows(s, "?sort=template&direction=desc&limit=10")
    assert [r["document_id"] for r in desc] == [c_id, a_id, b_id]

    walked, cursor, requests = [], None, 0
    while True:
        page = s["api"](
            "GET",
            "/documents?sort=template&limit=1"
            + (f"&cursor={cursor}" if cursor else ""),
        )
        requests += 1
        walked += page["items"]
        if not page["has_more"]:
            break
        cursor = page["next_cursor"]
    assert requests == 3
    assert [r["document_id"] for r in walked] == [b_id, a_id, c_id]
    assert all(isinstance(r["sort_value"], str) for r in walked)

    first = s["api"]("GET", "/documents?sort=review&direction=desc&limit=1")
    assert first["items"][0]["document_id"] == b_id
    assert first["items"][0]["review_pending"] == 1
    assert isinstance(first["items"][0]["sort_value"], int)
    rest = []
    cursor = first["next_cursor"]
    while cursor:
        page = s["api"](
            "GET", "/documents?sort=review&direction=desc&limit=1&cursor=" + cursor
        )
        rest += page["items"]
        cursor = page["next_cursor"]
    assert [r["review_pending"] for r in rest] == [0, 0]
    assert [r["document_id"] for r in rest] == sorted([a_id, c_id], reverse=True)
    assert [r["document_id"] for r in first["items"] + rest] == [
        r["document_id"] for r in rows(s, "?sort=review&direction=desc&limit=10")
    ]

    template_cursor = s["api"]("GET", "/documents?sort=template&limit=1")["next_cursor"]
    # 커서 범위 해시에 sort/direction이 들어 있어 다른 정렬에서는 거부된다(Problem 기본 422).
    error = s["api"](
        "GET", "/documents?sort=review&limit=1&cursor=" + template_cursor, code=422
    )["error"]
    assert error["code"] == "INVALID_CURSOR"
    error = s["api"]("GET", "/documents?sort=bogus", code=422)["error"]
    assert error["code"] == "VALIDATION_ERROR"

    filtered = rows(s, "?template=가나")
    assert [r["document_id"] for r in filtered] == [b_id]
    assert [t["name"] for t in filtered[0]["templates"]] == ["가나다"]


def test_documents_listing_never_opens_source_with_new_sorts(setup):
    s = setup
    original = s["service"].read

    def no_read(*args, **kwargs):
        raise AssertionError("document listing opened source")

    s["service"].read = no_read
    try:
        for query in (
            "?sort=template",
            "?sort=review&direction=desc",
            "?extraction_status=review",
            "?sort=template&direction=desc&template=공정",
        ):
            s["api"]("GET", "/documents" + query)
    finally:
        s["service"].read = original


def test_documents_sort_by_template_follows_the_template_filter(setup):
    s = setup
    b = register(s, "b.xlsx")
    a_id, b_id = s["doc"]["document_id"], b["document_id"]
    late = s["api"]("POST", "/templates", {"name": "하나", "definition": s["definition"]})
    mid = s["api"]("POST", "/templates", {"name": "나비", "definition": s["definition"]})
    apply_template(s, s["doc"], late, approved=False)
    apply_template(s, b, mid, approved=False)

    asc = rows(s, "?sort=template&limit=10")
    assert [(r["document_id"], r["sort_value"]) for r in asc] == [
        (a_id, "0공정온도"),
        (b_id, "0나비"),
    ]
    # template= 필터가 있으면 정렬 키는 필터에 맞는 첫 템플릿이다(a: 하나, b: 나비). templates[]는 이름순 그대로.
    filtered = rows(s, "?sort=template&template=나&limit=10")
    assert [(r["document_id"], r["sort_value"]) for r in filtered] == [
        (b_id, "0나비"),
        (a_id, "0하나"),
    ]
    assert [t["name"] for t in filtered[1]["templates"]] == ["공정온도", "하나"]
    desc = rows(s, "?sort=template&template=나&direction=desc&limit=10")
    assert [r["document_id"] for r in desc] == [a_id, b_id]

    walked, cursor = [], None
    while True:
        page = s["api"](
            "GET",
            "/documents?sort=template&template=나&limit=1"
            + (f"&cursor={cursor}" if cursor else ""),
        )
        walked += page["items"]
        if not page["has_more"]:
            break
        cursor = page["next_cursor"]
    assert [r["document_id"] for r in walked] == [b_id, a_id]
    # 날짜·상태 필터와 함께 써도 바인딩 순서가 맞는다.
    assert [
        r["document_id"]
        for r in rows(
            s,
            "?sort=template&template=나&date_from=2026-09-08&extraction_status=review&limit=10",
        )
    ] == [b_id, a_id]


def test_documents_listing_plan_is_index_backed_and_page_scoped(setup):
    s = setup
    from kg.v2.features import document_query, document_summary_query

    with s["service"].db.connect() as conn:
        for sort, template in (("name", ""), ("template", "공정"), ("review", "")):
            sql, params = document_query(
                "local-user", "", "", None, None, "", "", template, sort
            )
            plan = [
                r["detail"]
                for r in conn.execute(
                    "EXPLAIN QUERY PLAN "
                    + sql
                    + " ORDER BY sort_value ASC, document_id ASC LIMIT ?",
                    (*params, 31),
                )
            ]
            # 문서 테이블(d)만 훑고 mapping_revision·template_application 등은 인덱스로만 닿는다.
            # 말뭉치 전체 CTE(MATERIALIZE)는 페이지 SQL에 없다.
            assert not [d for d in plan if d.startswith("SCAN") and not d.startswith("SCAN d")], plan
            assert not [d for d in plan if d.startswith("MATERIALIZE")], plan
        sql, params = document_summary_query([s["doc"]["document_id"]])
        plan = [r["detail"] for r in conn.execute("EXPLAIN QUERY PLAN " + sql, params)]
        # 요약 SQL은 페이지 문서에서 출발하며 기본 테이블 전체 스캔이 없다: 페이지 범위 CTE(apps/cov/x/g)와
        # 재귀 CTE 자신(u), 파생 테이블만 훑는다.
        page_scoped = {"SCAN u", "SCAN apps", "SCAN cov", "SCAN x", "SCAN g"}
        scans = [
            d
            for d in plan
            if d.startswith("SCAN")
            and d not in page_scoped
            and not d.startswith("SCAN (subquery")
        ]
        assert not scans, plan
