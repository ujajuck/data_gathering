"""같은 양식 문서군 제안과 레시피(매핑) 이식: 서명 캐시, 점수·임계값, 검수 대기 이식, 시트 바인딩."""

import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from kg.v2.api import create_app
from kg.v2.db import uid
from kg.v2.spec import decimal


def process_workbook(path, second_sheet="공통 정보", base=100):
    wb = Workbook()
    ws = wb.active
    ws.title = "공정 기록"
    ws["A1"] = "공정 운전 기록"
    ws.merge_cells("A1:F1")
    ws["A3"], ws["B3"], ws["D3"], ws["E3"] = "배치", "공정", "온도", "시간"
    ws.merge_cells("B3:C3")
    for n, row in enumerate(range(4, 21)):
        ws.cell(row, 1, f"LOT-{n + 1:03d}")
        ws.cell(row, 2, base + n)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.cell(row, 4, base * 2 + n)
        ws.cell(row, 5, n)
    common = wb.create_sheet(second_sheet)
    common["A1"], common["B1"] = "온도 단위", "°C"
    wb.save(path)


def roster_workbook(path):
    """이름·자유 텍스트·식별자 열이 있는 명단. 라벨만 서명에 남고 레코드 값은 남지 않아야 한다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "명단"
    ws["A1"] = "직원 명단"
    ws.merge_cells("A1:D1")
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "사번", "이름", "점수", "비고"
    for n in range(3, 20):
        ws.cell(n, 1, f"EMP-{n:03d}")
        ws.cell(n, 2, f"홍길동{n}")
        ws.cell(n, 3, n * 3)
        ws.cell(n, 4, f"자유 메모 {n}번째")
    summary = wb.create_sheet("요약")
    summary["A1"], summary["B1"], summary["C1"] = "항목", "1월", "2월"
    for n, label in enumerate(("매출", "원가", "이익"), start=2):
        summary.cell(n, 1, label)
        summary.cell(n, 2, n)
        summary.cell(n, 3, n * 2)
    wb.save(path)


def sales_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "매출"
    ws["A1"], ws["B1"], ws["C1"] = "지역", "월", "매출액"
    for row in range(2, 15):
        ws.cell(row, 1, row * 10)
        ws.cell(row, 2, row)
        ws.cell(row, 3, row * 1000)
    wb.create_sheet("메모")["A1"] = "비고"
    wb.save(path)


@pytest.fixture(scope="module")
def workspace(tmp_path_factory):
    root = tmp_path_factory.mktemp("suggest")
    raw = root / "data/raw"
    raw.mkdir(parents=True)
    process_workbook(raw / "report_a.xlsx", base=100)
    process_workbook(raw / "report_b.xlsx", base=500)
    process_workbook(raw / "report_d.xlsx", second_sheet="단위", base=900)
    sales_workbook(raw / "report_c.xlsx")
    app = create_app(root, start_worker=False)
    with TestClient(app) as client:
        service = app.state.v2

        def request(method, url, body=None, code=200):
            response = client.request(method, "/api/v2" + url, json=body)
            assert response.status_code == code, response.text
            return response.json()

        def work(url, body):
            job = request("POST", url, {**body, "request_key": uid()}, 202)
            service.jobs.run_one()
            return request("GET", "/jobs/" + job["job_id"])

        def register(name):
            job = work("/documents/register", {"source_refs": [name]})
            assert job["state"] == "succeeded", job
            doc = job["result"]["documents"][0]
            doc["sheets"] = request("GET", f"/versions/{doc['version_id']}/sheets")[
                "items"
            ]
            doc["by_name"] = {s["name"]: s["sheet_id"] for s in doc["sheets"]}
            return doc

        docs = {k: register(f"report_{k}.xlsx") for k in ("a", "b", "c", "d")}
        kg = request(
            "POST",
            "/kg/import",
            {
                "concepts": [
                    {"concept_id": "temperature", "name": "공정온도", "level": 1},
                    {"concept_id": "ambient", "name": "주위온도", "level": 1},
                ]
            },
        )
        rule = {
            "rule_key": "temperature",
            "concept_id": "temperature",
            "record_spec": {"scope": ["process"], "key": {"column": "A"}},
            "selector": {
                "key": {
                    "areas": [
                        {"sheet_role": "main", "range": "B3:C3"},
                        {"sheet_role": "main", "range": "D3"},
                    ]
                },
                "value": {
                    "areas": [{"sheet_role": "main", "range": "B4:C20"}],
                    "cardinality": "list",
                    "axis": "down",
                    "element_layout": "one_per_row",
                },
                "unit": {"areas": [{"sheet_role": "common", "range": "B1"}]},
            },
            "value_spec": {"type": "decimal", "unit": "°C"},
        }
        template = request(
            "POST",
            "/templates",
            {
                "name": "공정 운전 기록",
                "definition": {
                    "kg_revision_id": kg["kg_revision_id"],
                    "sheet_roles": {
                        "main": {"cardinality": "one"},
                        "common": {"cardinality": "one"},
                    },
                    "rules": [rule],
                },
            },
        )
        app_a = request(
            "POST",
            "/applications",
            {
                "version_id": docs["a"]["version_id"],
                "template_version_id": template["template_version_id"],
                "bindings": {
                    "main": [docs["a"]["by_name"]["공정 기록"]],
                    "common": [docs["a"]["by_name"]["공통 정보"]],
                },
                "approved": True,
            },
        )["application_id"]
        first = request("GET", f"/applications/{app_a}/mappings")["items"][0]
        edited = json.loads(json.dumps(rule))
        edited["selector"]["value"]["areas"][0]["range"] = "B4:C18"
        head_a = request(
            "POST",
            f"/applications/{app_a}/mappings/{first['mapping_revision_id']}/revisions",
            {
                "expected_seq": first["edit_seq"],
                "effective_spec": edited,
                "concept_id": "temperature",
                "status": "approved",
                "reason": "값 영역 축소",
            },
        )
        app_c = request(
            "POST",
            "/applications",
            {
                "version_id": docs["c"]["version_id"],
                "template_version_id": template["template_version_id"],
                "bindings": {
                    "main": [docs["c"]["by_name"]["매출"]],
                    "common": [docs["c"]["by_name"]["메모"]],
                },
            },
        )["application_id"]
        yield {
            "api": request,
            "work": work,
            "register": register,
            "service": service,
            "client": client,
            "root": root,
            "docs": docs,
            "template": template,
            "app_a": app_a,
            "app_c": app_c,
            "head_a": head_a,
        }


def signature_of(s, version_id):
    with s["service"].db.connect() as conn:
        row = conn.execute(
            "SELECT * FROM version_signature WHERE document_version_id=?",
            (version_id,),
        ).fetchone()
    return dict(row) if row else None


def count(s, sql, params=()):
    with s["service"].db.connect() as conn:
        return conn.execute(sql, params).fetchone()[0]


def test_signature_cached_on_registration_and_listing_never_opens_sources(workspace):
    s = workspace
    for doc in s["docs"].values():
        row = signature_of(s, doc["version_id"])
        assert row and row["algorithm"] == "structure-v2"
        assert doc["signature"] == "ready"
    sheet = json.loads(signature_of(s, s["docs"]["b"]["version_id"])["signature_json"])[
        "sheets"
    ][0]
    assert sheet["name"] == "공정 기록" and sheet["name_norm"] == "공정 기록"
    # 헤더 성분은 라벨 셀만 담는다. 레코드 키(LOT-001 …)와 값은 서명에 남지 않는다.
    assert set(sheet["headers"]) == {"공정", "온도", "배치", "시간", "공정 운전 기록"}
    assert "lot-001" not in sheet["headers"]
    assert sheet["headers"] == sorted(sheet["headers"])
    for term in sheet["headers"]:
        with pytest.raises(Exception):
            decimal(term)
    assert {"A1:F1", "B3:C3", "B4:C4"} <= set(sheet["merges"])
    assert sheet["dims"]["rows"] == 20 and sheet["dims"]["cols"] == 6
    # 같은 양식에 데이터만 다른 두 문서(base 100/500)는 서명이 완전히 같다 = 데이터 값이 서명에 없다.
    assert (
        signature_of(s, s["docs"]["a"]["version_id"])["signature_sha256"]
        == signature_of(s, s["docs"]["b"]["version_id"])["signature_sha256"]
    )
    # 제안 목록은 대상 버전 권한만 확인하고, 메타데이터와 캐시된 서명만 사용하며 원본을 열지 않는다.
    original = s["service"].read
    ops = []

    def guarded(provider, principal, operation, payload, checkpoint=lambda: None):
        ops.append((operation, payload.get("source_ref")))
        if operation != "authorize":
            raise AssertionError("suggestion listing opened source")
        return original(provider, principal, operation, payload, checkpoint)

    s["service"].read = guarded
    try:
        result = s["api"]("GET", f"/versions/{s['docs']['b']['version_id']}/suggestions")
    finally:
        s["service"].read = original
    assert result["signature_status"] == "ready" and result["items"]
    assert ops == [("authorize", "report_b.xlsx")]


def test_suggestions_rank_by_score_and_apply_threshold(workspace):
    s = workspace
    vid = s["docs"]["b"]["version_id"]
    result = s["api"]("GET", f"/versions/{vid}/suggestions")
    assert result["signature_status"] == "ready"
    assert result["candidates"] == 2 and result["unsigned_candidates"] == 0
    assert [i["source_application_id"] for i in result["items"]] == [s["app_a"]]
    top = result["items"][0]
    assert top["score"] >= 0.95
    assert top["breakdown"]["sheet_names"]["score"] == 1
    assert top["breakdown"]["headers"]["score"] >= 0.9
    assert top["breakdown"]["merges"]["score"] >= 0.9
    assert top["source_document_name"] == "report_a.xlsx"
    assert top["template_name"] == "공정 운전 기록"
    assert top["template_version_id"] == s["template"]["template_version_id"]
    assert top["approved_rules"] == 1 and top["total_rules"] == 1
    assert top["unmatched_roles"] == []
    # 헤더·병합 성분은 개수만 설명하고 다른 문서의 셀 문자열은 응답에 싣지 않는다.
    headers = top["breakdown"]["headers"]
    assert headers["shared_count"] == headers["target_count"] == headers["source_count"] > 0
    assert "shared" not in headers and "shared" not in top["breakdown"]["merges"]
    assert top["breakdown"]["merges"]["shared_count"] > 0
    assert top["breakdown"]["sheet_names"]["shared"] == ["공정 기록", "공통 정보"]
    assert {m["role"]: m["target_sheet"] for m in top["matched_sheets"]} == {
        "main": "공정 기록",
        "common": "공통 정보",
    }
    assert all(m["target_sheet_id"] for m in top["matched_sheets"])
    everything = s["api"]("GET", f"/versions/{vid}/suggestions?threshold=0")
    assert [i["source_application_id"] for i in everything["items"]] == [
        s["app_a"],
        s["app_c"],
    ]
    assert everything["items"][1]["score"] < 0.5
    assert everything["items"][1]["unmatched_roles"] == ["common", "main"]
    assert everything["has_more"] is False and everything["next_cursor"] is None
    limited = s["api"]("GET", f"/versions/{vid}/suggestions?threshold=0&limit=1")
    assert len(limited["items"]) == 1 and limited["candidates"] == 2
    assert limited["has_more"] is True
    assert s["api"]("GET", f"/versions/{vid}/suggestions?limit=1")["has_more"] is False
    s["api"]("GET", "/versions/missing/suggestions", code=404)
    error = s["api"]("GET", f"/versions/{vid}/suggestions?threshold=2", code=422)
    assert error["error"]["code"] == "VALIDATION_ERROR"


def test_from_suggestion_creates_proposed_candidate_revisions_without_heads(workspace):
    s = workspace
    vid = s["docs"]["b"]["version_id"]
    before = count(s, "SELECT count(*) FROM mapping_revision WHERE application_id=?", (s["app_a"],))
    created = s["api"](
        "POST",
        f"/versions/{vid}/applications/from-suggestion",
        {"source_application_id": s["app_a"]},
    )
    new = created["application_id"]
    assert created["status"] == "proposed" and created["rules"] == 1
    assert created["template_version_id"] == s["template"]["template_version_id"]
    assert created["bindings"] == {
        "main": [s["docs"]["b"]["by_name"]["공정 기록"]],
        "common": [s["docs"]["b"]["by_name"]["공통 정보"]],
    }
    mappings = s["api"]("GET", f"/applications/{new}/mappings")["items"]
    assert len(mappings) == 1
    assert all(
        m["status"] == "proposed" and m["edit_seq"] == 0 and m["revision_no"] == 1
        for m in mappings
    )
    detail = s["api"]("GET", "/mappings/" + mappings[0]["mapping_revision_id"])
    assert detail["origin"] == "candidate" and detail["status"] == "proposed"
    assert detail["effective_spec"]["selector"]["value"]["areas"][0]["range"] == "B4:C18"
    assert detail["concept_id"] == "temperature"
    origin = detail["evidence"]["transplanted_from"]
    assert origin["mapping_revision_id"] == s["head_a"]["mapping_revision_id"]
    assert origin["was_head"] is True and origin["application_id"] == s["app_a"]
    assert detail["evidence"]["similarity"]["score"] >= 0.95
    assert detail["evidence"]["concept_dropped"] is False
    assert s["app_a"] in detail["reason"]
    assert count(s, "SELECT count(*) FROM mapping_head WHERE application_id=?", (new,)) == 0
    assert (
        count(s, "SELECT count(*) FROM mapping_revision WHERE application_id=?", (s["app_a"],))
        == before
    )
    # 자동 승인되지 않았으므로 추출은 검수를 요구한다.
    error = s["api"](
        "POST", f"/applications/{new}/extract", {"request_key": uid()}, code=409
    )
    assert error["error"]["code"] == "REVIEW_REQUIRED"
    approved = s["api"](
        "POST",
        f"/applications/{new}/mappings/{mappings[0]['mapping_revision_id']}/revisions",
        {
            "expected_seq": 0,
            "effective_spec": detail["effective_spec"],
            "concept_id": "temperature",
            "status": "approved",
        },
    )
    assert approved["edit_seq"] == 2 and approved["revision_no"] == 2
    job = s["work"](f"/applications/{new}/extract", {})
    assert job["state"] == "succeeded", job


def test_from_suggestion_reports_unmatched_roles_and_accepts_override(workspace):
    s = workspace
    vid = s["docs"]["d"]["version_id"]
    url = f"/versions/{vid}/applications/from-suggestion"
    error = s["api"]("POST", url, {"source_application_id": s["app_a"]}, code=409)
    assert error["error"]["code"] == "SHEET_UNMATCHED"
    assert "common" in error["error"]["message"] and "공통 정보" in error["error"]["message"]
    assert (
        count(s, "SELECT count(*) FROM template_application WHERE document_version_id=?", (vid,))
        == 0
    )
    unit = s["docs"]["d"]["by_name"]["단위"]
    created = s["api"](
        "POST",
        url,
        {"source_application_id": s["app_a"], "sheet_bindings": {"common": unit}},
    )
    assert created["bindings"] == {
        "main": [s["docs"]["d"]["by_name"]["공정 기록"]],
        "common": [unit],
    }
    bad = s["api"](
        "POST",
        url,
        {"source_application_id": s["app_a"], "sheet_bindings": {"common": "nope"}},
        code=422,
    )
    assert bad["error"]["code"] == "INVALID_SHEET_BINDING"
    dup = s["api"](
        "POST",
        url,
        {"source_application_id": s["app_a"], "sheet_bindings": {"common": [unit, unit]}},
        code=422,
    )
    assert dup["error"]["code"] == "INVALID_SHEET_BINDING"
    named = s["api"](
        "POST",
        url,
        {
            "source_application_id": s["app_a"],
            "sheet_bindings": {"common": [unit]},
            "name": "재검수",
        },
    )
    with s["service"].db.connect() as conn:
        assert (
            conn.execute(
                "SELECT scope_key FROM template_application WHERE application_id=?",
                (named["application_id"],),
            ).fetchone()[0]
            == "재검수"
        )
    conflict = s["api"](
        "POST",
        url,
        {
            "source_application_id": s["app_a"],
            "sheet_bindings": {"common": [unit]},
            "name": "재검수",
        },
        code=409,
    )
    assert conflict["error"]["code"] == "INTEGRITY_CONFLICT"
    same = s["api"](
        "POST",
        f"/versions/{s['docs']['a']['version_id']}/applications/from-suggestion",
        {"source_application_id": s["app_a"]},
        code=409,
    )
    assert same["error"]["code"] == "SAME_VERSION"
    s["api"]("POST", url, {"source_application_id": s["app_a"], "extra": 1}, code=422)
    s["api"]("POST", url, {"source_application_id": "missing"}, code=404)


def test_signature_endpoint_backfills_missing_signature(workspace):
    s = workspace
    vid = s["docs"]["b"]["version_id"]
    with s["service"].db.connect(write=True) as conn:
        conn.execute("DELETE FROM version_signature WHERE document_version_id=?", (vid,))
    missing = s["api"]("GET", f"/versions/{vid}/suggestions")
    assert missing["signature_status"] == "missing" and missing["items"] == []
    first = s["api"]("POST", f"/versions/{vid}/signature", {})
    assert first["status"] == "ready" and first["cached"] is False
    assert first["sheets"] == 2 and first["signature_sha256"]
    again = s["api"]("POST", f"/versions/{vid}/signature", {})
    assert again["cached"] is True
    restored = s["api"]("GET", f"/versions/{vid}/suggestions")
    assert restored["signature_status"] == "ready" and restored["items"]
    # 원본이 사라진 버전은 서명을 계산하지 못하지만 등록 자체는 유지된다.
    raw = s["root"] / "data/raw"
    process_workbook(raw / "report_e.xlsx", base=700)
    gone = s["register"]("report_e.xlsx")
    with s["service"].db.connect(write=True) as conn:
        conn.execute(
            "DELETE FROM version_signature WHERE document_version_id=?",
            (gone["version_id"],),
        )
    (raw / "report_e.xlsx").unlink()
    error = s["api"]("POST", f"/versions/{gone['version_id']}/signature", {}, code=404)
    assert error["error"]["code"] == "SOURCE_NOT_FOUND"
    assert s["api"]("GET", f"/versions/{gone['version_id']}/sheets")["items"]


def test_protected_reader_uses_viewport_fallback_for_signature(tmp_path, monkeypatch):
    monkeypatch.setenv("KG_V2_READER_FACTORY", "tests.v2_reader_fixture:factory")
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    protected = raw / "protected.bin"
    protected.write_bytes(b"DRM\0opaque-test-ciphertext")
    app = create_app(tmp_path, start_worker=False)
    with TestClient(app) as client:
        service = app.state.v2
        response = client.post(
            "/api/v2/documents/register",
            json={
                "provider": "protected-reader",
                "source_refs": ["opaque:protected-1"],
                "request_key": uid(),
            },
        )
        assert response.status_code == 202, response.text
        service.jobs.run_one()
        job = service.jobs.get(response.json()["job_id"], "local-user")
        assert job["state"] == "succeeded", job
        doc = job["result"]["documents"][0]
        assert doc["signature"] == "ready"
        with service.db.connect() as conn:
            row = conn.execute(
                "SELECT signature_json FROM version_signature WHERE document_version_id=?",
                (doc["version_id"],),
            ).fetchone()
        signature = json.loads(row["signature_json"])
        assert signature["sheets"][0]["name"] == "Protected"
        assert signature["sheets"][0]["headers"] == []
        assert signature["sheets"][0]["merges"] == []
        assert list(raw.iterdir()) == [protected]
        result = client.get(f"/api/v2/versions/{doc['version_id']}/suggestions").json()
        assert result["items"] == [] and result["signature_status"] == "ready"
        # 제안 목록도 대상 버전의 원본 권한을 확인한다.
        (tmp_path / "revoked").touch()
        denied = client.get(f"/api/v2/versions/{doc['version_id']}/suggestions")
        assert denied.status_code == 403, denied.text
        assert denied.json()["error"]["code"] == "ACCESS_DENIED"


def test_cli_parse_keeps_serve_default_and_adds_sign(workspace):
    from kg.v2.__main__ import parse_args
    from kg.v2.suggest import sign_versions

    serve = parse_args(["--ws", "x", "--port", "1"])
    assert serve.command == "serve" and serve.port == 1 and str(serve.ws) == "x"
    assert parse_args([]).command == "serve"
    sign = parse_args(["sign", "--ws", "x", "--version", "v"])
    assert sign.command == "sign" and sign.version == "v"
    s = workspace
    vid = s["docs"]["b"]["version_id"]
    with s["service"].db.connect(write=True) as conn:
        conn.execute("DELETE FROM version_signature WHERE document_version_id=?", (vid,))
    lines = list(sign_versions(s["service"], None, "local-user"))
    assert f"{vid} ready" in lines
    assert signature_of(s, vid)
    assert list(sign_versions(s["service"], vid, "local-user")) == [f"{vid} ready"]
    assert list(sign_versions(s["service"], "missing", "local-user")) == [
        "missing failed NOT_FOUND"
    ]
    # 템플릿을 연결한 이전 버전은 현재 버전이 아니어도 제안 후보이므로 함께 서명한다.
    raw = s["root"] / "data/raw"
    process_workbook(raw / "report_f.xlsx", base=300)
    first = s["register"]("report_f.xlsx")
    s["api"](
        "POST",
        "/applications",
        {
            "version_id": first["version_id"],
            "template_version_id": s["template"]["template_version_id"],
            "bindings": {
                "main": [first["by_name"]["공정 기록"]],
                "common": [first["by_name"]["공통 정보"]],
            },
        },
    )
    # 같은 문서의 새 버전(위치 변경). 이전 버전 원본은 그대로 남아 있어 서명을 다시 계산할 수 있다.
    process_workbook(raw / "report_f_moved.xlsx", base=300)
    job = s["work"](
        "/documents/register",
        {"source_refs": ["report_f_moved.xlsx"], "document_id": first["document_id"]},
    )
    assert job["state"] == "succeeded", job
    second = job["result"]["documents"][0]
    assert second["document_id"] == first["document_id"]
    assert second["version_id"] != first["version_id"]
    with s["service"].db.connect(write=True) as conn:
        conn.execute(
            "DELETE FROM version_signature WHERE document_version_id IN (?,?)",
            (first["version_id"], second["version_id"]),
        )
        # 이전 알고리즘으로 남은 캐시도 다시 계산한다.
        conn.execute(
            "INSERT INTO version_signature(document_version_id,algorithm,signature_json,signature_sha256,reader_revision,computed_at) VALUES (?,?,?,?,?,?)",
            (second["version_id"], "structure-v1", "{}", "x", "old", "2020-01-01T00:00:00+00:00"),
        )
    lines = list(sign_versions(s["service"], None, "local-user"))
    assert f"{first['version_id']} ready" in lines
    assert f"{second['version_id']} ready" in lines
    assert signature_of(s, first["version_id"])["algorithm"] == "structure-v2"
    assert signature_of(s, second["version_id"])["algorithm"] == "structure-v2"
    # 다시 실행하면 두 버전은 이미 서명이 있어 대상이 아니다(원본이 사라진 report_e만 실패로 남는다).
    again = list(sign_versions(s["service"], None, "local-user"))
    assert not any(first["version_id"] in line or second["version_id"] in line for line in again)
    assert all(line.endswith("failed SOURCE_NOT_FOUND") for line in again)


def test_openapi_from_suggestion_contract(workspace):
    schema = workspace["client"].get("/openapi.json").json()["components"]["schemas"]
    model = schema["FromSuggestionRequest"]
    assert model["required"] == ["source_application_id"]
    assert {"sheet_bindings", "name"} <= set(model["properties"])


def test_signature_keeps_labels_but_not_record_values(workspace):
    s = workspace
    roster_workbook(s["root"] / "data/raw/roster.xlsx")
    ops = []
    original = s["service"].read

    def counting(provider, principal, operation, payload, checkpoint=lambda: None):
        ops.append(operation)
        return original(provider, principal, operation, payload, checkpoint)

    s["service"].read = counting
    try:
        doc = s["register"]("roster.xlsx")
    finally:
        s["service"].read = original
    assert doc["signature"] == "ready"
    # 등록은 describe 한 번으로 서명까지 끝낸다(별도 authorize/signature Reader 호출 없음).
    assert ops == ["describe"]
    sheets = json.loads(signature_of(s, doc["version_id"])["signature_json"])["sheets"]
    roster, summary = sheets[0]["headers"], sheets[1]["headers"]
    assert set(roster) == {"직원 명단", "사번", "이름", "점수", "비고"}
    assert not any(t.startswith(("emp-", "홍길동", "자유 메모")) for t in roster)
    assert set(summary) == {"항목", "1월", "2월", "매출", "원가", "이익"}
    # 서명 응답 자체에도 레코드 값이 없다.
    dumped = json.dumps(sheets, ensure_ascii=False)
    assert "emp-" not in dumped and "홍길동" not in dumped and "메모" not in dumped
    # 같은 원본을 다시 등록하면 describe 한 번으로 캐시된 서명을 재사용한다.
    ops.clear()
    s["service"].read = counting
    try:
        again = s["register"]("roster.xlsx")
    finally:
        s["service"].read = original
    assert again["version_id"] == doc["version_id"] and again["signature"] == "ready"
    assert ops == ["describe"]


def test_signature_terms_rules_are_explainable():
    from kg.v2.readers import header_term, signature_terms

    assert header_term("lot-001") is None and header_term("No.12") is None
    assert header_term("2024-01-02") is None and header_term("2024년") is None
    assert header_term(" 온도 (°C) ") == "온도 (°c)" and header_term("1분기") == "1분기"
    assert header_term(12) is None and header_term("12.5") is None
    grid = [["제목"], [None], ["이름", "점수"]] + [
        [f"학생{n}", 80 + n] for n in range(12)
    ]
    # 블록 머리·라벨 행은 남고, 열 아래로 이어지는 값 옆 문자열(학생 이름)은 레코드로 제외한다.
    assert signature_terms(grid) == ["이름", "점수", "제목"]
    attributes = [["항목", "값"], ["매출", 1], ["원가", 2], ["이익", 3]]
    assert signature_terms(attributes) == ["값", "매출", "원가", "이익", "항목"]
    text_table = [["이름", "부서"]] + [[f"사람{n}", "부서" + str(n)] for n in range(5)]
    assert signature_terms(text_table) == ["부서", "이름"]
    assert signature_terms([]) == []
    merged = [[None, None], [None, "구분"], [None, 5]]
    assert signature_terms(merged, anchors={(2, 2)}) == ["구분"]


def test_register_survives_malformed_or_crashing_signature(tmp_path, monkeypatch):
    monkeypatch.setenv("KG_V2_READER_FACTORY", "tests.v2_reader_fixture:factory")
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    (raw / "protected.bin").write_bytes(b"DRM\0opaque-test-ciphertext")
    app = create_app(tmp_path, start_worker=False)
    with TestClient(app) as client:
        service = app.state.v2

        def register(provider, ref):
            response = client.post(
                "/api/v2/documents/register",
                json={"provider": provider, "source_refs": [ref], "request_key": uid()},
            )
            assert response.status_code == 202, response.text
            service.jobs.run_one()
            job = service.jobs.get(response.json()["job_id"], "local-user")
            assert job["state"] == "succeeded", job
            return job["result"]["documents"][0]

        # describe가 형식이 잘못된 서명을 돌려주는 제공자: 등록은 성공하고 서명만 실패로 남는다.
        doc = register("malformed-signature", "opaque:malformed-1")
        assert doc["signature"] == "failed:INVALID_READER_CONTRACT"
        assert client.get(f"/api/v2/versions/{doc['version_id']}/sheets").json()["items"]
        with service.db.connect() as conn:
            assert (
                conn.execute(
                    "SELECT count(*) FROM version_signature WHERE document_version_id=?",
                    (doc["version_id"],),
                ).fetchone()[0]
                == 0
            )
        listing = client.get(f"/api/v2/versions/{doc['version_id']}/suggestions").json()
        assert listing["signature_status"] == "missing"
        # 서명 계산 중의 예상하지 못한 예외도 이미 커밋된 등록을 실패시키지 않는다.
        import kg.v2.suggest as suggest

        def boom(*args, **kwargs):
            raise TypeError("reader returned garbage")

        monkeypatch.setattr(suggest, "store_signature", boom)
        crashed = register("protected-reader", "opaque:protected-2")
        assert crashed["signature"] == "failed:INTERNAL"
        assert client.get(f"/api/v2/versions/{crashed['version_id']}/sheets").json()["items"]


def test_from_suggestion_skips_rejected_revisions(workspace):
    s = workspace
    unit = s["docs"]["d"]["by_name"]["단위"]
    source = s["api"](
        "POST",
        "/applications",
        {
            "version_id": s["docs"]["d"]["version_id"],
            "template_version_id": s["template"]["template_version_id"],
            "bindings": {
                "main": [s["docs"]["d"]["by_name"]["공정 기록"]],
                "common": [unit],
            },
            "scope_key": "reject-check",
        },
    )["application_id"]
    first = s["api"]("GET", f"/applications/{source}/mappings")["items"][0]
    rejected = s["api"](
        "POST",
        f"/applications/{source}/mappings/{first['mapping_revision_id']}/revisions",
        {
            "expected_seq": 0,
            "effective_spec": s["api"]("GET", "/mappings/" + first["mapping_revision_id"])[
                "effective_spec"
            ],
            "concept_id": "temperature",
            "status": "rejected",
            "reason": "검수 거절",
        },
    )
    assert rejected["status"] == "rejected" and rejected["revision_no"] == 2
    target = s["docs"]["b"]["version_id"]
    url = f"/versions/{target}/applications/from-suggestion"
    created = s["api"](
        "POST",
        url,
        {
            "source_application_id": source,
            "sheet_bindings": {"common": [s["docs"]["b"]["by_name"]["공통 정보"]]},
            "name": "거절 건너뜀",
        },
    )
    mapping = s["api"]("GET", f"/applications/{created['application_id']}/mappings")["items"][0]
    origin = s["api"]("GET", "/mappings/" + mapping["mapping_revision_id"])["evidence"][
        "transplanted_from"
    ]
    # head가 없으면 거절되지 않은 최신 리비전(1, proposed)을 쓴다. 거절된 리비전 2는 후보가 되지 않는다.
    assert origin["mapping_revision_id"] == first["mapping_revision_id"]
    assert origin["revision_no"] == 1 and origin["status"] == "proposed"
    assert origin["was_head"] is False
    # 거절된 리비전만 있는 규칙은 이식할 수 없다.
    with s["service"].db.connect(write=True) as conn:
        only_rejected = uid()
        conn.execute(
            "INSERT INTO template_application(application_id,document_version_id,template_version_id,scope_key,created_at) VALUES (?,?,?,?,?)",
            (only_rejected, s["docs"]["d"]["version_id"], s["template"]["template_version_id"], "거절만", "2026-01-01T00:00:00+00:00"),
        )
        for role, sid in (("main", s["docs"]["d"]["by_name"]["공정 기록"]), ("common", unit)):
            conn.execute(
                "INSERT INTO application_sheet(application_id,document_version_id,role_key,ordinal,sheet_id) VALUES (?,?,?,?,?)",
                (only_rejected, s["docs"]["d"]["version_id"], role, 0, sid),
            )
        row = conn.execute(
            "SELECT * FROM mapping_revision WHERE mapping_revision_id=?",
            (rejected["mapping_revision_id"],),
        ).fetchone()
        conn.execute(
            "INSERT INTO mapping_revision(mapping_revision_id,application_id,document_version_id,template_version_id,rule_key,revision_no,kg_revision_id,concept_id,origin,status,effective_spec_json,evidence_json,created_by,reason,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                uid(),
                only_rejected,
                s["docs"]["d"]["version_id"],
                s["template"]["template_version_id"],
                row["rule_key"],
                1,
                row["kg_revision_id"],
                row["concept_id"],
                "manual",
                "rejected",
                row["effective_spec_json"],
                row["evidence_json"],
                "local-user",
                "거절",
                "2026-01-01T00:00:00+00:00",
            ),
        )
    error = s["api"](
        "POST",
        url,
        {
            "source_application_id": only_rejected,
            "sheet_bindings": {"common": [s["docs"]["b"]["by_name"]["공통 정보"]]},
        },
        code=409,
    )
    assert error["error"]["code"] == "SOURCE_REJECTED"
    assert (
        count(
            s,
            "SELECT count(*) FROM template_application WHERE document_version_id=? AND scope_key=?",
            (target, "거절만"),
        )
        == 0
    )
