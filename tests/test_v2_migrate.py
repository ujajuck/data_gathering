"""v1 kg.db → v2 워크스페이스 이관 CLI 검증: 새 ID·링크, proposed 매핑, 멱등, dry-run."""

import json
import sqlite3
from hashlib import sha256
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from kg.parsing import (
    add_version,
    assign,
    create_template,
    effective_mappings,
    run_parse,
    save_override,
)
from kg.store import KgStore
from kg.v2.__main__ import build_parser, main, normalize_argv
from kg.v2.api import create_app
from kg.v2.migrate import migrate
from kg.v2.service import Service


def _workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "190도"
    ws["A1"], ws["B1"] = "온도", 190
    ws["F7"], ws["G7"] = 42.1, 43.2
    wb.save(path)


@pytest.fixture
def v1_workspace(tmp_path):
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    path = raw / "coffee.xlsx"
    _workbook(path)
    store = KgStore(tmp_path / "data/kg/kg.db")
    store.upsert_concept(
        {
            "concept_id": "oven_temperature",
            "canonical_name": "오븐온도",
            "canonical_name_en": "Oven temperature",
            "description": "굽기 온도",
            "data_type": "numeric",
            "domain_level": "L1",
            "canonical_unit": "C",
        }
    )
    store.upsert_concept(
        {"concept_id": "weight", "canonical_name": "무게", "domain_level": "L2"}
    )
    store.upsert_concept(
        {"concept_id": "old_metric", "canonical_name": "구지표", "status": "DEPRECATED"}
    )
    store.add_alias("oven_temperature", "Oven Temp", "oven temp")
    store.add_relation("weight", "oven_temperature", "AFFECTS")
    store.add_relation("weight", "oven_temperature", "IS_A")
    store.add_relation("old_metric", "weight", "IS_A")
    store.upsert_unit("C", "temperature", 1.0, 0.0)

    store.upsert_document("doc-coffee", "coffee.xlsx", str(path))
    vid = store.add_version("doc-coffee", sha256(path.read_bytes()).hexdigest(), "test")
    store.upsert_document("doc-missing", "missing.xlsx", str(raw / "missing.xlsx"))
    missing_vid = store.add_version("doc-missing", "0" * 64, "test")

    create_template(store, "financier_recipe", "Financier Recipe", "financier")
    add_version(
        store,
        "financier_recipe",
        {
            "sheet_templates": [
                {
                    "name": "oven_test",
                    "match": {"names": ["190도"]},
                    "mappings": [
                        {
                            "key": "temperature",
                            "concept_id": "oven_temperature",
                            "source": {"key_search": ["온도"], "offset": {"row": 0, "col": 1}},
                            "type": "number",
                            "unit": "C",
                        },
                        {
                            "key": "weight",
                            "concept_id": "weight",
                            "source": {"range": "F7:F7"},
                            "type": "number",
                            "unit": "g",
                        },
                        {
                            "key": "legacy",
                            "concept_id": "old_metric",
                            "source": {"range": "B1"},
                            "type": "text",
                        },
                    ],
                }
            ]
        },
        "tester",
    )
    create_template(store, "broken", "Broken", None)
    add_version(
        store,
        "broken",
        {
            "sheet_templates": [
                {
                    "name": "any",
                    "match": {"names": ["190도"]},
                    "mappings": [{"key": "bad", "source": {"range": "A0"}, "type": "text"}],
                }
            ]
        },
        "tester",
    )
    create_template(store, "hdr", "Header Only", None)
    add_version(
        store,
        "hdr",
        {
            "sheet_templates": [
                {
                    "name": "by_header",
                    "match": {"headers": ["온도"]},
                    "mappings": [{"key": "temp", "source": {"range": "B1"}, "type": "number"}],
                }
            ]
        },
        "tester",
    )
    assign(store, "doc-coffee", vid, "financier_recipe", 1)
    assign(store, "doc-coffee", vid, "hdr", 1)
    weight = next(
        m for m in effective_mappings(store, vid, "financier_recipe") if m["mapping_key"] == "weight"
    )
    override = save_override(
        store, "doc-coffee", vid, weight["mapping_id"], {"range": "G7:G7"}, "column moved", "tester"
    )
    run_parse(store, "doc-coffee", vid, path, "financier_recipe")
    store.commit()
    store.close()
    return Workspace(tmp_path, {"vid": vid, "missing_vid": missing_vid, "override": override["override_id"]})


class Workspace(type(Path())):
    """tmp 경로에 v1 식별자를 붙여 테스트로 넘긴다."""

    def __new__(cls, path, ids):
        self = super().__new__(cls, path)
        self.ids = ids
        return self


def v2(tmp):
    conn = sqlite3.connect(tmp / "data/kg/v2.db")
    conn.row_factory = sqlite3.Row
    return conn


def v1(tmp):
    conn = sqlite3.connect(f"file:{(tmp / 'data/kg/kg.db').as_posix()}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def run(tmp, **kw):
    return migrate(tmp, tmp, **kw)


def links(report, entity, status=None):
    return [
        l for l in report["links"] if l["entity"] == entity and (status is None or l["status"] == status)
    ]


def stored_links(conn):
    out = {}
    for row in conn.execute("SELECT storage_ref FROM artifact WHERE policy_ref='v1-migration'"):
        link = json.loads(row["storage_ref"])
        out[(link["entity"], link["v1_id"])] = link
    return out


def test_kg_becomes_single_revision(v1_workspace):
    report = run(v1_workspace)
    assert report["counts"]["kg"]["migrated"] == 1
    assert report["counts"]["units"]["skipped"] == 1
    with v2(v1_workspace) as conn:
        assert conn.execute("SELECT count(*) FROM kg_revision").fetchone()[0] == 1
        rev = conn.execute("SELECT kg_revision_id FROM kg_revision").fetchone()[0]
        oven = dict(
            conn.execute(
                "SELECT * FROM domain_concept WHERE concept_id='oven_temperature'"
            ).fetchone()
        )
        assert (oven["name"], oven["level"], oven["value_type"], oven["canonical_unit"]) == (
            "오븐온도", 1, "numeric", "C",
        )
        assert oven["definition"] == "굽기 온도"
        assert (
            conn.execute("SELECT status FROM domain_concept WHERE concept_id='old_metric'").fetchone()[0]
            == "deprecated"
        )
        assert conn.execute(
            "SELECT count(*) FROM domain_alias WHERE concept_id='oven_temperature' AND alias_norm='oven temp'"
        ).fetchone()[0] == 1
        assert conn.execute(
            "SELECT count(*) FROM domain_edge WHERE from_concept_id='weight' AND to_concept_id='oven_temperature' AND relation_type='AFFECTS'"
        ).fetchone()[0] == 1
        # v1 IS_A(자식→부모, 레벨 차 1)는 v2 parent_of(부모→자식)로 방향을 뒤집어 옮긴다.
        assert conn.execute(
            "SELECT count(*) FROM domain_edge WHERE from_concept_id='oven_temperature' AND to_concept_id='weight' AND relation_type='parent_of'"
        ).fetchone()[0] == 1
        # old_metric(L1)→weight(L2)는 레벨 조건에 맞지 않아 IS_A 그대로 남긴다.
        assert conn.execute(
            "SELECT relation_type FROM domain_edge WHERE from_concept_id='old_metric' AND to_concept_id='weight'"
        ).fetchone()[0] == "IS_A"
        stored = stored_links(conn)
        kg_links = [l for (entity, _), l in stored.items() if entity == "kg"]
        assert len(kg_links) == 1 and kg_links[0]["v2_id"] == rev
        assert kg_links[0]["detail"]["dropped_fields"] == ["canonical_name_en"]
        assert kg_links[0]["detail"]["parent_of_from_is_a"] == 1


def test_document_gets_new_stable_id_and_link(v1_workspace):
    report = run(v1_workspace)
    vid = v1_workspace.ids["vid"]
    with v2(v1_workspace) as conn:
        docs = [dict(r) for r in conn.execute("SELECT * FROM document")]
        assert len(docs) == 1 and docs[0]["document_id"] != "doc-coffee"
        assert docs[0]["source_ref"] == "coffee.xlsx"
        stored = stored_links(conn)
        assert stored[("document", "doc-coffee")]["v2_id"] == docs[0]["document_id"]
        version_link = stored[("document_version", vid)]
        version = dict(
            conn.execute(
                "SELECT * FROM document_version WHERE document_version_id=?", (version_link["v2_id"],)
            ).fetchone()
        )
        with v1(v1_workspace) as old:
            expected = old.execute(
                "SELECT file_hash FROM document_version WHERE version_id=?", (vid,)
            ).fetchone()[0]
        assert version["content_sha256"] == expected
        assert version["document_id"] == docs[0]["document_id"]
        sheets = [r[0] for r in conn.execute("SELECT name FROM sheet WHERE document_version_id=?", (version["document_version_id"],))]
        assert sheets == ["190도"]
    assert report["counts"]["documents"]["migrated"] == 1
    assert report["counts"]["document_versions"]["migrated"] == 1


def test_missing_source_is_skipped_not_fabricated(v1_workspace):
    report = run(v1_workspace)
    missing = v1_workspace.ids["missing_vid"]
    assert {"entity": "document_version", "v1_id": missing, "reason": "source missing"} in report["skipped"]
    assert {"entity": "document", "v1_id": "doc-missing", "reason": "source missing"} in report["skipped"]
    with v2(v1_workspace) as conn:
        assert conn.execute("SELECT count(*) FROM document WHERE display_name='missing.xlsx'").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM document_version WHERE content_sha256=?", ("0" * 64,)).fetchone()[0] == 0


def test_template_is_converted_to_v2_grammar(v1_workspace):
    report = run(v1_workspace)
    with v2(v1_workspace) as conn:
        template = dict(conn.execute("SELECT * FROM template WHERE name='Financier Recipe'").fetchone())
        versions = [
            dict(r)
            for r in conn.execute(
                "SELECT * FROM template_version WHERE template_id=?", (template["template_id"],)
            )
        ]
        assert len(versions) == 1
        definition = json.loads(versions[0]["definition_json"])
        assert definition["sheet_roles"] == {"oven_test": {"cardinality": "one"}}
        rules = {r["rule_key"]: r for r in definition["rules"]}
        assert set(rules) == {"temperature", "weight", "legacy"}
        temperature = rules["temperature"]
        assert temperature["selector"]["key"]["areas"][0]["find"]["texts"] == ["온도"]
        assert temperature["selector"]["value"]["areas"][0]["relative"] == {
            "row": 0, "col": 1, "rows": 1, "cols": 1,
        }
        assert temperature["selector"]["value"]["cardinality"] == "scalar"
        assert temperature["concept_id"] == "oven_temperature"
        assert temperature["value_spec"]["unit"] == "C"
        weight = rules["weight"]
        assert weight["selector"]["value"]["areas"] == [{"sheet_role": "oven_test", "range": "F7:F7"}]
        assert weight["selector"]["key"]["areas"][0]["find"]["texts"][0] == "weight"
        assert "무게" in weight["selector"]["key"]["areas"][0]["find"]["texts"]
        assert weight["value_spec"]["type"] == "decimal"
        assert weight["value_spec"]["unit"] == "g"
        assert weight["concept_id"] == "weight"
        assert rules["legacy"]["concept_id"] is None
        rule_rows = {
            r["rule_key"]: dict(r)
            for r in conn.execute(
                "SELECT * FROM template_rule WHERE template_version_id=?", (versions[0]["template_version_id"],)
            )
        }
        assert rule_rows["legacy"]["concept_id"] is None and rule_rows["weight"]["concept_id"] == "weight"
        stored = stored_links(conn)
        link = stored[("template_version", "financier_recipe@1")]
        assert link["v2_id"] == versions[0]["template_version_id"]
        assert any("old_metric" in note for note in link["detail"]["notes"])
        assert link["detail"]["rules"] == {
            "oven_test/temperature": "temperature",
            "oven_test/weight": "weight",
            "oven_test/legacy": "legacy",
        }
        assert link["detail"]["inferred_keys"] == ["legacy", "weight"]
        assert stored[("template", "financier_recipe")]["v2_id"] == template["template_id"]
    assert report["counts"]["templates"]["migrated"] == 2
    assert report["counts"]["templates"]["needs_review"] == 1
    assert report["counts"]["template_versions"]["migrated"] == 2


def test_unconvertible_template_is_kept_for_review(v1_workspace):
    report = run(v1_workspace)
    link = next(l for l in links(report, "template_version") if l["v1_id"] == "broken@1")
    assert link["status"] == "needs_review" and link["v2_id"] is None
    assert "INVALID_RANGE" in link["detail"]["reason"]
    assert report["counts"]["template_versions"]["needs_review"] == 1
    template_link = next(l for l in links(report, "template") if l["v1_id"] == "broken")
    assert template_link["status"] == "needs_review"
    with v2(v1_workspace) as conn:
        assert conn.execute("SELECT count(*) FROM template WHERE name='Broken'").fetchone()[0] == 0
        sources = [
            json.loads(r[0])
            for r in conn.execute("SELECT storage_ref FROM artifact WHERE policy_ref='v1-migration-source'")
        ]
        assert len(sources) == 1 and sources[0]["v1_id"] == "broken@1"
        assert sources[0]["spec"]["sheet_templates"][0]["mappings"][0]["source"] == {"range": "A0"}
        assert stored_links(conn)[("template_version", "broken@1")]["detail"]["source_artifact_id"] == (
            conn.execute("SELECT artifact_id FROM artifact WHERE policy_ref='v1-migration-source'").fetchone()[0]
        )
    hdr = next(s for s in report["skipped"] if s["entity"] == "assignment" and s["v1_id"].endswith("@hdr"))
    assert "headers" in hdr["reason"]


def test_assignment_becomes_proposed_application(v1_workspace):
    report = run(v1_workspace)
    vid = v1_workspace.ids["vid"]
    assert report["counts"]["assignments"] == {
        "migrated": 1, "existing": 0, "skipped": 1, "needs_review": 0, "planned": 0,
    }
    with v2(v1_workspace) as conn:
        apps = [dict(r) for r in conn.execute("SELECT * FROM template_application")]
        assert len(apps) == 1 and apps[0]["scope_key"] == "v1-migration"
        stored = stored_links(conn)
        assert stored[("assignment", f"{vid}@financier_recipe")]["v2_id"] == apps[0]["application_id"]
        assert apps[0]["document_version_id"] == stored[("document_version", vid)]["v2_id"]
        assert apps[0]["template_version_id"] == stored[("template_version", "financier_recipe@1")]["v2_id"]
        bound = [
            dict(r)
            for r in conn.execute(
                "SELECT a.role_key,s.name FROM application_sheet a JOIN sheet s USING(sheet_id) WHERE a.application_id=?",
                (apps[0]["application_id"],),
            )
        ]
        assert bound == [{"role_key": "oven_test", "name": "190도"}]
        statuses = {r[0] for r in conn.execute("SELECT status FROM mapping_revision")}
        assert statuses == {"proposed"}
        assert conn.execute("SELECT count(*) FROM mapping_head").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM mapping_revision WHERE origin='template'").fetchone()[0] == 3
    app = create_app(v1_workspace, start_worker=False)
    with TestClient(app) as client:
        response = client.get("/api/v2/documents", params={"extraction_status": "review"})
        assert response.status_code == 200, response.text
        items = response.json()["items"]
        assert [i["display_name"] for i in items] == ["coffee.xlsx"]


def test_override_becomes_manual_proposed_revision(v1_workspace):
    report = run(v1_workspace)
    oid = v1_workspace.ids["override"]
    assert report["counts"]["overrides"]["migrated"] == 1
    with v2(v1_workspace) as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                "SELECT * FROM mapping_revision WHERE rule_key='weight' ORDER BY revision_no"
            )
        ]
        assert [r["revision_no"] for r in rows] == [1, 2]
        first, second = rows
        assert (second["origin"], second["status"]) == ("manual", "proposed")
        assert second["supersedes_id"] == first["mapping_revision_id"]
        assert second["concept_id"] == "weight"
        assert oid in second["reason"] and "column moved" in second["reason"]
        spec = json.loads(second["effective_spec_json"])
        assert spec["selector"]["value"]["areas"] == [{"sheet_role": "oven_test", "range": "G7:G7"}]
        assert spec["selector"]["key"] == json.loads(first["effective_spec_json"])["selector"]["key"]
        assert conn.execute("SELECT count(*) FROM mapping_head").fetchone()[0] == 0
        link = stored_links(conn)[("override", oid)]
        assert link["v2_id"] == second["mapping_revision_id"]
        assert link["detail"]["rule_key"] == "weight"


def test_values_are_not_copied_but_listed(v1_workspace):
    report = run(v1_workspace)
    vid = v1_workspace.ids["vid"]
    with v1(v1_workspace) as old:
        expected = old.execute(
            "SELECT count(*) FROM parsed_source WHERE value_json IS NOT NULL AND value_json<>'null'"
        ).fetchone()[0]
        payload = old.execute("SELECT count(*) FROM payload_value").fetchone()[0]
    assert expected == 3
    with v2(v1_workspace) as conn:
        assert conn.execute("SELECT count(*) FROM extraction_run").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM extracted_item").fetchone()[0] == 0
        app = conn.execute("SELECT application_id FROM template_application").fetchone()[0]
    items = report["re_extract_required"]
    assert len(items) == 1
    assert items[0]["sources"] == expected
    assert items[0]["v2_application_id"] == app
    assert (items[0]["v1_document_version"], items[0]["v1_template_id"]) == (vid, "financier_recipe")
    assert report["counts"]["values"] == {"parsed_sources": expected, "payload_values": payload, "copied": 0}


TABLES = (
    "kg_revision", "domain_concept", "document", "document_version", "sheet", "template",
    "template_version", "template_rule", "template_application", "application_sheet",
    "mapping_revision", "artifact",
)


def snapshot(tmp):
    with v2(tmp) as conn:
        return {t: conn.execute(f"SELECT count(*) FROM {t}").fetchone()[0] for t in TABLES}


def test_second_run_is_idempotent(v1_workspace):
    first = run(v1_workspace)
    before = snapshot(v1_workspace)
    second = run(v1_workspace)
    assert snapshot(v1_workspace) == before
    for key, counts in second["counts"].items():
        if key != "values":
            assert counts["migrated"] == 0, key
    for key in ("kg", "documents", "document_versions", "templates", "template_versions", "assignments", "overrides"):
        assert second["counts"][key]["existing"] > 0, key
    assert second["counts"]["template_versions"]["needs_review"] == 0
    assert len(second["skipped"]) == len(first["skipped"])
    assert {l["v2_id"] for l in links(second, "assignment")} == {l["v2_id"] for l in links(first, "assignment")}


def test_human_edits_are_never_overwritten(v1_workspace):
    run(v1_workspace)
    oid = v1_workspace.ids["override"]
    service = Service(v1_workspace)
    with v2(v1_workspace) as conn:
        app = conn.execute("SELECT application_id FROM template_application").fetchone()[0]
        latest = dict(
            conn.execute(
                "SELECT * FROM mapping_revision WHERE rule_key='weight' ORDER BY revision_no DESC LIMIT 1"
            ).fetchone()
        )
    human = service.revise(
        app, latest["mapping_revision_id"], 0, json.loads(latest["effective_spec_json"]),
        "weight", "approved", "검수", "human",
    )
    with v2(v1_workspace) as conn:
        conn.execute(
            "DELETE FROM artifact WHERE policy_ref='v1-migration' AND json_extract(storage_ref,'$.entity')='override'"
        )
        conn.commit()
        count = conn.execute("SELECT count(*) FROM mapping_revision WHERE rule_key='weight'").fetchone()[0]
    report = run(v1_workspace)
    skip = next(s for s in report["skipped"] if s["entity"] == "override" and s["v1_id"] == oid)
    assert "EDIT_CONFLICT" in skip["reason"]
    with v2(v1_workspace) as conn:
        head = dict(conn.execute("SELECT * FROM mapping_head WHERE rule_key='weight'").fetchone())
        assert head["mapping_revision_id"] == human["mapping_revision_id"]
        assert conn.execute("SELECT count(*) FROM mapping_revision WHERE rule_key='weight'").fetchone()[0] == count
        assert conn.execute("SELECT count(*) FROM mapping_head").fetchone()[0] == 1


def test_dry_run_writes_nothing(v1_workspace):
    report = run(v1_workspace, dry_run=True)
    assert not (v1_workspace / "data/kg/v2.db").exists()
    assert report["dry_run"] is True
    for entity in ("kg", "document", "document_version", "template_version", "assignment", "override"):
        planned = links(report, entity, "planned")
        assert planned, entity
        assert all(l["v2_id"] is None for l in planned)
    assert [l["status"] for l in links(report, "template_version") if l["v1_id"] == "broken@1"] == ["needs_review"]
    assert any(s["reason"] == "source missing" for s in report["skipped"])
    assert report["counts"]["assignments"]["planned"] == 1
    assert report["re_extract_required"][0]["v2_application_id"] is None
    real = run(v1_workspace)
    assert real["counts"]["documents"]["migrated"] == 1
    assert real["counts"]["assignments"]["migrated"] == 1
    assert real["counts"]["overrides"]["migrated"] == 1
    with v2(v1_workspace) as conn:
        assert conn.execute("SELECT count(*) FROM template_application").fetchone()[0] == 1


def test_cli_subcommands_and_serve_default(v1_workspace, capsys):
    assert normalize_argv(["--ws", "x", "--port", "1"]) == ["serve", "--ws", "x", "--port", "1"]
    args = build_parser().parse_args(normalize_argv([]))
    assert (args.command, args.port, args.host) == ("serve", 8010, "127.0.0.1")
    args = build_parser().parse_args(["migrate", "--ws", "a", "--from-ws", "b", "--dry-run"])
    assert (args.command, args.dry_run, args.raw, args.report) == ("migrate", True, None, None)
    out = v1_workspace / "out/report.json"
    code = main(["migrate", "--ws", str(v1_workspace), "--from-ws", str(v1_workspace), "--report", str(out)])
    assert code == 0
    report = json.loads(out.read_text(encoding="utf-8"))
    assert report["format"] == "v1-migration-report/1"
    assert report["counts"]["documents"]["migrated"] == 1
    summary = json.loads(capsys.readouterr().out)
    assert summary["report"] == str(out) and summary["dry_run"] is False
    empty = v1_workspace / "empty"
    assert main(["migrate", "--ws", str(empty), "--from-ws", str(empty)]) == 2
    assert "V1_DATABASE_MISSING" in capsys.readouterr().err
    assert not (empty / "data/kg/v2.db").exists()


def test_interrupted_run_recovers_without_duplicates(v1_workspace):
    # 링크를 쓰기 전에 중단된 실행을 흉내낸다: 배정·override 링크만 지우고 다시 실행한다.
    run(v1_workspace)
    before = snapshot(v1_workspace)
    with v2(v1_workspace) as conn:
        conn.execute(
            "DELETE FROM artifact WHERE policy_ref='v1-migration' AND json_extract(storage_ref,'$.entity') IN ('assignment','override')"
        )
        conn.commit()
    report = run(v1_workspace)
    after = snapshot(v1_workspace)
    assert {k: v for k, v in after.items() if k != "artifact"} == {
        k: v for k, v in before.items() if k != "artifact"
    }
    assignment = links(report, "assignment", "migrated")
    assert len(assignment) == 1 and assignment[0]["detail"].get("recovered") is True
    override = links(report, "override", "migrated")
    with v2(v1_workspace) as conn:
        latest = conn.execute(
            "SELECT mapping_revision_id FROM mapping_revision WHERE rule_key='weight' ORDER BY revision_no DESC LIMIT 1"
        ).fetchone()[0]
    assert [l["v2_id"] for l in override] == [latest]


def test_dry_run_against_existing_v2_reports_existing(v1_workspace):
    run(v1_workspace)
    before = snapshot(v1_workspace)
    report = run(v1_workspace, dry_run=True)
    assert snapshot(v1_workspace) == before
    for key in ("kg", "documents", "document_versions", "templates", "template_versions", "assignments", "overrides"):
        assert (report["counts"][key]["migrated"], report["counts"][key]["planned"]) == (0, 0), key
        assert report["counts"][key]["existing"] > 0, key


def test_source_copied_into_separate_v2_raw_is_found(v1_workspace, tmp_path):
    # v1의 절대경로는 이전 워크스페이스를 가리키지만, v2 raw에 같은 이름의 사본이 있으면 그것을 쓴다.
    import shutil

    ws2 = tmp_path / "v2ws"
    (ws2 / "data/raw").mkdir(parents=True)
    shutil.copy(v1_workspace / "data/raw/coffee.xlsx", ws2 / "data/raw/coffee.xlsx")
    report = migrate(ws2, v1_workspace)
    assert report["counts"]["documents"]["migrated"] == 1
    assert not any(s["reason"].startswith("source outside") for s in report["skipped"])
    with v2(ws2) as conn:
        version = conn.execute(
            "SELECT a.storage_ref FROM document_version v JOIN artifact a ON a.artifact_id=v.source_artifact_id JOIN document d USING(document_id) WHERE d.display_name='coffee.xlsx'"
        ).fetchone()
        assert version is not None and json.loads(version[0])["source_ref"] == "coffee.xlsx"
    assert {"entity": "document", "v1_id": "doc-missing", "reason": "source missing"} in report["skipped"]
