"""v3 E2E 작업 공간 시드(계약 §8) — openpyxl로 가상 문서를 만들고 실제 `kg.v3.service.Service`로 등록·승인·발행까지 돌린다.

python -m examples.schema_v3.demo --workspace /tmp/data-gathering-v3-demo
python -m kg.v3 serve --ws /tmp/data-gathering-v3-demo --port 8031

시나리오: 스키마 `공정 데이터 표준` → 프로파일 `공정데이터_A양식`(대표 문서 `공정데이터_2024_01.xlsx`로 승인) →
같은 양식 문서 3개(자동 적용·승인·발행), 앵커가 이동한 문서 1개(compatible → 검수), 다른 양식 1개(unmatched),
잠긴 파일 1개(locked), 이미지가 있는 시트 1개(`공정데이터_2024_03.xlsx`의 `첨부`). 사용자 원본은 읽지 않는다.
"""

from __future__ import annotations

import argparse
import io
import json
import struct
import zlib
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill

from kg.v3.service import Service

SCHEMA_KEY = "process_standard"
PROFILE_NAME = "공정데이터_A양식"
REFERENCE_DOCUMENT = "공정데이터_2024_01.xlsx"
IDENTICAL_DOCUMENTS = ("공정데이터_2024_01.xlsx", "공정데이터_2024_02.xlsx", "공정데이터_2024_03.xlsx")
SHIFTED_DOCUMENT = "공정데이터_2024_04_양식이동.xlsx"
OTHER_DOCUMENT = "품질검사_2024_05.xlsx"
LOCKED_DOCUMENT = "공정데이터_2024_06_잠김.xlsx"
IMAGE_DOCUMENT = "공정데이터_2024_03.xlsx"
MAIN_SHEET, COMMON_SHEET, IMAGE_SHEET = "공정 기록", "공통 정보", "첨부"
LOT_COUNT = 12

SCHEMA = {
    "format": "parsing-schema",
    "schema_version": "3.0",
    "schema_key": SCHEMA_KEY,
    "schema_name": "공정 데이터 표준",
    "description": "공정 기록 문서에서 추출하는 표준 필드",
    "fields": [
        {"field_key": "basic", "name": "기본 정보", "type": "group", "level": 1},
        {"field_key": "process", "name": "공정 정보", "type": "group", "level": 1},
        {"field_key": "result", "name": "결과 정보", "type": "group", "level": 1},
        {"field_key": "product_name", "name": "제품명", "type": "text", "level": 2, "parents": ["basic"], "aliases": ["제품", "Product"]},
        {"field_key": "recipe_name", "name": "레시피명", "type": "text", "level": 2, "parents": ["basic"], "aliases": ["레시피", "Recipe"]},
        {"field_key": "process_name", "name": "공정명", "type": "text", "level": 2, "parents": ["process"], "aliases": ["공정"]},
        {"field_key": "equipment", "name": "설비명", "type": "text", "level": 2, "parents": ["process"], "aliases": ["설비", "Equipment"]},
        {"field_key": "lot", "name": "배치", "type": "text", "level": 2, "parents": ["process"], "aliases": ["LOT", "배치번호"], "description": "반복 행의 업무 키"},
        {"field_key": "measured_at", "name": "측정일시", "type": "datetime", "level": 2, "parents": ["process"], "aliases": ["측정 시각"]},
        {"field_key": "temperature", "name": "온도", "type": "decimal", "unit": "°C", "level": 2, "parents": ["process"], "aliases": ["온도값", "Temp"], "related": ["pressure"], "description": "공정 설정 온도"},
        {"field_key": "pressure", "name": "압력", "type": "decimal", "unit": "bar", "level": 2, "parents": ["process"], "aliases": ["Pressure"]},
        {"field_key": "duration", "name": "시간", "type": "decimal", "unit": "min", "level": 2, "parents": ["process"], "aliases": ["소요 시간"]},
        {"field_key": "result_value", "name": "결과값", "type": "decimal", "level": 2, "parents": ["result"], "aliases": ["측정값"]},
        {"field_key": "verdict", "name": "판정", "type": "text", "level": 2, "parents": ["result"], "aliases": ["합부", "판정 결과"]},
    ],
}


def _scalar_rule(rule_key, rule_name, field_key, label):
    return {
        "rule_key": rule_key,
        "rule_name": rule_name,
        "field_key": field_key,
        "selector": {
            "key": {"areas": [{"sheet_role": "main", "find": {"texts": [label], "within": "A1:D10"}}]},
            "value": {"areas": [{"sheet_role": "main", "relative": {"row": 0, "col": 1}}]},
        },
        "value_spec": {"type": "text"},
    }


def _column_rule(rule_key, rule_name, field_key, anchor, value_type, unit_anchor=None, unit=None, key=None):
    rule = {
        "rule_key": rule_key,
        "rule_name": rule_name,
        "field_key": field_key,
        "selector": {
            "key": {"areas": [{"sheet_role": "main", "anchor": anchor}]},
            "value": {
                "areas": [{"sheet_role": "main", "relative": {"row": 1, "col": 0, "rows": 60, "cols": 1}}],
                "cardinality": "list",
                "axis": "down",
                "element_layout": "one_per_row",
                "stop": {"kind": "blank_run", "count": 2},
            },
        },
        "value_spec": {"type": value_type},
        "record_spec": {"scope": ["process-table"], "key": "physical_row"},
    }
    if unit_anchor:
        rule["selector"]["unit"] = {"areas": [{"sheet_role": "common", "relative": {"row": 0, "col": 1, "anchor": unit_anchor}}]}
        rule["value_spec"]["unit"] = unit
    if key:
        rule["relations"] = [{"to_rule": key, "kind": "same_row"}]
    return rule


PROFILE = {
    "format": "parsing-profile",
    "schema_version": "3.0",
    "profile_name": PROFILE_NAME,
    "schema_key": SCHEMA_KEY,
    "description": "공정 기록 A양식: 머리 정보 4개 + LOT 표(열마다 규칙 1개) + 공통 정보 시트의 단위",
    "sheet_roles": {
        "main": {"cardinality": "one", "match": {"name": MAIN_SHEET}},
        "common": {"cardinality": "one", "match": {"any_of": [{"name": COMMON_SHEET}, {"contains_text": {"texts": ["온도 단위"], "within": "A1:D10"}}]}},
    },
    "anchors": {
        "hdr_lot": {"sheet_role": "main", "find": {"regex": "^(배치|LOT)$", "within": "A1:H30"}},
        "hdr_time": {"sheet_role": "main", "find": {"texts": ["측정일시"], "within": "A1:H30"}},
        "hdr_temp": {"sheet_role": "main", "find": {"texts": ["온도"], "within": "A1:H30"}},
        "hdr_press": {"sheet_role": "main", "find": {"texts": ["압력"], "within": "A1:H30"}},
        "hdr_dur": {"sheet_role": "main", "find": {"texts": ["시간"], "within": "A1:H30"}},
        "hdr_result": {"sheet_role": "main", "find": {"texts": ["결과값"], "within": "A1:H30"}},
        "hdr_verdict": {"sheet_role": "main", "find": {"texts": ["판정"], "within": "A1:H30"}},
        "unit_temp": {"sheet_role": "common", "find": {"texts": ["온도 단위"], "within": "A1:D10"}},
        "unit_press": {"sheet_role": "common", "find": {"texts": ["압력 단위"], "within": "A1:D10"}},
        "unit_dur": {"sheet_role": "common", "find": {"texts": ["시간 단위"], "within": "A1:D10"}},
    },
    "rules": [
        _scalar_rule("product_name", "제품명", "product_name", "제품명"),
        _scalar_rule("recipe_name", "레시피명", "recipe_name", "레시피명"),
        _scalar_rule("process_name", "공정명", "process_name", "공정명"),
        _scalar_rule("equipment", "설비명", "equipment", "설비명"),
        _column_rule("lot", "배치", "lot", "hdr_lot", "text"),
        _column_rule("measured_at", "측정일시", "measured_at", "hdr_time", "datetime", key="lot"),
        _column_rule("temperature", "온도", "temperature", "hdr_temp", "decimal", "unit_temp", "°C", key="lot"),
        _column_rule("pressure", "압력", "pressure", "hdr_press", "decimal", "unit_press", "bar", key="lot"),
        _column_rule("duration", "시간", "duration", "hdr_dur", "decimal", "unit_dur", "min", key="lot"),
        _column_rule("result_value", "결과값", "result_value", "hdr_result", "decimal", key="lot"),
        _column_rule("verdict", "판정", "verdict", "hdr_verdict", "text", key="lot"),
    ],
}

HEADERS = ("LOT", "측정일시", "온도", "압력", "시간", "결과값", "판정")


def png_bytes(size=24, rgb=(37, 99, 235)):
    """의존성 없이 만드는 단색 PNG(렌더 asset 검증용)."""
    raw = b"".join(b"\x00" + bytes(rgb) * size for _ in range(size))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def build_process_workbook(path: Path, serial: int, table_shift=0, temp_offset=0.0, with_image=False):
    """A양식 문서. table_shift>0이면 표(앵커)가 아래로 이동해 compatible 매치가 된다."""
    wb = Workbook()
    ws = wb.active
    ws.title = MAIN_SHEET
    ws.merge_cells("A1:G1")
    ws["A1"] = "공정 데이터 기록 · 검증용 가상 데이터"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="2563EB")
    ws["A1"].alignment = Alignment(horizontal="center")
    for row, (label, value) in enumerate(
        (("제품명", f"제품-{serial:02d}"), ("레시피명", f"RCP-{100 + serial}"), ("공정명", "열처리"), ("설비명", f"FUR-{serial % 3 + 1}")), start=3
    ):
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, value)
    header_row = 8 + table_shift
    if table_shift:
        ws.cell(7, 1, "※ 양식 개정: 표 위치 변경")
    for col, text in enumerate(HEADERS, start=1):
        cell = ws.cell(header_row, col, text)
        cell.font = Font(bold=True, color="21564C")
        cell.fill = PatternFill("solid", fgColor="E8F0FF")
    base = datetime(2024, 1, serial, 9, 0)
    for n in range(LOT_COUNT):
        row = header_row + 1 + n
        ws.cell(row, 1, f"LOT-{serial:02d}-{n + 1:03d}")
        ws.cell(row, 2, base + timedelta(minutes=30 * n))
        ws.cell(row, 3, round(150 + (n % 5) * 2.5 + temp_offset, 2))
        ws.cell(row, 4, round(1.2 + (n % 4) * 0.05, 2))
        ws.cell(row, 5, 30 + (n % 3) * 5)
        ws.cell(row, 6, round(98.0 + (n % 7) * 0.3, 2))
        ws.cell(row, 7, "합격" if n % 6 else "불합격")
    for name, width in zip("ABCDEFG", (16, 20, 10, 10, 10, 12, 10)):
        ws.column_dimensions[name].width = width
    ws.freeze_panes = ws.cell(header_row + 1, 2)
    common = wb.create_sheet(COMMON_SHEET)
    common["A1"], common["B1"] = "온도 단위", "°C"
    common["A2"], common["B2"] = "압력 단위", "bar"
    common["A3"], common["B3"] = "시간 단위", "min"
    if with_image:
        attach = wb.create_sheet(IMAGE_SHEET)
        attach["A1"] = "설비 사진"
        attach.add_image(XlImage(io.BytesIO(png_bytes())), "B3")
    wb.save(path)
    return path


def build_other_workbook(path: Path):
    """다른 양식(시트 이름·헤더가 다름) → 프로파일 없음(unmatched)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "품질 검사"
    ws["A1"] = "품질 검사 성적서"
    ws["A3"], ws["B3"], ws["C3"] = "검사 항목", "규격", "측정치"
    for n in range(5):
        ws.cell(4 + n, 1, f"항목 {n + 1}")
        ws.cell(4 + n, 2, "10±1")
        ws.cell(4 + n, 3, 10 + n * 0.1)
    wb.save(path)
    return path


def build_locked_file(path: Path):
    # ZIP 매직(PK)이 아닌 첫 바이트 → 기본 Reader가 DRM_READER_REQUIRED(403)로 거부한다.
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"encrypted-workbook" * 8)
    return path


def write_documents(raw: Path):
    raw.mkdir(parents=True, exist_ok=True)
    for n, name in enumerate(IDENTICAL_DOCUMENTS, start=1):
        build_process_workbook(raw / name, n, with_image=name == IMAGE_DOCUMENT)
    build_process_workbook(raw / SHIFTED_DOCUMENT, 4, table_shift=3)
    build_other_workbook(raw / OTHER_DOCUMENT)
    build_locked_file(raw / LOCKED_DOCUMENT)


def mutate_document(root: Path, source_ref: str, temp_offset=0.5):
    """A양식 문서(원본 폴더 기준 상대 경로, 하위 폴더 가능)의 값(온도)을 바꿔 새 snapshot 시나리오(UC-5)를 만든다. 양식은 그대로."""
    path = Path(root) / "data/raw" / source_ref
    wb = load_workbook(path)
    ws = wb[MAIN_SHEET]
    for row in range(9, 9 + LOT_COUNT):
        cell = ws.cell(row, 3)
        cell.value = round(float(cell.value) + temp_offset, 2)
    ws["B4"] = "RCP-101-개정"
    wb.save(path)
    return source_ref


def mutate_first_document(root: Path, temp_offset=0.5):
    """`공정데이터_2024_01.xlsx`를 바꾼다(mutate_document 래퍼)."""
    return mutate_document(root, REFERENCE_DOCUMENT, temp_offset)


def write_heavy_document(root: Path, name="공정데이터_2024_09_대용량.xlsx", table_shift=3, rows=3000, cols=150):
    """E2E 도우미(작업 내역 JobBar 시나리오): A양식(기본은 표가 이동한 compatible 양식)에 큰 부속 시트를 붙인 문서를 원본 폴더에 만든다.

    Reader가 파일을 열 때마다 수 초가 걸려 등록·재파싱 작업이 '진행 중'으로 관찰될 만큼 오래 돈다. 시드 문서 집합은 바꾸지 않는다."""
    raw = Path(root) / "data/raw"
    raw.mkdir(parents=True, exist_ok=True)
    path = build_process_workbook(raw / name, 9, table_shift=table_shift)
    wb = load_workbook(path)
    ws = wb.create_sheet("부속 자료")
    for r in range(1, rows + 1):
        ws.append([r * c for c in range(1, cols + 1)])
    wb.save(path)
    return name


def write_wide_document(root: Path, name="공정데이터_2024_08_확장.xlsx", extra_lots=60, note_col=30):
    """E2E 도우미(Source Review 창 요청 시나리오): A양식 문서의 LOT 표를 60행 아래까지 늘리고 Z열 너머(기본 AD열)에 비고를 적어
    렌더 범위가 첫 창(A1:Z60)을 넘게 만든다 — 스크롤하면 뷰어가 두 번째 창(A61:… / AA1:…)을 요청한다. 시드 문서 집합은 바꾸지 않는다."""
    raw = Path(root) / "data/raw"
    raw.mkdir(parents=True, exist_ok=True)
    path = build_process_workbook(raw / name, 8)
    wb = load_workbook(path)
    ws = wb[MAIN_SHEET]
    base = datetime(2024, 1, 8, 9, 0)
    for n in range(LOT_COUNT, LOT_COUNT + extra_lots):
        row = 9 + n
        ws.cell(row, 1, f"LOT-08-{n + 1:03d}")
        ws.cell(row, 2, base + timedelta(minutes=30 * n))
        ws.cell(row, 3, round(150 + (n % 5) * 2.5, 2))
        ws.cell(row, 4, round(1.2 + (n % 4) * 0.05, 2))
        ws.cell(row, 5, 30 + (n % 3) * 5)
        ws.cell(row, 6, round(98.0 + (n % 7) * 0.3, 2))
        ws.cell(row, 7, "합격" if n % 6 else "불합격")
    ws.cell(2, note_col, "비고: 확장 열")
    ws.cell(9 + LOT_COUNT + extra_lots - 1, note_col, "마지막 행 비고")
    wb.save(path)
    return name


def seed(root: Path, principal="demo-user"):
    """작업 공간을 만들고 요약을 돌려준다(이미 시드된 작업 공간은 덮어쓰지 않는다)."""
    root = Path(root).resolve()
    raw = root / "data/raw"
    if (raw / REFERENCE_DOCUMENT).exists():
        raise SystemExit("이미 존재하는 샘플을 덮어쓰지 않습니다. 새 작업 공간을 지정하세요.")
    write_documents(raw)
    service = Service(root)
    try:
        schema = service.import_schema(SCHEMA)
        profile = service.import_profile(SCHEMA_KEY, PROFILE, principal=principal)
        # 대표 문서: 등록(unmatched) → 수동 적용 → 일괄 승인(추출·발행) → 프로파일 승인(+rematch 소급).
        first = _register(service, [REFERENCE_DOCUMENT], principal)[0]
        application = service.apply_profile(first["snapshot"]["snapshot_id"], profile["profile_id"], principal=principal)
        approved = service.approve_all(application["application_id"], reason="대표 문서 검수", principal=principal)
        if not approved["extraction"] or approved["extraction"]["state"] != "succeeded":
            raise SystemExit("대표 문서 추출에 실패했습니다: " + json.dumps(approved, ensure_ascii=False))
        reference = service.approve_profile(profile["profile_id"], application["application_id"], principal=principal)
        rematch = service.jobs.wait(reference["reparse_job"]["job_id"], 60, principal)
        others = _register(
            service, [*IDENTICAL_DOCUMENTS[1:], SHIFTED_DOCUMENT, OTHER_DOCUMENT, LOCKED_DOCUMENT], principal
        )
        documents = [first, *others]
        # 대표 문서는 등록 뒤 승인·발행됐으므로 최종 상태를 다시 읽는다.
        for d in documents:
            if d.get("document_id"):
                d["status"] = service.document(d["document_id"])["status"]
        first["applied"] = [{"application_id": application["application_id"], "profile_name": profile["profile_name"], "compatibility": "manual", "state": "published"}]
        return {
            "workspace": str(root),
            "schema": {"schema_key": schema["schema_key"], "current_rev": schema["current_rev"], "fields": schema["fields"]},
            "profile": {
                "profile_id": profile["profile_id"],
                "profile_name": profile["profile_name"],
                "status": reference["status"],
                "reference_application_id": application["application_id"],
                "rematch": rematch["state"],
            },
            "documents": [
                {
                    "document_id": d["document_id"],
                    "document_name": d["document_name"],
                    "status": d["status"],
                    "snapshot_id": (d.get("snapshot") or {}).get("snapshot_id"),
                    "applied": [{k: a[k] for k in ("application_id", "profile_name", "compatibility", "state")} for a in d.get("applied") or []],
                    "error": d.get("error"),
                }
                for d in documents
            ],
        }
    finally:
        service.close()


def _register(service, names, principal):
    job = service.register_documents(names, principal=principal, wait=120)
    if job["state"] != "succeeded":
        raise SystemExit("등록 작업 실패: " + json.dumps(job, ensure_ascii=False))
    return job["result"]["documents"]


def main(argv=None):
    parser = argparse.ArgumentParser(description="v3 데모 작업 공간 시드")
    parser.add_argument("--workspace", type=Path, required=True)
    args = parser.parse_args(argv)
    print(json.dumps(seed(args.workspace), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
