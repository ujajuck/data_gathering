"""실제 XLSX가 있는 재현용 작업 공간을 생성한다. 사용자 원본을 읽거나 복사하지 않는다.

python -m examples.schema_v2.runtime_demo --workspace /tmp/data-gathering-v2-demo
python -m kg.v2 --ws /tmp/data-gathering-v2-demo --port 8010
"""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

from kg.v2.db import uid
from kg.v2.service import Service


def seed(root):
    raw = root / "data/raw"
    raw.mkdir(parents=True, exist_ok=True)
    path = raw / "공정운전_샘플.xlsx"
    if path.exists():
        raise SystemExit(
            "이미 존재하는 샘플을 덮어쓰지 않습니다. 새 작업 공간을 지정하세요."
        )
    wb = Workbook()
    ws = wb.active
    ws.title = "공정 기록"
    ws.merge_cells("A1:F1")
    ws["A1"] = "공정 운전 기록 · 검증용 가상 데이터"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="087D70")
    ws.row_dimensions[1].height = 32
    ws["A3"], ws["B3"], ws["D3"] = "배치", "공정", "온도"
    ws.merge_cells("B3:C3")
    for c in ("A3", "B3", "D3"):
        ws[c].font = Font(bold=True, color="21564C")
        ws[c].fill = PatternFill("solid", fgColor="DEF0E9")
    for row in range(4, 69):
        ws.cell(row, 1, f"LOT-{row-3:03}")
        ws.cell(row, 2, str(150 + (row % 9) * 2.5))
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.cell(row, 5, row - 3)
        ws.cell(row, 6, 150 + (row % 9) * 2.5)
    for name in ("A", "B", "C", "D", "E", "F"):
        ws.column_dimensions[name].width = 16
    ws.freeze_panes = "B4"
    chart = LineChart()
    chart.title = "공정 온도 추이"
    chart.add_data(Reference(ws, min_col=6, min_row=4, max_row=12))
    ws.add_chart(chart, "H3")
    units = wb.create_sheet("공통 정보")
    units["A1"], units["B1"] = "온도 단위", "°C"
    units["A3"], units["B3"], units["C3"] = "측정방향", "세로", "가로 모두 지원"
    wb.save(path)
    service = Service(root)
    kg = service.import_kg(
        {
            "concepts": [
                {
                    "concept_id": "process_temperature",
                    "name": "공정온도",
                    "level": 1,
                    "canonical_unit": "°C",
                    "aliases": ["공정 온도", "Process Temp"],
                }
            ]
        },
        "demo-user",
    )
    doc = service.register(path.name, "local-xlsx", "demo-user")
    definition = {
        "kg_revision_id": kg["kg_revision_id"],
        "sheet_roles": {
            "main": {"cardinality": "one"},
            "common": {"cardinality": "one"},
        },
        "rules": [
            {
                "rule_key": "process_temperature",
                "concept_id": "process_temperature",
                "record_spec": {"scope": ["process-table"], "key": {"column": "A"}},
                "selector": {
                    "key": {
                        "areas": [
                            {"sheet_role": "main", "range": "B3:C3"},
                            {"sheet_role": "main", "range": "D3"},
                        ]
                    },
                    "value": {
                        "areas": [{"sheet_role": "main", "range": "B4:C68"}],
                        "cardinality": "list",
                        "axis": "down",
                        "element_layout": "one_per_row",
                        "stop": {"kind": "explicit_areas", "max_items": 10000},
                    },
                    "unit": {"areas": [{"sheet_role": "common", "range": "B1"}]},
                },
                "value_spec": {"type": "decimal", "unit": "°C"},
            }
        ],
    }
    template = service.create_template("공정 운전 기록", definition, "demo-user")
    with service.db.connect() as conn:
        sheets = {
            r["name"]: r["sheet_id"]
            for r in conn.execute(
                "SELECT * FROM sheet WHERE document_version_id=?", (doc["version_id"],)
            )
        }
    application = service.apply_template(
        doc["version_id"],
        template["template_version_id"],
        {"main": [sheets["공정 기록"]], "common": [sheets["공통 정보"]]},
        None,
        False,
        "demo-user",
    )
    print(
        json.dumps(
            {"workspace": str(root), **doc, **application}, ensure_ascii=False, indent=2
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", type=Path, required=True)
    seed(parser.parse_args().workspace.resolve())
