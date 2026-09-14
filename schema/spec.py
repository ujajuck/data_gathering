"""셀 주소·범위 파싱과 명시적 값 변환. 임의 Python/SQL 표현식을 실행하지 않는다."""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext

from openpyxl.utils.cell import get_column_letter, range_boundaries

from .db import Problem

MAX_ITEMS = 100000


def bounds(address):
    if (
        not isinstance(address, str)
        or len(address) > 40
        or not re.fullmatch(
            r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)?",
            address,
        )
    ):
        raise Problem("INVALID_RANGE", "A1 또는 B3:C20 형태의 유한 범위를 지정하세요.")
    c1, r1, c2, r2 = range_boundaries(address)
    if not (1 <= r1 <= r2 <= 1048576 and 1 <= c1 <= c2 <= 16384):
        raise Problem("INVALID_RANGE", "Excel 범위를 벗어났습니다.")
    return r1, c1, r2, c2


def address(r1, c1, r2=None, c2=None):
    r2, c2 = r2 or r1, c2 or c1
    a, b = f"{get_column_letter(c1)}{r1}", f"{get_column_letter(c2)}{r2}"
    return a if a == b else f"{a}:{b}"


def decimal(value):
    try:
        number = Decimal(str(value))
        if (
            not number.is_finite()
            or len(number.as_tuple().digits) > 1000
            or abs(number.as_tuple().exponent) > 1000
        ):
            raise InvalidOperation()
        return number
    except (InvalidOperation, ValueError):
        raise Problem(
            "INVALID_DECIMAL",
            "숫자로 변환할 수 없는 값이 있습니다. 원본과 변환 규칙을 확인하세요.",
        ) from None


def decimal_text(value):
    value = decimal(value)
    result = format(value, "f") if value else "0"
    return result.rstrip("0").rstrip(".") if "." in result else result


def typed(value, spec):
    target = spec.get("type", "text")
    if target == "decimal":
        value = decimal(value)
        normal = spec.get("normalization", {})
        if normal.get("operation") == "affine":
            with localcontext() as ctx:
                ctx.prec = 4096
                value = value * decimal(normal.get("factor", "1")) + decimal(
                    normal.get("offset", "0")
                )
        return target, decimal_text(value)
    if target == "boolean":
        if isinstance(value, bool):
            return target, "true" if value else "false"
        if str(value).casefold() not in ("true", "false", "1", "0"):
            raise Problem("INVALID_BOOLEAN", "참/거짓으로 변환할 수 없습니다.")
        return target, "true" if str(value).casefold() in ("true", "1") else "false"
    if target in ("date", "datetime"):
        try:
            if target == "date":
                return (
                    target,
                    (
                        value.date()
                        if isinstance(value, datetime)
                        else (
                            value
                            if isinstance(value, date)
                            else date.fromisoformat(str(value))
                        )
                    ).isoformat(),
                )
            return (
                target,
                (
                    value
                    if isinstance(value, datetime)
                    else datetime.fromisoformat(str(value))
                ).isoformat(),
            )
        except ValueError:
            raise Problem("INVALID_DATE", "날짜/시각 값이 유효하지 않습니다.") from None
    return "text", (
        str(value).strip()
        if spec.get("normalization", {}).get("operation") == "trim"
        else str(value)
    )
