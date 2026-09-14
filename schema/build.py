"""§4.10 데이터 빌드 — current_value(현재 snapshot의 발행 실행 값)를 행/열로 조립해 CSV·XLSX·SQLite로 내보낸다.

흐름: candidates(사용 가능 문서·제외 사유·필드별 값 있는 문서 수) → preview(50행) → build(산출물 + manifest).
행 조립은 문서 단위 스트리밍이다(한 문서의 값만 메모리에 두고 행을 바로 writer에 넘긴다). 문서 200개·행 20만 이하는 동기
(runtime_job kind='build' 이력 행), 초과는 작업(kind='build')이며 Service.handle_job이 execute_build로 넘긴다.
"""

from __future__ import annotations

import csv
import json
import os
import re
import sqlite3
import threading
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path


from .db import Problem, digest, dump, norm, now, rows
from .spec import decimal_text

SYNC_DOCUMENT_LIMIT = 200
SYNC_ROW_LIMIT = 200_000
PREVIEW_ROWS = 50
CONFLICT_LIMIT = 1000
XLSX_WIDTH_SAMPLE = 500
WRITE_BATCH = 500
EXPORT_DIR = "data/exports"
BUILD_KEY_RE = re.compile(r"^[0-9a-f]{16}$")
META_COLUMNS = ("_document", "_snapshot", "_record_key")
FORMATS = {
    "csv": ("data.csv", "text/csv; charset=utf-8"),
    "xlsx": ("data.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "sqlite": ("data.sqlite", "application/vnd.sqlite3"),
}
_UNITS_CACHE = {}


# ---------------------------------------------------------------------------- 입력 정리


def _request(request):
    """pydantic 모델 또는 dict → 기본값을 채운 dict."""
    if hasattr(request, "model_dump"):
        request = request.model_dump(mode="json")
    if not isinstance(request, dict):
        raise Problem("INVALID_REQUEST", "빌드 요청은 객체여야 합니다.")
    ids = request.get("document_ids")
    if not isinstance(ids, list) or not ids or any(not isinstance(i, str) or not i for i in ids):
        raise Problem("INVALID_REQUEST", "document_ids를 1개 이상 지정하세요.")
    columns = request.get("columns")
    if not isinstance(columns, list) or any(not isinstance(c, dict) for c in columns):
        raise Problem("INVALID_REQUEST", "columns는 객체 배열이어야 합니다.")
    out = {
        "document_ids": list(dict.fromkeys(ids)),
        "schema_key": request.get("schema_key"),
        "columns": [{"field_key": c.get("field_key"), "header": c.get("header"), "target_unit": c.get("target_unit") or None} for c in columns],
        "row_mode": request.get("row_mode") or "record",
        "format": request.get("format") or "xlsx",
    }
    if not isinstance(out["schema_key"], str) or not out["schema_key"]:
        raise Problem("INVALID_REQUEST", "schema_key를 지정하세요.")
    if out["row_mode"] not in ("record", "document"):
        raise Problem("INVALID_REQUEST", "row_mode는 record/document 중 하나여야 합니다.")
    if out["format"] not in FORMATS:
        raise Problem("INVALID_REQUEST", "format은 csv/xlsx/sqlite 중 하나여야 합니다.")
    return out


def load_units(root):
    """<ws>/config/units.yaml → UnitRegistry(없으면 None). 파일 mtime으로 캐시한다."""
    path = Path(root) / "config/units.yaml"
    if not path.is_file():
        return None
    stamp = path.stat().st_mtime_ns
    cached = _UNITS_CACHE.get(str(path))
    if cached and cached[0] == stamp:
        return cached[1]
    from .units import UnitRegistry

    registry = UnitRegistry.load(path)
    _UNITS_CACHE[str(path)] = (stamp, registry)
    return registry


def _stage_documents(conn, document_ids):
    """선택 문서를 TEMP 테이블에 올려 IN 목록 크기 제한 없이 조인한다."""
    conn.execute("CREATE TEMP TABLE IF NOT EXISTS build_docs (document_id TEXT PRIMARY KEY, ord INTEGER NOT NULL)")
    conn.execute("DELETE FROM build_docs")
    conn.executemany("INSERT INTO build_docs VALUES (?,?)", [(d, n) for n, d in enumerate(dict.fromkeys(document_ids))])


def _stage_scalars(conn, schema_id, field_ids):
    """value cardinality가 scalar인 리비전을 TEMP 테이블에 올린다. 엔진은 scalar 규칙에도 좌표 record_key(예: B55)를 주므로
    행 정체성은 '규칙이 scalar인가'로 판단해 그 값을 문서의 모든 행에 복제한다(§4.10)."""
    conn.execute("CREATE TEMP TABLE IF NOT EXISTS build_scalar (mapping_revision_id TEXT PRIMARY KEY)")
    conn.execute("DELETE FROM build_scalar")
    conn.execute(
        "INSERT OR IGNORE INTO build_scalar SELECT mr.mapping_revision_id FROM ("
        " SELECT DISTINCT cv.mapping_revision_id FROM current_value cv JOIN build_docs b ON b.document_id=cv.document_id"
        " JOIN parsing_application a ON a.application_id=cv.application_id WHERE a.schema_id=? AND cv.field_id IN (%s)) x"
        " JOIN mapping_revision mr ON mr.mapping_revision_id=x.mapping_revision_id"
        " WHERE json_extract(mr.effective_spec_json, '$.selector.value.cardinality')='scalar'" % ",".join("?" for _ in field_ids),
        (schema_id, *field_ids),
    )


# ---------------------------------------------------------------------------- 후보


def _document_reason(doc, apps):
    """§4.10 제외 사유. 발행된 application이 하나라도 있으면 사용 가능."""
    if doc.get("document_id") is None:
        return "not_found"
    if doc["status"] == "locked":
        return "locked"
    if not doc["current_snapshot_id"]:
        return "not_extracted"
    if not apps:
        return "unmatched"
    if any(a["published_run_id"] for a in apps):
        return None
    if any((a["unapproved"] or 0) > 0 or (a["inherited"] or 0) > 0 or not a["heads_total"] for a in apps):
        return "review_required"
    if any(a["last_run"] in ("failed", "cancelled") for a in apps):
        return "failed"
    return "not_extracted"


def _candidates(service, conn, document_ids, schema_key):
    _stage_documents(conn, document_ids)
    docs = rows(
        conn,
        "SELECT b.ord, b.document_id requested_id, d.document_id, d.document_name, d.status, d.current_snapshot_id, s.captured_at, s.revision_no "
        "FROM build_docs b LEFT JOIN document d ON d.document_id=b.document_id LEFT JOIN document_snapshot s ON s.snapshot_id=d.current_snapshot_id ORDER BY b.ord",
    )
    where, params = "a.snapshot_id=d.current_snapshot_id AND d.document_id IN (SELECT document_id FROM build_docs)", ()
    schema = None
    if schema_key:
        schema = conn.execute("SELECT * FROM parsing_schema WHERE schema_key=?", (schema_key,)).fetchone()
        if schema is None:
            raise Problem("UNKNOWN_SCHEMA", f"파싱 스키마 {schema_key!r}를 찾을 수 없습니다.", 404)
        where, params = where + " AND ps.schema_key=?", (schema_key,)
    aggs = service._application_aggregates(conn, where, params)
    by_doc = {}
    for a in aggs:
        by_doc.setdefault(a["document_id"], []).append(a)
    items, usable = [], []
    for d in docs:
        apps = by_doc.get(d["document_id"], [])
        reason = _document_reason(d, apps)
        published = [a for a in apps if a["published_run_id"]]
        item = {
            "document_id": d["document_id"] or d["requested_id"],
            "document_name": d["document_name"],
            "usable": reason is None,
            "snapshot": {"snapshot_id": d["current_snapshot_id"], "captured_at": d["captured_at"], "revision_no": d["revision_no"]} if d["current_snapshot_id"] else None,
            "status": d["status"],
        }
        if reason:
            item["reason"] = reason
        if published:
            first = published[0]
            item["profile"] = {"profile_id": first["profile_id"], "profile_name": first["profile_name"], "rev": first["profile_rev"], "application_id": first["application_id"], "run_id": first["published_run_id"], "schema_key": first["schema_key"]}
            item["applications"] = [
                {"application_id": a["application_id"], "run_id": a["published_run_id"], "profile_id": a["profile_id"], "profile_name": a["profile_name"], "rev": a["profile_rev"], "schema_key": a["schema_key"], "schema_name": a["schema_name"]}
                for a in published
            ]
            usable.append(item)
        items.append(item)
    return {"documents": items, "usable": usable, "schema": dict(schema) if schema else None}


def _fields(conn, schema_ids):
    """스키마별 활성·group 아닌 필드 + 선택 문서 중 present 값이 있는 문서 수."""
    if not schema_ids:
        return []
    marks = ",".join("?" for _ in schema_ids)
    counts = {
        r["field_id"]: r["n"]
        for r in conn.execute(
            f"SELECT cv.field_id, count(DISTINCT cv.document_id) n FROM current_value cv JOIN build_docs b ON b.document_id=cv.document_id "
            f"JOIN parsing_application a ON a.application_id=cv.application_id WHERE a.schema_id IN ({marks}) AND cv.value_state='present' GROUP BY cv.field_id",
            tuple(schema_ids),
        )
    }
    return [
        {"field_key": f["field_key"], "name": f["field_name"], "type": f["value_type"], "unit": f["canonical_unit"], "schema_key": f["schema_key"], "document_count": counts.get(f["field_id"], 0)}
        for f in conn.execute(
            f"SELECT f.field_id, f.field_key, f.field_name, f.value_type, f.canonical_unit, s.schema_key FROM parsing_field f JOIN parsing_schema s ON s.schema_id=f.schema_id "
            f"WHERE f.schema_id IN ({marks}) AND f.status='active' AND f.value_type<>'group' ORDER BY s.schema_key, f.ordinal, f.field_key",
            tuple(schema_ids),
        )
    ]


def candidates(service, document_ids, schema_key=None):
    """POST /builds/candidates → {documents[], summary{total, usable, excluded}, fields[], schemas[]}."""
    if not isinstance(document_ids, list) or not document_ids:
        raise Problem("INVALID_REQUEST", "document_ids를 1개 이상 지정하세요.")
    with service.db.connect() as conn:
        found = _candidates(service, conn, document_ids, schema_key)
        schemas = {}
        for item in found["usable"]:
            for a in item["applications"]:
                entry = schemas.setdefault(a["schema_key"], {"schema_key": a["schema_key"], "schema_name": a["schema_name"], "documents": set()})
                entry["documents"].add(item["document_id"])
        if found["schema"]:
            schema_ids = [found["schema"]["schema_id"]]
        else:
            schema_ids = [r["schema_id"] for r in conn.execute("SELECT schema_id, schema_key FROM parsing_schema WHERE schema_key IN (%s)" % ",".join("?" for _ in schemas), tuple(schemas))] if schemas else []
        fields = _fields(conn, schema_ids)
    return {
        "documents": found["documents"],
        "summary": {"total": len(found["documents"]), "usable": len(found["usable"]), "excluded": len(found["documents"]) - len(found["usable"])},
        "fields": fields,
        "schemas": [{"schema_key": s["schema_key"], "schema_name": s["schema_name"], "document_count": len(s["documents"])} for s in schemas.values()],
    }


# ---------------------------------------------------------------------------- 계획(열 검증·대상·build_key·행 수)


def _validate_columns(columns, fields, row_mode):
    problems, seen = [], {}
    for n, column in enumerate(columns):
        header = column.get("header")
        header = header.strip() if isinstance(header, str) else ""
        column["header"] = header
        key = norm(header)
        if not header:
            problems.append({"index": n, "field_key": column.get("field_key"), "header": header, "reason": "empty"})
        elif header.startswith("_"):
            problems.append({"index": n, "field_key": column.get("field_key"), "header": header, "reason": "reserved"})
        elif key in seen:
            problems.append({"index": n, "field_key": column.get("field_key"), "header": header, "reason": "duplicate"})
        else:
            seen[key] = n
    if row_mode == "document":
        # 문서 모드는 열마다 <header>_count 열이 붙으므로 그 이름과 겹치는 header도 중복이다.
        for n, column in enumerate(columns):
            if column["header"] and norm(column["header"] + "_count") in seen and seen[norm(column["header"] + "_count")] != n:
                problems.append({"index": seen[norm(column["header"] + "_count")], "field_key": column.get("field_key"), "header": column["header"] + "_count", "reason": "duplicate"})
    if problems:
        raise Problem("INVALID_HEADER", "출력 헤더가 비어 있거나 중복되었습니다.", 422, problems)
    unknown = [c["field_key"] for c in columns if c["field_key"] not in fields]
    if unknown:
        raise Problem("UNKNOWN_FIELD", "스키마에 없는 필드입니다: " + ", ".join(map(str, unknown)), 422, unknown)
    groups = [c["field_key"] for c in columns if fields[c["field_key"]]["value_type"] == "group"]
    if groups:
        raise Problem("GROUP_FIELD_TARGET", "묶음(group) 필드는 값이 없어 출력할 수 없습니다.", 422, groups)


def _plan(service, conn, request):
    """열 검증 + 후보 + 대상 실행 + build_key + 행 수. 빌드/미리보기/작업이 같은 계획을 쓴다."""
    fields = service.schema_fields(request["schema_key"], conn)
    if fields is None:
        raise Problem("UNKNOWN_SCHEMA", f"파싱 스키마 {request['schema_key']!r}를 찾을 수 없습니다.", 404)
    _validate_columns(request["columns"], fields, request["row_mode"])
    found = _candidates(service, conn, request["document_ids"], request["schema_key"])
    schema = found["schema"]
    columns = []
    for n, c in enumerate(request["columns"]):
        field = fields[c["field_key"]]
        columns.append({"index": n, "field_key": c["field_key"], "field_id": field["field_id"], "header": c["header"], "name": field["name"], "type": field["value_type"], "unit": c["target_unit"] or field["unit"], "target_unit": c["target_unit"], "field_unit": field["unit"]})
    usable = found["usable"]
    sources = [
        {"document_id": d["document_id"], "document_name": d["document_name"], "snapshot_id": d["snapshot"]["snapshot_id"], "application_id": a["application_id"], "run_id": a["run_id"], "profile": {"id": a["profile_id"], "name": a["profile_name"], "rev": a["rev"]}}
        for d in usable
        for a in d["applications"]
    ]
    excluded = [{"document_id": d["document_id"], "document_name": d["document_name"], "reason": d["reason"]} for d in found["documents"] if not d["usable"]]
    field_ids = sorted({c["field_id"] for c in columns})
    _stage_scalars(conn, schema["schema_id"], field_ids)
    if request["row_mode"] == "record" and usable:
        counted = {
            r["document_id"]: r["n"]
            for r in conn.execute(
                "SELECT cv.document_id, count(DISTINCT CASE WHEN bs.mapping_revision_id IS NULL THEN cv.record_key END) n FROM current_value cv JOIN build_docs b ON b.document_id=cv.document_id "
                "JOIN parsing_application a ON a.application_id=cv.application_id LEFT JOIN build_scalar bs ON bs.mapping_revision_id=cv.mapping_revision_id "
                "WHERE a.schema_id=? AND cv.field_id IN (%s) GROUP BY cv.document_id" % ",".join("?" for _ in field_ids),
                (schema["schema_id"], *field_ids),
            )
        }
        row_count = sum(max(1, counted.get(d["document_id"], 0)) for d in usable)
    else:
        row_count = len(usable)
    canonical = {
        "document_ids": request["document_ids"],
        "schema": {"key": schema["schema_key"], "rev": schema["current_rev"]},
        "columns": [{"field_key": c["field_key"], "header": c["header"], "target_unit": c["target_unit"]} for c in columns],
        "row_mode": request["row_mode"],
        "sources": [[s["document_id"], s["snapshot_id"], s["run_id"]] for s in sources],
    }
    return {
        "schema": {"key": schema["schema_key"], "rev": schema["current_rev"], "id": schema["schema_id"], "name": schema["schema_name"]},
        "columns": columns,
        "field_ids": field_ids,
        "usable": usable,
        "sources": sources,
        "excluded": excluded,
        "row_mode": request["row_mode"],
        "row_count": row_count,
        "build_key": digest(canonical)[:16],
    }


# ---------------------------------------------------------------------------- 행 조립


def _convert(units, value, target_unit, field_unit):
    """unit_normalized → target_unit. (text, unit, error) — error는 사용자에게 보이는 사유, None이면 변환 성공/불필요."""
    unit = value["unit_normalized"]
    if unit is None and target_unit == field_unit:
        return value["value_text"], target_unit, None
    same = unit == target_unit or (units is not None and units.normalize_unit(unit) == units.normalize_unit(target_unit))
    if same:
        return value["value_text"], target_unit, None
    if unit is None:
        return value["value_text"], None, "값에 단위가 없어 변환할 수 없습니다."
    if units is None:
        return value["value_text"], unit, "단위 변환표(config/units.yaml)가 없어 같은 단위만 출력할 수 있습니다."
    if value["value_type"] != "decimal":
        return value["value_text"], unit, "숫자(decimal) 값만 단위를 변환할 수 있습니다."
    source, target = units.normalize_unit(unit), units.normalize_unit(target_unit)
    shared = units.dimensions_of(source) & units.dimensions_of(target)
    if not shared:
        return value["value_text"], unit, f"{unit} → {target_unit} 변환식이 단위 변환표에 없습니다."
    dimension = sorted(shared)[0]
    # float 산술은 26.850000000000023 같은 잡음이 생기므로 변환표의 계수만 받아 Decimal로 계산한다.
    f_from, o_from = units.factor_offset(dimension, source)
    f_to, o_to = units.factor_offset(dimension, target)
    try:
        with localcontext() as ctx:
            ctx.prec = 15
            base = Decimal(value["value_text"]) * Decimal(repr(f_from)) + Decimal(repr(o_from))
            result = (base - Decimal(repr(o_to))) / Decimal(repr(f_to))
            return decimal_text(+result), target_unit, None
    except (InvalidOperation, ValueError, ArithmeticError, Problem):
        return value["value_text"], unit, "숫자로 해석할 수 없어 단위를 변환하지 못했습니다."


class _Assembler:
    """한 문서의 값 목록 → 행(셀) 목록. 중복(같은 source_identity_key)은 1개, 다른 출처는 첫 값 + conflicts."""

    def __init__(self, plan, units):
        self.columns = plan["columns"]
        self.row_mode = plan["row_mode"]
        self.units = units
        self.by_field = {}
        for column in self.columns:
            self.by_field.setdefault(column["field_id"], []).append(column)
        self.conflicts = []
        self.conflict_count = 0
        self.row_no = 0

    def _conflict(self, kind, doc, record_key, column, values, kept, message):
        self.conflict_count += 1
        if len(self.conflicts) >= CONFLICT_LIMIT:
            return
        self.conflicts.append(
            {
                "kind": kind,
                "row": self.row_no,
                "document_id": doc["document_id"],
                "document_name": doc["document_name"],
                "record_key": record_key,
                "field_key": column["field_key"],
                "header": column["header"],
                "message": message,
                "kept": kept,
                "values": [
                    {"value_id": v["value_id"], "text": v["value_text"], "unit": v["unit_normalized"], "application_id": v["application_id"], "rule_key": v["rule_key"], "sheet_name": v["sheet_name"], "range": v["range"]}
                    for v in values[:20]
                ],
            }
        )

    def _cell(self, doc, record_key, column, values):
        if not values:
            return None
        unique, seen = [], set()
        for v in values:
            if v["source_identity_key"] in seen:
                continue
            seen.add(v["source_identity_key"])
            unique.append(v)
        first = unique[0]
        if len(unique) > 1 and self.row_mode == "record":
            self._conflict("multiple_values", doc, record_key, column, unique, first["value_id"], "같은 행·필드에 서로 다른 출처의 값이 여럿입니다. 첫 값만 출력합니다.")
        state = first["value_state"]
        text, unit = (first["value_text"] if state == "present" else first["raw_text"] if state == "error" else "") or "", first["unit_normalized"]
        if state == "present" and column["target_unit"]:
            text, unit, error = _convert(self.units, first, column["target_unit"], column["field_unit"])
            if error:
                self._conflict("unit", doc, record_key, column, [first], first["value_id"], error)
        cell = {
            "text": text,
            "value_id": first["value_id"],
            "sheet_id": first["sheet_id"],
            "sheet_name": first["sheet_name"],
            "range": first["range"],
            "application_id": first["application_id"],
            "rule_key": first["rule_key"],
            "state": state,
            "type": first["value_type"],
            "unit": unit,
        }
        if self.row_mode == "document":
            cell["count"] = len(unique)
        return cell

    def rows(self, doc, values):
        """values: 한 문서의 current_value 행(실행·블록·순서). → [row]."""
        records, scalars = {}, {}
        for v in values:
            for column in self.by_field.get(v["field_id"], ()):
                bucket = scalars if v["scalar"] or v["record_key"] is None or self.row_mode == "document" else records.setdefault(v["record_key"], {})
                bucket.setdefault(column["index"], []).append(v)
        keys = list(records) or [None]
        out = []
        for record_key in keys:
            self.row_no += 1
            per_column = records.get(record_key, {}) if record_key is not None else {}
            cells = [self._cell(doc, record_key, column, per_column.get(column["index"], []) + scalars.get(column["index"], [])) for column in self.columns]
            out.append(
                {
                    "row_no": self.row_no,
                    "document": {"document_id": doc["document_id"], "document_name": doc["document_name"]},
                    "snapshot": {"snapshot_id": doc["snapshot_id"], "captured_at": doc["captured_at"]},
                    "record_key": record_key,
                    "cells": cells,
                }
            )
        return out


VALUE_SQL = (
    "SELECT b.ord, cv.document_id, cv.snapshot_id, cv.application_id, cv.profile_id, r.rule_key, cv.field_id, cv.value_id, cv.record_key, cv.item_index, "
    "cv.value_text, cv.display_text, cv.raw_text, cv.value_type, cv.value_state, cv.unit_normalized, cv.source_identity_key, "
    "sr.sheet_id, sh.sheet_name, sr.locator_key range, (bs.mapping_revision_id IS NOT NULL) scalar "
    "FROM build_docs b JOIN current_value cv ON cv.document_id=b.document_id JOIN parsing_application a ON a.application_id=cv.application_id "
    "LEFT JOIN build_scalar bs ON bs.mapping_revision_id=cv.mapping_revision_id "
    "JOIN mapping_revision mr ON mr.mapping_revision_id=cv.mapping_revision_id JOIN mapping m ON m.mapping_id=mr.mapping_id JOIN parsing_rule r ON r.rule_id=m.rule_id "
    "LEFT JOIN extracted_value_region er ON er.value_id=cv.value_id AND er.role='value' AND er.ordinal=0 "
    "LEFT JOIN extracted_value_region ei ON ei.value_id=cv.value_id AND ei.role='input' AND ei.ordinal=0 "
    "LEFT JOIN source_region sr ON sr.region_id=coalesce(er.region_id, ei.region_id) LEFT JOIN sheet sh ON sh.sheet_id=sr.sheet_id "
    "WHERE a.schema_id=? AND cv.field_id IN (%s) ORDER BY b.ord, cv.run_id, cv.group_key, cv.item_index, cv.value_id"
)


def _iterate_rows(conn, plan, units, checkpoint, limit=None):
    """문서 순서대로 값을 스트리밍하며 행을 만든다(문서 하나의 값만 메모리). limit 행에서 멈춘다. → (assembler, generator)."""
    assembler = _Assembler(plan, units)
    docs = {d["document_id"]: {"document_id": d["document_id"], "document_name": d["document_name"], "snapshot_id": d["snapshot"]["snapshot_id"], "captured_at": d["snapshot"]["captured_at"]} for d in plan["usable"]}
    order = [d["document_id"] for d in plan["usable"]]

    def generate():
        cursor = conn.execute(VALUE_SQL % ",".join("?" for _ in plan["field_ids"]), (plan["schema"]["id"], *plan["field_ids"]))
        state = {"pending": 0, "emitted": 0}

        def done():
            return bool(limit) and state["emitted"] >= limit

        def emit(doc_id, values):
            for row in assembler.rows(docs[doc_id], values):
                if done():
                    return
                state["emitted"] += 1
                yield row

        def through(doc_id, values):
            # 값이 하나도 없는 사용 가능 문서도 빈 행 1개를 낸다(문서 순서 유지).
            while state["pending"] < len(order) and order[state["pending"]] != doc_id:
                yield from emit(order[state["pending"]], [])
                state["pending"] += 1
            yield from emit(doc_id, values)
            state["pending"] += 1

        current, buffer = None, []
        try:
            for r in cursor:
                r = dict(r)
                if r["document_id"] not in docs:
                    continue
                if current is not None and r["document_id"] != current:
                    yield from through(current, buffer)
                    buffer = []
                    if done():
                        return
                    checkpoint(assembler.row_no, plan["row_count"])
                current = r["document_id"]
                buffer.append(r)
            if current is not None:
                yield from through(current, buffer)
            while state["pending"] < len(order) and not done():
                yield from emit(order[state["pending"]], [])
                state["pending"] += 1
        finally:
            cursor.close()

    return assembler, generate()


def _row_texts(row, columns, row_mode):
    texts = []
    for column, cell in zip(columns, row["cells"]):
        texts.append(cell["text"] if cell else "")
        if row_mode == "document":
            texts.append(cell["count"] if cell else 0)
    texts.extend([row["document"]["document_name"], (row["snapshot"]["captured_at"] or "")[:10], row["record_key"]])
    return texts


def _headers(columns, row_mode):
    out = []
    for column in columns:
        out.append(column["header"])
        if row_mode == "document":
            out.append(column["header"] + "_count")
    return out + list(META_COLUMNS)


# ---------------------------------------------------------------------------- 미리보기


def preview(service, request):
    """POST /builds/preview → {columns, rows[:50], row_count, excluded[], conflicts[](미리본 행 안의 충돌)}."""
    request = _request(request)
    units = load_units(service.root)
    with service.db.connect() as conn:
        plan = _plan(service, conn, request)
        assembler, generator = _iterate_rows(conn, plan, units, lambda *a, **k: None, limit=PREVIEW_ROWS)
        out_rows = list(generator)
    columns = [{"field_key": c["field_key"], "header": c["header"], "name": c["name"], "type": c["type"], "unit": c["unit"]} for c in plan["columns"]]
    return {
        "build_key": plan["build_key"],
        "schema": {"key": plan["schema"]["key"], "rev": plan["schema"]["rev"], "name": plan["schema"]["name"]},
        "row_mode": plan["row_mode"],
        "columns": columns,
        "headers": _headers(plan["columns"], plan["row_mode"]),
        "rows": out_rows,
        "row_count": plan["row_count"],
        "document_count": len(plan["usable"]),
        "excluded": plan["excluded"],
        "conflicts": assembler.conflicts,
        "sources": plan["sources"],
    }


# ---------------------------------------------------------------------------- 산출물 writer


FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def csv_safe(text):
    """원본 셀 텍스트가 스프레드시트에서 수식/DDE로 실행되지 않게 한다(=,+,-,@,탭,CR로 시작하면 작은따옴표 접두). 숫자('-5')는 그대로."""
    if text is None:
        return ""
    text = str(text)
    if text and text[0] in FORMULA_LEAD:
        try:
            Decimal(text.strip())
            return text
        except (InvalidOperation, ValueError):
            return "'" + text
    return text


class _CsvWriter:
    def __init__(self, path, columns, row_mode):
        self.columns, self.row_mode = columns, row_mode
        self.file = open(path, "w", encoding="utf-8-sig", newline="")
        self.writer = csv.writer(self.file, lineterminator="\r\n")
        self.writer.writerow([csv_safe(h) for h in _headers(columns, row_mode)])

    def write(self, row):
        self.writer.writerow([csv_safe(t) for t in _row_texts(row, self.columns, self.row_mode)])

    def close(self, manifest=None):
        self.file.close()


class _XlsxWriter:
    """write_only 워크북(대용량). 열 너비는 앞 500행 표본으로 정한 뒤 쓴다(write_only는 셀을 쓰기 전에만 너비를 정할 수 있다)."""

    def __init__(self, path, columns, row_mode):
        from openpyxl import Workbook

        self.path, self.columns, self.row_mode = path, columns, row_mode
        self.book = Workbook(write_only=True)
        self.sheet = self.book.create_sheet("data")
        self.headers = _headers(columns, row_mode)
        self.buffer, self.started = [], False

    def _text_cell(self, text):
        # '=...'로 시작하는 문자열은 openpyxl이 수식(data_type 'f')으로 저장하므로 문자열 타입을 강제한다.
        from openpyxl.cell import WriteOnlyCell

        cell = WriteOnlyCell(self.sheet, value=text)
        if cell.data_type == "f":
            cell.data_type = "s"
        return cell

    def _cells(self, row):
        out, texts = [], _row_texts(row, self.columns, self.row_mode)
        types = []
        for column in self.columns:
            types.append(column["type"])
            if self.row_mode == "document":
                types.append("count")
        types.extend(["text"] * len(META_COLUMNS))
        for kind, text in zip(types, texts):
            if kind == "count":
                out.append(int(text or 0))
            elif kind == "decimal" and text not in (None, ""):
                try:
                    out.append(Decimal(str(text)))
                except InvalidOperation:
                    out.append(text)
            elif kind == "boolean" and text in ("true", "false"):
                out.append(text == "true")
            else:
                out.append(self._text_cell("" if text is None else str(text)))
        return out

    def _start(self):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        widths = [len(h) for h in self.headers]
        for cells in self.buffer:
            for n, value in enumerate(cells):
                widths[n] = max(widths[n], min(60, len(str(getattr(value, "value", value)))))
        for n, width in enumerate(widths):
            self.sheet.column_dimensions[get_column_letter(n + 1)].width = min(60, max(8, width * 1.2 + 2))
        header = []
        for text in self.headers:
            cell = self._text_cell(text)
            cell.font = Font(bold=True)
            header.append(cell)
        self.sheet.append(header)
        for cells in self.buffer:
            self.sheet.append(cells)
        self.buffer, self.started = [], True

    def write(self, row):
        cells = self._cells(row)
        if self.started:
            self.sheet.append(cells)
            return
        self.buffer.append(cells)
        if len(self.buffer) >= XLSX_WIDTH_SAMPLE:
            self._start()

    def close(self, manifest=None):
        if not self.started:
            self._start()
        self.book.save(self.path)


class _SqliteWriter:
    """테이블 data: header 열 + <header>_count(문서 모드) + _source_<header>('sheet!range') + _document·_snapshot·_record_key."""

    def __init__(self, path, columns, row_mode):
        self.columns, self.row_mode = columns, row_mode
        Path(path).unlink(missing_ok=True)
        self.conn = sqlite3.connect(path)
        quote = lambda name: '"' + name.replace('"', '""') + '"'
        defs = ["_row_no INTEGER PRIMARY KEY"]
        for column in columns:
            defs.append(f"{quote(column['header'])} {'NUMERIC' if column['type'] == 'decimal' else 'TEXT'}")
            if row_mode == "document":
                defs.append(f"{quote(column['header'] + '_count')} INTEGER")
        defs.extend(f"{quote('_source_' + c['header'])} TEXT" for c in columns)
        defs.extend(f"{name} TEXT" for name in META_COLUMNS)
        self.conn.execute(f"CREATE TABLE data ({', '.join(defs)})")
        self.conn.execute("CREATE TABLE _manifest (json TEXT NOT NULL)")
        self.width = len(defs)
        self.batch = []

    def write(self, row):
        values = [row["row_no"]]
        for column, cell in zip(self.columns, row["cells"]):
            text = cell["text"] if cell else None
            if column["type"] == "decimal" and text not in (None, ""):
                try:
                    text = float(Decimal(str(text)))
                except (InvalidOperation, ValueError):
                    pass
            values.append(text)
            if self.row_mode == "document":
                values.append(cell["count"] if cell else 0)
        values.extend(f"{cell['sheet_name']}!{cell['range']}" if cell and cell["range"] else None for cell in row["cells"])
        values.extend(_row_texts(row, self.columns, self.row_mode)[-len(META_COLUMNS):])
        self.batch.append(values)
        if len(self.batch) >= WRITE_BATCH:
            self.flush()

    def flush(self):
        if self.batch:
            self.conn.executemany(f"INSERT INTO data VALUES ({','.join('?' for _ in range(self.width))})", self.batch)
            self.batch = []
        self.conn.commit()

    def close(self, manifest=None):
        self.flush()
        if manifest is not None:
            self.conn.execute("INSERT INTO _manifest VALUES (?)", (dump(manifest),))
            self.conn.commit()
        self.conn.close()


WRITERS = {"csv": _CsvWriter, "xlsx": _XlsxWriter, "sqlite": _SqliteWriter}


# ---------------------------------------------------------------------------- 빌드


def export_folder(service, build_key):
    if not BUILD_KEY_RE.match(build_key or ""):
        raise Problem("NOT_FOUND", "빌드를 찾을 수 없습니다.", 404)
    return service.root / EXPORT_DIR / build_key


def download_url(build_key, format):
    return f"/api/builds/{build_key}/download?format={format}"


def _read_manifest(folder):
    path = folder / "manifest.json"
    if not path.is_file():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except ValueError:
        return None


# 같은 build_key의 동기 빌드가 겹치면 같은 .tmp를 두 번 쓰므로 build_key별로 직렬화한다.
_BUILD_LOCKS, _BUILD_LOCKS_GUARD = {}, threading.Lock()


def _build_lock(build_key):
    with _BUILD_LOCKS_GUARD:
        return _BUILD_LOCKS.setdefault(build_key, threading.Lock())


def _existing(service, plan, format):
    """같은 입력(build_key)의 산출물이 이미 있으면 그대로 돌려준다."""
    folder = export_folder(service, plan["build_key"])
    manifest = _read_manifest(folder)
    filename = FORMATS[format][0]
    if manifest and format in (manifest.get("formats") or {}) and (folder / filename).is_file():
        return {"build_key": plan["build_key"], "download_url": download_url(plan["build_key"], format), "manifest": manifest, "path": str(folder / filename), "format": format, "row_count": manifest.get("row_count"), "reused": True}
    return None


def _execute(service, conn, plan, request, units, checkpoint):
    """행 스트리밍 → writer → manifest.json. 쓰기는 .tmp → os.replace."""
    format = request["format"]
    folder = export_folder(service, plan["build_key"])
    folder.mkdir(parents=True, exist_ok=True)
    filename = FORMATS[format][0]
    target, tmp = folder / filename, folder / (filename + ".tmp")
    assembler, generator = _iterate_rows(conn, plan, units, checkpoint)
    writer = WRITERS[format](tmp, plan["columns"], plan["row_mode"])
    written, finished = 0, False
    try:
        for row in generator:
            writer.write(row)
            written += 1
            if written % WRITE_BATCH == 0:
                checkpoint(written, plan["row_count"])
        previous = _read_manifest(folder) or {}
        manifest = {
            "build_key": plan["build_key"],
            "created_at": now(),
            "schema": {"key": plan["schema"]["key"], "rev": plan["schema"]["rev"], "name": plan["schema"]["name"]},
            "sources": plan["sources"],
            "columns": [{"field_key": c["field_key"], "header": c["header"], "type": c["type"], "unit": c["unit"]} for c in plan["columns"]],
            "row_mode": plan["row_mode"],
            "row_count": written,
            "document_count": len(plan["usable"]),
            "excluded": plan["excluded"],
            "conflicts": assembler.conflicts,
            "conflict_count": assembler.conflict_count,
            "format": format,
            "formats": {**(previous.get("formats") or {}), format: filename},
        }
        writer.close(manifest)
        finished = True
    finally:
        if not finished:
            try:
                writer.close()
            except Exception:
                pass
            tmp.unlink(missing_ok=True)
    os.replace(tmp, target)
    manifest_tmp = folder / "manifest.json.tmp"
    manifest_tmp.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(manifest_tmp, folder / "manifest.json")
    checkpoint(written, written, force=True)
    return {"build_key": plan["build_key"], "download_url": download_url(plan["build_key"], format), "manifest": manifest, "path": str(target), "format": format, "row_count": written, "reused": False}


def _label(plan, request):
    return f"{plan['schema']['name']} v{plan['schema']['rev']} · {len(plan['usable'])}문서 · {request['format'].upper()}"


def build(service, request, principal=None, wait=0):
    """POST /builds → {build_key, download_url, manifest, path}. 큰 입력은 작업(kind='build')이며 wait 안에 끝나지 않으면 manifest 없이 job을 돌려준다."""
    request = _request(request)
    principal = principal or service.principal
    units = load_units(service.root)
    # 계획과 실행은 같은 연결에서(TEMP 테이블·읽기 snapshot 공유).
    with service.db.connect() as conn:
        plan = _plan(service, conn, request)
        with _build_lock(plan["build_key"]):
            existing = _existing(service, plan, request["format"])
            if existing:
                return existing
            if len(plan["usable"]) > SYNC_DOCUMENT_LIMIT or plan["row_count"] > SYNC_ROW_LIMIT:
                job = service.jobs.submit("build", request, principal, f"{plan['build_key']}:{request['format']}", target_kind="build", target_id=plan["build_key"], label=_label(plan, request))
            else:
                job = None
                jid = service.jobs.record_sync("build", principal, "build", plan["build_key"], _label(plan, request), request)
                try:
                    result = _execute(service, conn, plan, request, units, lambda *a, **k: None)
                except Problem as exc:
                    service.jobs.finish_sync(jid, error=exc)
                    raise
                service.jobs.finish_sync(jid, {"build_key": result["build_key"], "download_url": result["download_url"], "row_count": result["row_count"]})
                return {**result, "job_id": jid}
    job = service.job_result(job, wait, principal)
    if job["state"] == "succeeded":
        return {**job["result"], "manifest": _read_manifest(export_folder(service, job["result"]["build_key"])), "path": str(export_folder(service, job["result"]["build_key"]) / FORMATS[request["format"]][0]), "job": job}
    if job["state"] in ("failed", "cancelled"):
        raise Problem(job["error_code"] or "BUILD_FAILED", job["error_message"] or "빌드에 실패했습니다.", 409 if job["state"] == "cancelled" else 422)
    return {"build_key": plan["build_key"], "download_url": download_url(plan["build_key"], request["format"]), "manifest": None, "path": None, "format": request["format"], "row_count": plan["row_count"], "job": job}


def execute_build(service, payload, principal, checkpoint):
    """작업 handler(kind='build'): 계획을 다시 세우고(값이 바뀌었을 수 있다) 산출물을 만든다. 결과는 요약만."""
    request = _request(payload)
    units = load_units(service.root)
    with service.db.connect() as conn:
        plan = _plan(service, conn, request)
        with _build_lock(plan["build_key"]):
            existing = _existing(service, plan, request["format"])
            if existing:
                return {"build_key": existing["build_key"], "download_url": existing["download_url"], "row_count": existing["row_count"], "format": request["format"], "reused": True}
            checkpoint(0, plan["row_count"], force=True)
            result = _execute(service, conn, plan, request, units, checkpoint)
    return {"build_key": result["build_key"], "download_url": result["download_url"], "row_count": result["row_count"], "format": request["format"], "reused": False}


run_build_job = execute_build


# ---------------------------------------------------------------------------- 다운로드


def manifest(service, build_key):
    """GET /builds/{build_key}/manifest."""
    found = _read_manifest(export_folder(service, build_key))
    if not found:
        raise Problem("NOT_FOUND", "빌드를 찾을 수 없습니다.", 404)
    return found


def output_path(service, build_key, format=None):
    """GET /builds/{build_key}/download → (path, filename, media_type). build_key는 16자 hex만 허용해 경로 조작을 막는다."""
    folder = export_folder(service, build_key)
    found = _read_manifest(folder)
    if not found:
        raise Problem("NOT_FOUND", "빌드를 찾을 수 없습니다.", 404)
    format = format or found.get("format") or "xlsx"
    if format not in FORMATS or format not in (found.get("formats") or {}):
        raise Problem("NOT_FOUND", "요청한 형식의 산출물이 없습니다.", 404)
    filename, media_type = FORMATS[format]
    path = folder / filename
    if not path.is_file():
        raise Problem("NOT_FOUND", "산출물 파일이 없습니다. 빌드를 다시 실행하세요.", 404)
    schema_key = re.sub(r"[^A-Za-z0-9_.\-]", "_", str((found.get("schema") or {}).get("key") or "data"))
    return path, f"{schema_key}-{build_key}.{format}", media_type
