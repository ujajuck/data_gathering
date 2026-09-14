"""Reader 계약(§3.2). 원본을 수정하지 않고, 경로는 Reader 안에서만 해석한다.

`XlsxReader`는 **평문 OOXML** 전용 기본 Reader다(권한 판정·원본 토큰 검증·구조 서명·매치·추출·렌더).
보호 문서는 운영자가 승인한 factory(`SCHEMA_READER_FACTORY`)의 Reader가 맡고, 그 Reader는 §3.5 해제 세션에서 얻은
평문 파일을 `plain_path`로 돌려주기만 하면 나머지 연산을 이 클래스에 그대로 위임할 수 있다.

컨테이너 판별은 `make_reader` **한 곳**에서만 한다(§3.5(2)) — `authorize`는 더 이상 `PK`를 보고 잠그지 않는다.
두 곳에서 판정하면 어댑터를 붙여도 계속 잠기는 화면이 남는다.
"""

from __future__ import annotations

import hashlib
import re
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from . import engine
from .db import Problem, norm
from .spec import decimal


def make_reader(root, provider, principal, source_ref=None):
    """컨테이너를 먼저 보고 Reader를 고른다(§3.5(2)).

    1. `provider != 'local-xlsx'` → 언제나 DRM Reader(어댑터). 없으면 403 `DRM_READER_REQUIRED`.
    2. `local-xlsx` + 평문 OOXML → `XlsxReader`.
    3. `local-xlsx` + 보호 문서 → DRM Reader. 없으면 403(문구가 무엇을 설정해야 하는지 말한다).
    4. `source_ref`가 없는 호출(방어) → `XlsxReader`.
    """
    from . import drm  # readers → drm 방향은 지연 import다(drm이 XlsxReader를 상속한다)

    root = Path(root)
    if provider == "local-xlsx":
        path = drm.source_path(root, source_ref) if source_ref is not None else None
        if path is None:
            # source_ref가 없거나 허용된 폴더에 그 파일이 없는 경우 — XlsxReader가 SOURCE_NOT_FOUND(404)로 말한다.
            return XlsxReader(root, principal)
        sniff = drm.sniff_path(path)
        if not sniff["protected"]:
            return XlsxReader(root, principal)
        if not drm.available():
            raise drm.reader_required(sniff["container"], source_ref)
    # 어댑터는 서버 설정만 읽는다. 요청 본문에서 모듈/함수 이름을 받지 않는다.
    return drm.load_factory()(root=root, provider=provider, principal=principal)


def release_sessions():
    """이 연산이 만든 해제본을 **연산이 끝나는 자리에서** 지운다(§3.5(3)).

    프로세스 종료(`atexit`)에 맡기면 forkserver 자식(`os._exit`)과 SIGTERM으로 죽는 경로에서 평문이 남는다.
    보호 문서를 한 번도 열지 않았으면 `schema.drm`이 아직 import되지도 않았으므로 아무 일도 하지 않는다."""
    module = sys.modules.get(__name__.rsplit(".", 1)[0] + ".drm")
    if module is not None:
        module.SESSIONS.release_all()


def file_hash(path):
    sha = hashlib.sha256()
    with path.open("rb") as source:
        for part in iter(lambda: source.read(1024 * 1024), b""):
            sha.update(part)
    return sha.hexdigest()


SIGNATURE_ROWS, SIGNATURE_COLS, SIGNATURE_SHEETS, SIGNATURE_TERMS = 30, 30, 64, 200


# 서명 헤더는 라벨로 판정한 셀만 담는다. 같은 열에서 이 값보다 길게 이어지는 값 옆 문자열은 레코드 값으로 본다.
SIGNATURE_COLUMN_RUN = 3
# lot-001, no.12, id_2024 처럼 접두어 뒤에 숫자가 이어지는 식별자 모양 문자열.
_ID_LIKE = re.compile(r"^[^\W\d_]{0,8}[-_./#:]?\d{2,}[^\W\d_]?$")


def cell_kind(value):
    """빈칸·숫자(숫자/날짜/숫자 문자열 등 문자열이 아닌 모든 값)·문자열을 구분한다. 숫자는 항상 데이터로 본다."""
    if value is None:
        return "empty"
    if not isinstance(value, str):
        return "number"
    if not value.strip():
        return "empty"
    try:
        decimal(value.strip())
    except Problem:
        return "text"
    return "number"


def header_term(value):
    """문자열 라벨의 정규형만 반환한다. 숫자·날짜·숫자 문자열과 식별자 모양 문자열은 데이터로 보고 제외한다."""
    if cell_kind(value) != "text":
        return None
    term = norm(value)
    if not term or len(term) > 64:
        return None
    compact = term.replace(" ", "")
    digits = sum(ch.isdigit() for ch in compact)
    if _ID_LIKE.match(compact) or digits * 2 > len(compact):
        return None
    return term


def signature_terms(grid, anchors=()):
    """창(grid)의 문자열 셀 중 라벨로 판정한 셀의 정규형을 정렬해 돌려준다.

    라벨 판정(하나라도 해당): 빈 행 다음의 첫 행(블록 머리), 병합 범위의 왼쪽 위 셀, 오른쪽/아래 이웃이 숫자인 셀,
    문자열만 있는 행(라벨 행)의 셀. 제외: 식별자·숫자 모양 문자열, 4행 이상 문자열만 이어지는 표의 둘째 행부터,
    같은 열에서 SIGNATURE_COLUMN_RUN칸을 넘겨 연속되는 값 옆 문자열(이름·자유 텍스트 같은 레코드 값).
    """
    anchors = {(int(r), int(c)) for r, c in anchors}
    height = len(grid)
    width = max((len(row) for row in grid), default=0)
    kinds = [
        [cell_kind(row[c]) if c < len(row) else "empty" for c in range(width)]
        for row in grid
    ]

    def kind(r, c):
        return kinds[r][c] if 0 <= r < height and 0 <= c < width else "empty"

    filled = [[k for k in row if k != "empty"] for row in kinds]
    block_head = [
        bool(filled[r]) and (r == 0 or not filled[r - 1]) for r in range(height)
    ]
    text_row = [
        len(filled[r]) >= 2 and all(k == "text" for k in filled[r])
        for r in range(height)
    ]
    label_row = [text_row[r] and _text_run_start(text_row, r) for r in range(height)]
    strong, weak = set(), {}
    for r in range(height):
        for c in range(width):
            if kinds[r][c] != "text" or header_term(grid[r][c]) is None:
                continue
            if block_head[r] or (r + 1, c + 1) in anchors or label_row[r]:
                strong.add((r, c))
            elif kind(r, c + 1) == "number" or kind(r + 1, c) == "number":
                weak.setdefault(c, []).append(r)
    chosen = set(strong)
    for c, rows in weak.items():
        # 같은 열에서 연속으로 이어지는 값 옆 문자열은 SIGNATURE_COLUMN_RUN칸까지만 라벨로 본다.
        runs, current = [], []
        for r in rows:
            if current and r != current[-1] + 1:
                runs.append(current)
                current = []
            current.append(r)
        runs.append(current)
        for group in runs:
            if len(group) <= SIGNATURE_COLUMN_RUN:
                chosen.update((r, c) for r in group)
    terms = {header_term(grid[r][c]) for r, c in chosen}
    return sorted(t for t in terms if t)[:SIGNATURE_TERMS]


def _text_run_start(text_row, r):
    # 문자열만 있는 행 묶음이 SIGNATURE_COLUMN_RUN행을 넘기면 첫 행만 라벨 행(표 머리)이고 나머지는 레코드다.
    start = r
    while start > 0 and text_row[start - 1]:
        start -= 1
    end = r
    while end + 1 < len(text_row) and text_row[end + 1]:
        end += 1
    return end - start + 1 <= SIGNATURE_COLUMN_RUN or r == start


def _sheets_of(wb):
    return [{"name": ws.title, "ordinal": n} for n, ws in enumerate(wb.worksheets)]


class XlsxReader:
    def __init__(self, root, principal):
        self.raw = (root / "data/raw").resolve()
        self.principal = principal

    def path(self, source_ref):
        """원본 경로. 토큰·크기 판정의 기준이고, 보호 문서라도 **원본**을 가리킨다."""
        path = (self.raw / source_ref).resolve()
        if not path.is_relative_to(self.raw) or not path.is_file():
            raise Problem(
                "SOURCE_NOT_FOUND", "허용된 원본 폴더에서 파일을 찾을 수 없습니다.", 404
            )
        return path

    def plain_path(self, source_ref, token):
        """평문으로 열 파일. 기본 Reader는 원본 그대로이고, DRM Reader가 해제본 경로로 덮어쓴다(§3.5(3))."""
        return self.path(source_ref)

    def authorize(self, source_ref, required="view"):
        self.path(source_ref)
        return {
            "can_view": True,
            "can_extract": True,
            "can_render_web": True,
            "can_cache_derivative": True,
            "native_render": False,
            "provider": "local-xlsx",
            "policy_revision": "local-filesystem-v1",
            "access_scope_key": self.principal,
            "expires_at": (datetime.now(timezone.utc) + timedelta(minutes=5)).isoformat(),
        }

    def _verify(self, source_ref, expected_token=None):
        """권한 재확인 + **원본** 해시. 연산 시작과 끝에 부른다(도중에 원본이 바뀌면 409).

        보호 문서도 토큰은 원본 해시다 — 해제본은 snapshot마다 달라질 수 있어 버전 기준이 될 수 없다."""
        self.authorize(source_ref)
        path = self.path(source_ref)
        if path.stat().st_size > 256 * 1024 * 1024:
            raise Problem(
                "SOURCE_SIZE_LIMIT",
                "기본 Reader의 256MB 원본 한도를 초과했습니다.",
                413,
            )
        token = file_hash(path)
        if expected_token is not None and token != expected_token:
            raise Problem(
                "SOURCE_VERSION_CHANGED",
                "원본이 변경되었습니다. 새 버전을 등록하세요.",
                409,
            )
        return token

    def _check(self, source_ref, expected_token=None):
        """(열어 읽을 파일, 원본 토큰). 보호 문서면 해제 세션의 평문 파일이 나온다(같은 snapshot은 한 번만 해제)."""
        token = self._verify(source_ref, expected_token)
        path = self.plain_path(source_ref, token)
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 30000
                or sum(i.file_size for i in entries) > 512 * 1024 * 1024
            ):
                raise Problem(
                    "EXPANDED_SIZE_LIMIT",
                    "압축 해제 예상 크기가 Reader 한도를 초과했습니다.",
                    413,
                )
        return path, token


    def describe(self, source_ref, profiles=()):
        """문서 메타·시트 목록·구조 서명(+ profiles가 있으면 matches[]). 워크북은 한 번만 열고 해시는 시작·끝 두 번만 계산한다."""
        try:
            return self._describe(source_ref, profiles)
        finally:
            release_sessions()

    def _describe(self, source_ref, profiles=()):
        path, token = self._check(source_ref)
        origin = self.path(source_ref)  # 이름·크기는 언제나 원본 기준(해제본 이름은 해시라 문서 이름이 아니다)
        capabilities = self.authorize(source_ref, "extract")
        wb = load_workbook(path, data_only=True, keep_links=False)
        try:
            sheets = [
                {
                    "name": s.title,
                    "ordinal": n,
                    "visibility": {"veryHidden": "very_hidden"}.get(s.sheet_state, s.sheet_state),
                    "estimated_rows": s.max_row,
                    "estimated_cols": s.max_column,
                }
                for n, s in enumerate(wb.worksheets)
            ]
            if len(sheets) > 2000:
                raise Problem("SHEET_LIMIT", "시트 수가 Reader 한도를 초과했습니다.", 413)
            result = {
                "token": token,
                "filename": origin.name,
                "author": wb.properties.creator,
                "authored_at": wb.properties.created.isoformat() if wb.properties.created else None,
                "excel_date_system": "1904" if wb.epoch.year == 1904 else "1900",
                "byte_size": origin.stat().st_size,
                "sheets": sheets,
                # 등록 경로가 별도 authorize 호출 없이 재사용할 수 있도록 추출 권한 기준으로 돌려준다.
                "capabilities": capabilities,
            }
            # 구조 서명 실패는 등록을 막지 않는다(필드 생략).
            try:
                result["signature"] = self._signature_of(wb, token)
            except Exception:
                pass
            if profiles:
                result["matches"] = self._matches(wb, profiles)
        finally:
            wb.close()
        if file_hash(origin) != token:
            raise Problem("SOURCE_VERSION_CHANGED", "읽는 동안 원본이 변경되었습니다.", 409)
        return result

    @staticmethod
    def _signature_of(wb, token, rows=SIGNATURE_ROWS, cols=SIGNATURE_COLS):
        """구조 서명을 이미 열린 워크북에서 계산한다(파일을 다시 열지 않는다)."""
        sheets = []
        for n, ws in enumerate(wb.worksheets[:SIGNATURE_SHEETS]):
            grid = [
                list(row)
                for row in ws.iter_rows(min_row=1, max_row=min(rows, ws.max_row or 1), min_col=1, max_col=min(cols, ws.max_column or 1), values_only=True)
            ]
            merged = [m for m in ws.merged_cells.ranges if m.min_row <= rows and m.min_col <= cols]
            sheets.append(
                {
                    "name": ws.title,
                    "ordinal": n,
                    "visibility": {"veryHidden": "very_hidden"}.get(ws.sheet_state, ws.sheet_state),
                    "dims": {"rows": ws.max_row, "cols": ws.max_column},
                    "headers": signature_terms(grid, {(m.min_row, m.min_col) for m in merged}),
                    "merges": sorted(str(m) for m in merged)[:SIGNATURE_TERMS],
                }
            )
        return {"token": token, "sheets": sheets}

    @staticmethod
    def _matches(wb, profiles):
        if not hasattr(wb, "worksheets"):
            # 경로를 받은 호출(match): 여기서 열고 닫는다.
            opened = load_workbook(wb, data_only=True, keep_links=False)
            try:
                return XlsxReader._matches(opened, profiles)
            finally:
                opened.close()
        sheets = _sheets_of(wb)
        out = []
        for profile in profiles:
            item = {"profile_id": profile.get("profile_id"), "profile_rev": profile.get("profile_rev")}
            try:
                item.update(engine.match_profile(profile["canonical"], wb, sheets, profile.get("reference"), profile.get("profile_rev")))
            except Problem as exc:
                item.update(bindings={}, compatibility="incompatible", match_signature=None, match_signature_json=None, missing=["*"], resolved={}, error=exc.code)
            out.append(item)
        return out

    def match(self, source_ref, expected_token, profiles):
        try:
            path, token = self._check(source_ref, expected_token)
            matches = self._matches(path, profiles)
            self._verify(source_ref, token)
            return matches
        finally:
            release_sessions()

    def match_specs(self, source_ref, expected_token, specs, bindings_hint):
        try:
            path, token = self._check(source_ref, expected_token)
            wb = load_workbook(path, data_only=True, keep_links=False)
            try:
                result = engine.match_specs(specs, wb, _sheets_of(wb), bindings_hint)
            finally:
                wb.close()
            self._verify(source_ref, token)
            return result
        finally:
            release_sessions()

    def extract(self, source_ref, expected_token, specs, bindings):
        # 스트림 연산의 finally는 소비자가 중간에 끊어도(GeneratorExit) 돈다 — 해제본이 남지 않는다.
        try:
            path, token = self._check(source_ref, expected_token)
            raw = load_workbook(path, data_only=False, keep_links=False)
            cached = load_workbook(path, data_only=True, keep_links=False)
            try:
                yield from engine.extract(raw, cached, bindings, specs)
            finally:
                raw.close()
                cached.close()
            self._verify(source_ref, token)
            yield {"type": "verified", "token": token}
        finally:
            release_sessions()

    def render(self, source_ref, expected_token, sheet_name, r1=1, c1=1, rows=2000, cols=200):
        try:
            path, token = self._check(source_ref, expected_token)
            # 렌더러는 별도 모듈(schema/render/renderer.py)이며 여기서 경로를 넘긴다(호출자는 경로를 모른다).
            from .render.renderer import render_events

            yield from render_events(path, sheet_name, token, r1, c1, rows, cols)
            self._verify(source_ref, token)
            yield {"type": "verified", "token": token}
        finally:
            release_sessions()
