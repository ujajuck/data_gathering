"""v3 Reader 계약(§3.2). `XlsxReader`는 v2를 상속해 region/area/_check/describe/signature를 재사용하고
`extract`를 v3 엔진으로 바꾸며 `match`·`match_specs`·`render`를 더한다. 경로는 Reader 안에서만 해석한다."""

from __future__ import annotations

import importlib
from pathlib import Path

from openpyxl import load_workbook

from kg.v2.db import Problem as V2Problem
from kg.v2.readers import SIGNATURE_COLS, SIGNATURE_ROWS, SIGNATURE_SHEETS, SIGNATURE_TERMS, signature_terms
from kg.v2.readers import XlsxReader as V2XlsxReader
from kg.v2.readers import file_hash  # noqa: F401  (v2 계약 테스트·서비스가 같은 이름을 쓴다)

from . import engine
from .db import Problem
from .jobs import env


def make_reader(root, provider, principal):
    if provider == "local-xlsx":
        return XlsxReader(Path(root), principal)
    factory = env("READER_FACTORY", "")
    if not factory or ":" not in factory:
        raise Problem("DRM_READER_REQUIRED", "이 문서의 보안 읽기 어댑터가 연결되지 않았습니다.", 403)
    module, function = factory.split(":", 1)
    # factory는 서버 설정만 읽는다. 요청 본문에서 모듈/함수 이름을 받지 않는다.
    return getattr(importlib.import_module(module), function)(root=Path(root), provider=provider, principal=principal)


def _sheets_of(wb):
    return [{"name": ws.title, "ordinal": n} for n, ws in enumerate(wb.worksheets)]


class XlsxReader(V2XlsxReader):
    def _check(self, source_ref, expected_token=None):
        try:
            return super()._check(source_ref, expected_token)
        except V2Problem as exc:
            # v2 Problem은 별개 클래스라 v3 서비스/격리 프로세스가 잡지 못한다.
            raise Problem(exc.code, exc.message, exc.status) from None

    def authorize(self, source_ref, required="view"):
        try:
            return super().authorize(source_ref, required)
        except V2Problem as exc:
            if exc.code == "DRM_READER_REQUIRED" and source_ref.lower().endswith(".xls"):
                # 구형 .xls(OLE2)와 암호화된 OOXML은 매직이 같아 구분되지 않는다(둘 다 PK가 아니다).
                # 코드·상태는 계약 §4.1 그대로 두고(잠김), 문구만 두 경우를 모두 알려 사용자가 헤매지 않게 한다.
                raise Problem(
                    exc.code,
                    "구형 .xls 형식이거나 암호화된 문서입니다. .xlsx로 저장한 뒤 등록하거나, 암호화 문서는 승인된 보안 읽기 어댑터로 접근하세요.",
                    exc.status,
                ) from None
            raise Problem(exc.code, exc.message, exc.status) from None

    def describe(self, source_ref, profiles=()):
        """v2 describe와 같은 결과(+ profiles가 있으면 matches[]). 워크북은 한 번만 열고 해시는 시작·끝 두 번만 계산한다."""
        path, token = self._check(source_ref)
        capabilities = self.authorize(source_ref, "extract")
        wb = load_workbook(path, data_only=True, keep_links=False)
        try:
            sheets = [
                {
                    "name": s.title,
                    "ordinal": n,
                    "visibility": {"veryHidden": "very_hidden"}.get(s.sheet_state, s.sheet_state),
                    "estimated_rows": s.max_row,
                    "estimated_cols": s.max_column,
                }
                for n, s in enumerate(wb.worksheets)
            ]
            if len(sheets) > 2000:
                raise Problem("SHEET_LIMIT", "시트 수가 Reader 한도를 초과했습니다.", 413)
            result = {
                "token": token,
                "filename": path.name,
                "author": wb.properties.creator,
                "authored_at": wb.properties.created.isoformat() if wb.properties.created else None,
                "excel_date_system": "1904" if wb.epoch.year == 1904 else "1900",
                "byte_size": path.stat().st_size,
                "sheets": sheets,
                # 등록 경로가 별도 authorize 호출 없이 재사용할 수 있도록 추출 권한 기준으로 돌려준다.
                "capabilities": capabilities,
            }
            # 구조 서명 실패는 등록을 막지 않는다(필드 생략).
            try:
                result["signature"] = self._signature_of(wb, token)
            except Exception:
                pass
            if profiles:
                result["matches"] = self._matches(wb, profiles)
        finally:
            wb.close()
        if file_hash(path) != token:
            raise Problem("SOURCE_VERSION_CHANGED", "읽는 동안 원본이 변경되었습니다.", 409)
        return result

    @staticmethod
    def _signature_of(wb, token, rows=SIGNATURE_ROWS, cols=SIGNATURE_COLS):
        """v2 `_signature`와 같은 결과를 이미 열린 워크북에서 계산한다(파일을 다시 열지 않는다)."""
        sheets = []
        for n, ws in enumerate(wb.worksheets[:SIGNATURE_SHEETS]):
            grid = [
                list(row)
                for row in ws.iter_rows(min_row=1, max_row=min(rows, ws.max_row or 1), min_col=1, max_col=min(cols, ws.max_column or 1), values_only=True)
            ]
            merged = [m for m in ws.merged_cells.ranges if m.min_row <= rows and m.min_col <= cols]
            sheets.append(
                {
                    "name": ws.title,
                    "ordinal": n,
                    "visibility": {"veryHidden": "very_hidden"}.get(ws.sheet_state, ws.sheet_state),
                    "dims": {"rows": ws.max_row, "cols": ws.max_column},
                    "headers": signature_terms(grid, {(m.min_row, m.min_col) for m in merged}),
                    "merges": sorted(str(m) for m in merged)[:SIGNATURE_TERMS],
                }
            )
        return {"token": token, "sheets": sheets}

    @staticmethod
    def _matches(wb, profiles):
        if not hasattr(wb, "worksheets"):
            # 경로를 받은 호출(match): 여기서 열고 닫는다.
            opened = load_workbook(wb, data_only=True, keep_links=False)
            try:
                return XlsxReader._matches(opened, profiles)
            finally:
                opened.close()
        sheets = _sheets_of(wb)
        out = []
        for profile in profiles:
            item = {"profile_id": profile.get("profile_id"), "profile_rev": profile.get("profile_rev")}
            try:
                item.update(engine.match_profile(profile["canonical"], wb, sheets, profile.get("reference"), profile.get("profile_rev")))
            except Problem as exc:
                item.update(bindings={}, compatibility="incompatible", match_signature=None, match_signature_json=None, missing=["*"], resolved={}, error=exc.code)
            out.append(item)
        return out

    def match(self, source_ref, expected_token, profiles):
        path, token = self._check(source_ref, expected_token)
        matches = self._matches(path, profiles)
        self._check(source_ref, token)
        return matches

    def match_specs(self, source_ref, expected_token, specs, bindings_hint):
        path, token = self._check(source_ref, expected_token)
        wb = load_workbook(path, data_only=True, keep_links=False)
        try:
            result = engine.match_specs(specs, wb, _sheets_of(wb), bindings_hint)
        finally:
            wb.close()
        self._check(source_ref, token)
        return result

    def extract(self, source_ref, expected_token, specs, bindings):
        path, token = self._check(source_ref, expected_token)
        raw = load_workbook(path, data_only=False, keep_links=False)
        cached = load_workbook(path, data_only=True, keep_links=False)
        try:
            yield from engine.extract(raw, cached, bindings, specs)
        finally:
            raw.close()
            cached.close()
        self._check(source_ref, token)
        yield {"type": "verified", "token": token}

    def render(self, source_ref, expected_token, sheet_name, r1=1, c1=1, rows=2000, cols=200):
        path, token = self._check(source_ref, expected_token)
        # 렌더러는 별도 모듈(kg/v3/render/renderer.py)이며 여기서 경로를 넘긴다(호출자는 경로를 모른다).
        from .render.renderer import render_events

        yield from render_events(path, sheet_name, token, r1, c1, rows, cols)
        self._check(source_ref, token)
        yield {"type": "verified", "token": token}
