"""openpyxl 간략 렌더(§3.2 `render` 스트림). 경로는 Reader가 넘기며 여기서는 파일을 열어 이벤트만 만든다.

이벤트 순서: meta → band(100행, 6MB 초과 시 50/25행) → image → (Reader가 verified를 덧붙인다).
셀·병합·이미지는 각각 한 번씩만 훑는다(O(셀)). 상한(2,000×200 / 비어 있지 않은 셀 200,000 / 직렬화 24MB)에
먼저 닿으면 밴드 경계에서 잘라 `truncated`·`rendered_bounds`로 알린다.
"""

from __future__ import annotations

import base64
import hashlib

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import column_index_from_string

from kg.v2.spec import bounds as v2_bounds
from kg.v2.db import Problem as V2Problem

from ..db import Problem, dump
from . import RENDERER_VERSION

MAX_ROWS, MAX_COLS = 2000, 200
MAX_CELLS = 200_000
MAX_BYTES = 24 * 1024 * 1024
BAND_ROWS = 100
BAND_BYTES = 6 * 1024 * 1024
IMAGE_BYTES = 2 * 1024 * 1024
# 작은 시트도 기본 창(A1:Z60)이 잘리지 않도록 기하는 최소 60행×26열을 준다(셀은 있는 것만).
MIN_ROWS, MIN_COLS = 60, 26
DEFAULT_COL_PX, DEFAULT_ROW_PX = 96, 20
IMAGE_EXT = {"png": "png", "jpeg": "jpeg", "jpg": "jpg", "gif": "gif"}


def _color(value, fallback):
    return (
        "#" + value.rgb[-6:].lower()
        if value is not None and value.type == "rgb" and isinstance(value.rgb, str)
        else fallback
    )


def _style(cell):
    font, alignment = cell.font, cell.alignment
    return {
        "background": _color(cell.fill.fgColor, "#ffffff") if cell.fill.patternType else "#ffffff",
        "color": _color(font.color, "#172b36"),
        "fontSize": round((font.sz or 11) * 96 / 72, 2),
        "fontFamily": font.name or "sans-serif",
        "fontWeight": 700 if font.b else 400,
        "fontStyle": "italic" if font.i else "normal",
        "textAlign": alignment.horizontal or "left",
        "verticalAlign": alignment.vertical or "bottom",
        "whiteSpace": "pre-wrap" if alignment.wrap_text else "nowrap",
    }


DEFAULT_STYLE = {
    "background": "#ffffff",
    "color": "#172b36",
    "fontSize": round(11 * 96 / 72, 2),
    "fontFamily": "Calibri",
    "fontWeight": 400,
    "fontStyle": "normal",
    "textAlign": "left",
    "verticalAlign": "bottom",
    "whiteSpace": "nowrap",
}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return str(value)


def _dimensions(ws, r1, c1, r2, c2):
    """열 너비·행 높이(px)를 원점(r1,c1)=0에서 누적한다. 설정된 치수만 순회하므로 O(치수 수 + 범위)."""
    widths, heights = {}, {}
    for letter, dim in ws.column_dimensions.items():
        lo = dim.min or column_index_from_string(letter)
        hi = dim.max or column_index_from_string(letter)
        for col in range(max(lo, c1), min(hi, c2) + 1):
            widths[col] = 0 if dim.hidden else max(12, (dim.width or 13) * 7 + 5)
    for row, dim in ws.row_dimensions.items():
        if r1 <= row <= r2:
            heights[row] = 0 if dim.hidden else (dim.height or 15) * 96 / 72
    columns, x = [], 0.0
    for c in range(c1, c2 + 1):
        w = widths.get(c, DEFAULT_COL_PX)
        columns.append({"index": c, "x": round(x, 2), "width": round(w, 2)})
        x += w
    rows, y = [], 0.0
    for r in range(r1, r2 + 1):
        h = heights.get(r, DEFAULT_ROW_PX)
        rows.append({"index": r, "y": round(y, 2), "height": round(h, 2)})
        y += h
    return rows, columns, round(x, 2), round(y, 2)


def _split_band(r1, r2, cells, size):
    """직렬화가 6MB를 넘는 밴드를 50행 → 25행으로 나눈다. 25행 이하는 더 나누지 않는다."""
    serialized = dump({"r1": r1, "r2": r2, "cells": cells})
    if len(serialized.encode()) <= BAND_BYTES or size <= 25:
        return [(r1, r2, cells, serialized)]
    half = 50 if size == BAND_ROWS else 25
    out = []
    for start in range(r1, r2 + 1, half):
        end = min(start + half - 1, r2)
        part = [c for c in cells if c["r1"] <= end and c["r2"] >= start]
        out.extend(_split_band(start, end, part, half))
    return out


def _freeze(ws):
    if not ws.freeze_panes:
        return None
    try:
        r, c, _, _ = v2_bounds(str(ws.freeze_panes))
    except V2Problem:
        return None
    return {"rows": r - 1, "cols": c - 1}


def render_events(path, sheet_name, token, r1=1, c1=1, rows=MAX_ROWS, cols=MAX_COLS):
    if not (isinstance(r1, int) and isinstance(c1, int) and r1 >= 1 and c1 >= 1):
        raise Problem("INVALID_RANGE", "렌더 시작 위치가 유효하지 않습니다.")
    rows = max(1, min(int(rows), MAX_ROWS))
    cols = max(1, min(int(cols), MAX_COLS))
    wb = load_workbook(path, data_only=True, keep_links=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise Problem("SHEET_NOT_FOUND", "시트를 찾을 수 없습니다.", 404)
        ws = wb[sheet_name]
        yield from _render_sheet(ws, token, r1, c1, rows, cols)
    finally:
        wb.close()


def _render_sheet(ws, token, r1, c1, rows, cols):
    cap_r2, cap_c2 = r1 + rows - 1, c1 + cols - 1
    estimated_rows, estimated_cols = ws.max_row or 0, ws.max_column or 0
    truncated = estimated_rows > cap_r2 or estimated_cols > cap_c2

    # 병합: 앵커 사전 한 번(셀마다 병합 목록을 훑지 않는다). 상한 범위와 교차하는 것만.
    anchors, merges = {}, []
    for merged in ws.merged_cells.ranges:
        if merged.max_row < r1 or merged.min_row > cap_r2 or merged.max_col < c1 or merged.min_col > cap_c2:
            continue
        anchors[(merged.min_row, merged.min_col)] = (merged.max_row, merged.max_col)

    styles, style_index = [DEFAULT_STYLE], {tuple(sorted(DEFAULT_STYLE.items())): 0}

    def style_of(cell):
        style = _style(cell)
        key = tuple(sorted(style.items()))
        index = style_index.get(key)
        if index is None:
            index = style_index[key] = len(styles)
            styles.append(style)
        return index

    # 셀: 비어 있는(값도 서식도 없는) 셀은 생략. 병합 앵커는 항상 담고 걸친 밴드마다 복제한다(창은 밴드 ≤3개만 읽으므로
    # 앵커가 다른 밴드에 있어도 찾을 수 있고, 창에서 (r1,c1)로 중복 제거).
    bands = {}
    source = getattr(ws, "_cells", None)
    iterator = (
        source.values()
        if source is not None
        else (cell for row in ws.iter_rows(min_row=r1, max_row=min(cap_r2, estimated_rows or r1)) for cell in row)
    )
    content_r2, content_c2 = r1, c1
    for cell in iterator:
        if isinstance(cell, MergedCell):
            continue
        r, c = cell.row, cell.column
        if r < r1 or r > cap_r2 or c < c1 or c > cap_c2:
            continue
        extent = anchors.get((r, c))
        text = _text(cell.value)
        if cell.has_style:
            s = style_of(cell)
        else:
            s = 0
        if not text and s == 0 and extent is None:
            continue
        r2, c2 = (min(extent[0], cap_r2), min(extent[1], cap_c2)) if extent else (r, c)
        item = {"r1": r, "c1": c, "r2": r2, "c2": c2, "text": text, "s": s}
        content_r2, content_c2 = max(content_r2, r2), max(content_c2, c2)
        for band in range((r - r1) // BAND_ROWS, (r2 - r1) // BAND_ROWS + 1):
            bands.setdefault(band, []).append(item)

    # 밴드를 순서대로 직렬화하며 셀 수·바이트 상한을 검사한다(상한에 닿으면 이전 밴드 경계에서 자른다).
    last_band = max(bands) if bands else -1
    emitted, cell_count, byte_count, capped = [], 0, 0, False
    for band in range(last_band + 1):
        b1 = r1 + band * BAND_ROWS
        b2 = min(b1 + BAND_ROWS - 1, cap_r2)
        cells = sorted(bands.get(band, []), key=lambda item: (item["r1"], item["c1"]))
        parts = _split_band(b1, b2, cells, BAND_ROWS)
        own = sum(1 for item in cells if item["r1"] >= b1)  # 복제된 앵커는 앵커 밴드에서만 센다
        size = sum(len(serialized.encode()) for _, _, _, serialized in parts)
        if emitted and (cell_count + own > MAX_CELLS or byte_count + size > MAX_BYTES):
            capped = True
            break
        cell_count += own
        byte_count += size
        emitted.extend(parts)
    if capped:
        truncated = True
        rendered_r2 = emitted[-1][1]
        # 잘린 경계 너머로 이어지는 병합은 경계에서 끊는다.
        for _, _, cells, _ in emitted:
            for item in cells:
                item["r2"] = min(item["r2"], rendered_r2)
    else:
        rendered_r2 = min(cap_r2, max(content_r2, r1 + MIN_ROWS - 1, estimated_rows))
    rendered_c2 = min(cap_c2, max(content_c2, c1 + MIN_COLS - 1, estimated_cols))
    # 밴드 경계는 렌더 범위 안으로 맞춘다(셀은 모두 rendered_r2 이하이므로 잘리는 것은 없다).
    emitted = [(b1, min(b2, rendered_r2), cells, s) for b1, b2, cells, s in emitted if b1 <= rendered_r2]
    rows_out, columns_out, width, height = _dimensions(ws, r1, c1, rendered_r2, rendered_c2)
    col_x = {col["index"]: col["x"] for col in columns_out}
    row_y = {row["index"]: row["y"] for row in rows_out}

    for (mr1, mc1), (mr2, mc2) in anchors.items():
        if mr1 <= rendered_r2 and mc1 <= rendered_c2:
            merges.append([mr1, mc1, min(mr2, rendered_r2), min(mc2, rendered_c2)])
    merges.sort()

    images = []
    for picture in getattr(ws, "_images", ()):
        anchor = getattr(picture.anchor, "_from", None)
        ext = IMAGE_EXT.get(str(getattr(picture, "format", "")).lower())
        if anchor is None or ext is None:
            continue
        col, row = anchor.col + 1, anchor.row + 1
        if col > rendered_c2 or row > rendered_r2 or col < c1 or row < r1:
            continue
        x = col_x[col] + anchor.colOff / 9525
        y = row_y[row] + anchor.rowOff / 9525
        try:
            data = picture._data()
        except Exception:
            continue
        if not data or len(data) > IMAGE_BYTES:
            continue
        images.append(
            {
                "type": "image",
                "asset_id": hashlib.sha256(data).hexdigest(),
                "ext": ext,
                "x": round(x, 2),
                "y": round(y, 2),
                "width": float(picture.width or 0),
                "height": float(picture.height or 0),
                "bytes": base64.b64encode(data).decode(),
            }
        )

    yield {
        "type": "meta",
        "renderer_version": RENDERER_VERSION,
        "mode": "simplified",
        "layout_revision": token,
        "sheet": ws.title,
        "r1": r1,
        "c1": c1,
        "rows": rows_out,
        "columns": columns_out,
        "styles": styles,
        "merges": merges,
        "freeze": _freeze(ws),
        "width": width,
        "height": height,
        "estimated_rows": estimated_rows,
        "estimated_cols": estimated_cols,
        "truncated": truncated,
        "rendered_bounds": {"r2": rendered_r2, "c2": rendered_c2},
        "band_rows": BAND_ROWS,
        "cell_count": cell_count,
    }
    for b1, b2, cells, _ in emitted:
        yield {"type": "band", "r1": b1, "r2": b2, "cells": cells}
    yield from images
