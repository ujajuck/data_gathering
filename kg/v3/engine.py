"""v3 실행 엔진(계약 §3.1·§3.3) — 열린 openpyxl 워크북 위에서 동작하는 순수 함수.

- resolve_area/resolve_anchor: range·find(texts|regex)·relative(anchor|key 기준)·이름 앵커(composite 포함) 해결.
- extract: effective_spec 목록 → 'group'/'values' 이벤트 스트림(v2 XlsxReader.extract의 리스트/행렬/병합/stop/결합/
  업무키 로직 + regex·앵커·relative.anchor·relations·split_delimiter).
- match_profile/match_specs: 시트 바인딩 + 모든 role의 영역 해결 + 매치 서명.
파일 경로·DB를 만지지 않는다. 텍스트 검색은 저장된 값(cached) 워크북에서 한다(수식 라벨도 결과값으로 비교).
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from decimal import localcontext

from openpyxl.utils.cell import column_index_from_string

from kg.v2.db import Problem as V2Problem
from kg.v2.db import norm
from kg.v2.readers import XlsxReader as _V2Reader
from kg.v2.spec import address, decimal, decimal_text, typed

from .db import Problem, dump
from .normalization import prepare
from .profile import DEFAULT_WITHIN, ROLES, bounds, compile_profile, compile_regex, relation_order

region = _V2Reader.region
area_of = _V2Reader.area
BATCH = 200
KEY_AREA_LIMIT = 10000


def regex_text(value):
    """regex 매칭 대상: NFKC + 공백 축약, casefold 없음, 앞 1,024자."""
    return " ".join(unicodedata.normalize("NFKC", str(value)[:1024]).split())


def _window(ws, within):
    # 시트 밖의 빈 셀을 만들지 않도록 실제 사용 영역으로 잘라 O(창)으로 훑는다.
    r1, c1, r2, c2 = bounds(within)
    return r1, c1, min(r2, ws.max_row or 1), min(c2, ws.max_column or 1)


def find_cells(ws, find):
    """find(texts|regex, within, occurrence?) → 병합 단위로 중복 제거한 region 목록(모든 히트)."""
    r1, c1, r2, c2 = _window(ws, find.get("within", DEFAULT_WITHIN))
    if "texts" in find:
        terms = {norm(t) for t in find["texts"]}

        def test(value):
            return norm(value) in terms

    else:
        pattern = compile_regex(find["regex"])

        def test(value):
            return pattern.search(regex_text(value)) is not None

    hits, seen = [], set()
    if r1 <= r2 and c1 <= c2:
        for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
            for cell in row:
                if cell.value is not None and test(cell.value):
                    found = region(ws, cell.row, cell.column)
                    if found["range"] not in seen:
                        seen.add(found["range"])
                        hits.append(found)
    if "occurrence" in find:
        index = find["occurrence"]
        hits = hits[index : index + 1] if isinstance(index, int) and index >= 0 else []
    return hits


def _single(ws, find, name):
    hits = find_cells(ws, find)
    if not hits:
        raise Problem("ANCHOR_NOT_FOUND", f"앵커 {name}: {ws.title} 시트에서 기준 셀을 찾을 수 없습니다.")
    if len(hits) > 1:
        raise Problem(
            "AMBIGUOUS_ANCHOR",
            f"앵커 {name}: {ws.title} 시트에 기준 셀이 여러 곳({', '.join(h['range'] for h in hits[:5])})입니다. occurrence를 지정하세요.",
        )
    return hits[0]


def resolve_anchor(wb, anchor, sheet_name, cache=None):
    """인라인 앵커({find} | {all_of:[{find},...]}) → region. composite는 멤버의 bounding box + members."""
    key = (dump(anchor), sheet_name)
    if cache is not None and key in cache:
        return cache[key]
    ws = wb[sheet_name]
    name = anchor.get("anchor_name") or "(이름 없음)"
    if "all_of" in anchor:
        members = [_single(ws, member["find"], name) for member in anchor["all_of"]]
        box = area_of(
            ws,
            (
                min(m["r1"] for m in members),
                min(m["c1"] for m in members),
                max(m["r2"] for m in members),
                max(m["c2"] for m in members),
            ),
        )
        box["members"] = members
        result = box
    else:
        result = _single(ws, anchor["find"], name)
    if cache is not None:
        cache[key] = result
    return result


def _is_anchor(area):
    return "all_of" in area or ("find" in area and "anchor_name" in area)


def resolve_area(wb, area, bindings, primary_role, primary_sheet, base=None, cache=None):
    """area → region 목록. base는 이 rule의 key.areas[0] 해결 영역(relative 기준, relative.anchor가 없을 때)."""
    role = area["sheet_role"]
    if role == primary_role:
        names = [primary_sheet]
    elif role in bindings:
        names = bindings[role]
    else:
        raise Problem("INVALID_SHEET_BINDING", f"시트 역할 {role}이 연결되지 않았습니다.")
    resolved = []
    for name in names:
        if name not in wb.sheetnames:
            raise Problem("SHEET_NOT_FOUND", f"시트 {name}을 찾을 수 없습니다.", 404)
        ws = wb[name]
        if "range" in area:
            resolved.append(area_of(ws, bounds(area["range"])))
        elif "relative" in area:
            spec = area["relative"]
            anchor = spec.get("anchor")
            origin = resolve_anchor(wb, anchor, name, cache) if isinstance(anchor, dict) else base
            if not origin:
                raise Problem("ANCHOR_NOT_FOUND", "상대 범위의 기준 영역을 해결할 수 없습니다.")
            r, c = origin["r1"] + spec.get("row", 0), origin["c1"] + spec.get("col", 0)
            if r < 1 or c < 1:
                raise Problem("INVALID_RANGE", "상대 범위가 시트 밖(1행/1열 앞)을 가리킵니다.")
            box = bounds(address(r, c, r + spec.get("rows", 1) - 1, c + spec.get("cols", 1) - 1))
            resolved.append(area_of(ws, box))
        elif _is_anchor(area):
            resolved.append(resolve_anchor(wb, area, name, cache))
        else:
            resolved.extend(find_cells(ws, area["find"]))
    return resolved


# ---------------------------------------------------------------------------- extract


class PairIndex:
    """relations 짝 찾기용 행/열 색인. 후보는 (item_index, 삽입 순서)로 정렬해 '첫 item'을 고른다."""

    def __init__(self):
        self.rows, self.cols, self.n = {}, {}, 0

    def add(self, item_index, record_key, box):
        entry = (item_index, self.n, record_key, box)
        self.n += 1
        for r in range(box["r1"], box["r2"] + 1):
            self.rows.setdefault(r, []).append(entry)
        for c in range(box["c1"], box["c2"] + 1):
            self.cols.setdefault(c, []).append(entry)

    @staticmethod
    def _first(candidates):
        return min(candidates, key=lambda e: e[:2]) if candidates else None

    def find(self, relation, box):
        kind = relation["kind"]
        if kind == "same_row":
            return self._first(self.rows.get(box["r1"], []))
        if kind == "same_column":
            return self._first(self.cols.get(box["c1"], []))
        r, c = box["r1"] + relation["row"], box["c1"] + relation["col"]
        return self._first([e for e in self.rows.get(r, []) if e[3]["c1"] <= c <= e[3]["c2"]])


def _key_labels(raw, cached, parts, separator):
    labels, seen = [], set()
    for part in parts:
        if (part["r2"] - part["r1"] + 1) * (part["c2"] - part["c1"] + 1) > KEY_AREA_LIMIT:
            raise Problem("KEY_AREA_LIMIT", "키 영역은 10,000셀 이하로 지정하세요.", 413)
        ws, cv = raw[part["sheet"]], cached[part["sheet"]]
        for row in ws.iter_rows(min_row=part["r1"], max_row=part["r2"], min_col=part["c1"], max_col=part["c2"]):
            for cell in row:
                found = region(ws, cell.row, cell.column)
                identity = (ws.title, found["range"])
                if identity in seen:
                    continue
                seen.add(identity)
                value = cv.cell(found["r1"], found["c1"]).value
                if value is not None:
                    labels.append(str(value))
    observed = separator.join(labels)
    if not observed.strip() or len(observed) > 32768:
        raise Problem("INVALID_KEY_TEXT", "키가 비어 있거나 너무 긴 영역입니다.")
    return observed


def _cell_state(cell, cached_cell):
    formula = str(cell.value) if cell.data_type == "f" else None
    value = cached_cell.value if formula else cell.value
    if formula and value is None:
        return formula, value, "error", "cached_missing"
    if value is None:
        return formula, value, "empty", "cached_unknown_age" if formula else "none"
    if cached_cell.data_type == "e":
        return formula, value, "error", "cached_unknown_age" if formula else "none"
    return formula, value, "present", "cached_unknown_age" if formula else "none"


def build_item(raw, cached, found, spec, selected):
    """값 셀 하나 → item(record_key 제외). 변환 실패는 예외가 아니라 value_state='error'다."""
    ws, cv = raw[found["sheet"]], cached[found["sheet"]]
    r, c = found["r1"], found["c1"]
    cell, cached_cell = ws.cell(r, c), cv.cell(r, c)
    formula, value, state, formula_state = _cell_state(cell, cached_cell)
    vspec = spec["value_spec"]
    target = vspec.get("type", "text")
    kind, text, inline_unit, ratio = target, None, None, False
    if state == "present":
        try:
            prepared, inline_unit, ratio = prepare(value, vspec.get("normalization") or {}, target)
            if prepared is None:
                state = "null"
            else:
                kind, text = typed(prepared, vspec)
        except (Problem, V2Problem):
            # 원문은 raw_text에 남기고 값만 오류로 표시한다. 실행 전체를 멈추지 않는다.
            state, kind, text = "error", target, None
    regions = {"value": [found], **{k: v for k, v in selected.items() if k in ("key", "unit", "context")}}
    region_unit = (
        " ".join(str(cached[a["sheet"]].cell(a["r1"], a["c1"]).value or "") for a in selected.get("unit", []))
        or None
    )
    if inline_unit and region_unit and norm(inline_unit) != norm(region_unit):
        raise Problem("UNIT_MISMATCH", f"값 셀의 단위는 {inline_unit}, 단위 영역의 단위는 {region_unit}입니다.")
    unit_raw = inline_unit or region_unit or vspec.get("source_unit") or vspec.get("unit")
    target_unit = vspec.get("unit") or unit_raw
    if ratio:
        if target_unit not in (None, "", "1"):
            raise Problem("UNIT_MISMATCH", f"비율 변환의 목표 단위는 없음 또는 1, 지정한 단위는 {target_unit}입니다.")
        target_unit = target_unit or None
    if norm(unit_raw or "") != norm(target_unit or "") and not ratio:
        affine = (vspec.get("normalization") or {}).get("operation") == "affine"
        if not affine or norm(vspec.get("source_unit") or "") != norm(unit_raw or ""):
            raise Problem(
                "UNIT_MISMATCH",
                f"원본 단위는 {unit_raw or '없음'}, 목표 단위는 {target_unit or '없음'}입니다. 단위 변환에는 원본 단위와 명시적 변환식을 지정하세요.",
            )
    return {
        "item_index": 0,
        "record_key": None,
        "raw_text": str(cell.value) if cell.value is not None else None,
        "display_text": str(value) if value is not None else "",
        "value_text": text,
        "value_type": kind if state == "present" else target,
        "value_state": state,
        "unit_raw": unit_raw,
        "unit_normalized": target_unit,
        "formula_state": formula_state,
        "derivation": None,
        "regions": regions,
    }


def _assign_record_key(item, spec, sheet_name, raw, cached, pair_indexes):
    """relations 짝(첫 relation 우선) → record_spec.key 순으로 record_key를 정한다."""
    box = item["regions"]["value"][0] if item["regions"].get("value") else item["regions"]["input"][0]
    for relation in spec.get("relations") or []:
        index = pair_indexes.get((relation["to_rule"], sheet_name))
        pair = index.find(relation, box) if index else None
        if pair:
            item["record_key"] = pair[2]
            item["regions"]["record_key"] = [pair[3]]
            item["derivation"] = f"relation:{relation['to_rule']}:{relation['kind']}"
            return
    key = spec["record_spec"].get("key", "coordinate")
    if key is None:
        return
    r, c = box["r1"], box["c1"]
    if isinstance(key, dict):
        ws, cv = raw[box["sheet"]], cached[box["sheet"]]
        kr, kc = (r, column_index_from_string(str(key["column"]))) if "column" in key else (int(key["row"]), c)
        keyregion = region(ws, kr, kc)
        value = cv.cell(keyregion["r1"], keyregion["c1"]).value
        if value is None:
            raise Problem("RECORD_KEY_MISSING", "업무키가 빈 항목이 있습니다.")
        item["record_key"], item["regions"]["record_key"] = str(value), [keyregion]
    else:
        item["record_key"] = f"r{r}" if key == "physical_row" else f"c{c}" if key == "physical_column" else address(r, c)


def _combine_scalar(parts, config, selected):
    if not parts:
        raise Problem("VALUE_NOT_FOUND", "단일 값을 찾을 수 없습니다.")
    if len(parts) == 1:
        return parts[0]
    combine = config.get("combine")
    if combine not in ("concat", "sum"):
        raise Problem("MULTIPLE_SCALAR_VALUES", "단일 값에 여러 셀이 선택되었습니다. 결합 방식을 지정하세요.")
    merged = dict(parts[0])
    values = [p["value_text"] for p in parts if p["value_state"] == "present"]
    try:
        with localcontext() as ctx:
            ctx.prec = 4096
            merged["value_text"] = (
                config.get("separator", " ").join(values)
                if combine == "concat"
                else decimal_text(sum((decimal(v) for v in values), decimal(0)))
            )
        merged["value_type"], merged["value_state"] = ("text" if combine == "concat" else "decimal"), "present"
    except V2Problem:
        merged["value_text"], merged["value_state"] = None, "error"
    if not values:
        merged["value_text"], merged["value_state"] = None, "null"
    merged["raw_text"], merged["formula_state"] = None, "none"
    merged["regions"] = {
        "input": [r for part in parts for r in part["regions"]["value"]],
        **{k: v for k, v in selected.items() if k in ("key", "unit", "context")},
    }
    return merged


def iter_group_items(raw, cached, spec, selected, sheet_name, pair_indexes, own_index):
    """한 group의 값 영역을 훑어 item 배치(≤ BATCH)를 낸다. v2 extract의 리스트/행렬/병합/stop 로직."""
    config = spec["selector"]["value"]
    cardinality, stop = config["cardinality"], config.get("stop") or {}
    max_items, blank_count = stop.get("max_items", 10000), stop.get("count", 1)
    blank_run = stop.get("kind") == "blank_run"
    right = config["axis"] in ("right", "column_major")
    layout = config.get("element_layout", "one_per_column" if right else "one_per_row")
    batch, index, pending, seen, scalar_parts = [], 0, [], set(), []
    stopped = False

    def emit(item):
        nonlocal index
        item["item_index"] = index
        _assign_record_key(item, spec, sheet_name, raw, cached, pair_indexes)
        own_index.add(index, item["record_key"], item["regions"]["value"][0])
        batch.append(item)
        index += 1

    for selected_area in selected["value"]:
        ws = raw[selected_area["sheet"]]
        r1, c1, r2, c2 = (selected_area[k] for k in ("r1", "c1", "r2", "c2"))
        steps = (
            ([(r, c) for r in range(r1, r2 + 1)] for c in range(c1, c2 + 1))
            if right
            else ([(r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1))
        )
        for step in steps:
            regions = []
            for r, c in step:
                found = region(ws, r, c)
                identity = (ws.title, found["range"])
                if identity not in seen:
                    regions.append(found)
                    seen.add(identity)
            if not regions:
                continue
            if cardinality == "list" and layout != "each_cell":
                filled = [x for x in regions if ws.cell(x["r1"], x["c1"]).value is not None]
                if len(filled) > 1:
                    raise Problem("MULTIPLE_VALUES_PER_STEP", "한 행/열에 여러 값이 있습니다. 행렬 또는 셀별 읽기를 지정하세요.")
                regions = filled or regions[:1]
            for found in regions:
                item = build_item(raw, cached, found, spec, selected)
                if cardinality == "scalar":
                    scalar_parts.append(item)
                    if len(scalar_parts) > max_items:
                        raise Problem("ITEM_LIMIT", "최대 항목 수를 초과했습니다.", 413)
                    continue
                if blank_run and item["value_state"] == "empty":
                    pending.append(item)
                    if len(pending) >= blank_count:
                        stopped = True
                        break
                    continue
                for emitted in [*pending, item]:
                    emit(emitted)
                pending = []
                if index > max_items:
                    raise Problem("ITEM_LIMIT", "최대 항목 수를 초과했습니다. 영역을 나누세요.", 413)
                if len(batch) >= BATCH:
                    yield batch
                    batch = []
            if stopped:
                break
        if stopped:
            pending = []
            break
    if cardinality == "scalar":
        merged = _combine_scalar(scalar_parts, config, selected)
        merged["item_index"] = 0
        _assign_record_key(merged, spec, sheet_name, raw, cached, pair_indexes)
        own_index.add(0, merged["record_key"], merged["regions"].get("value", merged["regions"].get("input"))[0])
        batch.append(merged)
    else:
        for item in pending:
            emit(item)
        if index > max_items:
            raise Problem("ITEM_LIMIT", "빈칸을 포함한 최대 항목 수를 초과했습니다.", 413)
    if batch:
        yield batch


def iter_groups(wb, spec, bindings, cache):
    """rule의 primary 시트마다 키(앵커)를 찾고 모든 role 영역을 해결한다 → (sheet_name, anchor, selected)."""
    selectors = spec["selector"]
    first = selectors["key"]["areas"][0]
    primary = first["sheet_role"]
    rule_key = spec["rule_key"]
    if primary not in bindings:
        raise Problem("INVALID_SHEET_BINDING", f"{rule_key}: 시트 역할 {primary}이 연결되지 않았습니다.")
    for name in bindings[primary]:
        found = resolve_area(wb, first, bindings, primary, name, None, cache)
        if not found:
            raise Problem("KEY_NOT_FOUND", f"{rule_key}: 키를 찾을 수 없습니다.")
        if len(found) > 1 and selectors["key"].get("repeat", "once") != "each":
            raise Problem("AMBIGUOUS_KEY", f"{rule_key}: 같은 키가 여러 곳에 있습니다. 반복 또는 위치를 지정하세요.")
        for anchor in found:
            # composite 키는 멤버별로 기록한다(KEY_AREA_LIMIT·mapping_region role=key).
            selected = {"key": list(anchor.get("members") or [anchor])}
            for part in selectors["key"]["areas"][1:]:
                parts = resolve_area(wb, part, bindings, primary, name, anchor, cache)
                if not parts:
                    raise Problem("KEY_NOT_FOUND", f"{rule_key}: 키의 일부 영역을 찾을 수 없습니다.")
                selected["key"].extend(parts)
            for role in ("value", "unit", "context"):
                if role in selectors:
                    selected[role] = []
                    for area in selectors[role]["areas"]:
                        resolved = resolve_area(wb, area, bindings, primary, name, anchor, cache)
                        if not resolved:
                            raise Problem("AREA_NOT_FOUND", f"{rule_key}: 지정한 {role} 영역의 일부를 찾을 수 없습니다.")
                        selected[role].extend(resolved)
            if not selected.get("value"):
                raise Problem("VALUE_NOT_FOUND", f"{rule_key}: 값 영역을 찾을 수 없습니다.")
            yield name, anchor, selected


def _observed_key(raw, cached, anchor, selected, separator):
    members = anchor.get("members")
    if not members:
        return _key_labels(raw, cached, selected["key"], separator)
    texts = [str(cached[m["sheet"]].cell(m["r1"], m["c1"]).value or "") for m in members]
    rest = selected["key"][len(members) :]
    if rest:
        texts.append(_key_labels(raw, cached, rest, separator))
    observed = separator.join(texts)
    if not observed.strip() or len(observed) > 32768:
        raise Problem("INVALID_KEY_TEXT", "키가 비어 있거나 너무 긴 영역입니다.")
    return observed


def extract(raw, cached, bindings, specs):
    """effective_spec 목록 → group/values 이벤트. relations의 to_rule을 먼저 실행한다(위상 정렬)."""
    specs = list(specs)
    by_key = {s["rule_key"]: s for s in specs}
    if not isinstance(bindings, dict) or any(not isinstance(v, list) or not v for v in bindings.values()):
        raise Problem("INVALID_SHEET_BINDING", "시트 역할마다 시트를 하나 이상 연결하세요.")
    cache, pair_indexes = {}, {}
    for rule_key in relation_order(specs):
        spec = by_key[rule_key]
        separator = spec["selector"]["key"].get("separator", " ")
        for name, anchor, selected in iter_groups(cached, spec, bindings, cache):
            observed = _observed_key(raw, cached, anchor, selected, separator)
            group_key = f"{rule_key}|{name}|{anchor['range']}"
            yield {
                "type": "group",
                "group_key": group_key,
                "rule_key": rule_key,
                "observed_key": observed,
                "primary_sheet": name,
                "anchor_locator": anchor["range"],
                "regions": {role: selected[role] for role in ROLES if role in selected},
            }
            own = pair_indexes.setdefault((rule_key, name), PairIndex())
            for batch in iter_group_items(raw, cached, spec, selected, name, pair_indexes, own):
                yield {"type": "values", "group_key": group_key, "items": batch}


# ---------------------------------------------------------------------------- match


def _sheet_terms(wb, name, within, cache):
    key = ("terms", name, within)
    if key not in cache:
        ws = wb[name]
        r1, c1, r2, c2 = _window(ws, within)
        terms = set()
        if r1 <= r2 and c1 <= c2:
            for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True):
                terms.update(norm(v) for v in row if v is not None)
        cache[key] = terms
    return cache[key]


def matcher_hits(match, sheet, wb, cache):
    kind, value = next(iter(match.items()))
    if kind == "name":
        return sheet["name"] == value
    if kind == "name_regex":
        return compile_regex(value).search(sheet["name"]) is not None
    if kind == "ordinal":
        return sheet.get("ordinal") == value
    if kind == "contains_text":
        terms = {norm(t) for t in value["texts"]}
        present = _sheet_terms(wb, sheet["name"], value.get("within", "A1:AD30"), cache) & terms
        return present == terms if value.get("mode", "any") == "all" else bool(present)
    return any(matcher_hits(item, sheet, wb, cache) for item in value)


def bind_sheets(sheet_roles, wb, sheets_meta=None, cache=None):
    """sheet_roles[].match → bindings{role:[sheet_name]}(워크북 순서), missing[{role, code}]."""
    cache = {} if cache is None else cache
    names = set(wb.sheetnames)
    sheets = [s for s in (sheets_meta or [{"name": n, "ordinal": i} for i, n in enumerate(wb.sheetnames)]) if s["name"] in names]
    bindings, missing = {}, []
    for role, definition in sheet_roles.items():
        match = definition.get("match")
        if not match:
            missing.append({"role": role, "code": "MATCH_REQUIRED"})
            continue
        matched = [s["name"] for s in sheets if matcher_hits(match, s, wb, cache)]
        if not matched:
            missing.append({"role": role, "code": "SHEET_NOT_FOUND"})
        elif definition.get("cardinality", "one") == "one" and len(matched) != 1:
            missing.append({"role": role, "code": "AMBIGUOUS_SHEET"})
        else:
            bindings[role] = matched
    return bindings, missing


def locator(found):
    if found.get("members"):
        return [f"{m['sheet']}!{m['range']}" for m in found["members"]]
    return f"{found['sheet']}!{found['range']}"


def _flatten(locators):
    if not locators:
        return None
    return locators[0] if len(locators) == 1 else locators


def resolve_rules(specs, wb, bindings, cache=None):
    """모든 rule의 모든 role 영역을 해결한다 → (resolved{rule_key:{role: locator}}, missing[], signature_rows)."""
    cache = {} if cache is None else cache
    resolved, missing, rows = {}, [], []
    for spec in specs:
        rule_key, selectors = spec["rule_key"], spec["selector"]
        first = selectors["key"]["areas"][0]
        primary = first["sheet_role"]
        found_all = {role: [] for role in ROLES if role in selectors}
        problem = None
        if primary not in bindings:
            problem = ("key", "SHEET_NOT_FOUND")
        for name in bindings.get(primary, []):
            try:
                anchors = resolve_area(wb, first, bindings, primary, name, None, cache)
                if not anchors:
                    raise Problem("KEY_NOT_FOUND", "키를 찾을 수 없습니다.")
                if len(anchors) > 1 and selectors["key"].get("repeat", "once") != "each":
                    raise Problem("AMBIGUOUS_KEY", "같은 키가 여러 곳에 있습니다.")
            except Problem as exc:
                problem = ("key", exc.code)
                break
            for anchor in anchors:
                found_all["key"].append(locator(anchor))
                for role in ROLES:
                    if role not in selectors:
                        continue
                    areas = selectors[role]["areas"][1:] if role == "key" else selectors[role]["areas"]
                    for area in areas:
                        try:
                            parts = resolve_area(wb, area, bindings, primary, name, anchor, cache)
                        except Problem as exc:
                            problem = (role, exc.code)
                            break
                        if not parts:
                            problem = (role, "KEY_NOT_FOUND" if role == "key" else "VALUE_NOT_FOUND" if role == "value" else "AREA_NOT_FOUND")
                            break
                        found_all[role].extend(locator(p) for p in parts)
                    if problem:
                        break
                if problem:
                    break
            if problem:
                break
        if problem:
            # 시트 역할이 바인딩되지 않은 경우는 missing[{role}]이 이미 말하므로 규칙마다 반복하지 않는다.
            if primary in bindings and problem[1] != "INVALID_SHEET_BINDING":
                missing.append({"rule_key": rule_key, "role": problem[0], "code": problem[1]})
            continue
        resolved[rule_key] = {role: _flatten(found_all[role]) for role in found_all}
        rows.append([rule_key, primary, *(_flatten(found_all.get(role) or []) for role in ROLES)])
    rows.sort(key=lambda row: row[0])
    return resolved, missing, rows


def signature_of(rows):
    return hashlib.sha256(dump(rows).encode()).hexdigest()


def _judge(bindings, resolved, missing, rows, reference=None, profile_rev=None):
    signature = signature_of(rows)
    if missing:
        compatibility = "incompatible"
    elif (
        reference
        and reference.get("signature") == signature
        and (profile_rev is None or reference.get("profile_rev") == profile_rev)
    ):
        compatibility = "identical"
    else:
        compatibility = "compatible"
    return {
        "bindings": bindings,
        "compatibility": compatibility,
        "match_signature": signature,
        "match_signature_json": rows,
        "missing": missing,
        "resolved": resolved,
    }


def match_profile(canonical, wb, sheets_meta=None, reference=None, profile_rev=None):
    """§3.3 매치 판정. reference={signature, profile_rev}; profile_rev는 현재 rev(identical 판정에 필요)."""
    cache = {}
    specs = list(compile_profile(canonical).values())
    bindings, missing = bind_sheets(canonical["sheet_roles"], wb, sheets_meta, cache)
    resolved, rule_missing, rows = resolve_rules(specs, wb, bindings, cache)
    return _judge(bindings, resolved, missing + rule_missing, rows, reference, profile_rev)


def match_specs(specs, wb, sheets_meta=None, bindings_hint=None):
    """이전 헤드의 effective_spec 목록과 바인딩 힌트로 새 snapshot을 판정한다(§4.4 승계)."""
    cache = {}
    names = set(wb.sheetnames)
    bindings, missing = {}, []
    for role, sheet_names in (bindings_hint or {}).items():
        present = [n for n in (sheet_names or []) if n in names]
        if present:
            bindings[role] = present
        else:
            missing.append({"role": role, "code": "SHEET_NOT_FOUND"})
    resolved, rule_missing, rows = resolve_rules(list(specs), wb, bindings, cache)
    return _judge(bindings, resolved, missing + rule_missing, rows)
