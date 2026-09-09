"""템플릿 명세 검증과 명시적 값 변환. 임의 Python/SQL 표현식을 실행하지 않는다."""

from __future__ import annotations

import copy
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext

from openpyxl.utils.cell import get_column_letter, range_boundaries

from .db import Problem, dump

MAX_ITEMS = 100000


def bounds(address):
    if (
        not isinstance(address, str)
        or len(address) > 40
        or not re.fullmatch(
            r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)?",
            address,
        )
    ):
        raise Problem("INVALID_RANGE", "A1 또는 B3:C20 형태의 유한 범위를 지정하세요.")
    c1, r1, c2, r2 = range_boundaries(address)
    if not (1 <= r1 <= r2 <= 1048576 and 1 <= c1 <= c2 <= 16384):
        raise Problem("INVALID_RANGE", "Excel 범위를 벗어났습니다.")
    return r1, c1, r2, c2


def address(r1, c1, r2=None, c2=None):
    r2, c2 = r2 or r1, c2 or c1
    a, b = f"{get_column_letter(c1)}{r1}", f"{get_column_letter(c2)}{r2}"
    return a if a == b else f"{a}:{b}"


def decimal(value):
    try:
        number = Decimal(str(value))
        if (
            not number.is_finite()
            or len(number.as_tuple().digits) > 1000
            or abs(number.as_tuple().exponent) > 1000
        ):
            raise InvalidOperation()
        return number
    except (InvalidOperation, ValueError):
        raise Problem(
            "INVALID_DECIMAL",
            "숫자로 변환할 수 없는 값이 있습니다. 원본과 변환 규칙을 확인하세요.",
        ) from None


def decimal_text(value):
    value = decimal(value)
    result = format(value, "f") if value else "0"
    return result.rstrip("0").rstrip(".") if "." in result else result


def validate_template(spec):
    if not isinstance(spec, dict) or len(dump(spec).encode()) > 512000:
        raise Problem("INVALID_TEMPLATE", "템플릿은 512KB 이하 JSON 객체여야 합니다.")
    out = copy.deepcopy(spec)
    if out.get("format", "json") != "json":
        raise Problem(
            "UNSUPPORTED_TEMPLATE_FORMAT",
            "현재 실행기는 JSON 템플릿을 지원합니다. Python 확장은 별도 등록이 필요합니다.",
        )
    roles = out.get("sheet_roles")
    if not isinstance(roles, dict) or not 1 <= len(roles) <= 16:
        raise Problem("INVALID_TEMPLATE", "시트 역할을 1~16개 지정하세요.")
    for role, definition in roles.items():
        if not isinstance(role, str) or not role or not isinstance(definition, dict):
            raise Problem("INVALID_TEMPLATE", "시트 역할 정의가 유효하지 않습니다.")
        if definition.get("cardinality", "one") not in ("one", "many"):
            raise Problem(
                "INVALID_TEMPLATE", "시트 역할은 one/many 중 하나여야 합니다."
            )
    rules = out.get("rules")
    if not isinstance(rules, list) or not 1 <= len(rules) <= 200:
        raise Problem("INVALID_TEMPLATE", "추출 규칙을 1~200개 지정하세요.")
    seen = set()
    for rule in rules:
        key = rule.get("rule_key") if isinstance(rule, dict) else None
        if not isinstance(key, str) or not key or len(key) > 128 or key in seen:
            raise Problem(
                "INVALID_RULE", "규칙 이름은 비어 있지 않고 서로 달라야 합니다."
            )
        seen.add(key)
        validate_rule(rule, roles)
    out["format"], out["schema_version"] = "json", "2.0"
    return out


def validate_rule(rule, roles):
    selectors = rule.get("selector", {})
    if not isinstance(selectors, dict):
        raise Problem("INVALID_SELECTOR", "키와 값의 영역 정의가 필요합니다.")
    for role in ("key", "value", "unit", "context"):
        source = selectors.get(role)
        if source is None and role not in ("key", "value"):
            continue
        if (
            not isinstance(source, dict)
            or not isinstance(source.get("areas"), list)
            or not 1 <= len(source["areas"]) <= 32
        ):
            raise Problem("INVALID_SELECTOR", f"{role} 영역은 1~32개여야 합니다.")
        for area in source["areas"]:
            if not isinstance(area, dict) or area.get("sheet_role") not in roles:
                raise Problem(
                    "INVALID_SHEET_ROLE", "영역이 선언되지 않은 시트 역할을 참조합니다."
                )
            kinds = [k for k in ("range", "find", "relative") if k in area]
            if len(kinds) != 1:
                raise Problem(
                    "INVALID_SELECTOR",
                    "영역은 range/find/relative 중 하나로 지정하세요.",
                )
            if "range" in area:
                r1, c1, r2, c2 = bounds(area["range"])
                if (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_ITEMS:
                    raise Problem(
                        "RANGE_LIMIT", "영역당 100,000셀 이하로 나누어 지정하세요.", 413
                    )
            elif "find" in area:
                find = area["find"]
                if (
                    not isinstance(find, dict)
                    or not isinstance(find.get("texts"), list)
                    or not 1 <= len(find["texts"]) <= 50
                    or any(
                        not isinstance(t, str) or not 1 <= len(t) <= 512
                        for t in find["texts"]
                    )
                ):
                    raise Problem("INVALID_SELECTOR", "검색할 키 표현을 지정하세요.")
                if "occurrence" in find and (
                    type(find["occurrence"]) is not int or find["occurrence"] < 0
                ):
                    raise Problem(
                        "INVALID_SELECTOR", "검색 일치 순서는 0 이상의 정수여야 합니다."
                    )
                r1, c1, r2, c2 = bounds(find.get("within", "A1:AZ100"))
                if (r2 - r1 + 1) * (c2 - c1 + 1) > 10000:
                    raise Problem(
                        "RANGE_LIMIT", "키 검색 범위는 10,000셀 이하로 지정하세요.", 413
                    )
            else:
                relative = area["relative"]
                if not isinstance(relative, dict) or any(
                    type(relative.get(k, default)) is not int
                    for k, default in (("row", 0), ("col", 0), ("rows", 1), ("cols", 1))
                ):
                    raise Problem("INVALID_SELECTOR", "상대 위치는 정수여야 합니다.")
                if (
                    not (
                        1 <= relative.get("rows", 1) <= MAX_ITEMS
                        and 1 <= relative.get("cols", 1) <= 16384
                    )
                    or relative.get("rows", 1) * relative.get("cols", 1) > MAX_ITEMS
                ):
                    raise Problem("RANGE_LIMIT", "상대 범위가 너무 큽니다.", 413)
    value = selectors["value"]
    if selectors["key"].get("repeat", "once") not in ("once", "each"):
        raise Problem("INVALID_REPEAT", "키 반복은 once/each 중 하나여야 합니다.")
    if (
        value.get("merge_policy", "anchor_once") != "anchor_once"
        or value.get("blank_policy", "preserve") != "preserve"
    ):
        raise Problem(
            "UNSUPPORTED_POLICY",
            "병합 셀은 anchor_once, 빈칸은 preserve 정책을 지원합니다.",
        )
    if value.get("element_layout", "each_cell") not in (
        "each_cell",
        "one_per_row",
        "one_per_column",
    ):
        raise Problem(
            "INVALID_LAYOUT",
            "리스트 항목은 each_cell/one_per_row/one_per_column 중 하나여야 합니다.",
        )
    cardinality, axis = value.setdefault("cardinality", "scalar"), value.setdefault(
        "axis", "none"
    )
    allowed = {
        "scalar": {"none"},
        "list": {"down", "right"},
        "matrix": {"row_major", "column_major"},
    }
    if cardinality not in allowed or axis not in allowed[cardinality]:
        raise Problem(
            "INVALID_DIRECTION", "값 모양과 가로/세로 방향이 일치하지 않습니다."
        )
    if cardinality == "list" and (
        (value.get("element_layout") == "one_per_row" and axis != "down")
        or (value.get("element_layout") == "one_per_column" and axis != "right")
    ):
        raise Problem(
            "INVALID_LAYOUT", "행별 값은 down, 열별 값은 right 방향으로 지정하세요."
        )
    if value.get("combine", "ordered_union") not in ("ordered_union", "concat", "sum"):
        raise Problem(
            "UNSUPPORTED_COMBINE",
            "현재 결합은 ordered_union/concat/sum을 지원합니다. 시트 간 조인은 통합 DB에서 지정하세요.",
        )
    if (
        cardinality != "scalar"
        and value.get("combine", "ordered_union") != "ordered_union"
    ):
        raise Problem(
            "UNSUPPORTED_COMBINE", "concat/sum 결합은 단일 값에서 지정하세요."
        )
    stop = value.setdefault("stop", {"kind": "explicit_areas", "max_items": 10000})
    if (
        stop.get("kind") not in ("explicit_areas", "blank_run")
        or type(stop.get("max_items", 10000)) is not int
        or not 1 <= stop.get("max_items", 10000) <= MAX_ITEMS
    ):
        raise Problem("INVALID_STOP", "종료 조건과 최대 항목 수를 지정하세요.")
    if stop.get("kind") == "blank_run" and (
        type(stop.get("count", 1)) is not int or not 1 <= stop.get("count", 1) <= 100
    ):
        raise Problem("INVALID_STOP", "빈칸 종료 수는 1~100이어야 합니다.")
    spec = rule.setdefault("value_spec", {"type": "text"})
    if spec.get("type", "text") not in (
        "decimal",
        "text",
        "boolean",
        "date",
        "datetime",
    ):
        raise Problem(
            "INVALID_TYPE",
            "지원하는 값 타입은 decimal/text/boolean/date/datetime입니다.",
        )
    if spec.get("formula_policy", "cached_only") != "cached_only":
        raise Problem(
            "UNSUPPORTED_FORMULA", "원본 수정 없이 저장된 수식 결과만 읽습니다."
        )
    normal = spec.setdefault("normalization", {"operation": "identity", "version": "1"})
    if normal.get("operation", "identity") not in (
        "identity",
        "trim",
        "affine",
        "pipeline",
    ):
        raise Problem("UNSUPPORTED_NORMALIZER", "등록되지 않은 변환입니다.")
    if normal.get("operation") == "pipeline":
        from .normalization import validate_pipeline

        validate_pipeline(normal, spec.get("type", "text"))
    if (
        normal.get("operation") == "affine"
        and spec.get("type", "text") != "decimal"
        or normal.get("operation") == "trim"
        and spec.get("type", "text") != "text"
    ):
        raise Problem(
            "NORMALIZER_TYPE_MISMATCH", "affine은 decimal, trim은 text에 적용하세요."
        )
    if normal.get("operation") == "affine":
        decimal(normal.get("factor", "1"))
        decimal(normal.get("offset", "0"))
    records = rule.setdefault(
        "record_spec", {"scope": [rule["rule_key"]], "key": "coordinate"}
    )
    if (
        not isinstance(records.get("scope"), list)
        or not 1 <= len(records["scope"]) <= 8
        or any(
            not isinstance(s, str) or not 1 <= len(s) <= 128 for s in records["scope"]
        )
    ):
        raise Problem(
            "INVALID_RECORD_SCOPE", "표/반복 블록을 구분할 레코드 범위를 지정하세요."
        )
    key = records.get("key", "coordinate")
    if not (
        key in ("coordinate", "physical_row", "physical_column")
        if isinstance(key, str)
        else isinstance(key, dict) and ("column" in key or "row" in key)
    ):
        raise Problem(
            "INVALID_RECORD_KEY", "물리 행/열/좌표 또는 업무키 행/열을 지정하세요."
        )
    if isinstance(key, dict):
        if len(key) != 1 or (
            "row" in key
            and (type(key["row"]) is not int or not 1 <= key["row"] <= 1048576)
        ):
            raise Problem("INVALID_RECORD_KEY", "업무키 행/열 하나를 지정하세요.")
        if "column" in key:
            bounds(str(key["column"]) + "1")


def typed(value, spec):
    target = spec.get("type", "text")
    if target == "decimal":
        value = decimal(value)
        normal = spec.get("normalization", {})
        if normal.get("operation") == "affine":
            with localcontext() as ctx:
                ctx.prec = 4096
                value = value * decimal(normal.get("factor", "1")) + decimal(
                    normal.get("offset", "0")
                )
        return target, decimal_text(value)
    if target == "boolean":
        if isinstance(value, bool):
            return target, "true" if value else "false"
        if str(value).casefold() not in ("true", "false", "1", "0"):
            raise Problem("INVALID_BOOLEAN", "참/거짓으로 변환할 수 없습니다.")
        return target, "true" if str(value).casefold() in ("true", "1") else "false"
    if target in ("date", "datetime"):
        try:
            if target == "date":
                return (
                    target,
                    (
                        value.date()
                        if isinstance(value, datetime)
                        else (
                            value
                            if isinstance(value, date)
                            else date.fromisoformat(str(value))
                        )
                    ).isoformat(),
                )
            return (
                target,
                (
                    value
                    if isinstance(value, datetime)
                    else datetime.fromisoformat(str(value))
                ).isoformat(),
            )
        except ValueError:
            raise Problem("INVALID_DATE", "날짜/시각 값이 유효하지 않습니다.") from None
    return "text", (
        str(value).strip()
        if spec.get("normalization", {}).get("operation") == "trim"
        else str(value)
    )
