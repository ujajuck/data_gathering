"""원본 저장을 하지 않는 Reader 계약. DRM 제공자는 운영자가 승인한 factory로 등록한다."""

from __future__ import annotations

import base64
import hashlib
import importlib
import io
import os
import zipfile
from bisect import bisect_right
from decimal import localcontext
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter

from .db import Problem, norm
from .spec import MAX_ITEMS, address, bounds, decimal, decimal_text, typed


def make_reader(root, provider, principal):
    if provider == "local-xlsx":
        return XlsxReader(Path(root), principal)
    factory = os.environ.get("KG_V2_READER_FACTORY", "")
    if not factory or ":" not in factory:
        raise Problem(
            "DRM_READER_REQUIRED",
            "이 문서의 보안 읽기 어댑터가 연결되지 않았습니다.",
            403,
        )
    module, function = factory.split(":", 1)
    # factory는 서버 설정만 읽는다. 요청 본문에서 모듈/함수 이름을 받지 않는다.
    return getattr(importlib.import_module(module), function)(
        root=Path(root), provider=provider, principal=principal
    )


def file_hash(path):
    sha = hashlib.sha256()
    with path.open("rb") as source:
        for part in iter(lambda: source.read(1024 * 1024), b""):
            sha.update(part)
    return sha.hexdigest()


class XlsxReader:
    def __init__(self, root, principal):
        self.raw = (root / "data/raw").resolve()
        self.principal = principal

    def path(self, source_ref):
        path = (self.raw / source_ref).resolve()
        if not path.is_relative_to(self.raw) or not path.is_file():
            raise Problem(
                "SOURCE_NOT_FOUND", "허용된 원본 폴더에서 파일을 찾을 수 없습니다.", 404
            )
        return path

    def authorize(self, source_ref, required="view"):
        path = self.path(source_ref)
        with path.open("rb") as source:
            magic = source.read(4)
        if magic[:2] != b"PK":
            raise Problem(
                "DRM_READER_REQUIRED",
                "암호화 문서는 승인된 보안 읽기 어댑터로 접근해야 합니다.",
                403,
            )
        return {
            "can_view": True,
            "can_extract": True,
            "can_render_web": True,
            "can_cache_derivative": True,
            "native_render": False,
            "provider": "local-xlsx",
            "policy_revision": "local-filesystem-v1",
            "access_scope_key": self.principal,
            "expires_at": (
                datetime.now(timezone.utc) + timedelta(minutes=5)
            ).isoformat(),
        }

    def _check(self, source_ref, expected_token=None):
        self.authorize(source_ref)
        path = self.path(source_ref)
        if path.stat().st_size > 256 * 1024 * 1024:
            raise Problem(
                "SOURCE_SIZE_LIMIT",
                "기본 Reader의 256MB 원본 한도를 초과했습니다.",
                413,
            )
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 30000
                or sum(i.file_size for i in entries) > 512 * 1024 * 1024
            ):
                raise Problem(
                    "EXPANDED_SIZE_LIMIT",
                    "압축 해제 예상 크기가 Reader 한도를 초과했습니다.",
                    413,
                )
        token = file_hash(path)
        if expected_token is not None and token != expected_token:
            raise Problem(
                "SOURCE_VERSION_CHANGED",
                "원본이 변경되었습니다. 새 버전을 등록하세요.",
                409,
            )
        return path, token

    def describe(self, source_ref):
        path, token = self._check(source_ref)
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheets = [
                {
                    "name": s.title,
                    "ordinal": n,
                    "visibility": {"veryHidden": "very_hidden"}.get(
                        s.sheet_state, s.sheet_state
                    ),
                    "estimated_rows": s.max_row,
                    "estimated_cols": s.max_column,
                }
                for n, s in enumerate(wb.worksheets)
            ]
            if len(sheets) > 2000:
                raise Problem(
                    "SHEET_LIMIT", "시트 수가 Reader 한도를 초과했습니다.", 413
                )
            result = {
                "token": token,
                "filename": path.name,
                "author": wb.properties.creator,
                "authored_at": (
                    wb.properties.created.isoformat() if wb.properties.created else None
                ),
                "excel_date_system": "1904" if wb.epoch.year == 1904 else "1900",
                "byte_size": path.stat().st_size,
                "sheets": sheets,
                "capabilities": self.authorize(source_ref),
            }
        finally:
            wb.close()
        if file_hash(path) != token:
            raise Problem(
                "SOURCE_VERSION_CHANGED", "읽는 동안 원본이 변경되었습니다.", 409
            )
        return result

    @staticmethod
    def region(ws, r, c):
        # 같은 행의 병합 목록을 재사용하고 열 이진 탐색으로 항목마다 전체 목록을 훑지 않는다.
        if not hasattr(ws, "_v2_merge_rows"):
            ws._v2_merge_rows = {}
        if r not in ws._v2_merge_rows:
            if len(ws._v2_merge_rows) >= 128:
                ws._v2_merge_rows.pop(next(iter(ws._v2_merge_rows)))
            spans = sorted(
                (m for m in ws.merged_cells.ranges if m.min_row <= r <= m.max_row),
                key=lambda m: m.min_col,
            )
            ws._v2_merge_rows[r] = ([m.min_col for m in spans], spans)
        starts, spans = ws._v2_merge_rows[r]
        index = bisect_right(starts, c) - 1
        if index >= 0:
            merged = spans[index]
            if c <= merged.max_col:
                return {
                    "sheet": ws.title,
                    "range": str(merged),
                    "r1": merged.min_row,
                    "c1": merged.min_col,
                    "r2": merged.max_row,
                    "c2": merged.max_col,
                    "merge_anchor_r": merged.min_row,
                    "merge_anchor_c": merged.min_col,
                }
        return {
            "sheet": ws.title,
            "range": address(r, c),
            "r1": r,
            "c1": c,
            "r2": r,
            "c2": c,
        }

    @staticmethod
    def area(ws, box):
        r1, c1, r2, c2 = box
        return {
            "sheet": ws.title,
            "range": address(r1, c1, r2, c2),
            "r1": r1,
            "c1": c1,
            "r2": r2,
            "c2": c2,
        }

    def resolve(self, wb, area, bindings, primary_role, primary_sheet, anchor=None):
        names = (
            [primary_sheet]
            if area["sheet_role"] == primary_role
            else bindings[area["sheet_role"]]
        )
        resolved = []
        for name in names:
            ws = wb[name]
            if "range" in area:
                resolved.append(self.area(ws, bounds(area["range"])))
            elif "relative" in area:
                if not anchor:
                    raise Problem(
                        "ANCHOR_REQUIRED", "상대 범위에는 키 앵커가 필요합니다."
                    )
                spec = area["relative"]
                r, c = anchor["r1"] + spec.get("row", 0), anchor["c1"] + spec.get(
                    "col", 0
                )
                box = bounds(
                    address(
                        r, c, r + spec.get("rows", 1) - 1, c + spec.get("cols", 1) - 1
                    )
                )
                resolved.append(self.area(ws, box))
            else:
                find = area["find"]
                r1, c1, r2, c2 = bounds(find.get("within", "A1:AZ100"))
                terms = {norm(t) for t in find["texts"]}
                hits, seen = [], set()
                for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                    for cell in row:
                        if cell.value is not None and norm(cell.value) in terms:
                            region = self.region(ws, cell.row, cell.column)
                            if region["range"] not in seen:
                                seen.add(region["range"])
                                hits.append(region)
                if "occurrence" in find:
                    index = find["occurrence"]
                    hits = (
                        hits[index : index + 1]
                        if isinstance(index, int) and index >= 0
                        else []
                    )
                resolved.extend(hits)
        return resolved

    def extract(self, source_ref, expected_token, rules, bindings):
        path, token = self._check(source_ref, expected_token)
        raw = load_workbook(path, data_only=False, keep_links=False)
        cached = load_workbook(path, data_only=True, keep_links=False)
        try:
            for rule in rules:
                selectors = rule["selector"]
                first = selectors["key"]["areas"][0]
                primary = first["sheet_role"]
                for name in bindings[primary]:
                    found = self.resolve(raw, first, bindings, primary, name)
                    repeat = selectors["key"].get("repeat", "one")
                    if not found:
                        raise Problem(
                            "KEY_NOT_FOUND",
                            f"{rule['rule_key']}: 키를 찾을 수 없습니다.",
                        )
                    if len(found) > 1 and repeat != "each":
                        raise Problem(
                            "AMBIGUOUS_KEY",
                            f"{rule['rule_key']}: 같은 키가 여러 곳에 있습니다. 반복 또는 위치를 지정하세요.",
                        )
                    for anchor in found:
                        selected = {"key": [anchor]}
                        for part in selectors["key"]["areas"][1:]:
                            parts = self.resolve(
                                raw, part, bindings, primary, name, anchor
                            )
                            if not parts:
                                raise Problem(
                                    "KEY_NOT_FOUND",
                                    "키의 일부 영역을 찾을 수 없습니다.",
                                )
                            selected["key"].extend(parts)
                        for role in ("value", "unit", "context"):
                            if role in selectors:
                                selected[role] = []
                                for area in selectors[role]["areas"]:
                                    resolved = self.resolve(
                                        raw, area, bindings, primary, name, anchor
                                    )
                                    if not resolved:
                                        raise Problem(
                                            "AREA_NOT_FOUND",
                                            f"{rule['rule_key']}: 지정한 {role} 영역의 일부를 찾을 수 없습니다.",
                                        )
                                    selected[role].extend(resolved)
                        if not selected["value"]:
                            raise Problem(
                                "VALUE_NOT_FOUND",
                                f"{rule['rule_key']}: 값 영역을 찾을 수 없습니다.",
                            )
                        series = f"{rule['rule_key']}:{name}:{anchor['range']}"
                        labels, key_anchors = [], set()
                        for part in selected["key"]:
                            if (part["r2"] - part["r1"] + 1) * (
                                part["c2"] - part["c1"] + 1
                            ) > 10000:
                                raise Problem(
                                    "KEY_AREA_LIMIT",
                                    "키 영역은 10,000셀 이하로 지정하세요.",
                                    413,
                                )
                            ws, cv = raw[part["sheet"]], cached[part["sheet"]]
                            for row in ws.iter_rows(
                                min_row=part["r1"],
                                max_row=part["r2"],
                                min_col=part["c1"],
                                max_col=part["c2"],
                            ):
                                for cell in row:
                                    area = self.region(ws, cell.row, cell.column)
                                    identity = (ws.title, area["range"])
                                    if identity in key_anchors:
                                        continue
                                    key_anchors.add(identity)
                                    value = cv.cell(area["r1"], area["c1"]).value
                                    if value is not None:
                                        labels.append(str(value))
                        observed = selectors["key"].get("separator", "").join(labels)
                        if not observed.strip() or len(observed) > 32768:
                            raise Problem(
                                "INVALID_KEY_TEXT",
                                "키가 비어 있거나 너무 긴 영역입니다.",
                            )
                        config = selectors["value"]
                        yield {
                            "type": "series",
                            "key": series,
                            "rule_key": rule["rule_key"],
                            "observed_key": observed,
                            "primary_sheet": name,
                            "scope": rule["record_spec"]["scope"],
                            "cardinality": config["cardinality"],
                            "axis": config["axis"],
                            "regions": selected,
                        }
                        batch, index, pending_blanks, seen = [], 0, [], set()
                        scalar_parts = []
                        for selected_area in selected["value"]:
                            ws, cv = (
                                raw[selected_area["sheet"]],
                                cached[selected_area["sheet"]],
                            )
                            r1, c1, r2, c2 = (
                                selected_area[k] for k in ("r1", "c1", "r2", "c2")
                            )
                            right = config["axis"] in ("right", "column_major")
                            steps = (
                                [
                                    [(r, c) for r in range(r1, r2 + 1)]
                                    for c in range(c1, c2 + 1)
                                ]
                                if right
                                else [
                                    [(r, c) for c in range(c1, c2 + 1)]
                                    for r in range(r1, r2 + 1)
                                ]
                            )
                            for step in steps:
                                regions = []
                                for r, c in step:
                                    region = self.region(ws, r, c)
                                    key = (ws.title, region["range"])
                                    if key not in seen:
                                        regions.append(region)
                                        seen.add(key)
                                if not regions:
                                    continue
                                if (
                                    config["cardinality"] == "list"
                                    and config.get(
                                        "element_layout",
                                        (
                                            "one_per_row"
                                            if not right
                                            else "one_per_column"
                                        ),
                                    )
                                    != "each_cell"
                                ):
                                    filled = [
                                        r
                                        for r in regions
                                        if ws.cell(r["r1"], r["c1"]).value is not None
                                    ]
                                    if len(filled) > 1:
                                        raise Problem(
                                            "MULTIPLE_VALUES_PER_STEP",
                                            "한 행/열에 여러 값이 있습니다. 행렬 또는 셀별 읽기를 지정하세요.",
                                        )
                                    regions = filled or regions[:1]
                                for region in regions:
                                    item = self._item(
                                        raw, cached, region, rule, selected, index
                                    )
                                    if config["cardinality"] == "scalar":
                                        scalar_parts.append(item)
                                        if len(scalar_parts) > config["stop"].get(
                                            "max_items", 10000
                                        ):
                                            raise Problem(
                                                "ITEM_LIMIT",
                                                "최대 항목 수를 초과했습니다.",
                                                413,
                                            )
                                        continue
                                    stop = config["stop"]
                                    if (
                                        stop["kind"] == "blank_run"
                                        and item["value_state"] == "blank"
                                    ):
                                        pending_blanks.append(item)
                                        if len(pending_blanks) >= stop.get("count", 1):
                                            break
                                        continue
                                    for emitted in [*pending_blanks, item]:
                                        emitted["item_index"] = index
                                        batch.append(emitted)
                                        index += 1
                                    pending_blanks = []
                                    if index > stop.get("max_items", 10000):
                                        raise Problem(
                                            "ITEM_LIMIT",
                                            "최대 항목 수를 초과했습니다. 영역을 나누세요.",
                                            413,
                                        )
                                    if len(batch) >= 100:
                                        yield {
                                            "type": "items",
                                            "series": series,
                                            "items": batch,
                                        }
                                        batch = []
                                if config["stop"]["kind"] == "blank_run" and len(
                                    pending_blanks
                                ) >= config["stop"].get("count", 1):
                                    break
                            if config["stop"]["kind"] == "blank_run" and len(
                                pending_blanks
                            ) >= config["stop"].get("count", 1):
                                pending_blanks = []
                                break
                        if config["cardinality"] == "scalar":
                            if not scalar_parts:
                                raise Problem(
                                    "VALUE_NOT_FOUND", "단일 값을 찾을 수 없습니다."
                                )
                            if len(scalar_parts) > 1:
                                combine = config.get("combine")
                                if combine not in ("concat", "sum"):
                                    raise Problem(
                                        "MULTIPLE_SCALAR_VALUES",
                                        "단일 값에 여러 셀이 선택되었습니다. 결합 방식을 지정하세요.",
                                    )
                                merged = dict(scalar_parts[0])
                                values = [
                                    p["value_text"]
                                    for p in scalar_parts
                                    if p["value_state"] == "present"
                                ]
                                with localcontext() as ctx:
                                    ctx.prec = 4096
                                    merged["value_text"] = (
                                        config.get("separator", "").join(values)
                                        if combine == "concat"
                                        else decimal_text(
                                            sum(
                                                (decimal(v) for v in values), decimal(0)
                                            )
                                        )
                                    )
                                merged["value_type"] = (
                                    "text" if combine == "concat" else "decimal"
                                )
                                merged["value_state"] = "present"
                                if not values:
                                    (
                                        merged["value_type"],
                                        merged["value_text"],
                                        merged["value_state"],
                                    ) = ("null", None, "blank")
                                (
                                    merged["raw_type"],
                                    merged["raw_text"],
                                    merged["formula_text"],
                                    merged["formula_state"],
                                ) = ("derived", None, None, "none")
                                merged["regions"] = {
                                    "input": [
                                        r
                                        for part in scalar_parts
                                        for r in part["regions"]["value"]
                                    ],
                                    **{
                                        k: v
                                        for k, v in selected.items()
                                        if k in ("unit", "context")
                                    },
                                }
                                scalar_parts = [merged]
                            batch.extend(scalar_parts)
                        else:
                            for item in pending_blanks:
                                item["item_index"] = index
                                batch.append(item)
                                index += 1
                            if index > config["stop"].get("max_items", 10000):
                                raise Problem(
                                    "ITEM_LIMIT",
                                    "빈칸을 포함한 최대 항목 수를 초과했습니다.",
                                    413,
                                )
                        if batch:
                            yield {"type": "items", "series": series, "items": batch}
        finally:
            raw.close()
            cached.close()
        self._check(source_ref, token)
        yield {"type": "verified", "token": token}

    def _item(self, raw, cached, region, rule, selected, index):
        ws, cv = raw[region["sheet"]], cached[region["sheet"]]
        r, c = region["r1"], region["c1"]
        cell, cached_cell = ws.cell(r, c), cv.cell(r, c)
        formula = str(cell.value) if cell.data_type == "f" else None
        value = cached_cell.value if formula else cell.value
        value_state = (
            "unavailable"
            if formula and value is None
            else (
                "blank"
                if value is None
                else "excel_error" if cached_cell.data_type == "e" else "present"
            )
        )
        kind, text = (
            typed(value, rule["value_spec"])
            if value_state == "present"
            else ("null", None)
        )
        if formula and value is None:
            raise Problem(
                "FORMULA_CACHE_MISSING",
                "수식의 저장된 결과가 없습니다. 원본 제공자에서 계산한 뒤 다시 등록하세요.",
            )
        if value_state == "excel_error":
            raise Problem(
                "EXCEL_ERROR",
                "원본 값에 Excel 오류가 있습니다. 원본과 매핑을 확인하세요.",
            )
        records = rule["record_spec"]
        key = records.get("key", "coordinate")
        regions = {
            "value": [region],
            **{k: v for k, v in selected.items() if k in ("unit", "context")},
        }
        if isinstance(key, dict):
            kr, kc = (
                (r, column_index_from_string(str(key["column"])))
                if "column" in key
                else (int(key["row"]), c)
            )
            keyregion = self.region(ws, kr, kc)
            record_value = cv.cell(keyregion["r1"], keyregion["c1"]).value
            if record_value is None:
                raise Problem("RECORD_KEY_MISSING", "업무키가 빈 항목이 있습니다.")
            record_key = str(record_value)
            regions["record_key"] = [keyregion]
        else:
            record_key = (
                f"r{r}"
                if key == "physical_row"
                else f"c{c}" if key == "physical_column" else address(r, c)
            )
        spec = rule["value_spec"]
        unit_raw = " ".join(
            str(cached[a["sheet"]].cell(a["r1"], a["c1"]).value or "")
            for a in selected.get("unit", [])
        ) or spec.get("source_unit", spec.get("unit"))
        target_unit = spec.get("unit", unit_raw)
        if norm(unit_raw or "") != norm(target_unit or ""):
            if spec.get("normalization", {}).get("operation") != "affine" or norm(
                spec.get("source_unit", "")
            ) != norm(unit_raw or ""):
                raise Problem(
                    "UNIT_MISMATCH",
                    "원본 단위와 목표 단위가 다릅니다. 원본 단위와 명시적 변환식을 지정하세요.",
                )
        return {
            "item_index": index,
            "record_key": record_key,
            "row_ordinal": r - 1,
            "col_ordinal": c - 1,
            "raw_type": "formula" if formula else type(cell.value).__name__,
            "raw_text": str(cell.value) if cell.value is not None else None,
            "display_text": str(value) if value is not None else "",
            "value_type": kind,
            "value_text": text,
            "value_state": value_state,
            "formula_text": formula,
            "formula_state": "cached_unknown_age" if formula else "none",
            "unit_raw": unit_raw,
            "unit_normalized": target_unit,
            "regions": regions,
        }

    def viewport(self, source_ref, expected_token, sheet, r1, c1, rows=40, cols=12):
        if not 1 <= rows <= 100 or not 1 <= cols <= 30 or rows * cols > 3000:
            raise Problem(
                "VIEWPORT_LIMIT", "표시 범위는 100행·30열·3,000셀 이내여야 합니다.", 413
            )
        bounds(address(r1, c1, r1 + rows - 1, c1 + cols - 1))
        path, token = self._check(source_ref, expected_token)
        wb = load_workbook(path, data_only=True, keep_links=False)
        try:
            if sheet not in wb.sheetnames:
                raise Problem("SHEET_NOT_FOUND", "시트를 찾을 수 없습니다.", 404)
            ws = wb[sheet]
            r2, c2 = r1 + rows - 1, c1 + cols - 1
            widths, heights = {}, {}
            for letter, dim in ws.column_dimensions.items():
                for col in range(
                    dim.min or column_index_from_string(letter),
                    (dim.max or column_index_from_string(letter)) + 1,
                ):
                    widths[col] = (
                        0 if dim.hidden else max(12, (dim.width or 13) * 7 + 5)
                    )
            for row, dim in ws.row_dimensions.items():
                heights[row] = 0 if dim.hidden else (dim.height or 15) * 96 / 72

            def px(index, values, default):
                return (index - 1) * default + sum(
                    v - default for k, v in values.items() if k < index
                )

            origin_x, origin_y = px(c1, widths, 96), px(r1, heights, 20)
            columns = [
                {
                    "index": c,
                    "label": get_column_letter(c),
                    "x": px(c, widths, 96) - origin_x,
                    "width": widths.get(c, 96),
                }
                for c in range(c1, c2 + 1)
            ]
            rowdims = [
                {
                    "index": r,
                    "y": px(r, heights, 20) - origin_y,
                    "height": heights.get(r, 20),
                }
                for r in range(r1, r2 + 1)
            ]

            def geometry(region):
                x, y = px(region["c1"], widths, 96), px(region["r1"], heights, 20)
                return {
                    "x": x - origin_x,
                    "y": y - origin_y,
                    "width": px(region["c2"] + 1, widths, 96) - x,
                    "height": px(region["r2"] + 1, heights, 20) - y,
                }

            cells, seen = [], set()
            for row in range(r1, r2 + 1):
                for col in range(c1, c2 + 1):
                    region = self.region(ws, row, col)
                    if region["range"] in seen:
                        continue
                    seen.add(region["range"])
                    cell = ws.cell(region["r1"], region["c1"])

                    def color(value, fallback):
                        return (
                            "#" + value.rgb[-6:]
                            if value
                            and value.type == "rgb"
                            and isinstance(value.rgb, str)
                            else fallback
                        )

                    cells.append(
                        {
                            **region,
                            **geometry(region),
                            "text": str(cell.value) if cell.value is not None else "",
                            "style": {
                                "background": (
                                    color(cell.fill.fgColor, "#ffffff")
                                    if cell.fill.patternType
                                    else "#ffffff"
                                ),
                                "color": color(cell.font.color, "#172b36"),
                                "fontSize": (cell.font.sz or 11) * 96 / 72,
                                "fontFamily": cell.font.name or "sans-serif",
                                "fontWeight": 700 if cell.font.b else 400,
                                "fontStyle": "italic" if cell.font.i else "normal",
                                "textAlign": cell.alignment.horizontal or "left",
                                "whiteSpace": (
                                    "pre-wrap" if cell.alignment.wrap_text else "nowrap"
                                ),
                            },
                        }
                    )
            images = []
            image_bytes = 0
            for picture in ws._images:
                anchor = getattr(picture.anchor, "_from", None)
                if not anchor:
                    continue
                x = px(anchor.col + 1, widths, 96) + anchor.colOff / 9525 - origin_x
                y = px(anchor.row + 1, heights, 20) + anchor.rowOff / 9525 - origin_y
                if (
                    x > px(c2 + 1, widths, 96) - origin_x
                    or y > px(r2 + 1, heights, 20) - origin_y
                    or x + picture.width < 0
                    or y + picture.height < 0
                ):
                    continue
                data = picture._data()
                if (
                    len(data) > 2 * 1024 * 1024
                    or image_bytes + len(data) > 4 * 1024 * 1024
                ):
                    continue
                image_bytes += len(data)
                mime = {
                    "png": "image/png",
                    "jpeg": "image/jpeg",
                    "jpg": "image/jpeg",
                    "gif": "image/gif",
                }.get(picture.format)
                if mime:
                    images.append(
                        {
                            "x": x,
                            "y": y,
                            "width": picture.width,
                            "height": picture.height,
                            "data_url": f"data:{mime};base64,"
                            + base64.b64encode(data).decode(),
                        }
                    )
            result = {
                "mode": "simplified",
                "fidelity": "간략 보기: 차트·도형·일부 서식은 원본 렌더 어댑터에서 확인해야 합니다.",
                "sheet": sheet,
                "r1": r1,
                "c1": c1,
                "r2": r2,
                "c2": c2,
                "rows": rowdims,
                "columns": columns,
                "width": px(c2 + 1, widths, 96) - origin_x,
                "height": px(r2 + 1, heights, 20) - origin_y,
                "cells": cells,
                "images": images,
                "layout_revision": token,
                "estimated_rows": ws.max_row,
                "estimated_cols": ws.max_column,
            }
        finally:
            wb.close()
        self._check(source_ref, token)
        return result
