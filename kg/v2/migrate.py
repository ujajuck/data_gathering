"""v1 kg.db → v2 워크스페이스 이관 (docs/design/db-schema-v2.md §12).

v1은 읽기 전용(mode=ro)으로만 열고, v2 쓰기는 Service를 통해서만 한다.
문서는 해시/이름으로 병합하지 않고 새 안정 ID를 만들어 artifact 행으로 대응을 기록한다.
매핑은 항상 proposed로 만들며(§4.10) 값(parsed_source/payload_value)은 옮기지 않는다.
"""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import sys
from pathlib import Path

from .db import Problem, digest, dump, insert, now, uid
from .spec import bounds, validate_template

V1_DB = "data/kg/kg.db"
LINK_SCHEMA = "v1-migration/1"
LINK_POLICY = "v1-migration"
SOURCE_POLICY = "v1-migration-source"
SCOPE_KEY = "v1-migration"
PROVIDER = "local-xlsx"
STATUSES = ("migrated", "existing", "skipped", "needs_review", "planned")
COUNT_KEYS = {
    "kg": "kg",
    "units": "units",
    "document": "documents",
    "document_version": "document_versions",
    "template": "templates",
    "template_version": "template_versions",
    "assignment": "assignments",
    "override": "overrides",
}


def shape(address):
    r1, c1, r2, c2 = bounds(address)
    if (r1, c1) == (r2, c2):
        return {"cardinality": "scalar", "axis": "none"}
    if c1 == c2:
        return {"cardinality": "list", "axis": "down", "element_layout": "each_cell"}
    if r1 == r2:
        return {"cardinality": "list", "axis": "right", "element_layout": "each_cell"}
    return {"cardinality": "matrix", "axis": "row_major"}


def _texts(values):
    out = []
    for value in values:
        text = str(value).strip() if value is not None else ""
        if text and len(text) <= 512 and text not in out:
            out.append(text)
    return out[:50]


def _loads(text, default):
    try:
        value = json.loads(text) if text else default
    except ValueError:
        return default
    return value if isinstance(value, type(default)) else default


def convert_template(sheets, kg_revision_id, active, names):
    """v1 sheet_templates → v2 JSON 문법. 표현 불가하면 UNCONVERTIBLE_TEMPLATE."""

    def fail(reason):
        raise Problem("UNCONVERTIBLE_TEMPLATE", reason)

    if not 1 <= len(sheets) <= 16:
        fail(f"sheet templates must be 1~16, got {len(sheets)}")
    total = sum(len(s["mappings"]) for s in sheets)
    if not 1 <= total <= 200:
        fail(f"mappings must be 1~200, got {total}")
    keys = [m["mapping_key"] for s in sheets for m in s["mappings"]]
    unique = len(keys) == len(set(keys))
    detail = {"roles": {}, "rules": {}, "inferred_keys": [], "notes": []}
    roles, rules = {}, []
    for sheet in sheets:
        name, matcher = sheet["name"], sheet["matcher"]
        if not isinstance(name, str) or not name:
            fail("sheet template name is empty")
        listed = matcher.get("names")
        cardinality = (
            "one"
            if isinstance(listed, list)
            and len(listed) == 1
            and not matcher.get("name_regex")
            and not matcher.get("headers")
            else "many"
        )
        roles[name] = {"cardinality": cardinality}
        detail["roles"][name] = {"matcher": matcher, "cardinality": cardinality}
        for mapping in sheet["mappings"]:
            key = mapping["mapping_key"]
            rule_key = key if unique else f"{name}.{key}"
            if not rule_key or len(rule_key) > 128:
                fail(f"mapping {name}/{key}: rule key is empty or longer than 128")
            detail["rules"][f"{name}/{key}"] = rule_key
            source = mapping["source"]
            if not isinstance(source, dict):
                fail(f"mapping {key}: source is not an object")
            if source.get("range"):
                address = source["range"]
                try:
                    value = {
                        "areas": [{"sheet_role": name, "range": address}],
                        **shape(address),
                    }
                except Problem as exc:
                    fail(f"mapping {key}: {exc.code}: {exc.message}")
                canonical, aliases = names.get(mapping.get("concept_id"), (None, []))
                texts = _texts([key, canonical, *aliases])
                if not texts:
                    fail(f"mapping {key}: no key text to search")
                selector = {
                    "key": {"areas": [{"sheet_role": name, "find": {"texts": texts}}]},
                    "value": value,
                }
                detail["inferred_keys"].append(rule_key)
            elif source.get("key_search"):
                terms = source["key_search"]
                texts = _texts(terms if isinstance(terms, list) else [terms])
                if not texts:
                    fail(f"mapping {key}: key_search is empty")
                offset = source.get("offset") or {}
                if not isinstance(offset, dict):
                    fail(f"mapping {key}: offset is not an object")
                try:
                    row, col = int(offset.get("row", 0)), int(offset.get("col", 1))
                except (TypeError, ValueError):
                    fail(f"mapping {key}: offset must be integers")
                selector = {
                    "key": {"areas": [{"sheet_role": name, "find": {"texts": texts}}]},
                    "value": {
                        "areas": [
                            {
                                "sheet_role": name,
                                "relative": {"row": row, "col": col, "rows": 1, "cols": 1},
                            }
                        ],
                        "cardinality": "scalar",
                        "axis": "none",
                    },
                }
                # v1은 시트 전체를 훑었지만 v2의 기본 검색 창은 A1:AZ100이다.
                detail["notes"].append(
                    f"{rule_key}: key search limited to A1:AZ100 in v2 (v1 scanned the whole sheet); verify anchor"
                )
            else:
                fail(f"mapping {key}: source has neither range nor key_search")
            value_spec = {
                "type": {"number": "decimal", "text": "text"}.get(
                    mapping.get("value_type"), "text"
                )
            }
            if mapping.get("unit"):
                value_spec["unit"] = mapping["unit"]
            target = (mapping.get("normalization") or {}).get("target_unit")
            if target and target != mapping.get("unit"):
                detail["notes"].append(
                    f"{rule_key}: unit conversion {mapping.get('unit')}→{target} needs explicit affine; review"
                )
            concept = mapping.get("concept_id")
            if concept and concept not in active:
                detail["notes"].append(
                    f"{rule_key}: concept {concept} missing/deprecated in v2 KG"
                )
                concept = None
            rules.append(
                {
                    "rule_key": rule_key,
                    "concept_id": concept or None,
                    "selector": selector,
                    "value_spec": value_spec,
                }
            )
    definition = {
        "kg_revision_id": kg_revision_id,
        "format": "json",
        "sheet_roles": roles,
        "rules": rules,
    }
    try:
        validate_template(definition)
    except Problem as exc:
        fail(f"{exc.code}: {exc.message}")
    return definition, detail


class Migration:
    def __init__(self, ws, from_ws, raw=None, dry_run=False, principal="v1-migration"):
        self.ws = Path(ws)
        self.from_ws = Path(from_ws)
        self.v1_path = self.from_ws / V1_DB
        self.raw = Path(raw) if raw else self.from_ws / "data/raw"
        self.v2_raw = (self.ws / "data/raw").resolve()
        self.principal = principal
        self.dry_run = dry_run
        self.kg = None
        self.kg_planned = False
        self.active = set()
        self.names = {}
        self.links = {}
        self.service = None
        self.v1 = None
        self.report = {
            "format": "v1-migration-report/1",
            "source_db": str(self.v1_path.resolve()),
            "workspace": str(self.ws.resolve()),
            "raw_dir": str(self.raw.resolve()),
            "dry_run": dry_run,
            "principal": principal,
            "started_at": now(),
            "finished_at": None,
            "counts": {key: dict.fromkeys(STATUSES, 0) for key in COUNT_KEYS.values()},
            "skipped": [],
            "links": [],
            "re_extract_required": [],
        }
        self.report["counts"]["values"] = {
            "parsed_sources": 0,
            "payload_values": 0,
            "copied": 0,
        }

    # ------------------------------------------------------------ plumbing --
    @staticmethod
    def _readonly(path):
        conn = sqlite3.connect(f"file:{Path(path).as_posix()}?mode=ro", uri=True)
        conn.row_factory = sqlite3.Row
        return conn

    def _v2(self):
        # 적용 시 Service 연결, dry-run에서는 기존 v2.db가 있을 때만 읽기 전용 연결.
        if self.service is not None:
            return self.service.db.connect()
        path = self.ws / "data/kg/v2.db"
        from contextlib import closing, nullcontext

        if not path.is_file():
            return nullcontext(None)
        return closing(self._readonly(path))

    def _load_links(self):
        with self._v2() as conn:
            if conn is None:
                return
            # 같은 (entity, v1_id)의 나중 링크가 이전 needs_review 링크를 대체한다.
            for row in conn.execute(
                "SELECT artifact_id,storage_ref FROM artifact WHERE kind='manifest' AND policy_ref=? ORDER BY created_at,artifact_id",
                (LINK_POLICY,),
            ):
                link = _loads(row["storage_ref"], {})
                if link.get("schema") != LINK_SCHEMA or "entity" not in link:
                    continue
                link["artifact_id"] = row["artifact_id"]
                self.links[(link["entity"], str(link["v1_id"]))] = link

    def _count(self, entity, status):
        self.report["counts"][COUNT_KEYS[entity]][status] += 1

    def _link(self, entity, v1_id, v2_id, status, detail=None):
        detail = detail or {}
        if self.dry_run and status == "migrated":
            status = "planned"
        link = {
            "schema": LINK_SCHEMA,
            "source_db": self.report["source_db"],
            "entity": entity,
            "v1_id": str(v1_id),
            "v2_id": v2_id,
            "status": status,
            "detail": detail,
            "migrated_at": now(),
        }
        if status in ("migrated", "needs_review") and not self.dry_run:
            text = dump(link)
            with self.service.db.connect(write=True) as conn:
                insert(
                    conn,
                    "artifact",
                    artifact_id=uid(),
                    kind="manifest",
                    storage_ref=text,
                    sha256=hashlib.sha256(text.encode()).hexdigest(),
                    byte_size=len(text.encode()),
                    media_type="application/json",
                    policy_ref=LINK_POLICY,
                    created_at=now(),
                )
        self.links[(entity, str(v1_id))] = link
        self.report["links"].append(
            {
                "entity": entity,
                "v1_id": str(v1_id),
                "v2_id": v2_id,
                "status": status,
                "detail": detail,
            }
        )
        self._count(entity, status)
        return link

    def _existing(self, entity, v1_id):
        link = self.links.get((entity, str(v1_id)))
        if link is None:
            return None
        self.report["links"].append(
            {
                "entity": entity,
                "v1_id": str(v1_id),
                "v2_id": link.get("v2_id"),
                "status": "existing",
                "detail": link.get("detail") or {},
            }
        )
        self._count(entity, "existing")
        return link

    def _skip(self, entity, v1_id, reason):
        self.report["skipped"].append(
            {"entity": entity, "v1_id": str(v1_id), "reason": reason}
        )
        self._count(entity, "skipped")

    def _source_artifact(self, entity, v1_id, spec, reason):
        if self.dry_run:
            return None
        text = dump({"entity": entity, "v1_id": v1_id, "spec": spec, "reason": reason})
        artifact_id = uid()
        with self.service.db.connect(write=True) as conn:
            insert(
                conn,
                "artifact",
                artifact_id=artifact_id,
                kind="manifest",
                storage_ref=text,
                sha256=hashlib.sha256(text.encode()).hexdigest(),
                byte_size=len(text.encode()),
                media_type="application/json",
                policy_ref=SOURCE_POLICY,
                created_at=now(),
            )
        return artifact_id

    # --------------------------------------------------------------- run ----
    def run(self):
        if not self.v1_path.is_file():
            raise Problem(
                "V1_DATABASE_MISSING", f"v1 kg.db가 없습니다: {self.v1_path}", 404
            )
        self.v1 = self._readonly(self.v1_path)
        tables = {
            r[0] for r in self.v1.execute("SELECT name FROM sqlite_master WHERE type='table'")
        }
        required = {"domain_concept", "domain_relation", "domain_alias", "document", "document_version"}
        if not required <= tables:
            self.v1.close()
            raise Problem(
                "V1_DATABASE_INVALID",
                f"v1 kg.db 스키마가 아닙니다 (없는 테이블: {', '.join(sorted(required - tables))}): {self.v1_path}",
            )
        if not self.dry_run:
            from .service import Service

            self.service = Service(self.ws)
        try:
            self._load_links()
            self._kg()
            self._documents()
            self._templates()
            self._assignments()
            self._values()
        finally:
            self.v1.close()
        self.report["finished_at"] = now()
        from .contracts import MigrationReport

        return MigrationReport(**self.report).model_dump()

    # ---------------------------------------------------------------- kg ----
    def _kg(self):
        v1 = self.v1
        concepts = [
            dict(r)
            for r in v1.execute(
                "SELECT concept_id,canonical_name,canonical_name_en,description,concept_type,data_type,domain_level,canonical_unit,unit_dimension,status FROM domain_concept ORDER BY concept_id"
            )
        ]
        aliases = {}
        for row in v1.execute(
            "SELECT concept_id,alias_text FROM domain_alias ORDER BY concept_id,alias_norm"
        ):
            aliases.setdefault(row["concept_id"], []).append(row["alias_text"])
        relations = [
            list(r)
            for r in v1.execute(
                "SELECT source_concept_id,target_concept_id,relation_type FROM domain_relation ORDER BY source_concept_id,target_concept_id,relation_type"
            )
        ]
        units = v1.execute("SELECT count(*) FROM unit").fetchone()[0]
        if units:
            self._skip("units", "unit", "v2 has no unit table; units.yaml stays the source")
        notes, definition = [], {"concepts": [], "relations": []}
        levels = {}
        for c in concepts:
            match = re.fullmatch(r"L?(\d+)", str(c["domain_level"] or "L1").strip())
            level = int(match.group(1)) if match else 1
            if not match:
                notes.append(
                    f"level '{c['domain_level']}' of {c['concept_id']} not parsable; set to 1"
                )
            if level < 1:
                level = 1
            levels[c["concept_id"]] = level
            status = (
                "deprecated" if str(c["status"] or "").upper() == "DEPRECATED" else "active"
            )
            definition["concepts"].append(
                {
                    "concept_id": c["concept_id"],
                    "name": c["canonical_name"],
                    "definition": c["description"] or c["canonical_name"],
                    "level": level,
                    "value_type": c["data_type"],
                    "canonical_unit": c["canonical_unit"],
                    "status": status,
                    "aliases": aliases.get(c["concept_id"], []),
                }
            )
            self.names[c["concept_id"]] = (
                c["canonical_name"],
                aliases.get(c["concept_id"], []),
            )
            if status == "active":
                self.active.add(c["concept_id"])
        ids = {c["concept_id"] for c in concepts}
        translated = 0
        for source, target, kind in relations:
            if source == target or source not in ids or target not in ids:
                notes.append(f"relation {source}→{target} {kind} dropped (self edge or unknown endpoint)")
                continue
            # v1 IS_A는 자식→부모, v2 parent_of는 부모→자식이며 레벨이 정확히 1 커야 한다(설계 §4.9).
            # 그 조건을 만족하는 IS_A만 방향을 뒤집어 parent_of로 옮기고, 나머지는 그대로 복사한다.
            if kind == "IS_A" and levels[source] == levels[target] + 1:
                definition["relations"].append([target, source, "parent_of"])
                translated += 1
                continue
            if kind == "IS_A":
                notes.append(
                    f"IS_A {source}(L{levels[source]})→{target}(L{levels[target]}) kept verbatim: level gap is not 1"
                )
            definition["relations"].append([source, target, kind])
        if not definition["concepts"]:
            self._skip("kg", "kg", "no concepts in v1")
            return
        dropped = sorted(
            {
                field
                for c in concepts
                for field in ("canonical_name_en", "concept_type", "unit_dimension")
                if c[field]
            }
        )
        v1_id = digest(definition)
        existing = self._existing("kg", v1_id)
        if existing:
            self.kg = existing["v2_id"]
            return
        detail = {
            "concepts": len(definition["concepts"]),
            "aliases": sum(len(c["aliases"]) for c in definition["concepts"]),
            "relations": len(definition["relations"]),
            "parent_of_from_is_a": translated,
            "dropped_fields": dropped,
            "notes": notes,
        }
        if self.dry_run:
            with self._v2() as conn:
                found = (
                    conn.execute(
                        "SELECT kg_revision_id FROM kg_revision WHERE content_sha256=?",
                        (v1_id,),
                    ).fetchone()
                    if conn is not None
                    else None
                )
            if found:
                self.kg = found[0]
                self._link("kg", v1_id, self.kg, "existing", detail)
            else:
                self.kg = None
                self._link("kg", v1_id, None, "migrated", detail)
            self.kg_planned = True
            return
        result = self.service.import_kg(definition, self.principal)
        self.kg = result["kg_revision_id"]
        detail["revision_no"] = result["revision_no"]
        with self.service.db.connect() as conn:
            self.active = {
                r[0]
                for r in conn.execute(
                    "SELECT concept_id FROM domain_concept WHERE kg_revision_id=? AND status='active'",
                    (self.kg,),
                )
            }
        self._link("kg", v1_id, self.kg, "migrated", detail)

    # --------------------------------------------------------- documents ----
    def _locate(self, filename, filepath):
        # v1의 절대경로는 대개 v2 raw 밖(이전 워크스페이스)을 가리키므로, v2 raw에 놓인 사본을
        # 먼저 찾고 절대경로는 마지막 후보로만 쓴다. 바이트 해시 검증은 호출부가 수행한다.
        candidates = []
        path = Path(filepath) if filepath else None
        if path is not None and not path.is_absolute():
            candidates += [self.v2_raw / path, self.raw / path]
        if path is not None:
            candidates += [self.v2_raw / path.name, self.raw / path.name]
        if filename:
            candidates += [self.v2_raw / filename, self.raw / filename]
        if path is not None and path.is_absolute():
            candidates.append(path)
        seen = set()
        for candidate in candidates:
            if candidate in seen:
                continue
            seen.add(candidate)
            if candidate.is_file():
                return candidate
        return None

    def _link_duplicates(self, duplicates, v2_version_id, actual, registrable):
        for dup in duplicates:
            self._link(
                "document_version",
                dup["version_id"],
                v2_version_id,
                "migrated",
                {
                    "file_hash": actual,
                    "parser_version": dup["parser_version"],
                    "v1_parsed_at": dup["parsed_at"],
                    "same_bytes_as": registrable["version_id"],
                },
            )

    def _documents(self):
        from .readers import file_hash

        for doc in self.v1.execute(
            "SELECT document_id,filename,filepath,file_type,current_version FROM document ORDER BY document_id"
        ):
            doc = dict(doc)
            versions = [
                dict(r)
                for r in self.v1.execute(
                    "SELECT version_id,document_id,file_hash,parser_version,parsed_at FROM document_version WHERE document_id=? ORDER BY parsed_at,version_id",
                    (doc["document_id"],),
                )
            ]
            doc_link = self.links.get(("document", doc["document_id"]))
            if not versions:
                self._skip("document", doc["document_id"], "no v1 versions")
                continue
            pending = [v for v in versions if ("document_version", v["version_id"]) not in self.links]
            for v in versions:
                if v not in pending:
                    self._existing("document_version", v["version_id"])
            if doc_link and not pending:
                self._existing("document", doc["document_id"])
                continue
            path = self._locate(doc["filename"], doc["filepath"])
            if path is None:
                self._skip_document(doc, pending, doc_link, "source missing")
                continue
            resolved = path.resolve()
            if not resolved.is_relative_to(self.v2_raw):
                self._skip_document(
                    doc, pending, doc_link, f"source outside workspace raw: {resolved}"
                )
                continue
            source_ref = str(resolved.relative_to(self.v2_raw))
            actual = file_hash(resolved)
            matching = [v for v in pending if v["file_hash"] == actual]
            for v in pending:
                if v["file_hash"] != actual:
                    self._skip(
                        "document_version",
                        v["version_id"],
                        f"source changed (v1 hash {v['file_hash'][:8]}… ≠ current {actual[:8]}…)",
                    )
            # 같은 바이트를 가리키는 v1 버전이 여럿이면(A→B→A 이력) 한 번만 등록하고 전부 그 v2 버전에 연결한다.
            registrable = matching[-1] if matching else None
            duplicates = matching[:-1]
            if registrable is None:
                if not doc_link:
                    self._skip("document", doc["document_id"], "no version matches current bytes")
                else:
                    self._existing("document", doc["document_id"])
                continue
            if doc_link:
                self._existing("document", doc["document_id"])
            if self.dry_run:
                from openpyxl import load_workbook

                try:
                    wb = load_workbook(resolved, read_only=True, keep_links=False)
                    sheets = list(wb.sheetnames)
                    wb.close()
                except Exception as exc:  # 잠긴/손상 파일은 등록 시 Reader가 거부한다.
                    self._skip("document_version", registrable["version_id"], f"unreadable: {exc}")
                    if not doc_link:
                        self._skip("document", doc["document_id"], "unreadable source")
                    continue
                if not doc_link:
                    self._link(
                        "document",
                        doc["document_id"],
                        None,
                        "migrated",
                        {"source_ref": source_ref, "filename": doc["filename"]},
                    )
                self._link(
                    "document_version",
                    registrable["version_id"],
                    None,
                    "migrated",
                    {
                        "file_hash": actual,
                        "parser_version": registrable["parser_version"],
                        "v1_parsed_at": registrable["parsed_at"],
                        "v2_document_id": None,
                        "sheets": sheets,
                    },
                )
                self._link_duplicates(duplicates, None, actual, registrable)
                continue
            with self.service.db.connect() as conn:
                found = conn.execute(
                    "SELECT 1 FROM document WHERE provider=? AND source_ref=?",
                    (PROVIDER, source_ref),
                ).fetchone()
            try:
                result = self.service.register(source_ref, PROVIDER, self.principal)
            except Problem as exc:
                self._skip(
                    "document_version", registrable["version_id"], f"{exc.code}: {exc.message}"
                )
                if not doc_link:
                    self._skip("document", doc["document_id"], f"{exc.code}: {exc.message}")
                continue
            if not doc_link:
                self._link(
                    "document",
                    doc["document_id"],
                    result["document_id"],
                    "migrated",
                    {
                        "source_ref": source_ref,
                        "filename": doc["filename"],
                        "existing_v2_document": bool(found or result.get("unchanged")),
                    },
                )
            self._link(
                "document_version",
                registrable["version_id"],
                result["version_id"],
                "migrated",
                {
                    "file_hash": actual,
                    "parser_version": registrable["parser_version"],
                    "v1_parsed_at": registrable["parsed_at"],
                    "v2_document_id": result["document_id"],
                },
            )
            self._link_duplicates(duplicates, result["version_id"], actual, registrable)

    def _skip_document(self, doc, pending, doc_link, reason):
        for v in pending:
            self._skip("document_version", v["version_id"], reason)
        if doc_link:
            self._existing("document", doc["document_id"])
        else:
            self._skip("document", doc["document_id"], reason)

    # --------------------------------------------------------- templates ----
    def _v1_sheets(self, template_id, version):
        sheets = []
        for st in self.v1.execute(
            "SELECT sheet_template_id,name,matcher_json,ordinal FROM sheet_template WHERE template_id=? AND template_version=? ORDER BY ordinal",
            (template_id, version),
        ):
            mappings = [
                {
                    "mapping_id": m["mapping_id"],
                    "mapping_key": m["mapping_key"],
                    "concept_id": m["concept_id"],
                    "document_kg_node": m["document_kg_node"],
                    "source": _loads(m["source_json"], {}),
                    "value_type": m["value_type"],
                    "unit": m["unit"],
                    "normalization": _loads(m["normalization_json"], {}),
                }
                for m in self.v1.execute(
                    "SELECT mapping_id,mapping_key,concept_id,document_kg_node,source_json,value_type,unit,normalization_json FROM template_mapping WHERE sheet_template_id=? ORDER BY mapping_key",
                    (st["sheet_template_id"],),
                )
            ]
            sheets.append(
                {
                    "name": st["name"],
                    "matcher": _loads(st["matcher_json"], {}),
                    "mappings": mappings,
                }
            )
        return sheets

    def _recover_template_version(self, name, definition, template_id):
        # 링크를 기록하기 전에 중단된 이전 실행이 만든 같은 정의의 버전이 있으면 새로 만들지 않고 재사용한다.
        spec = validate_template(definition)
        with self.service.db.connect() as conn:
            row = conn.execute(
                "SELECT tv.template_id,tv.template_version_id,tv.revision_no FROM template_version tv JOIN template t USING(template_id) WHERE tv.definition_sha256=? AND tv.created_by=? AND t.name=? AND t.created_by=? ORDER BY tv.created_at DESC LIMIT 1",
                (digest(spec), self.principal, name, self.principal),
            ).fetchone()
        if row is None or (template_id and row["template_id"] != template_id):
            return None
        return {
            "template_id": row["template_id"],
            "template_version_id": row["template_version_id"],
            "revision_no": row["revision_no"],
        }

    def _templates(self):
        for t in self.v1.execute(
            "SELECT template_id,name,target_document_kg,lifecycle,created_at FROM parsing_template ORDER BY template_id"
        ):
            t = dict(t)
            tid = t["template_id"]
            prior = self.links.get(("template", tid))
            template_link = prior if prior and prior.get("v2_id") else None
            if template_link:
                self._existing("template", tid)
            v2_template_id = template_link["v2_id"] if template_link else None
            versions = [
                dict(r)
                for r in self.v1.execute(
                    "SELECT template_id,version,spec_json,created_at,created_by FROM parsing_template_version WHERE template_id=? ORDER BY version",
                    (tid,),
                )
            ]
            for v in versions:
                v1_id = f"{tid}@{v['version']}"
                if self._existing("template_version", v1_id):
                    continue
                if self.kg is None and not self.kg_planned:
                    self._skip("template_version", v1_id, "no KG revision")
                    continue
                sheets = self._v1_sheets(tid, v["version"])
                try:
                    definition, detail = convert_template(
                        sheets, self.kg, self.active, self.names
                    )
                except Problem as exc:
                    spec = _loads(v["spec_json"], {}) or {"sheet_templates": sheets}
                    artifact_id = self._source_artifact(
                        "template_version", v1_id, spec, exc.message
                    )
                    self._link(
                        "template_version",
                        v1_id,
                        None,
                        "needs_review",
                        {
                            "reason": exc.message,
                            "source_artifact_id": artifact_id,
                            "v1_version": v["version"],
                        },
                    )
                    continue
                detail.update(
                    {
                        "v1_version": v["version"],
                        "created_by": v["created_by"],
                        "v1_created_at": v["created_at"],
                    }
                )
                if self.dry_run:
                    detail["v2_template_id"] = None
                    if not template_link:
                        template_link = self._link(
                            "template",
                            tid,
                            None,
                            "migrated",
                            {
                                "name": t["name"],
                                "lifecycle": t["lifecycle"],
                                "target_document_kg": t["target_document_kg"],
                                "v1_created_at": t["created_at"],
                            },
                        )
                    self._link("template_version", v1_id, None, "migrated", detail)
                    continue
                result = self._recover_template_version(t["name"], definition, v2_template_id)
                if result:
                    detail["recovered"] = True
                else:
                    try:
                        result = self.service.create_template(
                            t["name"], definition, self.principal, template_id=v2_template_id
                        )
                    except sqlite3.IntegrityError as exc:
                        self._skip("template_version", v1_id, f"INTEGRITY_CONFLICT: {exc}")
                        continue
                    except Problem as exc:
                        self._skip("template_version", v1_id, f"{exc.code}: {exc.message}")
                        continue
                v2_template_id = result["template_id"]
                if not template_link:
                    template_link = self._link(
                        "template",
                        tid,
                        v2_template_id,
                        "migrated",
                        {
                            "name": t["name"],
                            "lifecycle": t["lifecycle"],
                            "target_document_kg": t["target_document_kg"],
                            "v1_created_at": t["created_at"],
                        },
                    )
                detail.update(
                    {
                        "v2_revision_no": result["revision_no"],
                        "v2_template_id": v2_template_id,
                    }
                )
                self._link(
                    "template_version", v1_id, result["template_version_id"], "migrated", detail
                )
            if template_link:
                continue
            if prior:
                self._existing("template", tid)
            elif versions:
                self._link(
                    "template", tid, None, "needs_review", {"reason": "no convertible version"}
                )
            else:
                self._skip("template", tid, "no versions")

    # ------------------------------------------------------- assignments ----
    def _sheets_of(self, dv):
        if not dv.get("v2_id"):
            return [
                {"sheet_id": None, "name": name, "ordinal": n}
                for n, name in enumerate((dv.get("detail") or {}).get("sheets") or [])
            ]
        with self._v2() as conn:
            if conn is None:
                return []
            return [
                dict(r)
                for r in conn.execute(
                    "SELECT sheet_id,name,ordinal FROM sheet WHERE document_version_id=? ORDER BY ordinal",
                    (dv["v2_id"],),
                )
            ]

    @staticmethod
    def _bind(roles, sheets):
        bindings, names = {}, {}
        for role, info in roles.items():
            matcher = info.get("matcher") or {}
            listed, regex = matcher.get("names") or [], matcher.get("name_regex")
            if not isinstance(listed, list):
                listed = [listed]
            if not listed and not regex:
                raise Problem(
                    "UNBINDABLE_ROLE", "sheet matcher requires cell content (headers)"
                )
            try:
                matched = [
                    s
                    for s in sheets
                    if (listed and s["name"] in listed)
                    or (regex and re.search(regex, s["name"]))
                ]
            except re.error as exc:
                raise Problem("UNBINDABLE_ROLE", f"role {role}: invalid name_regex ({exc})")
            if info.get("cardinality") == "one" and len(matched) != 1:
                raise Problem(
                    "UNBINDABLE_ROLE", f"role {role} matched {len(matched)} sheets"
                )
            if not matched:
                raise Problem("UNBINDABLE_ROLE", f"role {role} matched no sheet")
            bindings[role] = [s["sheet_id"] for s in matched]
            names[role] = [s["name"] for s in matched]
        return bindings, names

    def _assignments(self):
        rows = [
            dict(r)
            for r in self.v1.execute(
                "SELECT document_id,document_version,template_id,template_version,status,assigned_at FROM document_template_assignment ORDER BY document_id,document_version,template_id"
            )
        ]
        for a in rows:
            v1_id = f"{a['document_version']}@{a['template_id']}"
            dv = self.links.get(("document_version", a["document_version"]))
            tv = self.links.get(
                ("template_version", f"{a['template_id']}@{a['template_version']}")
            )
            existing = self._existing("assignment", v1_id)
            if existing:
                if dv and tv and tv.get("status") != "needs_review":
                    self._overrides(a, existing["v2_id"], tv, dv)
                continue
            if not dv or (not dv.get("v2_id") and not self.dry_run):
                self._skip("assignment", v1_id, "document version not migrated")
                continue
            if not tv:
                self._skip("assignment", v1_id, "template version not migrated")
                continue
            if tv.get("status") == "needs_review":
                self._skip("assignment", v1_id, "template version needs review")
                continue
            roles = (tv.get("detail") or {}).get("roles") or {}
            try:
                bindings, names = self._bind(roles, self._sheets_of(dv))
            except Problem as exc:
                self._skip("assignment", v1_id, exc.message)
                continue
            detail = {
                "v1_status": a["status"],
                "assigned_at": a["assigned_at"],
                "bindings": names,
                "v2_document_version_id": dv.get("v2_id"),
                "v2_template_version_id": tv.get("v2_id"),
            }
            if self.dry_run:
                self._link("assignment", v1_id, None, "migrated", detail)
                self._overrides(a, None, tv, dv, names)
                continue
            app = None
            try:
                app = self.service.apply_template(
                    dv["v2_id"], tv["v2_id"], bindings, SCOPE_KEY, False, self.principal
                )["application_id"]
            except sqlite3.IntegrityError:
                with self.service.db.connect() as conn:
                    found = conn.execute(
                        "SELECT application_id FROM template_application WHERE document_version_id=? AND template_version_id=? AND scope_key=?",
                        (dv["v2_id"], tv["v2_id"], SCOPE_KEY),
                    ).fetchone()
                if not found:
                    self._skip("assignment", v1_id, "INTEGRITY_CONFLICT: application exists")
                    continue
                app = found[0]
                detail["recovered"] = True
            except Problem as exc:
                self._skip("assignment", v1_id, f"{exc.code}: {exc.message}")
                continue
            self._link("assignment", v1_id, app, "migrated", detail)
            self._overrides(a, app, tv, dv, names)

    def _overrides(self, a, app, tv, dv, names=None):
        if names is None:
            names = ((self.links.get(("assignment", f"{a['document_version']}@{a['template_id']}")) or {}).get("detail") or {}).get("bindings") or {}
        rules = (tv.get("detail") or {}).get("rules") or {}
        seen = set()
        for o in self.v1.execute(
            "SELECT o.override_id,o.override_source_json,o.status,o.reason,o.created_by,o.updated_at,st.name sheet_name,tm.mapping_key FROM document_override o JOIN template_mapping tm ON tm.mapping_id=o.template_mapping_id JOIN sheet_template st ON st.sheet_template_id=tm.sheet_template_id WHERE o.document_version=? AND st.template_id=? ORDER BY o.updated_at DESC,o.override_id DESC",
            (a["document_version"], a["template_id"]),
        ):
            o = dict(o)
            oid = o["override_id"]
            if self._existing("override", oid):
                continue
            if o["status"] not in ("APPROVED", "CONFLICT"):
                self._skip("override", oid, f"v1 status {o['status']}")
                continue
            slot = (o["sheet_name"], o["mapping_key"])
            if slot in seen:
                self._skip("override", oid, "superseded by a newer override")
                continue
            seen.add(slot)
            rule_key = rules.get(f"{o['sheet_name']}/{o['mapping_key']}")
            if not rule_key:
                self._skip("override", oid, "mapping not in converted template")
                continue
            role = o["sheet_name"]
            src = _loads(o["override_source_json"], {})
            try:
                value = {
                    "areas": [{"sheet_role": role, "range": src.get("range")}],
                    **shape(src.get("range")),
                }
            except Problem as exc:
                self._skip("override", oid, f"{exc.code}: {exc.message}")
                continue
            if src.get("sheet") and src["sheet"] not in (names.get(role) or []):
                self._skip("override", oid, "override sheet not bound to role")
                continue
            detail = {
                "v2_application_id": app,
                "rule_key": rule_key,
                "v1_status": o["status"],
                "v1_updated_at": o["updated_at"],
            }
            if self.dry_run:
                self._link("override", oid, None, "migrated", detail)
                continue
            with self.service.db.connect() as conn:
                latest = conn.execute(
                    "SELECT mapping_revision_id,concept_id,effective_spec_json,reason,created_by FROM mapping_revision WHERE application_id=? AND rule_key=? ORDER BY revision_no DESC LIMIT 1",
                    (app, rule_key),
                ).fetchone()
            if latest is None:
                self._skip("override", oid, "rule has no mapping revision")
                continue
            reason = f"v1 override {oid}: {o['reason'] or ''}".strip()
            if latest["created_by"] != self.principal:
                # 사람이 v2에서 이미 손댄 규칙(proposed/rejected 포함)은 덮어쓰지 않는다 (§6).
                self._skip("override", oid, "EDIT_CONFLICT: human revision exists")
                continue
            if latest["reason"] == reason:
                # 링크 기록 전에 중단된 이전 실행의 리비전을 재사용한다.
                self._link("override", oid, latest["mapping_revision_id"], "migrated", detail)
                continue
            spec = json.loads(latest["effective_spec_json"])
            spec.setdefault("selector", {})["value"] = value
            try:
                new = self.service.revise(
                    app,
                    latest["mapping_revision_id"],
                    0,
                    spec,
                    latest["concept_id"],
                    "proposed",
                    reason,
                    self.principal,
                )
            except Problem as exc:
                self._skip("override", oid, f"{exc.code}: {exc.message}")
                continue
            self._link("override", oid, new["mapping_revision_id"], "migrated", detail)

    # ------------------------------------------------------------ values ----
    def _values(self):
        total = 0
        for r in self.v1.execute(
            "SELECT r.parse_run_id,r.document_id,r.document_version,r.template_id,r.template_version,r.status,count(s.parsed_source_id) sources FROM parse_run r LEFT JOIN parsed_source s ON s.parse_run_id=r.parse_run_id AND s.value_json IS NOT NULL AND s.value_json<>'null' GROUP BY r.parse_run_id ORDER BY r.started_at,r.parse_run_id"
        ):
            link = self.links.get(("assignment", f"{r['document_version']}@{r['template_id']}"))
            total += r["sources"]
            self.report["re_extract_required"].append(
                {
                    "v1_parse_run_id": r["parse_run_id"],
                    "v1_document_id": r["document_id"],
                    "v1_document_version": r["document_version"],
                    "v1_template_id": r["template_id"],
                    "v1_template_version": r["template_version"],
                    "v1_status": r["status"],
                    "sources": r["sources"],
                    "v2_application_id": link.get("v2_id") if link else None,
                }
            )
        payload = self.v1.execute("SELECT count(*) FROM payload_value").fetchone()[0]
        self.report["counts"]["values"] = {
            "parsed_sources": total,
            "payload_values": payload,
            "copied": 0,
        }


def migrate(ws, from_ws, raw=None, dry_run=False, principal="v1-migration"):
    return Migration(ws, from_ws, raw, dry_run, principal).run()


def run_cli(args):
    try:
        report = migrate(args.ws, args.from_ws, args.raw, args.dry_run, args.principal)
    except Problem as exc:
        print(f"{exc.code}: {exc.message}", file=sys.stderr)
        return 2
    except sqlite3.Error as exc:
        print(f"V1_DATABASE_ERROR: {exc}", file=sys.stderr)
        return 2
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(
        json.dumps(
            {
                "dry_run": report["dry_run"],
                "counts": report["counts"],
                "skipped": len(report["skipped"]),
                "re_extract_required": len(report["re_extract_required"]),
                "report": str(args.report) if args.report else None,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0
